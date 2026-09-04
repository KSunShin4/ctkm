# -*- coding: utf-8 -*-
"""User-facing workbook layout helpers for the CTKM exporter.

This module changes presentation only. It never normalizes, infers, or replaces
source business values such as KEY Channel, Channel, item name, UoM, campaign,
dates, or prices.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd


TOTAL_SHEET_NAME = "TỔNG HỢP - TEM"
ACTION_SHEET_NAMES = {"Gắn tem", "Tháo tem", "Chờ gắn", "Cần review", TOTAL_SHEET_NAME}

USER_ACTION_COLUMNS = [
    "key_channel", "channel", "item_id", "item_name", "promo_uom",
    "pog_code", "location_id", "facing",
    "scheduled_action", "action", "ngay_gan_tem", "ngay_thao_tem", "so_ngay_con_lai",
    "campaign_name", "campaign_detail", "promo_from", "promo_to",
    "normal_price", "promo_price", "program_source", "reason",
    "product_type", "store_id",
    "location_completeness", "location_status", "location_needs_review",
    "processing_status", "processed_by", "processed_at", "staff_note",
]

HEADER_LABELS = {
    "key_channel": "KEY Channel",
    "channel": "Channel",
    "item_id": "Mã sản phẩm",
    "item_name": "Tên sản phẩm",
    "promo_uom": "Đơn vị sản phẩm (Promo UoM)",
    "pog_code": "POG",
    "location_id": "Vị trí",
    "facing": "Facing",
    "scheduled_action": "Action theo lịch",
    "action": "Action",
    "ngay_gan_tem": "Ngày gắn",
    "ngay_thao_tem": "Ngày tháo",
    "campaign_name": "Campaign",
    "campaign_detail": "Scheme detail",
    "promo_from": "From",
    "promo_to": "To",
    "normal_price": "Giá thường",
    "promo_price": "Giá KM",
    "program_source": "Nguồn",
    "reason": "Lý do / Cảnh báo",
}

_COLUMN_WIDTHS = {
    "key_channel": 18,
    "channel": 24,
    "item_id": 14,
    "item_name": 38,
    "promo_uom": 20,
    "pog_code": 16,
    "location_id": 12,
    "facing": 10,
    "scheduled_action": 18,
    "action": 18,
    "ngay_gan_tem": 14,
    "ngay_thao_tem": 14,
    "so_ngay_con_lai": 14,
    "campaign_name": 28,
    "campaign_detail": 42,
    "promo_from": 14,
    "promo_to": 14,
    "normal_price": 14,
    "promo_price": 14,
    "program_source": 16,
    "reason": 36,
    "product_type": 16,
    "store_id": 14,
}

_HEADER_GROUPS = {
    "key_channel": "1F4E78",
    "channel": "1F4E78",
    "item_id": "2F75B5",
    "item_name": "2F75B5",
    "promo_uom": "2F75B5",
    "pog_code": "5B9BD5",
    "location_id": "5B9BD5",
    "facing": "5B9BD5",
    "scheduled_action": "70AD47",
    "action": "70AD47",
    "ngay_gan_tem": "70AD47",
    "ngay_thao_tem": "70AD47",
    "so_ngay_con_lai": "70AD47",
    "campaign_name": "4472C4",
    "campaign_detail": "4472C4",
    "promo_from": "4472C4",
    "promo_to": "4472C4",
    "normal_price": "4472C4",
    "promo_price": "4472C4",
    "program_source": "7F8C8D",
    "reason": "C65911",
}


def reorder_user_action_columns(df: pd.DataFrame, technical_columns: list[str]) -> pd.DataFrame:
    """Put operational columns first and technical/audit columns last."""
    if df is None or df.empty:
        return df
    technical = set(technical_columns)
    ordered = [column for column in USER_ACTION_COLUMNS if column in df.columns]
    middle = [column for column in df.columns if column not in ordered and column not in technical]
    audit = [column for column in technical_columns if column in df.columns]
    result = df.loc[:, [*ordered, *middle, *audit]].copy()
    sort_cols = [
        column for column in ("store_id", "pog_code", "location_id", "item_id")
        if column in result.columns
    ]
    if sort_cols:
        result = result.sort_values(sort_cols, kind="stable", na_position="last")
    return result.reset_index(drop=True)


def build_total_action_frame(
    df_fefo_actions: pd.DataFrame | None,
    *,
    unique_work_tracking_rows,
    technical_columns: list[str],
) -> pd.DataFrame:
    """Build the main user view from report-scope tasks only."""
    if df_fefo_actions is None:
        return pd.DataFrame()
    actions = df_fefo_actions.copy()
    if "in_report_scope" in actions.columns:
        actions = actions[actions["in_report_scope"].fillna(False).astype(bool)].copy()
    tasks = unique_work_tracking_rows(actions)
    return reorder_user_action_columns(tasks, technical_columns)


def apply_user_sheet_style(ws, sheet_name: str) -> None:
    """Apply the approved compact, grouped header layout to user action sheets."""
    if ws is None or sheet_name not in ACTION_SHEET_NAMES:
        return

    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    # Freeze KEY Channel + Channel and the header row.
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions if ws.max_row >= 1 and ws.max_column >= 1 else None
    ws.sheet_view.showGridLines = False

    raw_headers = [str(cell.value or "") for cell in ws[1]]
    for index, raw_header in enumerate(raw_headers, start=1):
        fill_color = _HEADER_GROUPS.get(raw_header)
        if fill_color:
            cell = ws.cell(row=1, column=index)
            cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)
            cell.font = Font(bold=True, color="FFFFFF")
        width = _COLUMN_WIDTHS.get(raw_header)
        if width:
            ws.column_dimensions[get_column_letter(index)].width = width


def apply_header_labels(exporter_module) -> None:
    """Use concise labels while retaining raw source values unchanged."""
    for canonical, label in HEADER_LABELS.items():
        exporter_module._BILINGUAL_HEADERS[canonical] = label


def add_total_action_sheet(
    output_path: str | Path,
    df_fefo_actions: pd.DataFrame | None,
    *,
    exporter_module,
    unique_work_tracking_rows,
    technical_columns: list[str],
) -> None:
    """Add one first-sheet overview without changing existing audit sheets."""
    if df_fefo_actions is None:
        return

    from openpyxl import load_workbook
    from openpyxl.utils.dataframe import dataframe_to_rows

    total = build_total_action_frame(
        df_fefo_actions,
        unique_work_tracking_rows=unique_work_tracking_rows,
        technical_columns=technical_columns,
    )
    if total is None:
        return

    path = Path(output_path)
    workbook = load_workbook(path)
    if TOTAL_SHEET_NAME in workbook.sheetnames:
        del workbook[TOTAL_SHEET_NAME]
    ws = workbook.create_sheet(TOTAL_SHEET_NAME, 0)

    if total.empty and len(total.columns) == 0:
        ws.append(["Ghi chú"])
        ws.append(["Không có công việc CTKM trong phạm vi báo cáo."])
    else:
        for row in dataframe_to_rows(total, index=False, header=True):
            ws.append(list(row))

    # Reuse the existing exporter styling first, then apply the approved layout.
    exporter_module._apply_professional_style(ws, TOTAL_SHEET_NAME, total)
    exporter_module._hide_technical_columns(ws, action_sheet=True)
    apply_user_sheet_style(ws, TOTAL_SHEET_NAME)
    exporter_module._apply_bilingual_headers(ws)
    ws.sheet_properties.tabColor = "1F4E78"

    workbook.save(path)
    workbook.close()
