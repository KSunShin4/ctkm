# -*- coding: utf-8 -*-
"""All-promotion reconciliation surface for the final CTKM workbook.

Historically ``Tong hop`` and ``POG không có CTKM`` were driven by the Monthly
CTKM ↔ POG reconciliation only.  That is useful as an audit, but it makes a
Clearance-only SKU look missing even when the operational action pipeline has a
valid promotion for that product.

This contract keeps the Monthly audit explicitly, while making the user-facing
coverage surfaces mean what their names say:

* ``Tong hop`` = POG products covered by any in-scope KM (Monthly or Clearance)
* ``POG không có KM`` = POG products with no in-scope KM at all
* ``POG không có Monthly CTKM`` = the historical Monthly-only audit

A hidden ``POG không có CTKM`` compatibility sheet is retained temporarily for
older validators/frontends.  It mirrors the canonical all-KM gap, so it cannot
reintroduce the misleading Monthly-only meaning.
"""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from modules.item_code_identity import build_transition_identity, canonical_series


_CONTRACT_MARKER = "_ctkm_all_km_reconciliation_v1"
_MONTHLY_SHEET = "POG không có Monthly CTKM"
_ALL_KM_GAP_SHEET = "POG không có KM"
_LEGACY_SHEET = "POG không có CTKM"


def _blank(value) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(value, str) and not value.strip()


def _comparison_identity(core_module, frame: pd.DataFrame | None):
    base = frame if frame is not None else pd.DataFrame()
    return build_transition_identity(
        base,
        key_func=core_module.NM.make_item_key,
        is_sku=core_module.RC.is_sku,
    )


def _item_series(core_module, frame: pd.DataFrame | None, identity) -> pd.Series:
    if frame is None or frame.empty:
        return pd.Series(dtype=object)
    return canonical_series(
        frame,
        identity,
        key_func=core_module.NM.make_item_key,
    )


def _item_keys(core_module, frame: pd.DataFrame | None, identity) -> set[str]:
    series = _item_series(core_module, frame, identity)
    return {str(value) for value in series.dropna() if str(value).strip()}


def _scoped_actions(actions: pd.DataFrame | None) -> pd.DataFrame:
    if actions is None or actions.empty:
        return pd.DataFrame(columns=[] if actions is None else actions.columns)
    scoped = actions.copy()
    if "in_report_scope" in scoped.columns:
        scoped = scoped[scoped["in_report_scope"].fillna(False).astype(bool)].copy()
    if "pog_applicable" in scoped.columns:
        scoped = scoped[~scoped["pog_applicable"].fillna(True).eq(False)].copy()
    return scoped.reset_index(drop=True)


def _source_label(values: Iterable[object]) -> str:
    seen = []
    for value in values:
        text = str(value or "").strip().upper()
        if not text or text in seen:
            continue
        seen.append(text)
    ordered = [value for value in ("CTKM", "CLEARANCE") if value in seen]
    ordered.extend(value for value in seen if value not in ordered)
    return " + ".join(ordered) if ordered else "KM"


def _source_and_action_maps(core_module, actions: pd.DataFrame, identity):
    if actions.empty:
        return {}, {}
    keys = _item_series(core_module, actions, identity).reset_index(drop=True)
    sources: dict[str, list[object]] = defaultdict(list)
    representative: dict[str, dict] = {}
    for position, (_, row) in enumerate(actions.reset_index(drop=True).iterrows()):
        key = str(keys.iloc[position]) if position < len(keys) and pd.notna(keys.iloc[position]) else ""
        if not key:
            continue
        sources[key].append(row.get("program_source"))
        representative.setdefault(key, row.to_dict())
    return {key: _source_label(values) for key, values in sources.items()}, representative


def _filter_selected_columns(df: pd.DataFrame, selected_columns) -> pd.DataFrame:
    if df is None or df.empty:
        return df
    selected = {str(value).strip().lower() for value in (selected_columns or []) if str(value).strip()}
    if not selected:
        return df
    keep = []
    for column in df.columns:
        text = str(column).strip().lower()
        base = text.split("__", 1)[0]
        if text in selected or base in selected:
            keep.append(column)
    for required in ("item_id", "program_source", "store_id", "pog_code", "location_id", "channel"):
        if required in df.columns and required not in keep:
            keep.append(required)
    return df.loc[:, keep].copy() if keep else df


def build_all_km_reconciliation(
    core_module,
    df_matched: pd.DataFrame | None,
    df_pog_only_monthly: pd.DataFrame | None,
    df_fefo_actions: pd.DataFrame | None,
    *,
    selected_output_columns=None,
):
    """Return (all-KM matched, all-KM gap, diagnostics).

    Monthly matched rows remain the baseline.  A Monthly-unmatched POG row is
    promoted into ``Tong hop`` when an in-scope CTKM/Clearance action proves that
    the same canonical SKU has a promotion.  The POG row itself supplies the
    physical store/POG/location fields, so Clearance-only products do not lose
    their shelf location merely because Monthly reconciliation did not match.
    """
    matched = df_matched.copy() if df_matched is not None else pd.DataFrame()
    pog_only = (
        df_pog_only_monthly.copy()
        if df_pog_only_monthly is not None
        else pd.DataFrame()
    )
    actions = _scoped_actions(df_fefo_actions)
    identity = _comparison_identity(core_module, actions)

    matched_keys = _item_keys(core_module, matched, identity)
    pog_keys_series = _item_series(core_module, pog_only, identity)
    pog_keys = {str(value) for value in pog_keys_series.dropna() if str(value).strip()}
    action_keys = _item_keys(core_module, actions, identity)
    source_by_key, action_by_key = _source_and_action_maps(core_module, actions, identity)

    if not matched.empty:
        if "program_source" not in matched.columns:
            matched["program_source"] = "CTKM"
        else:
            matched["program_source"] = matched["program_source"].where(
                ~matched["program_source"].apply(_blank), "CTKM"
            )

    if pog_only.empty:
        covered_monthly_gaps = pog_only.copy()
        all_km_gap = pog_only.copy()
    else:
        covered_mask = pog_keys_series.astype(str).isin(action_keys)
        covered_monthly_gaps = pog_only.loc[covered_mask].copy()
        all_km_gap = pog_only.loc[~covered_mask].copy()

        covered_keys = pog_keys_series.loc[covered_mask].astype(str).to_list()
        covered_monthly_gaps["program_source"] = [
            source_by_key.get(key, "KM") for key in covered_keys
        ]
        covered_monthly_gaps["coverage_basis"] = "KM ngoài Monthly vẫn phủ SKU POG"

        enrich_columns = (
            "campaign_name", "campaign_detail", "promo_from", "promo_to",
            "normal_price", "promo_price", "key_channel", "channel",
            "product_type", "promotion_id",
        )
        for column in enrich_columns:
            values = []
            for key, (_, row) in zip(covered_keys, covered_monthly_gaps.iterrows()):
                current = row.get(column) if column in covered_monthly_gaps.columns else None
                if not _blank(current):
                    values.append(current)
                    continue
                values.append(action_by_key.get(key, {}).get(column))
            if column in covered_monthly_gaps.columns:
                covered_monthly_gaps[column] = values
            elif any(not _blank(value) for value in values):
                covered_monthly_gaps[column] = values

    # Monthly matched and Monthly-unmatched are disjoint by reconciliation
    # contract, but transition aliases can make their canonical keys converge.
    # Do not append a second coverage row for a SKU already represented by the
    # Monthly matched surface.
    if not covered_monthly_gaps.empty:
        covered_keys = _item_series(core_module, covered_monthly_gaps, identity).astype(str)
        covered_monthly_gaps = covered_monthly_gaps.loc[~covered_keys.isin(matched_keys)].copy()

    all_km_matched = pd.concat([matched, covered_monthly_gaps], ignore_index=True, sort=False)
    all_km_matched = _filter_selected_columns(all_km_matched, selected_output_columns)
    all_km_gap = _filter_selected_columns(all_km_gap, selected_output_columns)

    all_pog_keys = matched_keys | pog_keys
    covered_pog_keys = all_pog_keys & action_keys
    all_km_keys = _item_keys(core_module, all_km_matched, identity)
    missing_from_tong_hop = covered_pog_keys - all_km_keys

    diagnostics = {
        "monthly_matched_items": len(matched_keys),
        "monthly_gap_items": len(pog_keys),
        "all_km_covered_items": len(covered_pog_keys),
        "all_km_gap_items": len(all_pog_keys - action_keys),
        "items_added_to_tong_hop": len(all_km_keys - matched_keys),
        "items_missing_from_tong_hop": len(missing_from_tong_hop),
        "missing_keys": tuple(sorted(missing_from_tong_hop)),
    }
    return all_km_matched, all_km_gap, diagnostics


def _excel_value(value):
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, (list, tuple, set, dict)):
        return str(value)
    return value


def _style_sheet(ws) -> None:
    thin = Side(border_style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    alt_fill = PatternFill(fill_type="solid", fgColor="F7F9FC")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row_index in range(2, ws.max_row + 1):
        for column_index in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_index, column=column_index)
            if cell.value is None:
                continue
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
            if row_index % 2 == 0:
                cell.fill = alt_fill
        ws.row_dimensions[row_index].height = 22
    for column_index in range(1, ws.max_column + 1):
        values = [ws.cell(row=row, column=column_index).value for row in range(1, ws.max_row + 1)]
        width = min(60, max(12, max((len(str(value)) for value in values if value is not None), default=10) + 2))
        ws.column_dimensions[get_column_letter(column_index)].width = width
    ws.freeze_panes = "A2"
    if ws.max_row and ws.max_column:
        ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.tabColor = "4F81BD"


def _write_frame(workbook, sheet_name: str, frame: pd.DataFrame | None, note: str, *, index: int | None = None):
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    ws = workbook.create_sheet(sheet_name, index if index is not None else len(workbook.sheetnames))
    data = frame if frame is not None and not frame.empty else pd.DataFrame({"Ghi chu": [note]})
    ws.append([str(column) for column in data.columns])
    for row in data.itertuples(index=False, name=None):
        ws.append([_excel_value(value) for value in row])
    _style_sheet(ws)
    return ws


def _postprocess_workbook(
    output_path,
    all_km_matched: pd.DataFrame,
    all_km_gap: pd.DataFrame,
    monthly_pog_only: pd.DataFrame | None,
):
    path = Path(output_path)
    if not path.is_file():
        return
    workbook = load_workbook(path)

    tong_index = workbook.sheetnames.index("Tong hop") if "Tong hop" in workbook.sheetnames else 1
    _write_frame(
        workbook,
        "Tong hop",
        all_km_matched,
        "Không có SKU POG nào có KM phù hợp trong dữ liệu/tháng đang xét.",
        index=tong_index,
    )

    old_index = workbook.sheetnames.index(_LEGACY_SHEET) if _LEGACY_SHEET in workbook.sheetnames else min(tong_index + 1, len(workbook.sheetnames))
    if _LEGACY_SHEET in workbook.sheetnames:
        del workbook[_LEGACY_SHEET]

    _write_frame(
        workbook,
        _ALL_KM_GAP_SHEET,
        all_km_gap,
        "Không có SKU POG nào thiếu KM phù hợp trong dữ liệu/tháng đang xét.",
        index=old_index,
    )
    _write_frame(
        workbook,
        _MONTHLY_SHEET,
        monthly_pog_only,
        "Không có SKU POG nào thiếu CTKM Monthly phù hợp trong dữ liệu/tháng đang xét.",
        index=old_index + 1,
    )

    # Temporary hidden compatibility alias for old validators/frontends.  It
    # mirrors canonical all-KM semantics, never the old Monthly-only meaning.
    compat = _write_frame(
        workbook,
        _LEGACY_SHEET,
        all_km_gap,
        "Không có SKU POG nào thiếu KM phù hợp trong dữ liệu/tháng đang xét.",
    )
    compat.sheet_state = "hidden"

    workbook.save(path)


def install_all_km_reconciliation_contract(core_module, facade_module=None) -> None:
    """Wrap the current exporter before runtime/final-output hooks are reapplied."""
    current = core_module.EX.export_final_report
    if getattr(current, _CONTRACT_MARKER, False):
        return

    def export_with_all_km_reconciliation(*args, **kwargs):
        df_matched = kwargs.get("df_matched")
        if df_matched is None and len(args) > 1:
            df_matched = args[1]
        df_pog_only = kwargs.get("df_pog_only")
        if df_pog_only is None and len(args) > 3:
            df_pog_only = args[3]
        summary = kwargs.get("reconciliation_summary")
        if summary is None and len(args) > 6:
            summary = args[6]
        selected = kwargs.get("selected_output_columns")
        if selected is None and len(args) > 8:
            selected = args[8]
        actions = kwargs.get("df_fefo_actions")
        if actions is None and len(args) > 9:
            actions = args[9]

        all_km_matched, all_km_gap, diagnostics = build_all_km_reconciliation(
            core_module,
            df_matched,
            df_pog_only,
            actions,
            selected_output_columns=selected,
        )
        if isinstance(summary, dict):
            summary["So item KM+POG thieu trong Tong hop"] = diagnostics["items_missing_from_tong_hop"]
            summary["So item bo sung vao Tong hop tu KM ngoai Monthly"] = diagnostics["items_added_to_tong_hop"]

        output_path = current(*args, **kwargs)
        _postprocess_workbook(output_path, all_km_matched, all_km_gap, df_pog_only)
        return output_path

    setattr(export_with_all_km_reconciliation, _CONTRACT_MARKER, True)
    setattr(export_with_all_km_reconciliation, "__wrapped__", current)
    core_module.EX.export_final_report = export_with_all_km_reconciliation
