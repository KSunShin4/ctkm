# -*- coding: utf-8 -*-
"""Exporter entrypoint with business-first layout and explicit channel fidelity."""
from __future__ import annotations

import importlib
import importlib.util
from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook

_CORE_NAME = "_ctkm_exporter_core"
_CORE_PATH = Path(__file__).with_name("exporter_core.py")

if _CORE_NAME in sys.modules:
    _core = sys.modules[_CORE_NAME]
else:
    _spec = importlib.util.spec_from_file_location(_CORE_NAME, _CORE_PATH)
    if _spec is None or _spec.loader is None:
        raise ImportError(f"Không nạp được exporter core: {_CORE_PATH}")
    _core = importlib.util.module_from_spec(_spec)
    sys.modules[_CORE_NAME] = _core
    _spec.loader.exec_module(_core)

_performance = importlib.import_module("modules.ctkm_performance")
_original_hide = _core._hide_technical_columns
_original_shelf_colors = _core._apply_alternating_shelf_colors
_original_style = _core._apply_professional_style
_original_action_inputs = _core._configure_action_inputs
_original_confirmation_inputs = _core._configure_confirmation_inputs
_original_guide_sheet = _core._create_guide_sheet
_original_export_final_report = _core.export_final_report

# Comparison-only keys remain available for audit/code but should not clutter
# operational/confirmation sheets. Raw source/new codes stay visible instead.
for _field in ("new_item_key", "transition_match_key"):
    if _field not in _performance.TECHNICAL_ACTION_COLUMNS:
        _performance.TECHNICAL_ACTION_COLUMNS.append(_field)

# Exact user-facing order approved for the workbook mockup. Any additional
# operational diagnostics are retained after this core block and technical IDs
# remain in the audit tail.
_APPROVED_ACTION_COLUMNS = [
    "key_channel", "channel", "item_id", "item_name", "promo_uom",
    "pog_code", "location_id", "facing",
    "action", "ngay_gan_tem", "ngay_thao_tem",
    "campaign_name", "campaign_detail", "promo_from", "promo_to",
    "normal_price", "promo_price", "program_source", "reason",
    "product_type", "store_id",
]
_existing_user_columns = list(_performance.USER_ACTION_COLUMNS)
_performance.USER_ACTION_COLUMNS[:] = [
    *_APPROVED_ACTION_COLUMNS,
    *[column for column in _existing_user_columns if column not in _APPROVED_ACTION_COLUMNS],
]

# Additional operational context stays available but never breaks the first five
# columns or the approved 21-column front view.
for _field in (
    "new_item_id", "item_code_transition_status", "matched_pog_item_id",
    "pog_match_code_type", "pog_applicable", "promotion_count_at_position",
    "has_multiple_promotions_at_position", "pog_month_status",
):
    if _field not in _performance.USER_ACTION_COLUMNS:
        _performance.USER_ACTION_COLUMNS.append(_field)

_core._prepare_action_sheet = _performance.prepare_action_sheet_user_first
_core._filter_selected_columns = _performance.filter_selected_columns_keep_channels
_core._add_work_tracking_columns = _performance.unique_work_tracking_rows


_TOTAL_SHEET_NAME = "TỔNG HỢP - TEM"
_USER_ACTION_SHEETS = {"Gắn tem", "Tháo tem", "Chờ gắn", "Cần review"}
_USER_TABLE_SHEETS = {*_USER_ACTION_SHEETS, _TOTAL_SHEET_NAME}
_RECONCILIATION_SHEETS = {
    "Tong hop", "CTKM thieu vi tri ke", "POG cần kiểm tra",
    "POG không có CTKM", "POG cần review", "Xác nhận CTKM Clearance",
}
_WARNING_SHEETS = {"Canh bao trung thoi gian"}
_AUDIT_SHEETS = {"Audit hành động", "Tom tat", "Mapping log"}

# Fields intentionally hidden on task sheets so the first screen matches the
# approved mockup. They remain in the workbook and can be unhidden for audit.
_DEFAULT_HIDDEN_ACTION_FIELDS = {
    "scheduled_action", "so_ngay_con_lai",
    "new_item_id", "item_code_transition_status", "matched_pog_item_id",
    "pog_match_code_type", "pog_applicable", "promotion_count_at_position",
    "has_multiple_promotions_at_position", "pog_month_status",
    "location_completeness", "location_status", "location_needs_review",
}

_COLUMN_WIDTHS = {
    "key_channel": 18,
    "channel": 24,
    "item_id": 14,
    "item_name": 38,
    "item_name__ctkm": 38,
    "item_name__pog": 38,
    "promo_uom": 20,
    "pog_code": 16,
    "location_id": 12,
    "facing": 10,
    "action": 16,
    "ngay_gan_tem": 14,
    "ngay_thao_tem": 14,
    "campaign_name": 28,
    "campaign_detail": 42,
    "promo_from": 14,
    "promo_to": 14,
    "normal_price": 14,
    "promo_price": 14,
    "program_source": 15,
    "reason": 36,
    "product_type": 16,
    "store_id": 14,
    "processing_status": 16,
    "processed_by": 18,
    "processed_at": 20,
    "staff_note": 30,
    "new_item_id": 16,
    "item_code_transition_status": 22,
    "location_status": 22,
    "location_completeness": 20,
    "pog_month_status": 18,
    "trang_thai_doi_chieu": 28,
    "quan_he_thoi_gian": 30,
    "nguon_chuong_trinh": 18,
    "xac_nhan_ap_dung": 22,
    "ghi_chu_xac_nhan": 34,
    "campaign_a": 28,
    "campaign_b": 28,
    "period_a": 22,
    "period_b": 22,
    "channel_a": 18,
    "channel_b": 18,
    "conflict_type": 24,
}

_HEADER_GROUPS = {
    "key_channel": "1F4E78",
    "channel": "1F4E78",
    "item_id": "2F75B5",
    "item_name": "2F75B5",
    "item_name__ctkm": "2F75B5",
    "item_name__pog": "2F75B5",
    "promo_uom": "2F75B5",
    "pog_code": "5B9BD5",
    "location_id": "5B9BD5",
    "facing": "5B9BD5",
    "action": "70AD47",
    "ngay_gan_tem": "70AD47",
    "ngay_thao_tem": "70AD47",
    "campaign_name": "4472C4",
    "campaign_detail": "4472C4",
    "promo_from": "4472C4",
    "promo_to": "4472C4",
    "normal_price": "4472C4",
    "promo_price": "4472C4",
    "program_source": "7F8C8D",
    "reason": "C65911",
    "product_type": "7F8C8D",
    "store_id": "7F8C8D",
    "processing_status": "C79B00",
    "processed_by": "C79B00",
    "processed_at": "C79B00",
    "staff_note": "C79B00",
    "location_status": "5B9BD5",
    "location_completeness": "5B9BD5",
    "pog_month_status": "5B9BD5",
    "nguon_chuong_trinh": "7F8C8D",
    "trang_thai_doi_chieu": "00A6A6",
    "quan_he_thoi_gian": "00A6A6",
    "xac_nhan_ap_dung": "00A6A6",
    "ghi_chu_xac_nhan": "00A6A6",
    "campaign_a": "C79B00",
    "campaign_b": "C79B00",
    "period_a": "C79B00",
    "period_b": "C79B00",
    "channel_a": "C79B00",
    "channel_b": "C79B00",
    "conflict_type": "C65911",
    "is_blocking": "C65911",
}

_DISPLAY_LABELS = {
    "key_channel": "KEY Channel",
    "channel": "Channel",
    "item_id": "Mã sản phẩm",
    "item_name": "Tên sản phẩm",
    "promo_uom": "Đơn vị sản phẩm (Promo UoM)",
    "pog_code": "POG",
    "location_id": "Vị trí",
    "facing": "Facing",
    "action": "Công việc",
    "ngay_gan_tem": "Ngày gắn",
    "ngay_thao_tem": "Ngày tháo",
    "campaign_name": "Campaign",
    "campaign_detail": "Scheme detail",
    "promo_from": "From",
    "promo_to": "To",
    "normal_price": "Giá thường",
    "promo_price": "Giá KM",
    "program_source": "Nguồn",
    "reason": "Lý do / Cảnh báo",
    "product_type": "Product type",
    "store_id": "Cửa hàng",
    "processing_status": "Đã xử lý",
    "processed_by": "Người thực hiện",
    "processed_at": "Thời gian xử lý",
    "staff_note": "Ghi chú tại quầy",
    "new_item_id": "Mã mới (Item code 1 NEW / New P code)",
    "item_code_transition_status": "Trạng thái chuyển mã",
    "matched_pog_item_id": "Mã ITEMID thực tế trên POG",
    "pog_match_code_type": "POG khớp theo mã nào?",
    "pog_applicable": "Có áp dụng vị trí POG?",
    "promotion_count_at_position": "Số CTKM tại cùng vị trí",
    "has_multiple_promotions_at_position": "Nhiều CTKM cùng vị trí?",
}


def _apply_runtime_header_labels() -> None:
    _core._BILINGUAL_HEADERS.update(_DISPLAY_LABELS)


def _set_price_and_date_formats(ws, raw_headers: list[str]) -> None:
    for index, header in enumerate(raw_headers, start=1):
        if header in {"normal_price", "promo_price"}:
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row=row_idx, column=index).number_format = "#,##0"
        elif header in {"promo_from", "promo_to", "ngay_gan_tem", "ngay_thao_tem"}:
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row=row_idx, column=index).number_format = "dd/mm/yyyy"


def _apply_user_output_style(ws, sheet_name: str, df=None):
    """Style workbook views to match the approved visual hierarchy."""
    _original_style(ws, sheet_name, df)
    if ws is None:
        return

    if sheet_name in _USER_TABLE_SHEETS:
        ws.freeze_panes = "C2"
    elif sheet_name == "Xác nhận CTKM Clearance":
        ws.freeze_panes = "D2"
    elif sheet_name in _RECONCILIATION_SHEETS | _WARNING_SHEETS | _AUDIT_SHEETS:
        ws.freeze_panes = "A2"

    if sheet_name in _USER_TABLE_SHEETS | _RECONCILIATION_SHEETS | _WARNING_SHEETS | _AUDIT_SHEETS:
        ws.auto_filter.ref = ws.dimensions if ws.max_row >= 1 and ws.max_column >= 1 else None
        ws.sheet_view.showGridLines = False

    raw_headers = [str(cell.value or "") for cell in ws[1]]
    for index, raw_header in enumerate(raw_headers, start=1):
        width = _COLUMN_WIDTHS.get(raw_header)
        if width:
            ws.column_dimensions[_core.get_column_letter(index)].width = width
        fill_color = _HEADER_GROUPS.get(raw_header)
        if fill_color:
            cell = ws.cell(row=1, column=index)
            cell.fill = _core.PatternFill(fill_type="solid", fgColor=fill_color)
            cell.font = _core.Font(bold=True, color="FFFFFF")

    _set_price_and_date_formats(ws, raw_headers)

    if sheet_name == _TOTAL_SHEET_NAME:
        ws.sheet_properties.tabColor = "1F4E78"
    elif sheet_name == "Canh bao trung thoi gian" and "is_blocking" in raw_headers:
        blocking_col = raw_headers.index("is_blocking") + 1
        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=blocking_col).value
            is_blocking = value is True or str(value).strip().lower() in {"true", "1", "yes"}
            fill = _core.PatternFill(
                fill_type="solid", fgColor="FCE8E6" if is_blocking else "FFF4E5"
            )
            for col_idx in range(1, ws.max_column + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill


def _configure_action_inputs(ws) -> None:
    """Protect source cells but keep normal Excel filtering/sorting available."""
    _original_action_inputs(ws)
    if ws is None:
        return
    # In openpyxl SheetProtection, True means the feature is locked.
    ws.protection.autoFilter = False
    ws.protection.sort = False


def _configure_confirmation_inputs(ws) -> None:
    _original_confirmation_inputs(ws)
    if ws is None:
        return
    ws.protection.autoFilter = False
    ws.protection.sort = False


def _create_guide_sheet(*args, **kwargs) -> None:
    """Keep HƯỚNG DẪN first and describe the approved channel-first workflow."""
    _original_guide_sheet(*args, **kwargs)
    workbook = args[0] if args else kwargs.get("workbook")
    if workbook is None or "HƯỚNG DẪN" not in workbook.sheetnames:
        return
    ws = workbook["HƯỚNG DẪN"]
    instructions = [
        "1. Mở TỔNG HỢP - TEM để xem toàn bộ công việc hoặc vào từng sheet thao tác.",
        "2. Lọc KEY Channel trước, sau đó lọc Channel theo đúng value gốc từ file công ty.",
        "3. Lọc tiếp Cửa hàng → POG → Vị trí để đi thao tác tại quầy.",
        "4. POG chỉ cung cấp vị trí; không phản ánh tem đã gắn thực tế hay chưa.",
        "5. Cập nhật các ô màu vàng ở sheet thao tác; không sửa dữ liệu nguồn.",
    ]
    for offset, instruction in enumerate(instructions, start=5):
        ws.cell(offset, 4, instruction)


def _hide_technical_columns(ws, *, action_sheet=False):
    _performance.hide_technical_columns(
        ws,
        get_column_letter=_core.get_column_letter,
        original_hide=_original_hide,
        action_sheet=action_sheet,
    )
    if ws is None or not action_sheet:
        return
    for index, cell in enumerate(ws[1], start=1):
        raw = str(cell.value or "").strip()
        if raw in _DEFAULT_HIDDEN_ACTION_FIELDS:
            ws.column_dimensions[_core.get_column_letter(index)].hidden = True


def _truthy(value) -> bool:
    return value is True or str(value).strip().lower() in {"true", "1", "yes"}


def _apply_alternating_shelf_colors(ws, df):
    """Keep shelf grouping, then flag POG review and same-position promotions."""
    _original_shelf_colors(ws, df)
    if ws is None or df is None or df.empty:
        return

    headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]

    collision_fields = [
        field for field in (
            "promotion_count_at_position", "has_multiple_promotions_at_position",
        )
        if field in headers
    ]
    if collision_fields and "has_multiple_promotions_at_position" in df.columns:
        warning_fill = _core.PatternFill(fill_type="solid", fgColor="FFF2CC")
        warning_font = _core.Font(bold=True, color="7F6000")
        collision_columns = [headers.index(field) + 1 for field in collision_fields]
        for excel_row, (_, row) in enumerate(df.iterrows(), start=2):
            if not _truthy(row.get("has_multiple_promotions_at_position")):
                continue
            for column in collision_columns:
                cell = ws.cell(row=excel_row, column=column)
                cell.fill = warning_fill
                cell.font = warning_font

    if "location_needs_review" not in df.columns:
        return
    review_fields = [
        field for field in (
            "location_completeness", "pog_month_status",
            "location_status", "location_needs_review",
        )
        if field in headers
    ]
    if not review_fields:
        return

    review_fill = _core.PatternFill(fill_type="solid", fgColor="FFC7CE")
    review_font = _core.Font(bold=True, color="9C0006")
    column_indexes = [headers.index(field) + 1 for field in review_fields]

    for excel_row, (_, row) in enumerate(df.iterrows(), start=2):
        if not _truthy(row.get("location_needs_review")):
            continue
        for column in column_indexes:
            cell = ws.cell(row=excel_row, column=column)
            cell.fill = review_fill
            cell.font = review_font


def _build_total_action_frame(df_fefo_actions: pd.DataFrame | None) -> pd.DataFrame:
    """Build the 21-column mockup view without changing source values."""
    columns = list(_APPROVED_ACTION_COLUMNS)
    if df_fefo_actions is None:
        return pd.DataFrame(columns=columns)
    actions = df_fefo_actions.copy()
    if "in_report_scope" in actions.columns:
        actions = actions[actions["in_report_scope"].fillna(False).astype(bool)].copy()
    actions = _performance.unique_work_tracking_rows(actions)
    if "action" not in actions.columns:
        actions["action"] = pd.NA
    if "scheduled_action" in actions.columns:
        actions["action"] = actions["action"].where(
            actions["action"].notna() & actions["action"].astype(str).str.strip().ne(""),
            actions["scheduled_action"],
        )
    for column in columns:
        if column not in actions.columns:
            actions[column] = pd.NA
    result = actions.loc[:, columns].copy()
    sort_cols = [
        column for column in ("store_id", "pog_code", "location_id", "item_id")
        if column in result.columns
    ]
    if sort_cols:
        result = result.sort_values(sort_cols, kind="stable", na_position="last")
    return result.reset_index(drop=True)


def _write_total_action_sheet(workbook, df_fefo_actions: pd.DataFrame | None) -> None:
    total = _build_total_action_frame(df_fefo_actions)
    if _TOTAL_SHEET_NAME in workbook.sheetnames:
        del workbook[_TOTAL_SHEET_NAME]
    insert_index = 1 if "HƯỚNG DẪN" in workbook.sheetnames else 0
    ws = workbook.create_sheet(_TOTAL_SHEET_NAME, insert_index)

    for column_index, canonical in enumerate(_APPROVED_ACTION_COLUMNS, start=1):
        ws.cell(row=1, column=column_index, value=canonical)
    for row_index, (_, row) in enumerate(total.iterrows(), start=2):
        for column_index, canonical in enumerate(_APPROVED_ACTION_COLUMNS, start=1):
            value = row[canonical]
            try:
                missing = bool(pd.isna(value))
            except (TypeError, ValueError):
                missing = False
            ws.cell(row=row_index, column=column_index, value=None if missing else value)

    _core._apply_professional_style(ws, _TOTAL_SHEET_NAME, total)
    _core._apply_bilingual_headers(ws)
    ws.sheet_properties.tabColor = "1F4E78"


def _finalize_workbook_views(
    output_path,
    metadata,
    df_fefo_actions: pd.DataFrame | None,
) -> None:
    """Reuse one final workbook pass for POG metadata and the mockup overview."""
    if metadata is None and df_fefo_actions is None:
        return

    workbook = load_workbook(output_path)
    _write_total_action_sheet(workbook, df_fefo_actions)

    if metadata is not None and not metadata.empty:
        sheet_name = "POG metadata"
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]
        ws = workbook.create_sheet(sheet_name)
        columns = list(metadata.columns)
        ws.append(columns)
        for _, row in metadata.iterrows():
            values = []
            for column in columns:
                value = row[column]
                try:
                    missing = bool(_core.pd.isna(value))
                except (TypeError, ValueError):
                    missing = False
                values.append(None if missing else value)
            ws.append(values)
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
        ws.sheet_state = "hidden"

    preferred_order = [
        "HƯỚNG DẪN", _TOTAL_SHEET_NAME,
        "Gắn tem", "Tháo tem", "Chờ gắn", "Cần review",
        "CTKM thieu vi tri ke", "POG cần kiểm tra", "Tong hop",
        "POG không có CTKM", "POG cần review",
        "Xác nhận CTKM Clearance", "Canh bao trung thoi gian",
        "Audit hành động", "Tom tat", "Mapping log", "POG metadata",
    ]
    ordered = [workbook[name] for name in preferred_order if name in workbook.sheetnames]
    ordered_names = {sheet.title for sheet in ordered}
    ordered.extend(sheet for sheet in workbook.worksheets if sheet.title not in ordered_names)
    workbook._sheets = ordered

    workbook.save(output_path)
    workbook.close()


def _export_final_report_with_pog_metadata(*args, **kwargs):
    _apply_runtime_header_labels()
    df_pog_review = kwargs.get("df_pog_review")
    if df_pog_review is None and len(args) > 11:
        df_pog_review = args[11]
    df_fefo_actions = kwargs.get("df_fefo_actions")
    if df_fefo_actions is None and len(args) > 9:
        df_fefo_actions = args[9]

    metadata = None
    if df_pog_review is not None:
        metadata = getattr(df_pog_review, "attrs", {}).get("pog_metadata")

    output_path = _original_export_final_report(*args, **kwargs)
    _finalize_workbook_views(output_path, metadata, df_fefo_actions)
    return output_path


_core._hide_technical_columns = _hide_technical_columns
_core._apply_professional_style = _apply_user_output_style
_core._configure_action_inputs = _configure_action_inputs
_core._configure_confirmation_inputs = _configure_confirmation_inputs
_core._create_guide_sheet = _create_guide_sheet
_core._apply_alternating_shelf_colors = _apply_alternating_shelf_colors
_apply_runtime_header_labels()

for _name in dir(_core):
    if not _name.startswith("__"):
        globals()[_name] = getattr(_core, _name)

export_final_report = _export_final_report_with_pog_metadata
_prepare_action_sheet = _core._prepare_action_sheet
_filter_selected_columns = _core._filter_selected_columns
_add_work_tracking_columns = _core._add_work_tracking_columns
_hide_technical_columns = _core._hide_technical_columns
_apply_professional_style = _core._apply_professional_style
_configure_action_inputs = _core._configure_action_inputs
_configure_confirmation_inputs = _core._configure_confirmation_inputs
_create_guide_sheet = _core._create_guide_sheet
_apply_alternating_shelf_colors = _core._apply_alternating_shelf_colors
_BILINGUAL_HEADERS = _core._BILINGUAL_HEADERS
