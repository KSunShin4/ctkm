# -*- coding: utf-8 -*-
"""Validate manually entered item/batch/expiration triples against an Excel export."""
from __future__ import annotations

from io import BytesIO
import re
from typing import BinaryIO
import zipfile

import openpyxl
import pandas as pd

REQUIRED_COLUMNS = ("ITEMID", "BATCHNUMER", "EXPDATE")
OPTIONAL_COLUMNS = ("ITEMNAME", "INVENTUNIT", "ONHANDNEAREXPIRATION")
MAX_EXPIRATION_ROWS = 250_000
MAX_XLSX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024


def _text(value) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _date_text(value) -> str:
    if value is None or pd.isna(value):
        return ""
    parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)
    return "" if pd.isna(parsed) else parsed.strftime("%d/%m/%Y")


def _source_bytes(source: bytes | BinaryIO) -> bytes:
    if isinstance(source, bytes):
        return source
    position = None
    try:
        position = source.tell()
    except (AttributeError, OSError):
        pass
    raw = source.read()
    if position is not None:
        try:
            source.seek(position)
        except (AttributeError, OSError):
            pass
    return raw


def _sanitize_xlsx_bytes(raw: bytes) -> bytes:
    if not raw:
        raise ValueError("File hạn dùng đang rỗng.")
    if not zipfile.is_zipfile(BytesIO(raw)):
        raise ValueError(
            "File hạn dùng không phải XLSX hợp lệ; file có thể bị hỏng, "
            "đặt sai đuôi hoặc đang dùng định dạng XLS/XLSB."
        )

    try:
        with zipfile.ZipFile(BytesIO(raw), "r") as archive:
            infos = archive.infolist()
            total_uncompressed = sum(info.file_size for info in infos)
            if total_uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValueError(
                    "File hạn dùng giải nén quá lớn để xử lý an toàn. "
                    "Hãy xuất lại file với đúng vùng dữ liệu cần kiểm tra."
                )
            names = {info.filename for info in infos}
            if "xl/workbook.xml" not in names:
                raise ValueError("File hạn dùng thiếu cấu trúc workbook XLSX bắt buộc.")
            workbook_xml = archive.read("xl/workbook.xml").decode("utf-8", errors="ignore")
            if "<definedNames>" not in workbook_xml:
                return raw
            cleaned_workbook_xml = re.sub(
                r"<definedNames>.*?</definedNames>",
                "",
                workbook_xml,
                flags=re.DOTALL,
            ).encode("utf-8")
            members = {
                info.filename: (
                    cleaned_workbook_xml
                    if info.filename == "xl/workbook.xml"
                    else archive.read(info.filename)
                )
                for info in infos
            }
    except zipfile.BadZipFile as exc:
        raise ValueError("File hạn dùng có cấu trúc ZIP/XLSX bị hỏng.") from exc

    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return output.getvalue()


def load_expiration_export(source: bytes | BinaryIO) -> tuple[pd.DataFrame, str]:
    """Return the first sheet containing the three expected export columns."""
    raw = _source_bytes(source)
    safe = _sanitize_xlsx_bytes(raw)
    workbook = None
    try:
        workbook = openpyxl.load_workbook(
            BytesIO(safe),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            if worksheet.max_column and worksheet.max_column > 2048:
                worksheet.reset_dimensions()
            rows = worksheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            if header_row is None:
                continue
            headers = [str(value).strip() if value is not None else "" for value in header_row]
            normalized_headers = [header.upper() for header in headers]
            lookup = {
                normalized: index
                for index, normalized in enumerate(normalized_headers)
                if normalized
            }
            if not all(column in lookup for column in REQUIRED_COLUMNS):
                continue

            selected_names = list(REQUIRED_COLUMNS) + [
                column for column in OPTIONAL_COLUMNS if column in lookup
            ]
            duplicate_selected = sorted({
                column
                for column in selected_names
                if normalized_headers.count(column) > 1
            })
            if duplicate_selected:
                raise ValueError(
                    "Sheet hạn dùng có tên cột bị trùng: " + ", ".join(duplicate_selected)
                )

            records = []
            for row in rows:
                record = [
                    row[lookup[column]] if lookup[column] < len(row) else None
                    for column in selected_names
                ]
                if all(value is None or str(value).strip() == "" for value in record):
                    continue
                records.append(record)
                if len(records) > MAX_EXPIRATION_ROWS:
                    raise ValueError(
                        f"Sheet hạn dùng vượt quá {MAX_EXPIRATION_ROWS:,} dòng dữ liệu."
                    )

            selected = pd.DataFrame(records, columns=selected_names)
            selected["ITEMID"] = selected["ITEMID"].map(_text).str.upper()
            selected["BATCHNUMER"] = selected["BATCHNUMER"].map(_text)
            selected["EXPDATE"] = selected["EXPDATE"].map(_date_text)
            for column in ("ITEMNAME", "INVENTUNIT"):
                if column in selected:
                    selected[column] = selected[column].map(_text)
            return selected, sheet_name
    except ValueError:
        raise
    except Exception as exc:
        message = str(exc).lower()
        if "encrypt" in message or "password" in message or "agile" in message:
            raise ValueError(
                "File hạn dùng đang được mã hóa/đặt mật khẩu. Hãy bỏ mật khẩu rồi tải lại."
            ) from exc
        raise ValueError(
            f"Không thể mở file hạn dùng; file có thể bị hỏng hoặc không đúng chuẩn XLSX: {exc}"
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()

    raise ValueError(
        "Không tìm thấy sheet có đủ các cột ITEMID, BATCHNUMER và EXPDATE."
    )


def check_entries(frame: pd.DataFrame, raw_entries: str) -> list[dict]:
    rows = [line.strip() for line in raw_entries.splitlines() if line.strip()]
    if not rows:
        raise ValueError("Hãy nhập ít nhất một dòng theo mẫu MÃ_P;SỐ_LÔ;DD/MM/YYYY.")

    results = []
    item_index = frame.attrs.get("item_index")
    if item_index is None:
        item_index = {
            item_id: positions.tolist()
            for item_id, positions in frame.groupby("ITEMID", sort=False).groups.items()
        }
        frame.attrs["item_index"] = item_index
    for line_number, raw in enumerate(rows, start=1):
        parts = [part.strip() for part in raw.split(";")]
        base = {"line": line_number, "input": raw}
        if len(parts) != 3 or not all(parts):
            results.append({**base, "status": "invalid", "message": "Cần đúng 3 phần, ngăn cách bằng dấu ;"})
            continue

        item_id, batch_number, raw_date = parts
        item_id = item_id.upper()
        exp_date = _date_text(raw_date)
        normalized = f"{item_id};{batch_number};{exp_date or raw_date}"
        detail = {**base, "item_id": item_id, "batch_number": batch_number,
                  "exp_date": exp_date or raw_date, "normalized": normalized}
        if not item_id.startswith("P"):
            results.append({**detail, "status": "invalid", "message": "Mã sản phẩm phải bắt đầu bằng P"})
            continue
        if not exp_date:
            results.append({**detail, "status": "invalid", "message": "Hạn dùng phải có dạng DD/MM/YYYY"})
            continue

        positions = item_index.get(item_id)
        if not positions:
            results.append({**detail, "status": "not_found", "message": "Không tìm thấy mã sản phẩm trong file"})
            continue
        item_rows = frame.loc[positions]
        first_item = item_rows.iloc[0]
        related_rows = []
        for _, record in item_rows.drop_duplicates(["BATCHNUMER", "EXPDATE"]).iterrows():
            related_rows.append({
                "batch_number": _text(record["BATCHNUMER"]),
                "exp_date": _text(record["EXPDATE"]),
                "on_hand": _text(record.get("ONHANDNEAREXPIRATION", "")),
            })
        related_rows.sort(
            key=lambda row: pd.to_datetime(row["exp_date"], format="%d/%m/%Y", errors="coerce"),
            reverse=True,
        )
        product_info = {
            "item_name": _text(first_item.get("ITEMNAME", "")),
            "inventory_unit": _text(first_item.get("INVENTUNIT", "")),
        }
        batch_rows = item_rows[item_rows["BATCHNUMER"] == batch_number]
        if batch_rows.empty:
            results.append({**detail, **product_info, "status": "batch_mismatch",
                            "message": "Số lô không có trong file", "related_rows": related_rows})
            continue
        if not (batch_rows["EXPDATE"] == exp_date).any():
            dates = sorted({value for value in batch_rows["EXPDATE"] if value})
            suffix = f" (trong file: {', '.join(dates)})" if dates else ""
            results.append({**detail, **product_info, "status": "date_mismatch",
                            "message": "Hạn dùng không khớp" + suffix,
                            "related_rows": related_rows})
            continue
        matched_record = batch_rows[batch_rows["EXPDATE"] == exp_date].iloc[0]
        results.append({
            **detail, **product_info,
            "on_hand": _text(matched_record.get("ONHANDNEAREXPIRATION", "")),
            "status": "matched", "message": "Khớp dữ liệu", "related_rows": related_rows,
        })
    return results
