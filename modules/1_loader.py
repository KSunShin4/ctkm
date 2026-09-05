# -*- coding: utf-8 -*-
"""Loader entrypoint with conservative workbook recovery.

The existing loader logic is preserved in ``loader_core.py``. Normal XLSX files
are opened byte-for-byte first; defined-name stripping is only a fallback when
openpyxl cannot parse the original workbook.
"""
from __future__ import annotations

import importlib.util
import io
from pathlib import Path
import re
import sys

import openpyxl

_CORE_NAME = "modules.loader_core"
_CORE_PATH = Path(__file__).with_name("loader_core.py")

if _CORE_NAME in sys.modules:
    _core = sys.modules[_CORE_NAME]
else:
    _spec = importlib.util.spec_from_file_location(_CORE_NAME, _CORE_PATH)
    if _spec is None or _spec.loader is None:
        raise ImportError(f"Không nạp được loader core: {_CORE_PATH}")
    _core = importlib.util.module_from_spec(_spec)
    sys.modules[_CORE_NAME] = _core
    _spec.loader.exec_module(_core)

for _name in dir(_core):
    if not _name.startswith("__"):
        globals()[_name] = getattr(_core, _name)


def _open_workbook_safely(filepath: str):
    """Open the original workbook first and repair only after a parser failure."""
    _core._validate_xlsx_container(filepath)
    original_error = None
    try:
        workbook = openpyxl.load_workbook(
            filepath, data_only=True, read_only=True, keep_links=False
        )
    except Exception as exc:  # openpyxl/lxml expose several parser exception types
        original_error = exc
        try:
            safe = _core._strip_broken_defined_names(filepath)
            workbook = openpyxl.load_workbook(
                io.BytesIO(safe), data_only=True, read_only=True, keep_links=False
            )
        except Exception as repaired_exc:
            message = f"{original_error} {repaired_exc}".lower()
            if "encrypt" in message or "password" in message or "agile" in message:
                raise ValueError(
                    f"File Excel '{Path(filepath).name}' đang được mã hóa/đặt mật khẩu. "
                    "Hãy bỏ mật khẩu rồi tải lại."
                ) from repaired_exc
            raise ValueError(
                f"Không thể mở file Excel '{Path(filepath).name}'. File có thể bị hỏng "
                f"hoặc không đúng chuẩn XLSX: {repaired_exc}"
            ) from repaired_exc

    if not workbook.sheetnames:
        workbook.close()
        raise ValueError(f"File Excel '{Path(filepath).name}' không có sheet nào.")
    return workbook


def _snapshot_month_candidate(parts, *, month_index: int, day_index: int):
    """Infer one month only when every DATEID date supports the same format."""
    if not parts:
        return None
    months = []
    years = []
    for values in parts:
        month = int(values[month_index])
        day = int(values[day_index])
        year = int(values[2])
        if not (1 <= month <= 12 and 1 <= day <= 31 and 1900 <= year <= 9999):
            return None
        months.append(month)
        years.append(year)
    if len(set(months)) != 1 or len(set(years)) != 1:
        return None
    return f"{years[0]:04d}-{months[0]:02d}"


def _extract_pog_snapshot_month_from_values(row_values):
    """Read DATEID metadata without guessing an ambiguous day/month order.

    Pharmacity's real export currently uses ``DATEID 9/1/2026 - 9/30/2026``
    (MM/DD/YYYY). Some Excel exports can instead render DD/MM/YYYY. We inspect the
    whole date range: if only one interpretation is valid and stays in one month,
    use it; if both interpretations disagree, return ``None`` so the pipeline
    warns instead of silently assigning the wrong POG month.
    """
    pattern = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")
    for value in row_values:
        if value is None:
            continue
        text = str(value)
        lower = text.lower()
        if "applied filter" not in lower and "dateid" not in lower:
            continue
        parts = pattern.findall(text)
        if not parts:
            continue

        month_first = _snapshot_month_candidate(parts, month_index=0, day_index=1)
        day_first = _snapshot_month_candidate(parts, month_index=1, day_index=0)
        candidates = {candidate for candidate in (month_first, day_first) if candidate}
        if len(candidates) == 1:
            return next(iter(candidates))
        if len(candidates) > 1:
            # Both formats are valid but imply different months: do not guess.
            return None
    return None


# load_workbook_raw resolves these helpers in loader_core globals.
_core._open_workbook_safely = _open_workbook_safely
_core._extract_pog_snapshot_month_from_values = _extract_pog_snapshot_month_from_values
load_workbook_raw = _core.load_workbook_raw
inspect_workbook_sheets = _core.inspect_workbook_sheets
