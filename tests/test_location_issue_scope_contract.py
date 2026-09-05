from openpyxl import load_workbook
import pandas as pd

import main


def _action(item_id, source, *, location="Không tìm thấy SKU trong POG", in_scope=True):
    return {
        "task_id": f"TASK-{source}-{item_id}",
        "promotion_id": f"PROMO-{source}-{item_id}",
        "pog_position_id": f"POGPOS-{item_id}",
        "item_id": item_id,
        "item_key": item_id,
        "item_name": f"Product {item_id}",
        "program_source": source,
        "campaign_name": "Campaign",
        "campaign_detail": "Offer",
        "promo_from": pd.Timestamp("2026-09-01"),
        "promo_to": pd.Timestamp("2026-09-30"),
        "scheduled_action": "gắn tem",
        "action": "gắn tem",
        "reason": "trong_thoi_gian_ap_dung",
        "in_report_scope": in_scope,
        "location_completeness": location,
        "location_status": location,
        "source_row_id": f"{source}:Monthly:2",
        "source_row_ids": f"{source}:Monthly:2",
        "source_row_count": 1,
        "duplicate_status": "Dòng duy nhất",
    }


def _actions():
    return pd.DataFrame([
        _action("P1001", "CTKM"),
        _action("P2001", "CLEARANCE"),
        _action("P3001", "CTKM", location="Đủ vị trí"),
        _action("P4001", "CLEARANCE", in_scope=False),
    ])


def test_location_summary_distinguishes_monthly_and_clearance_sources():
    summary = {
        "So dong CTKM thieu vi tri ke": 99,
        "So item CTKM thieu vi tri ke": 99,
    }

    main._align_location_summary_with_program_sources(summary, _actions())

    assert summary["So dong KM thieu vi tri POG"] == 2
    assert summary["So item KM thieu vi tri POG"] == 2
    assert summary["So dong CTKM Monthly thieu vi tri POG"] == 1
    assert summary["So item CTKM Monthly thieu vi tri POG"] == 1
    assert summary["So dong Clearance thieu vi tri POG"] == 1
    assert summary["So item Clearance thieu vi tri POG"] == 1

    # Historical API keys stay as total-KM aliases for old frontend/jobs.
    assert summary["So dong CTKM thieu vi tri ke"] == 2
    assert summary["So item CTKM thieu vi tri ke"] == 2


def test_export_shows_generic_location_metrics_and_preserves_legacy_sheet_name(tmp_path):
    summary = {
        "Thang bao cao": "2026-09",
        "Ngay nghiep vu": pd.Timestamp("2026-09-15"),
        "So dong CTKM thieu vi tri ke": 99,
        "So item CTKM thieu vi tri ke": 99,
    }
    output = tmp_path / "location-scope.xlsx"

    main.EX.export_final_report(
        output,
        pd.DataFrame(),
        pd.DataFrame(),
        pd.DataFrame(),
        [],
        [],
        summary,
        df_fefo_actions=_actions(),
    )

    workbook = load_workbook(output, data_only=True)

    # Keep the historical tab name as a compatibility surface for now.
    assert "CTKM thieu vi tri ke" in workbook.sheetnames
    location_sheet = workbook["CTKM thieu vi tri ke"]
    all_values = [cell.value for row in location_sheet.iter_rows() for cell in row]
    assert "CTKM" in all_values
    assert "CLEARANCE" in all_values
    assert "P1001" in all_values
    assert "P2001" in all_values
    assert "P3001" not in all_values
    assert "P4001" not in all_values

    tom_tat = workbook["Tom tat"]
    summary_rows = {
        tom_tat.cell(row=row, column=1).value: (tom_tat.cell(row=row, column=2).value, row)
        for row in range(2, tom_tat.max_row + 1)
    }
    assert summary_rows["So dong KM thieu vi tri POG"][0] == 2
    assert summary_rows["So dong CTKM Monthly thieu vi tri POG"][0] == 1
    assert summary_rows["So dong Clearance thieu vi tri POG"][0] == 1

    # Misleading legacy aliases remain in the workbook only for compatibility
    # and are hidden from the normal summary view.
    legacy_row = summary_rows["So dong CTKM thieu vi tri ke"][1]
    assert tom_tat.row_dimensions[legacy_row].hidden is True

    guide_values = [cell.value for row in workbook["HƯỚNG DẪN"].iter_rows() for cell in row]
    assert "Thiếu vị trí KM" in guide_values
