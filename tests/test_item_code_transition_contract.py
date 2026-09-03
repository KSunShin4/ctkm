from openpyxl import load_workbook
import pandas as pd

import main


BUSINESS_DATE = pd.Timestamp("2026-09-15")


def _raw_monthly(rows):
    return pd.DataFrame(rows)


def _monthly_row(old_code, new_code, name, campaign="FULL MONTH"):
    return {
        "KEY Channel": "All channel",
        "Channel": "ALL CHANNEL",
        "Item code 1 (NEW)": new_code,
        "Item code 2": old_code,
        "Product Name": name,
        "Promo UoM": "Tuyp",
        "Normal": 100000,
        "Promo": 80000,
        "Scheme detail for customer": "Discount 20%",
        "From": pd.Timestamp("2026-09-01"),
        "To": pd.Timestamp("2026-09-30"),
        "Campaign": campaign,
        "Product type": "H&B",
    }


def _process_monthly(rows):
    raw = _raw_monthly(rows)
    mapping = main._infer_standard_fact_mapping(raw)
    warnings = []
    result = main._process(
        "TỔNG HỢP CTKM UPDATE",
        raw,
        main.SCHEMA["fact_fields"],
        [],
        True,
        warnings,
        mapping_override=mapping,
        source_excel_rows=list(range(3, 3 + len(raw))),
        source_prefix="CTKM",
    )
    return result, warnings


def _pog(item_id, location_id, row_id):
    return {
        "item_id": item_id,
        "store_id": "SGPMC1885",
        "pog_code": "DERMO2.W",
        "location_id": location_id,
        "facing": 1,
        "pog_snapshot_month": "2026-09",
        "source_row_id": row_id,
    }


def test_monthly_keeps_source_item_id_and_exposes_exact_new_item_code():
    processed, warnings = _process_monthly([
        _monthly_row("P09544", "P30519", "La Roche Posay Effaclar Duo+M"),
    ])

    assert warnings == []
    assert processed.loc[0, "item_id"] == "P09544"
    assert processed.loc[0, "item_key"] == "P09544"
    assert processed.loc[0, "new_item_id"] == "P30519"
    assert processed.loc[0, "new_item_key"] == "P30519"
    assert processed.loc[0, "item_code_transition_status"] == "Có mã mới - đối chiếu song song"
    # Original source column is still present for byte/source audit.
    assert processed.loc[0, "extra__Item code 1 (NEW)"] == "P30519"


def test_clearance_new_p_code_is_preserved_as_same_derived_contract():
    raw = pd.DataFrame([{
        "Kênh áp dụng": "All Channel",
        "New P code": "P30648",
        "Item code": "P27878",
        "Product Name": "Kem Duoc Lieu Thuan Moc",
        "Promo UoM": "Tuyp",
        "Normal": 199000,
        "Promo": 109450,
        "Scheme detail for customer": "Discount 45%",
        "From": pd.Timestamp("2026-09-01"),
        "To": pd.Timestamp("2026-09-30"),
        "Campaign Detail": "Clearance Isstop 1-3-4",
        "Product type": "H&B",
    }])
    mapping = main._infer_standard_fact_mapping(raw)
    processed = main._process(
        "TỔNG HỢP CTKM UPDATE", raw, main.SCHEMA["fact_fields"], [], True, [],
        mapping_override=mapping,
        source_excel_rows=[3], source_prefix="CLEARANCE",
    )

    assert processed.loc[0, "item_id"] == "P27878"
    assert processed.loc[0, "new_item_id"] == "P30648"
    assert processed.loc[0, "new_item_key"] == "P30648"
    assert processed.loc[0, "extra__New P code"] == "P30648"


def test_new_code_pog_matches_source_promotion_without_changing_promotion_identity():
    processed, _ = _process_monthly([
        _monthly_row("P09544", "P30519", "La Roche Posay Effaclar Duo+M"),
    ])
    pog_new = pd.DataFrame([_pog("P30519", "5", "POG:Export:342")])

    joined = main.JN.join_ctkm_pog(processed, pog_new)
    assert len(joined.df_matched) == 1
    assert joined.df_ctkm_only.empty
    assert joined.df_pog_only.empty
    assert joined.df_matched.loc[0, "item_id"] == "P09544"
    assert joined.df_matched.loc[0, "item_id__pog"] == "P30519"
    assert joined.df_matched.loc[0, "item_key"] == "P09544"

    actions_new = main.RC.build_fefo_action_frame(
        processed, "2026-09", pog_new, business_date=BUSINESS_DATE,
    )
    row = actions_new.iloc[0]
    assert row["item_id"] == "P09544"
    assert row["new_item_id"] == "P30519"
    assert row["matched_pog_item_id"] == "P30519"
    assert row["pog_match_code_type"] == "Khớp mã mới"
    assert row["location_completeness"] == "Đủ vị trí"
    assert bool(row["location_needs_review"]) is False

    # If the POG snapshot still uses the source/old code at the same position,
    # promotion/task identity remains stable across the code migration.
    pog_old = pd.DataFrame([_pog("P09544", "5", "POG:Export:342")])
    actions_old = main.RC.build_fefo_action_frame(
        processed, "2026-09", pog_old, business_date=BUSINESS_DATE,
    )
    assert actions_old.loc[0, "promotion_id"] == row["promotion_id"]
    assert actions_old.loc[0, "task_id"] == row["task_id"]
    assert actions_old.loc[0, "pog_match_code_type"] == "Khớp mã nguồn"


def test_source_code_pog_still_matches_when_new_code_exists():
    processed, _ = _process_monthly([
        _monthly_row("P21290", "P30623", "Balance Serum Vitamin C"),
    ])
    pog = pd.DataFrame([_pog("P21290", "4", "POG:Export:372")])

    actions = main.RC.build_fefo_action_frame(
        processed, "2026-09", pog, business_date=BUSINESS_DATE,
    )
    assert len(actions) == 1
    assert actions.loc[0, "matched_pog_item_id"] == "P21290"
    assert actions.loc[0, "pog_match_code_type"] == "Khớp mã nguồn"
    assert actions.loc[0, "location_completeness"] == "Đủ vị trí"


def test_old_and_new_pog_positions_are_both_preserved_for_one_promotion():
    processed, _ = _process_monthly([
        _monthly_row("P1001", "P9001", "Transition product"),
    ])
    pog = pd.DataFrame([
        _pog("P1001", "1", "POG:Export:10"),
        _pog("P9001", "2", "POG:Export:11"),
    ])

    actions = main.RC.build_fefo_action_frame(
        processed, "2026-09", pog, business_date=BUSINESS_DATE,
    )

    assert len(actions) == 2
    assert actions["promotion_id"].nunique() == 1
    assert actions["task_id"].nunique() == 2
    assert set(actions["matched_pog_item_id"]) == {"P1001", "P9001"}
    assert set(actions["pog_match_code_type"]) == {"Khớp mã nguồn", "Khớp mã mới"}
    assert set(actions["location_completeness"]) == {"Đủ vị trí"}


def test_ambiguous_replacement_mapping_is_not_silently_applied():
    processed, _ = _process_monthly([
        _monthly_row("P1001", "P9001", "Product A", campaign="A"),
        _monthly_row("P1002", "P9001", "Product B", campaign="B"),
    ])
    pog = pd.DataFrame([_pog("P9001", "7", "POG:Export:20")])

    joined = main.JN.join_ctkm_pog(processed, pog)
    assert joined.df_matched.empty
    assert len(joined.df_ctkm_only) == 2
    assert len(joined.df_pog_only) == 1

    actions = main.RC.build_fefo_action_frame(
        processed, "2026-09", pog, business_date=BUSINESS_DATE,
    )
    assert set(actions["item_code_transition_status"]) == {"Mơ hồ - không tự remap"}
    assert set(actions["location_completeness"]) == {"Không tìm thấy SKU trong POG"}
    assert set(actions["location_needs_review"]) == {True}


def test_export_shows_source_new_and_matched_pog_codes(tmp_path):
    processed, _ = _process_monthly([
        _monthly_row("P09544", "P30519", "La Roche Posay Effaclar Duo+M"),
    ])
    actions = main.RC.build_fefo_action_frame(
        processed,
        "2026-09",
        pd.DataFrame([_pog("P30519", "5", "POG:Export:342")]),
        business_date=BUSINESS_DATE,
    )
    output = tmp_path / "item-code-transition.xlsx"
    summary = {
        "Thang bao cao": "2026-09",
        "Ngay nghiep vu": BUSINESS_DATE,
        "So dong CTKM thieu vi tri ke": 0,
        "So item CTKM thieu vi tri ke": 0,
        "So dong POG khong co CTKM phu hop": 0,
        "So item POG khong co CTKM phu hop": 0,
    }

    main.EX.export_final_report(
        output_path=output,
        df_matched=pd.DataFrame(),
        df_ctkm_only=pd.DataFrame(),
        df_pog_only=pd.DataFrame(),
        overlap_warnings=[],
        mapping_logs=[],
        reconciliation_summary=summary,
        df_fefo_actions=actions,
        df_pog_review=pd.DataFrame(),
    )

    workbook = load_workbook(output, data_only=True)
    sheet = workbook["Gắn tem"]
    headers = [cell.value for cell in sheet[1]]
    assert main.EX._BILINGUAL_HEADERS["item_id"] in headers
    assert main.EX._BILINGUAL_HEADERS["new_item_id"] in headers
    assert main.EX._BILINGUAL_HEADERS["matched_pog_item_id"] in headers
    assert main.EX._BILINGUAL_HEADERS["pog_match_code_type"] in headers

    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert "P09544" in values
    assert "P30519" in values
    assert "Khớp mã mới" in values
