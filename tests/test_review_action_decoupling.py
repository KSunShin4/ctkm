import pandas as pd
from openpyxl import load_workbook

import main


BUSINESS_DATE = pd.Timestamp("2026-09-15")


def _promo(campaign_name, *, channel="OFFLINE", key_channel="Offline", source="CTKM", **overrides):
    row = {
        "item_id": "P1001",
        "item_name": "Product",
        "promo_uom": "Hộp",
        "campaign_name": campaign_name,
        "campaign_detail": "Offer",
        "promo_from": pd.Timestamp("2026-09-01"),
        "promo_to": pd.Timestamp("2026-09-30"),
        "normal_price": 100_000,
        "promo_price": 90_000,
        "key_channel": key_channel,
        "channel": channel,
        "program_source": source,
        "source_sheet": "Monthly",
    }
    row.update(overrides)
    return row


def _actions(frame, warnings=None):
    return main.RC.build_fefo_action_frame(
        pd.DataFrame(frame),
        "2026-09",
        None,
        business_date=BUSINESS_DATE,
        overlap_warnings=warnings,
    )


def test_same_key_channel_but_different_channel_is_not_blocking():
    frame = pd.DataFrame([
        _promo("A", channel="OFFLINE", key_channel="Offline"),
        _promo("B", channel="COMBO ĐI SINH", key_channel="Offline"),
    ])

    warnings = main.RC.detect_campaign_time_overlaps(frame)
    actions = _actions(frame.to_dict("records"), warnings)

    assert len(warnings) == 1
    assert warnings[0].is_blocking is False
    assert warnings[0].conflict_type == "potential_cross_channel_overlap"
    assert set(actions["scheduled_action"]) == {"gắn tem"}
    assert not actions["review_required"].astype(bool).any()


def test_same_scope_overlap_keeps_action_and_adds_review_warning():
    frame = pd.DataFrame([
        _promo("A", promo_price=90_000),
        _promo("B", promo_price=80_000),
    ])

    warnings = main.RC.detect_campaign_time_overlaps(frame)
    actions = _actions(frame.to_dict("records"), warnings)

    assert len(warnings) == 1
    assert warnings[0].is_blocking is True
    assert set(actions["scheduled_action"]) == {"gắn tem"}
    assert actions["review_required"].astype(bool).all()
    assert not actions["review_blocks_action"].astype(bool).any()
    assert set(actions["review_severity"]) == {"warning"}
    assert actions["review_reason"].str.contains("xung_dot_chuong_trinh").all()


def test_ctkm_clearance_overlap_keeps_action_but_requires_review():
    frame = pd.DataFrame([
        _promo("Monthly", channel="OFFLINE", source="CTKM"),
        _promo("Clearance", channel="WEBAPP", key_channel="Online", source="CLEARANCE"),
    ])

    warnings = main.RC.detect_campaign_time_overlaps(frame)
    actions = _actions(frame.to_dict("records"), warnings)

    assert len(warnings) == 1
    assert warnings[0].is_blocking is True
    assert set(actions["scheduled_action"]) == {"gắn tem"}
    assert actions["review_required"].astype(bool).all()
    assert not actions["review_blocks_action"].astype(bool).any()


def test_intrinsic_data_error_still_blocks_action():
    actions = _actions([
        _promo("Broken", promo_from=pd.NaT),
    ])

    assert actions.loc[0, "scheduled_action"] == "review"
    assert bool(actions.loc[0, "review_required"]) is True
    assert bool(actions.loc[0, "review_blocks_action"]) is True
    assert actions.loc[0, "review_severity"] == "blocking"


def test_review_warning_is_present_in_action_and_review_sheets(tmp_path):
    frame = pd.DataFrame([
        _promo("A", promo_price=90_000),
        _promo("B", promo_price=80_000),
    ])
    warnings = main.RC.detect_campaign_time_overlaps(frame)
    actions = _actions(frame.to_dict("records"), warnings)
    output = tmp_path / "review-action.xlsx"

    main.EX.export_final_report(
        output,
        pd.DataFrame(), pd.DataFrame(), pd.DataFrame(),
        warnings, [],
        {"Thang bao cao": "2026-09", "Ngay nghiep vu": BUSINESS_DATE},
        df_fefo_actions=actions,
    )

    workbook = load_workbook(output, data_only=False)
    assert workbook["Gắn tem"].max_row == 3
    assert workbook["Cần review"].max_row == 3

    review_headers = [cell.value for cell in workbook["Cần review"][1]]
    assert "Cần review?" in review_headers
    assert "Review chặn action?" in review_headers
    assert "Lý do review" in review_headers

    review_values = [cell.value for row in workbook["Cần review"] for cell in row]
    assert "gắn tem" in review_values
    assert any(isinstance(value, str) and "xung_dot_chuong_trinh" in value for value in review_values)
