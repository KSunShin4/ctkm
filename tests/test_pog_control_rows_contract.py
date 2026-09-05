from openpyxl import load_workbook
import pandas as pd

import main


def _frames():
    ctkm = pd.DataFrame([
        {"item_id": "P1001", "item_key": "P1001", "campaign_name": "September"},
    ])
    pog = pd.DataFrame([
        {
            "item_id": "P1001", "item_key": "P1001", "store_id": "S1",
            "pog_code": "A", "location_id": "1", "source_row_id": "POG:Export:2",
            "pog_snapshot_month": "2026-09",
        },
        {
            # Real invalid business row: must remain reviewable.
            "item_id": None, "item_key": None, "store_id": "S1",
            "pog_code": "BROKEN", "location_id": "9", "source_row_id": "POG:Export:1616",
            "pog_snapshot_month": "2026-09",
        },
        {
            # Pharmacity export summary/footer row; not a business SKU.
            "item_id": "Total", "item_key": "TOTAL", "store_id": None,
            "pog_code": None, "location_id": None, "source_row_id": "POG:Export:1617",
            "pog_snapshot_month": "2026-09",
        },
        {
            # Pharmacity snapshot metadata stored in the first source column.
            "item_id": None, "item_key": None,
            "store_id": (
                "Applied filters:\nFIXELNAME does not contain 'NONPOG'\n"
                "MODELAPPLYSTORE is VTB\nDATEID 9/1/2026 - 9/30/2026"
            ),
            "pog_code": None, "location_id": None, "source_row_id": "POG:Export:1618",
            "pog_snapshot_month": "2026-09",
        },
    ])
    return ctkm, pog


def test_pog_control_rows_do_not_distort_reconciliation_or_missing_item_review():
    ctkm, pog = _frames()

    joined = main.JN.join_ctkm_pog(ctkm, pog)

    assert len(joined.df_matched) == 1
    assert joined.df_matched.loc[0, "item_id"] == "P1001"

    # Total/DATEID are control rows, not POG-only business positions.
    assert joined.df_pog_only.empty

    # The genuinely broken business row remains reviewable exactly as before.
    assert len(joined.df_pog_review) == 1
    assert joined.df_pog_review.loc[0, "source_row_id"] == "POG:Export:1616"
    assert joined.df_pog_review.loc[0, "reason"] == "thieu_ma_san_pham_pog"

    metadata = joined.df_pog_review.attrs.get("pog_metadata")
    assert metadata is not None
    assert len(metadata) == 2
    assert set(metadata["pog_control_kind"]) == {"summary_total", "snapshot_metadata"}
    assert set(metadata["source_row_id"]) == {"POG:Export:1617", "POG:Export:1618"}
    assert metadata.loc[
        metadata["pog_control_kind"].eq("summary_total"), "item_id"
    ].iloc[0] == "Total"
    assert "DATEID 9/1/2026 - 9/30/2026" in metadata.loc[
        metadata["pog_control_kind"].eq("snapshot_metadata"), "store_id"
    ].iloc[0]


def test_export_keeps_pog_control_rows_in_hidden_metadata_audit(tmp_path):
    ctkm, pog = _frames()
    joined = main.JN.join_ctkm_pog(ctkm, pog)
    output = tmp_path / "pog-control-rows.xlsx"
    summary = {
        "Thang bao cao": "2026-09",
        "Ngay nghiep vu": pd.Timestamp("2026-09-15"),
        "Thang snapshot POG": "2026-09",
        "So dong POG khong co CTKM phu hop": 0,
        "So item POG khong co CTKM phu hop": 0,
    }

    main.EX.export_final_report(
        output_path=output,
        df_matched=joined.df_matched,
        df_ctkm_only=joined.df_ctkm_only,
        df_pog_only=joined.df_pog_only,
        overlap_warnings=[],
        mapping_logs=[],
        reconciliation_summary=summary,
        df_fefo_actions=pd.DataFrame(),
        df_pog_review=joined.df_pog_review,
    )

    workbook = load_workbook(output, data_only=True)

    # User-facing review sheet contains only the real invalid POG row.
    review = workbook["POG cần review"]
    review_values = [cell.value for row in review.iter_rows() for cell in row]
    assert "POG:Export:1616" in review_values
    assert "POG:Export:1617" not in review_values
    assert "POG:Export:1618" not in review_values
    assert "Total" not in review_values

    # Control rows remain source-auditable without cluttering operational review.
    assert "POG metadata" in workbook.sheetnames
    metadata_sheet = workbook["POG metadata"]
    assert metadata_sheet.sheet_state == "hidden"
    metadata_values = [cell.value for row in metadata_sheet.iter_rows() for cell in row]
    assert "POG:Export:1617" in metadata_values
    assert "POG:Export:1618" in metadata_values
    assert "Total" in metadata_values
    assert "summary_total" in metadata_values
    assert "snapshot_metadata" in metadata_values
    assert any(
        isinstance(value, str) and "DATEID 9/1/2026 - 9/30/2026" in value
        for value in metadata_values
    )
