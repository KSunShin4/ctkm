from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from backend import main_api  # installs production scanner layers
from backend import tag_audit_business_date_contract as business_date
from backend import tag_audit_contract
from backend import tag_audit_progress_contract as progress_contract


HEADERS = [
    "Phạm vi áp dụng (KEY Channel)", "Kênh chương trình (Channel)",
    "Mã sản phẩm", "Tên sản phẩm", "Đơn vị sản phẩm (Promo UoM)",
    "POG", "Vị trí", "Facing", "Công việc", "Ngày gắn", "Ngày tháo",
    "Campaign", "Scheme detail", "From", "To", "Giá thường", "Giá KM",
    "Nguồn", "Lý do / Cảnh báo", "Product type", "Cửa hàng",
]


def _workbook_bytes(
    *, report_date="2026-09-05", report_month="2026-09",
    promo_from="2026-09-01", promo_to="2026-09-30",
    source_sheet="Gắn tem", include_guide=True,
):
    workbook = Workbook()
    workbook.remove(workbook.active)
    if include_guide:
        guide = workbook.create_sheet("HƯỚNG DẪN")
        guide.append(["DANH SÁCH THAO TÁC TEM KHUYẾN MÃI"])
        guide.append(["Tháng báo cáo", report_month])
        guide.append(["Ngày nghiệp vụ", report_date])
    for name in ("Gắn tem", "Tháo tem", "Chờ gắn"):
        ws = workbook.create_sheet(name)
        ws.append(HEADERS)
    workbook[source_sheet].append([
        "All channel", "ALL CHANNEL", "P30766", "Sản phẩm thử nghiệm", "Hộp",
        "POG-A", "1", 1, "snapshot action", promo_from, "2026-10-01",
        "FULL MONTH", "GIẢM 25%", promo_from, promo_to,
        225000, 168750, "CTKM", "", "NORMAL", "S1",
    ])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _observed(*, promo_from="2026-09-01", promo_to="2026-09-30"):
    return {
        "item_id": "P30766",
        "item_id_candidates": ["P30766"],
        "barcode": "",
        "product_name": "SAN PHAM THU NGHIEM",
        "promo_text": "GIAM 25%",
        "normal_price": 225000,
        "promo_price": 168750,
        "promo_from": promo_from,
        "promo_to": promo_to,
        "raw_text": "P30766",
        "confidence": 0.95,
    }


def test_production_installs_business_date_contract():
    assert tag_audit_contract._business_date_contract_v1 is True


def test_parser_preserves_exported_report_business_date_as_technical_metadata():
    rows = tag_audit_contract._parse_action_workbook(BytesIO(_workbook_bytes()))

    assert len(rows) == 1
    assert rows[0]["_report_business_date"] == "2026-09-05"
    assert rows[0]["_report_month"] == "2026-09"
    assert rows[0]["item_id"] == "P30766"
    assert rows[0]["channel"] == "ALL CHANNEL"
    assert rows[0]["expected_action"] == "gan_tem"


def test_same_month_old_report_date_is_visible_but_not_blocking(monkeypatch):
    monkeypatch.setattr(business_date, "current_vn_business_date", lambda: "2026-09-06")
    rows = tag_audit_contract._parse_action_workbook(BytesIO(_workbook_bytes()))

    payload = tag_audit_contract._context_payload({
        "source": "uploaded-workbook",
        "job_id": None,
        "month": None,
        "rows": rows,
    })

    assert payload["business_date"] == "2026-09-05"
    assert payload["scanner_business_date"] == "2026-09-06"
    assert payload["business_date_stale"] is True
    assert payload["business_date_blocked"] is False
    assert payload["scanner_reclassifies_actions"] is True
    assert payload["report_month"] == "2026-09"


def test_same_workbook_next_day_active_promo_scans_as_gan_tem(monkeypatch):
    monkeypatch.setattr(business_date, "current_vn_business_date", lambda: "2026-09-06")
    rows = tag_audit_contract._parse_action_workbook(BytesIO(_workbook_bytes()))
    original_action = rows[0]["expected_action"]

    result = tag_audit_contract.compare_observation(rows, _observed())

    assert result["verdict"]["code"] == "correct"
    assert result["expected"]["expected_action"] == "gan_tem"
    assert result["progress_eligible"] is True
    assert rows[0]["expected_action"] == original_action  # source/cache row not mutated


def test_promo_that_ended_yesterday_is_derived_thao_tem(monkeypatch):
    monkeypatch.setattr(business_date, "current_vn_business_date", lambda: "2026-09-06")
    rows = tag_audit_contract._parse_action_workbook(BytesIO(_workbook_bytes(
        promo_to="2026-09-05", source_sheet="Gắn tem",
    )))

    result = tag_audit_contract.compare_observation(
        rows, _observed(promo_to="2026-09-05")
    )

    assert result["verdict"]["code"] == "remove_tag"
    assert result["expected"]["expected_action"] == "thao_tem"
    assert result["progress_eligible"] is False


def test_future_promo_is_derived_cho_gan(monkeypatch):
    monkeypatch.setattr(business_date, "current_vn_business_date", lambda: "2026-09-06")
    rows = tag_audit_contract._parse_action_workbook(BytesIO(_workbook_bytes(
        promo_from="2026-09-07", source_sheet="Gắn tem",
    )))

    result = tag_audit_contract.compare_observation(
        rows, _observed(promo_from="2026-09-07")
    )

    assert result["verdict"]["code"] == "early_tag"
    assert result["expected"]["expected_action"] == "cho_gan"
    assert result["progress_eligible"] is False


def test_missing_promo_date_becomes_review_not_snapshot_truth(monkeypatch):
    monkeypatch.setattr(business_date, "current_vn_business_date", lambda: "2026-09-06")
    rows = tag_audit_contract._parse_action_workbook(BytesIO(_workbook_bytes(
        promo_from="", source_sheet="Gắn tem",
    )))

    result = tag_audit_contract.compare_observation(rows, _observed(promo_from=""))

    assert result["verdict"]["code"] == "review_business_date"
    assert result["expected"]["expected_action"] == "review"
    assert result["progress_eligible"] is False


def test_checklist_denominator_uses_current_date_not_snapshot_sheet(monkeypatch):
    monkeypatch.setattr(business_date, "current_vn_business_date", lambda: "2026-09-06")
    active = tag_audit_contract._parse_action_workbook(BytesIO(_workbook_bytes()))[0]
    ended = dict(active, item_id="P30002", promo_to="2026-09-05", expected_action="gan_tem")
    future = dict(active, item_id="P30003", promo_from="2026-09-07", expected_action="gan_tem")

    items = progress_contract.checklist_rows(tag_audit_contract, [future, ended, active])

    assert [item["item_id"] for item in items] == ["P30766"]


def test_current_date_reclassification_is_row_order_independent(monkeypatch):
    monkeypatch.setattr(business_date, "current_vn_business_date", lambda: "2026-09-06")
    active = tag_audit_contract._parse_action_workbook(BytesIO(_workbook_bytes()))[0]
    second = dict(active, item_id="P30002", campaign_name="SECOND")

    first = progress_contract.context_id(tag_audit_contract, [active, second])
    second_order = progress_contract.context_id(tag_audit_contract, [second, active])

    assert first == second_order


def test_known_other_report_month_still_blocks(monkeypatch):
    monkeypatch.setattr(business_date, "current_vn_business_date", lambda: "2026-09-06")
    rows = tag_audit_contract._parse_action_workbook(BytesIO(_workbook_bytes(
        report_month="2026-08",
    )))

    payload = tag_audit_contract._context_payload({
        "source": "uploaded-workbook", "job_id": None, "month": None, "rows": rows,
    })
    assert payload["business_date_blocked"] is True

    with pytest.raises(HTTPException) as exc_info:
        tag_audit_contract.compare_observation(rows, _observed())
    assert exc_info.value.status_code == 409
    assert "2026-08" in str(exc_info.value.detail)


def test_workbook_without_metadata_is_fail_closed(monkeypatch):
    monkeypatch.setattr(business_date, "current_vn_business_date", lambda: "2026-09-06")
    rows = tag_audit_contract._parse_action_workbook(
        BytesIO(_workbook_bytes(include_guide=False))
    )

    payload = tag_audit_contract._context_payload({
        "source": "uploaded-workbook", "job_id": None, "month": None, "rows": rows,
    })
    assert payload["business_date"] is None
    assert payload["business_date_blocked"] is True
    assert payload["business_date_block_reason"] == "missing_report_metadata"
    assert payload["scanner_reclassifies_actions"] is False

    with pytest.raises(HTTPException) as exc_info:
        tag_audit_contract.compare_observation(rows, _observed())
    assert exc_info.value.status_code == 409
    assert "metadata" in str(exc_info.value.detail).lower()


def test_incomplete_report_metadata_is_fail_closed(monkeypatch):
    monkeypatch.setattr(business_date, "current_vn_business_date", lambda: "2026-09-06")
    raw = _workbook_bytes(report_date="", report_month="2026-09")
    rows = tag_audit_contract._parse_action_workbook(BytesIO(raw))

    payload = tag_audit_contract._context_payload({
        "source": "uploaded-workbook", "job_id": None, "month": None, "rows": rows,
    })

    assert payload["business_date_blocked"] is True
    assert payload["business_date_block_reason"] == "missing_report_metadata"
    assert payload["scanner_reclassifies_actions"] is False
    with pytest.raises(HTTPException) as exc_info:
        tag_audit_contract.compare_observation(rows, _observed())
    assert exc_info.value.status_code == 409
