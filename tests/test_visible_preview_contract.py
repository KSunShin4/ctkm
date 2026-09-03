from openpyxl import Workbook

from backend.preview_contract import build_visible_preview, sheet_kind


def test_preview_matches_final_workbook_visibility(tmp_path):
    path = tmp_path / "report.xlsx"
    wb = Workbook()

    guide = wb.active
    guide.title = "HƯỚNG DẪN"
    guide["A1"] = "DANH SÁCH THAO TÁC"

    action = wb.create_sheet("Gắn tem")
    action.append(["Mã sản phẩm", "task_id", "Kênh chương trình"])
    action.append(["P1001", "TASK-1", "Offline"])
    action.append(["P1002", "TASK-2", "Online"])
    action.append(["P1003", "TASK-3", "All Channel"])
    action.column_dimensions["B"].hidden = True
    action.row_dimensions[3].hidden = True

    empty_review = wb.create_sheet("Cần review")
    empty_review.append(["Ghi chu"])
    empty_review.append(["Không có CTKM cần review."])

    summary = wb.create_sheet("Tom tat")
    summary.append(["Chi so", "Gia tri"])
    summary.append(["Metric visible", 2])
    summary.append(["Legacy hidden", 99])
    summary.row_dimensions[3].hidden = True

    metadata = wb.create_sheet("POG metadata")
    metadata.append(["ITEMID", "pog_control_kind"])
    metadata.append(["", "snapshot_metadata"])
    metadata.sheet_state = "hidden"

    mapping = wb.create_sheet("Mapping log")
    mapping.append(["field", "source"])
    mapping.append(["item_id", "ITEMID"])
    mapping.sheet_state = "hidden"

    wb.save(path)
    wb.close()

    preview = build_visible_preview(path, rows=1)["sheets"]

    assert "HƯỚNG DẪN" not in preview
    assert "POG metadata" not in preview
    assert "Mapping log" not in preview

    assert preview["Gắn tem"]["sheet_kind"] == "action"
    assert preview["Gắn tem"]["headers"] == ["Mã sản phẩm", "Kênh chương trình"]
    assert preview["Gắn tem"]["total_rows"] == 2
    assert preview["Gắn tem"]["rows"] == [["P1001", "Offline"]]

    # Exporter _safe() keeps a human-readable placeholder row in Excel, but it is
    # not a real review task and therefore must show zero in the preview badge.
    assert preview["Cần review"]["total_rows"] == 0
    assert preview["Cần review"]["rows"] == []

    # Hidden legacy summary aliases stay out of web preview exactly as in Excel.
    assert preview["Tom tat"]["sheet_kind"] == "summary"
    assert preview["Tom tat"]["total_rows"] == 1
    assert preview["Tom tat"]["rows"] == [["Metric visible", "2"]]


def test_sheet_kind_contract_is_explicit_for_user_facing_groups():
    assert sheet_kind("Gắn tem") == "action"
    assert sheet_kind("Tong hop") == "reconciliation"
    assert sheet_kind("Audit hành động") == "audit"
    assert sheet_kind("Canh bao trung thoi gian") == "audit"
    assert sheet_kind("Xác nhận CTKM Clearance") == "confirmation"
    assert sheet_kind("Tom tat") == "summary"
    assert sheet_kind("Other") == "data"
