import importlib.util
from pathlib import Path
import re
import sys
import zipfile

import pytest


def _module():
    path = Path(__file__).resolve().parents[1] / "modules" / "10_tem_r.py"
    spec = importlib.util.spec_from_file_location("tem_r_hardening", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _basic_workbook(path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "ton"])
    sheet.append(["S1", "G1", "1/1", 1, "P1", "One", 2])
    sheet.append(["S1", "G1", "1/1", 2, "P2", "Two", 0])
    workbook.save(path)


def test_tem_r_always_disables_external_link_parsing(tmp_path, monkeypatch):
    tem_r = _module()
    import openpyxl

    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _basic_workbook(source)

    real_load_workbook = openpyxl.load_workbook
    calls = []

    def recording_load_workbook(*args, **kwargs):
        calls.append(dict(kwargs))
        return real_load_workbook(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", recording_load_workbook)

    inspected = tem_r.inspect_tem_r_workbook(source)
    tem_r.create_tem_r_workbook(source, output, stock_column="ton")

    assert inspected["sheet_name"] == "Export"
    assert len(calls) == 3
    assert all(call.get("keep_links") is False for call in calls)


def test_selected_sheet_finds_header_only_inside_that_sheet(tmp_path):
    tem_r = _module()
    from openpyxl import Workbook, load_workbook

    source = tmp_path / "multi-sheet.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()

    export = workbook.active
    export.title = "Export"
    export.append(["STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "ton"])
    export.append(["WRONG", "WRONG", "WRONG", 1, "BAD", "Wrong sheet", 9])

    chosen = workbook.create_sheet("Chosen")
    chosen.append(["Báo cáo Tem R"])
    chosen.append([])
    chosen.append(["STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "ton"])
    chosen.append(["S1", "G1", "1/1", 1, "P1", "One", 2])
    chosen.append(["S1", "G1", "1/1", 2, "P2", "Two", 0])
    workbook.save(source)

    plan = tem_r.create_tem_r_workbook(
        source,
        output,
        stock_column="ton",
        sheet_name="Chosen",
        header_row=None,
    )

    assert plan.rows["_excel_row"].tolist() == [4, 5]
    assert plan.rows.loc[1, "TEM R"] == "P1"
    result = load_workbook(output, data_only=True)
    assert result["Chosen"].cell(5, 8).value == "P1"
    assert result["Export"].max_column == 7


def test_formula_stock_with_cached_numeric_value_is_rejected(tmp_path):
    tem_r = _module()
    from openpyxl import Workbook

    source = tmp_path / "formula-stock.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "ton"])
    sheet.append(["S1", "G1", "1/1", 1, "P1", "One", "=2"])
    sheet.append(["S1", "G1", "1/1", 2, "P2", "Two", 0])
    workbook.save(source)

    # openpyxl does not calculate formulas. Inject the cached value that a real
    # Excel export may contain. Tem R must not trust that cached stock value.
    with zipfile.ZipFile(source, "r") as zin:
        members = {name: zin.read(name) for name in zin.namelist()}
    patched, replacements = re.subn(
        br'(<c r="G2"[^>]*><f>2</f><v>)(?:[^<]*)(</v></c>)',
        br'\g<1>2\g<2>',
        members["xl/worksheets/sheet1.xml"],
    )
    assert replacements == 1
    members["xl/worksheets/sheet1.xml"] = patched
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, content in members.items():
            zout.writestr(name, content)

    with pytest.raises(ValueError, match="đang chứa công thức Excel") as exc_info:
        tem_r.create_tem_r_workbook(source, output, stock_column="ton")

    message = str(exc_info.value)
    assert "ton" in message
    assert "dòng 2" in message
    assert "Paste Values" in message
    assert not output.exists()


def test_header_scan_ignores_inflated_used_range(tmp_path):
    tem_r = _module()
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    source = tmp_path / "inflated-range.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "ton"])
    sheet.append(["S1", "G1", "1/1", 1, "P1", "One", 2])
    # Formatting the last Excel column makes max_column = 16384 even though the
    # actual Tem R table has only seven columns.
    sheet["XFD1"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    workbook.save(source)

    result = tem_r.inspect_tem_r_workbook(source)

    assert result["sheet_name"] == "Export"
    assert result["columns"] == [
        "STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "ton"
    ]
    assert result["suggested_stock_column"] == "ton"
