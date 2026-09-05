import importlib.util
from pathlib import Path
import sys

from openpyxl import Workbook


def _loader_module():
    import modules.config_loader

    path = Path(__file__).resolve().parents[1] / "modules" / "1_loader.py"
    spec = importlib.util.spec_from_file_location("modules.loader_recovery_test", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_normal_workbook_is_opened_without_defined_name_repair(tmp_path, monkeypatch):
    source = tmp_path / "normal.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ITEMID", "Campaign", "From", "To"])
    sheet.append(["P1", "A", "2026-09-01", "2026-09-30"])
    workbook.save(source)

    loader = _loader_module()
    called = False

    def unexpected_repair(_filepath):
        nonlocal called
        called = True
        raise AssertionError("normal workbook must not be rewritten in memory")

    monkeypatch.setattr(loader._core, "_strip_broken_defined_names", unexpected_repair)
    loaded = loader.load_workbook_raw(str(source), selected_sheet_names=["Sheet"])

    assert called is False
    assert loaded.sheets["Sheet"].df.iloc[0]["ITEMID"] == "P1"


def test_defined_name_repair_is_used_only_after_original_open_fails(tmp_path, monkeypatch):
    import openpyxl

    source = tmp_path / "fallback.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ITEMID", "Campaign", "From", "To"])
    sheet.append(["P1", "A", "2026-09-01", "2026-09-30"])
    workbook.save(source)

    loader = _loader_module()
    real_load = openpyxl.load_workbook
    calls = []

    def flaky_load(target, *args, **kwargs):
        calls.append(target)
        if len(calls) == 1:
            raise ValueError("simulated original workbook parser failure")
        return real_load(target, *args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", flaky_load)
    loaded = loader.load_workbook_raw(str(source), selected_sheet_names=["Sheet"])

    assert len(calls) == 2
    assert loaded.sheets["Sheet"].df.iloc[0]["ITEMID"] == "P1"
