import importlib.util
from pathlib import Path
import sys

import pandas as pd
from openpyxl import Workbook


def _performance_module():
    root = Path(__file__).resolve().parents[1]
    path = root / "modules" / "ctkm_performance.py"
    spec = importlib.util.spec_from_file_location("ctkm_excel_layout_test", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _exporter_module():
    root = Path(__file__).resolve().parents[1]
    path = root / "modules" / "9_exporter.py"
    spec = importlib.util.spec_from_file_location("ctkm_exporter_layout_test", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _normalize(value):
    return " ".join(str(value).strip().lower().split())


def test_key_channel_and_channel_are_both_preserved():
    perf = _performance_module()
    frame = pd.DataFrame({
        "KEY Channel": ["OFFLINE", "ONLINE"],
        "Channel": ["COMBO ĐI SINH", "MẸ & BÉ"],
        "Item code": ["P1", "P2"],
    })
    mapping = {"KEY Channel": "channel", "Item code": "item_id"}

    result = perf.preserve_both_channel_mappings(
        mapping, frame, normalize_header=_normalize
    )

    assert result["KEY Channel"] == "key_channel"
    assert result["Channel"] == "channel"


def test_action_sheet_starts_with_approved_user_columns():
    perf = _performance_module()
    frame = pd.DataFrame({
        "task_id": ["T1"],
        "promotion_id": ["P1"],
        "pog_position_id": ["L1"],
        "item_id": ["P123"],
        "item_name": ["Thuốc A"],
        "promo_uom": ["Hộp"],
        "key_channel": ["OFFLINE"],
        "channel": ["MẸ & BÉ"],
        "campaign_name": ["CTKM A"],
        "store_id": ["S1"],
        "pog_code": ["POG1"],
        "location_id": [1],
        "scheduled_action": ["gắn tem"],
        "processing_status": ["Chưa xử lý"],
    })

    result = perf.prepare_action_sheet_user_first(frame)
    columns = list(result.columns)

    assert columns[:5] == [
        "key_channel", "channel", "item_id", "item_name", "promo_uom"
    ]
    assert columns.index("item_id") < columns.index("task_id")
    assert columns.index("key_channel") < columns.index("promotion_id")
    assert columns.index("channel") < columns.index("pog_position_id")


def test_selected_output_pins_both_channels_before_identifier():
    perf = _performance_module()
    frame = pd.DataFrame({
        "item_id": ["P1"],
        "item_name": ["Thuốc A"],
        "key_channel": ["ALL CHANNEL - EXCEPT MKP"],
        "channel": ["COMBO ĐI SINH"],
        "campaign_name": ["CTKM A"],
    })

    result = perf.filter_selected_columns_keep_channels(frame, ["item_name"])

    assert list(result.columns) == ["key_channel", "channel", "item_id", "item_name"]


def test_selected_output_does_not_reorder_columns_by_checkbox_click_sequence():
    perf = _performance_module()
    frame = pd.DataFrame({
        "item_id": ["P1"],
        "item_name": ["Thuốc A"],
        "campaign_name": ["CTKM A"],
        "key_channel": ["Offline"],
        "channel": ["OFFLINE"],
    })

    result = perf.filter_selected_columns_keep_channels(
        frame, ["campaign_name", "item_name"]
    )

    assert list(result.columns) == [
        "key_channel", "channel", "item_id", "item_name", "campaign_name"
    ]


def test_protected_action_sheet_still_allows_filter_and_sort():
    exporter = _exporter_module()
    workbook = Workbook()
    ws = workbook.active
    ws.append(["processing_status", "processed_by", "processed_at", "staff_note"])
    ws.append(["Chưa xử lý", "", "", ""])

    exporter._configure_action_inputs(ws)

    assert ws.protection.sheet is True
    assert ws.protection.autoFilter is False
    assert ws.protection.sort is False


def test_guide_sheet_is_created_first_and_mentions_channel_first_workflow():
    exporter = _exporter_module()
    workbook = Workbook()
    del workbook[workbook.sheetnames[0]]
    empty = pd.DataFrame()

    exporter._create_guide_sheet(
        workbook, empty, empty, empty, empty, empty,
        {"Thang bao cao": "2026-09", "Ngay nghiep vu": "2026-09-04"},
        empty, empty, empty,
    )

    assert workbook.sheetnames[0] == "HƯỚNG DẪN"
    text = " ".join(str(cell.value or "") for row in workbook["HƯỚNG DẪN"] for cell in row)
    assert "Lọc KEY Channel trước" in text
    assert "sau đó lọc Channel" in text


def test_total_action_frame_matches_approved_21_column_mockup():
    exporter = _exporter_module()
    actions = pd.DataFrame([{
        "key_channel": "All channel",
        "channel": "All Channel",
        "item_id": "P1001",
        "item_name": "Thuốc A",
        "promo_uom": "Hộp",
        "pog_code": "A01",
        "location_id": "1",
        "facing": 2,
        "scheduled_action": "gắn tem",
        "ngay_gan_tem": pd.Timestamp("2026-09-01"),
        "ngay_thao_tem": pd.Timestamp("2026-10-01"),
        "campaign_name": "FULL MONTH",
        "campaign_detail": "Giảm 10%",
        "promo_from": pd.Timestamp("2026-09-01"),
        "promo_to": pd.Timestamp("2026-09-30"),
        "normal_price": 100000,
        "promo_price": 90000,
        "program_source": "CTKM",
        "reason": "",
        "product_type": "SKU",
        "store_id": "SGPMC1885",
        "in_report_scope": True,
        "task_id": "T1",
    }])

    total = exporter._build_total_action_frame(actions)

    assert list(total.columns) == exporter._APPROVED_ACTION_COLUMNS
    assert len(total.columns) == 21
    assert total.iloc[0]["key_channel"] == "All channel"
    assert total.iloc[0]["channel"] == "All Channel"
    assert total.iloc[0]["action"] == "gắn tem"


def test_total_sheet_is_inserted_after_guide_and_uses_mockup_headers():
    exporter = _exporter_module()
    workbook = Workbook()
    workbook.active.title = "HƯỚNG DẪN"
    workbook.create_sheet("Gắn tem")
    actions = pd.DataFrame([{
        "key_channel": "Offline",
        "channel": "OFFLINE",
        "item_id": "P1",
        "item_name": "Thuốc A",
        "promo_uom": "Hộp",
        "action": "gắn tem",
    }])

    exporter._write_total_action_sheet(workbook, actions)

    assert workbook.sheetnames[:3] == ["HƯỚNG DẪN", "TỔNG HỢP - TEM", "Gắn tem"]
    ws = workbook["TỔNG HỢP - TEM"]
    headers = [cell.value for cell in ws[1]]
    assert headers[:5] == [
        "KEY Channel", "Channel", "Mã sản phẩm", "Tên sản phẩm",
        "Đơn vị sản phẩm (Promo UoM)",
    ]
    assert ws.freeze_panes == "C2"
    assert ws.auto_filter.ref is not None



def test_selected_output_always_pins_product_name_and_uom_when_available():
    perf = _performance_module()
    frame = pd.DataFrame({
        "key_channel": ["Offline"],
        "channel": ["OFFLINE"],
        "item_id": ["P1"],
        "item_name": ["Thuốc A"],
        "promo_uom": ["Hộp"],
        "campaign_name": ["Campaign A"],
    })

    result = perf.filter_selected_columns_keep_channels(frame, ["campaign_name"])

    assert list(result.columns) == [
        "key_channel", "channel", "item_id", "item_name", "promo_uom", "campaign_name"
    ]
