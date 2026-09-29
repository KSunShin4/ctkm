import asyncio
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from backend import main_api
from backend.auxiliary_workflows_contract import (
    _build_inventory_inspection,
    _resolve_tem_r_selection,
    validate_tem_r_workbook,
)


TEM_R_COLUMNS = [
    "STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "ton",
]
TEM_R_MAPPING = {
    "STOREID": "STOREID",
    "POGCODE": "POGCODE",
    "FIXELID": "FIXELID",
    "LOCATIONID": "LOCATIONID",
    "ITEMID": "ITEMID",
    "ITEMNAME": "ITEMNAME",
}


def _save_tem_r_source(path, *, second_sheet=False, control_only=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append(TEM_R_COLUMNS)
    if control_only:
        ws.append(["", "", "", "", "Total", "", 0])
    else:
        ws.append(["S1", "P01", "F01", 1, "P1001", "Có tồn", 2])
        ws.append(["S1", "P01", "F01", 2, "P1002", "Hết tồn", 0])
    if second_sheet:
        other = wb.create_sheet("Other")
        other.append(TEM_R_COLUMNS)
        other.append(["S1", "P02", "F02", 1, "P2001", "Khác", 1])
    wb.save(path)
    wb.close()


def _xlsx_bytes(builder):
    stream = BytesIO()
    wb = Workbook()
    builder(wb)
    wb.save(stream)
    wb.close()
    return stream.getvalue()


class FakeUpload:
    def __init__(self, filename, content):
        self.filename = filename
        self._content = content

    async def read(self, size):
        return self._content[:size]


def test_tem_r_preflight_resolves_exact_selection_and_rejects_stale_configuration(tmp_path):
    source = tmp_path / "tem_r.xlsx"
    _save_tem_r_source(source)

    sheet, header, mapping, _inspection = _resolve_tem_r_selection(
        main_api._tem_r_mod,
        source,
        stock_column="ton",
        column_mapping=TEM_R_MAPPING,
        sheet_name="Export",
        header_row=1,
    )
    assert (sheet, header) == ("Export", 1)
    assert mapping == TEM_R_MAPPING

    with pytest.raises(ValueError, match="không còn khớp file hiện tại"):
        _resolve_tem_r_selection(
            main_api._tem_r_mod,
            source,
            stock_column="ton",
            column_mapping=TEM_R_MAPPING,
            sheet_name="Export",
            header_row=2,
        )

    with pytest.raises(ValueError, match="đồng thời được dùng"):
        _resolve_tem_r_selection(
            main_api._tem_r_mod,
            source,
            stock_column="ITEMID",
            column_mapping=TEM_R_MAPPING,
            sheet_name="Export",
            header_row=1,
        )


def test_tem_r_preflight_requires_sheet_when_multiple_and_rejects_control_only(tmp_path):
    multi = tmp_path / "multi.xlsx"
    _save_tem_r_source(multi, second_sheet=True)
    with pytest.raises(ValueError, match="nhiều sheet dữ liệu"):
        _resolve_tem_r_selection(
            main_api._tem_r_mod,
            multi,
            stock_column="ton",
            column_mapping=TEM_R_MAPPING,
            sheet_name=None,
            header_row=None,
        )

    control = tmp_path / "control.xlsx"
    _save_tem_r_source(control, control_only=True)
    with pytest.raises(ValueError, match="không có dòng sản phẩm ITEMID"):
        _resolve_tem_r_selection(
            main_api._tem_r_mod,
            control,
            stock_column="ton",
            column_mapping=TEM_R_MAPPING,
            sheet_name="Export",
            header_row=1,
        )


def test_tem_r_final_workbook_must_match_plan_and_audit_counts(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _save_tem_r_source(source)

    plan = main_api._tem_r_mod.create_tem_r_workbook(
        source,
        output,
        max_location_distance=1,
        stock_column="ton",
        column_mapping=TEM_R_MAPPING,
        sheet_name="Export",
        header_row=1,
    )
    validate_tem_r_workbook(output, plan, sheet_name="Export", header_row=1)

    wb = load_workbook(output)
    del wb["Tem R ngoại lệ"]
    wb.save(output)
    wb.close()
    with pytest.raises(RuntimeError, match="thiếu sheet bắt buộc"):
        validate_tem_r_workbook(output, plan, sheet_name="Export", header_row=1)


def test_inspect_tem_r_always_removes_temporary_directory(monkeypatch, tmp_path):
    content = _xlsx_bytes(lambda wb: (
        setattr(wb.active, "title", "Export"),
        wb.active.append(TEM_R_COLUMNS),
        wb.active.append(["S1", "P01", "F01", 1, "P1001", "Prod", 1]),
    ))
    monkeypatch.setattr(main_api, "_TMP_DIR", tmp_path)

    result = asyncio.run(main_api.inspect_tem_r_file(
        tem_r_file=FakeUpload("tem_r.xlsx", content),
        _admin="admin",
    ))
    assert result["sheet_name"] == "Export"
    assert not list(tmp_path.glob("inspect_tem_r_*"))


def test_invalid_tem_r_selection_fails_before_queue_submission(monkeypatch, tmp_path):
    content = _xlsx_bytes(lambda wb: (
        setattr(wb.active, "title", "Export"),
        wb.active.append(TEM_R_COLUMNS),
        wb.active.append(["S1", "P01", "F01", 1, "P1001", "Prod", 1]),
    ))
    monkeypatch.setattr(main_api, "_TMP_DIR", tmp_path / "tmp")
    monkeypatch.setattr(main_api, "_OUT_DIR", tmp_path / "out")
    main_api._TMP_DIR.mkdir()
    main_api._OUT_DIR.mkdir()
    submitted = []
    monkeypatch.setattr(main_api, "_submit_job", lambda *args: submitted.append(args))

    with pytest.raises(main_api.HTTPException) as exc:
        asyncio.run(main_api.enqueue_tem_r(
            tem_r_file=FakeUpload("tem_r.xlsx", content),
            max_location_distance=1,
            stock_column="missing-stock",
            column_mapping=__import__("json").dumps(TEM_R_MAPPING),
            sheet_name="Export",
            header_row=1,
            _admin="admin",
        ))
    assert exc.value.status_code == 422
    assert submitted == []
    assert not list(main_api._TMP_DIR.glob("tem_r_*"))


def test_expiration_api_requires_exact_dd_mm_yyyy_input_format():
    frame = pd.DataFrame({
        "ITEMID": ["P1001"],
        "BATCHNUMER": ["B1"],
        "EXPDATE": ["01/01/2030"],
    })
    exact = main_api._expiration_mod.check_entries(frame, "P1001;B1;01/01/2030")
    iso = main_api._expiration_mod.check_entries(frame, "P1001;B1;2030-01-01")
    dashed = main_api._expiration_mod.check_entries(frame, "P1001;B1;01-01-2030")

    assert exact[0]["status"] == "matched"
    assert iso[0]["status"] == "invalid"
    assert dashed[0]["status"] == "invalid"
    assert iso[0]["message"] == "Hạn dùng phải có dạng DD/MM/YYYY"


def test_expiration_user_file_errors_are_422_and_route_stays_public():
    content = _xlsx_bytes(lambda wb: (
        setattr(wb.active, "title", "Wrong"),
        wb.active.append(["A", "B", "C"]),
        wb.active.append([1, 2, 3]),
    ))
    with pytest.raises(main_api.HTTPException) as exc:
        asyncio.run(main_api.check_expiration(
            expiration_file=FakeUpload("expiration.xlsx", content),
            entries="P1001;B1;01/01/2030",
        ))
    assert exc.value.status_code == 422

    route = next(
        route for route in main_api.app.routes
        if getattr(route, "path", None) == "/api/check-expiration"
    )
    dependency_calls = [getattr(dep, "call", None) for dep in route.dependant.dependencies]
    assert main_api.require_admin not in dependency_calls



def test_inventory_inspection_uses_store_level_stock_without_summing_repeated_pog_rows(tmp_path):
    source = tmp_path / "inventory.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append(TEM_R_COLUMNS)
    ws.append(["S1", "POG-A", "F1", 1, "P1001", "Product 1", 3])
    ws.append(["S1", "POG-B", "F2", 2, "P1001", "Product 1", 3])
    ws.append(["S1", "POG-C", "F3", "", "P2001", "Product 2", 0])
    ws.append(["S1", "POG-D", "F4", 4, "P3001", "Product 3", "#N/A"])
    wb.save(source)
    wb.close()

    result = _build_inventory_inspection(
        main_api._tem_r_mod,
        source,
        stock_column="ton",
        column_mapping=TEM_R_MAPPING,
        sheet_name="Export",
        header_row=1,
    )
    by_item = {row["item_id"]: row for row in result["rows"]}

    assert by_item["P1001"]["stock"] == 3
    assert by_item["P1001"]["stock_values"] == [3]
    assert by_item["P1001"]["position_count"] == 2
    assert by_item["P1001"]["status"] == "ok"

    assert by_item["P2001"]["stock"] == 0
    assert by_item["P2001"]["status"] == "location_missing"
    assert by_item["P3001"]["stock"] is None
    assert by_item["P3001"]["status"] == "stock_missing"


def test_inventory_inspection_flags_conflicting_store_stock_instead_of_choosing_or_summing(tmp_path):
    source = tmp_path / "inventory-conflict.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append(TEM_R_COLUMNS)
    ws.append(["S1", "POG-A", "F1", 1, "P1001", "Product 1", 3])
    ws.append(["S1", "POG-B", "F2", 2, "P1001", "Product 1", 5])
    wb.save(source)
    wb.close()

    result = _build_inventory_inspection(
        main_api._tem_r_mod,
        source,
        stock_column="ton",
        column_mapping=TEM_R_MAPPING,
        sheet_name="Export",
        header_row=1,
    )
    row = result["rows"][0]
    assert row["stock"] is None
    assert row["stock_values"] == [3, 5]
    assert row["position_count"] == 2
    assert row["status"] == "stock_conflict"


def test_inventory_inspect_route_is_store_facing_and_cleans_temporary_file(monkeypatch, tmp_path):
    content = _xlsx_bytes(lambda wb: (
        setattr(wb.active, "title", "Export"),
        wb.active.append(TEM_R_COLUMNS),
        wb.active.append(["S1", "P01", "F01", 1, "P1001", "Prod", 2]),
    ))
    monkeypatch.setattr(main_api, "_TMP_DIR", tmp_path)

    result = asyncio.run(main_api.inspect_inventory(
        tem_r_file=FakeUpload("inventory.xlsx", content),
        stock_column="ton",
        column_mapping=__import__("json").dumps(TEM_R_MAPPING),
        sheet_name="Export",
        header_row=1,
    ))
    assert result["total_items"] == 1
    assert result["rows"][0]["stock"] == 2
    assert not list(tmp_path.glob("inspect_inventory_*"))
