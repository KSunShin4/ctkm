import pandas as pd
from openpyxl import Workbook

from backend import main_api


def _rgb(cell) -> str:
    value = cell.fill.fgColor.rgb or ""
    return str(value)[-6:].upper()


def test_export_speed_contract_is_installed_on_shared_exporter_core():
    exporter = main_api._main_mod._core.EX
    exporter_core = exporter._core
    assert main_api.app.state.export_speed_contract_v1 is True
    assert getattr(exporter_core._apply_professional_style, "_ctkm_export_speed_contract_v1", False)
    assert getattr(exporter_core._apply_alternating_shelf_colors, "_ctkm_export_speed_contract_v1", False)


def test_single_pass_style_preserves_shelf_review_and_collision_precedence():
    exporter_core = main_api._main_mod._core.EX._core
    wb = Workbook()
    ws = wb.active
    ws.title = "Gắn tem"
    headers = [
        "store_id", "pog_code", "action", "location_completeness",
        "location_needs_review", "promotion_count_at_position",
        "has_multiple_promotions_at_position",
    ]
    ws.append(headers)
    ws.append(["S1", "P1", "gắn tem", "Đủ vị trí", False, 1, False])
    ws.append(["S1", "P1", "tháo tem", "Đủ vị trí", False, 1, False])
    ws.append(["S2", "P2", "review", "Thiếu một phần thông tin vị trí", True, 2, True])

    frame = pd.DataFrame([
        dict(zip(headers, row))
        for row in [
            ["S1", "P1", "gắn tem", "Đủ vị trí", False, 1, False],
            ["S1", "P1", "tháo tem", "Đủ vị trí", False, 1, False],
            ["S2", "P2", "review", "Thiếu một phần thông tin vị trí", True, 2, True],
        ]
    ])

    exporter_core._apply_professional_style(ws, "Gắn tem", frame)
    exporter_core._apply_alternating_shelf_colors(ws, frame)

    # Shelf grouping is the final row background and therefore overrides action
    # green/orange fills, exactly as the historical second shelf pass did.
    assert _rgb(ws["A2"]) == "EAF3FF"
    assert _rgb(ws["A3"]) == "EAF3FF"
    assert _rgb(ws["A4"]) == "FFF4E5"

    # Review and same-position diagnostics intentionally override shelf color on
    # their targeted diagnostic fields.
    assert _rgb(ws["D4"]) == "FFC7CE"
    assert _rgb(ws["E4"]) == "FFC7CE"
    assert _rgb(ws["F4"]) == "FFF2CC"
    assert _rgb(ws["G4"]) == "FFF2CC"

    assert ws.column_dimensions["A"].width >= 12
    assert ws.row_dimensions[2].height == 22
    wb.close()
