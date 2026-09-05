import asyncio
from io import BytesIO
import json
import threading

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from backend import main_api
from backend.auxiliary_workflows_contract import (
    _resolve_tem_r_selection,
    validate_tem_r_workbook,
)
from backend.runtime_contract import ReservedHeavyJobQueue


MAPPING = {
    "STOREID": "STOREID",
    "POGCODE": "POGCODE",
    "FIXELID": "FIXELID",
    "LOCATIONID": "LOCATIONID",
    "ITEMID": "ITEMID",
    "ITEMNAME": "ITEMNAME",
}


class FakeUpload:
    def __init__(self, filename, content):
        self.filename = filename
        self._content = content

    async def read(self, size):
        return self._content[:size]


def _workbook_bytes(headers, rows, *, sheet="Export"):
    stream = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def _write_bytes(path, content):
    path.write_bytes(content)
    return path


def test_tem_r_preflight_uses_physical_header_columns_when_blank_gaps_exist(tmp_path):
    headers = [
        "STOREID", None, "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "ton",
    ]
    rows = [["S1", None, "P01", "F01", 1, "P1001", "Prod", 1]]
    source = _write_bytes(tmp_path / "gap.xlsx", _workbook_bytes(headers, rows))

    sheet, header, mapping, _inspection = _resolve_tem_r_selection(
        main_api._tem_r_mod,
        source,
        stock_column="ton",
        column_mapping=MAPPING,
        sheet_name="Export",
        header_row=1,
    )
    assert sheet == "Export"
    assert header == 1
    assert mapping["ITEMID"] == "ITEMID"


def test_tem_r_final_validator_detects_saved_cell_tampering(tmp_path):
    headers = ["STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "ton"]
    rows = [
        ["S1", "P01", "F01", 1, "P1001", "Có tồn", 2],
        ["S1", "P01", "F01", 2, "P1002", "Hết tồn", 0],
    ]
    source = _write_bytes(tmp_path / "source.xlsx", _workbook_bytes(headers, rows))
    output = tmp_path / "output.xlsx"
    plan = main_api._tem_r_mod.create_tem_r_workbook(
        source,
        output,
        max_location_distance=1,
        stock_column="ton",
        column_mapping=MAPPING,
        sheet_name="Export",
        header_row=1,
    )
    validate_tem_r_workbook(output, plan, sheet_name="Export", header_row=1)

    wb = load_workbook(output)
    ws = wb["Export"]
    tem_r_col = [cell.value for cell in ws[1]].index("TEM R") + 1
    ws.cell(3, tem_r_col).value = "P9999"
    wb.save(output)
    wb.close()

    with pytest.raises(RuntimeError, match="không khớp kế hoạch"):
        validate_tem_r_workbook(output, plan, sheet_name="Export", header_row=1)


def test_legacy_tem_r_creates_completed_downloadable_job_and_cleans_temp(monkeypatch, tmp_path):
    headers = ["STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "ton"]
    rows = [
        ["S1", "P01", "F01", 1, "P1001", "Có tồn", 2],
        ["S1", "P01", "F01", 2, "P1002", "Hết tồn", 0],
    ]
    content = _workbook_bytes(headers, rows)
    tmp_dir, out_dir = tmp_path / "tmp", tmp_path / "out"
    tmp_dir.mkdir(); out_dir.mkdir()
    monkeypatch.setattr(main_api, "_TMP_DIR", tmp_dir)
    monkeypatch.setattr(main_api, "_OUT_DIR", out_dir)
    monkeypatch.setattr(main_api, "_heavy_queue", ReservedHeavyJobQueue(2))

    result = asyncio.run(main_api.run_tem_r(
        tem_r_file=FakeUpload("tem_r.xlsx", content),
        max_location_distance=1,
        stock_column="ton",
        column_mapping=json.dumps(MAPPING),
        sheet_name="Export",
        header_row=1,
        _admin="admin",
    ))
    job_id = result["job_id"]
    try:
        job = main_api._jobs[job_id]
        assert job["status"] == "completed"
        assert job["summary"]["tem_r_total_rows"] == 2
        assert main_api._job_output_ready(job) is True
        assert not list(tmp_dir.glob(f"tem_r_{job_id}"))
    finally:
        job = main_api._jobs.pop(job_id, None)
        if job:
            __import__("pathlib").Path(job["output_path"]).unlink(missing_ok=True)


def test_expiration_empty_item_source_is_rejected_instead_of_reporting_not_found(monkeypatch):
    content = _workbook_bytes(
        ["ITEMID", "BATCHNUMER", "EXPDATE"],
        [],
        sheet="Export",
    )
    frame = pd.DataFrame(columns=["ITEMID", "BATCHNUMER", "EXPDATE"])
    monkeypatch.setattr(
        main_api,
        "_load_expiration_cached",
        lambda _content: (frame, "Export", False),
    )
    monkeypatch.setattr(main_api, "_expiration_slots", threading.BoundedSemaphore(1))

    with pytest.raises(main_api.HTTPException) as exc:
        asyncio.run(main_api.check_expiration(
            expiration_file=FakeUpload("expiration.xlsx", content),
            entries="P1001;B1;01/01/2030",
        ))
    assert exc.value.status_code == 422
    assert "không có dòng sản phẩm ITEMID" in str(exc.value.detail)


def test_expiration_overload_fails_fast_without_using_threadpool_waiters(monkeypatch):
    content = _workbook_bytes(
        ["ITEMID", "BATCHNUMER", "EXPDATE"],
        [["P1001", "B1", "01/01/2030"]],
        sheet="Export",
    )

    class FullSemaphore:
        def acquire(self, blocking=True):
            assert blocking is False
            return False

        def release(self):
            raise AssertionError("release must not run when acquire failed")

    monkeypatch.setattr(main_api, "_expiration_slots", FullSemaphore())
    called = []
    monkeypatch.setattr(main_api, "_load_expiration_cached", lambda content: called.append(content))

    with pytest.raises(main_api.HTTPException) as exc:
        asyncio.run(main_api.check_expiration(
            expiration_file=FakeUpload("expiration.xlsx", content),
            entries="P1001;B1;01/01/2030",
        ))
    assert exc.value.status_code == 503
    assert exc.value.headers.get("Retry-After") == "2"
    assert called == []
