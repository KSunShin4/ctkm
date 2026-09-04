import importlib.util
from pathlib import Path
import sys
import types
import zipfile

import pytest


def _loader_module():
    root = Path(__file__).resolve().parents[1]
    modules_dir = root / "modules"
    if "modules" not in sys.modules:
        pkg = types.ModuleType("modules")
        pkg.__path__ = [str(modules_dir)]
        sys.modules["modules"] = pkg

    config_spec = importlib.util.spec_from_file_location(
        "modules.config_loader", modules_dir / "config_loader.py"
    )
    config_module = importlib.util.module_from_spec(config_spec)
    sys.modules[config_spec.name] = config_module
    config_spec.loader.exec_module(config_module)

    spec = importlib.util.spec_from_file_location(
        "modules.input_reader_hardening_test", modules_dir / "1_loader.py"
    )
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _tem_r_module():
    root = Path(__file__).resolve().parents[1]
    path = root / "modules" / "10_tem_r.py"
    spec = importlib.util.spec_from_file_location("tem_r_input_reader_test", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize("is_pog_file", [False, True])
def test_shared_loader_disables_external_link_parsing(tmp_path, monkeypatch, is_pog_file):
    from openpyxl import Workbook
    import openpyxl

    loader = _loader_module()
    source = tmp_path / ("pog.xlsx" if is_pog_file else "ctkm-clearance.xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["STOREID", "POGCODE", "ITEMID", "LOCATIONID"])
    sheet.append(["S1", "POG1", "P1", 1])
    workbook.save(source)

    real_load_workbook = openpyxl.load_workbook
    calls = []

    def recording_load_workbook(*args, **kwargs):
        calls.append(dict(kwargs))
        return real_load_workbook(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", recording_load_workbook)
    loaded = loader.load_workbook_raw(
        str(source), is_pog_file=is_pog_file, selected_sheet_names=["Export"]
    )

    assert len(loaded.sheets["Export"].df) == 1
    assert calls
    assert calls[0].get("keep_links") is False
    assert calls[0].get("read_only") is True
    assert calls[0].get("data_only") is True


def test_loader_rejects_zip_that_is_not_an_xlsx(tmp_path):
    loader = _loader_module()
    source = tmp_path / "fake.xlsx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("hello.txt", "not a workbook")

    with pytest.raises(ValueError, match="thiếu cấu trúc XLSX bắt buộc"):
        loader.load_workbook_raw(str(source))


def test_header_detection_prefers_business_headers_over_denser_banner_row(tmp_path):
    from openpyxl import Workbook

    loader = _loader_module()
    source = tmp_path / "monthly.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TỔNG HỢP CTKM UPDATE"
    sheet.append(["alpha", "beta", "gamma", "delta", "epsilon", "zeta"])
    sheet.append(["Item code 2", "Product Name", "From", "To", None, None])
    sheet.append(["P1", "Product 1", "2026-09-01", "2026-09-30", None, None])
    workbook.save(source)

    loaded = loader.load_workbook_raw(str(source), selected_sheet_names=[sheet.title])
    result = loaded.sheets[sheet.title]

    assert result.header_row_index == 1
    assert list(result.df.columns[:4]) == ["Item code 2", "Product Name", "From", "To"]
    assert result.df.iloc[0]["Item code 2"] == "P1"


def test_business_row_containing_total_word_is_not_dropped(tmp_path):
    from openpyxl import Workbook

    loader = _loader_module()
    source = tmp_path / "clearance.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TỔNG HỢP CTKM UPDATE"
    sheet.append(["Item code", "Product Name", "Campaign", "From", "To"])
    sheet.append(["P100", "Total Care Shampoo", "Total Beauty Promo", "2026-09-01", "2026-09-30"])
    workbook.save(source)

    loaded = loader.load_workbook_raw(str(source), selected_sheet_names=[sheet.title])
    frame = loaded.sheets[sheet.title].df

    assert len(frame) == 1
    assert frame.iloc[0]["Item code"] == "P100"
    assert frame.iloc[0]["Product Name"] == "Total Care Shampoo"


def test_pog_sparse_filter_row_sets_snapshot_but_is_not_business_data(tmp_path):
    from openpyxl import Workbook

    loader = _loader_module()
    source = tmp_path / "pog.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["STOREID", "POGCODE", "ITEMID", "LOCATIONID"])
    sheet.append(["S1", "POG1", "P1", 1])
    sheet.append(["Applied Filter DateID 9/1/2026 - 9/30/2026", None, None, None])
    workbook.save(source)

    loaded = loader.load_workbook_raw(
        str(source), is_pog_file=True, selected_sheet_names=["Export"]
    )

    assert loaded.pog_snapshot_month == "2026-09"
    assert len(loaded.sheets["Export"].df) == 1
    assert loaded.sheets["Export"].df.iloc[0]["ITEMID"] == "P1"


def test_tem_r_inspection_already_skips_external_links(tmp_path, monkeypatch):
    from openpyxl import Workbook
    import openpyxl

    tem_r = _tem_r_module()
    source = tmp_path / "stock.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append([
        "STOREID", "POGCODE", "FIXELID", "LOCATIONID",
        "ITEMID", "ITEMNAME", "ton",
    ])
    sheet.append(["S1", "POG1", "1/1", 1, "P1", "Product 1", 2])
    workbook.save(source)

    real_load_workbook = openpyxl.load_workbook
    calls = []

    def recording_load_workbook(*args, **kwargs):
        calls.append(dict(kwargs))
        return real_load_workbook(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", recording_load_workbook)
    result = tem_r.inspect_tem_r_workbook(source)

    assert result["sheet_name"] == "Export"
    assert result["suggested_item_column"] == "ITEMID"
    assert result["suggested_stock_column"] == "ton"
    assert calls
    assert calls[0].get("keep_links") is False
    assert calls[0].get("read_only") is True
    assert calls[0].get("data_only") is True
