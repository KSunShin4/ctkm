from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

import main
from backend.final_output_contract import (
    _LEGACY_HIDDEN_SUMMARY_KEYS,
    _REQUIRED_CTKM_SHEETS,
    validate_ctkm_workbook,
)
from backend.mapping_result_contract import install_mapping_result_contract
from backend.workbook_semantics import visible_data_row_count
from modules.mapping_validation import validate_manual_mapping


SUMMARY = {
    "Thang bao cao": "2026-09",
    "So dong khop (Matched)": 0,
    "So item duy nhat khop": 0,
    "So dong KM thieu vi tri POG": 0,
    "So item KM thieu vi tri POG": 0,
    "So dong CTKM thieu vi tri ke": 0,
    "So item CTKM thieu vi tri ke": 0,
    "So dong can kiem tra vi tri": 0,
    "So canh bao trung thoi gian": 0,
    "So dong cho gan": 0,
    "So dong can review": 0,
    "So dong POG khong co KM phu hop": 0,
    "So item POG khong co KM phu hop": 0,
    "So dong POG khong co CTKM phu hop": 0,
    "So item POG khong co CTKM phu hop": 0,
}


def _build_minimal_valid_report(path: Path):
    wb = Workbook()
    guide = wb.active
    guide.title = "HƯỚNG DẪN"

    for name in sorted(_REQUIRED_CTKM_SHEETS - {"HƯỚNG DẪN", "Tom tat", "Mapping log"}):
        ws = wb.create_sheet(name)
        ws.append(["Ghi chu"])
        note = (
            "Khong phat hien campaign nao trung thoi gian."
            if name == "Canh bao trung thoi gian"
            else "Không có dữ liệu nghiệp vụ."
        )
        ws.append([note])
        if name == "POG không có CTKM":
            ws.sheet_state = "hidden"

    summary_ws = wb.create_sheet("Tom tat")
    summary_ws.append(["Chi so", "Gia tri"])
    for key, value in SUMMARY.items():
        summary_ws.append([key, value])
        if key in _LEGACY_HIDDEN_SUMMARY_KEYS:
            summary_ws.row_dimensions[summary_ws.max_row].hidden = True

    mapping = wb.create_sheet("Mapping log")
    mapping.append(["Ghi chu"])
    mapping.append(["Chua co mapping log."])
    mapping.sheet_state = "hidden"

    cards = [
        ("Gắn tem", "Gắn tem"),
        ("Tháo tem", "Tháo tem"),
        ("Chờ gắn", "Chờ gắn"),
        ("Cần review", "Cần review"),
        ("Thiếu vị trí KM", "CTKM thieu vi tri ke"),
        ("Review vị trí", "POG cần kiểm tra"),
        ("POG cần review", "POG cần review"),
        ("Xác nhận CTKM / Clearance", "Xác nhận CTKM Clearance"),
    ]
    for row, (label, _target) in enumerate(cards, start=6):
        guide.cell(row=row, column=1, value=label)
        guide.cell(row=row, column=2, value=0)

    wb.save(path)
    wb.close()


def test_overlap_no_warning_note_is_zero_visible_rows(tmp_path):
    path = tmp_path / "report.xlsx"
    _build_minimal_valid_report(path)
    wb = load_workbook(path)
    try:
        assert visible_data_row_count(wb["Canh bao trung thoi gian"]) == 0
        assert visible_data_row_count(wb["Gắn tem"]) == 0
    finally:
        wb.close()


def test_final_workbook_contract_checks_summary_guide_and_required_sheets(tmp_path):
    path = tmp_path / "report.xlsx"
    _build_minimal_valid_report(path)
    validate_ctkm_workbook(path, dict(SUMMARY))

    wb = load_workbook(path)
    wb["HƯỚNG DẪN"]["B6"] = 1
    wb.save(path)
    wb.close()
    with pytest.raises(RuntimeError, match="HƯỚNG DẪN không khớp"):
        validate_ctkm_workbook(path, dict(SUMMARY))

    _build_minimal_valid_report(path)
    wb = load_workbook(path)
    del wb["Cần review"]
    wb.save(path)
    wb.close()
    with pytest.raises(RuntimeError, match="thiếu sheet bắt buộc"):
        validate_ctkm_workbook(path, dict(SUMMARY))


def test_final_workbook_contract_rejects_api_excel_and_sheet_count_mismatch(tmp_path):
    path = tmp_path / "report.xlsx"
    _build_minimal_valid_report(path)
    bad_summary = dict(SUMMARY)
    bad_summary["So dong can review"] = 2
    with pytest.raises(RuntimeError, match="không khớp summary"):
        validate_ctkm_workbook(path, bad_summary)

    _build_minimal_valid_report(path)
    wb = load_workbook(path)
    summary_ws = wb["Tom tat"]
    for row in range(2, summary_ws.max_row + 1):
        if summary_ws.cell(row=row, column=1).value == "So dong can review":
            summary_ws.cell(row=row, column=2, value=1)
            break
    wb.save(path)
    wb.close()
    sheet_mismatch_summary = dict(SUMMARY)
    sheet_mismatch_summary["So dong can review"] = 1
    with pytest.raises(RuntimeError, match="không khớp sheet 'Cần review'"):
        validate_ctkm_workbook(path, sheet_mismatch_summary)


def test_final_workbook_contract_rejects_legacy_canonical_drift_and_visible_legacy_rows(tmp_path):
    path = tmp_path / "report.xlsx"
    _build_minimal_valid_report(path)
    drifted = dict(SUMMARY)
    drifted["So dong CTKM thieu vi tri ke"] = 1
    with pytest.raises(RuntimeError, match="legacy .* lệch canonical"):
        validate_ctkm_workbook(path, drifted)

    _build_minimal_valid_report(path)
    wb = load_workbook(path)
    ws = wb["Tom tat"]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "So dong POG khong co CTKM phu hop":
            ws.row_dimensions[row].hidden = False
            break
    wb.save(path)
    wb.close()
    with pytest.raises(RuntimeError, match="phải ẩn legacy metric"):
        validate_ctkm_workbook(path, dict(SUMMARY))


def _install_business_presence_contract(mapping):
    def original(_sheet, _df, _schema, logs):
        logs.append({"source": "standard", "mapping": dict(mapping)})
        return dict(mapping)

    fake_core_main = SimpleNamespace(
        _get_or_map=original,
        _ctkm_main_facade_baseline={"get_or_map": original},
    )
    fake_main = SimpleNamespace(
        _core=fake_core_main,
        NM=main.NM,
        JN=main.JN,
        _separate_fact_control_rows=main._separate_fact_control_rows,
        cfg=main.cfg,
        MC=main.MC,
        CM=main.CM,
        _infer_standard_fact_mapping=lambda frame: dict(mapping),
        _infer_standard_pog_mapping=lambda frame: dict(mapping),
    )
    fake_api = SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace()),
        _main_mod=fake_main,
    )
    install_mapping_result_contract(fake_api, validate_manual_mapping)
    return fake_core_main


def test_control_only_fact_sheet_is_rejected_but_real_missing_item_row_is_reviewable():
    mapping = {
        "Item": "item_id",
        "Campaign": "campaign_name",
        "From": "promo_from",
        "To": "promo_to",
    }
    fake_core = _install_business_presence_contract(mapping)
    schema = main.SCHEMA["fact_fields"]

    control_only = pd.DataFrame([
        {"Item": None, "Campaign": 'Combo "A" Total', "From": None, "To": None},
        {"Item": None, "Campaign": "Grand Total", "From": None, "To": None},
    ])
    with pytest.raises(ValueError, match="chỉ chứa subtotal/Total/control"):
        fake_core._get_or_map("Detail", control_only, schema, [])

    malformed_business = pd.DataFrame([
        {
            "Item": None,
            "Campaign": "Promo thật thiếu mã",
            "From": "2026-09-01",
            "To": "2026-09-30",
        }
    ])
    assert fake_core._get_or_map("Detail", malformed_business, schema, []) == mapping


def test_control_only_pog_sheet_is_rejected_but_real_bad_pog_row_remains_reviewable():
    mapping = {"Item": "item_id", "Store": "store_id"}
    fake_core = _install_business_presence_contract(mapping)
    schema = main.SCHEMA["pog_fields"]

    total_only = pd.DataFrame([{"Item": "Total", "Store": None}])
    with pytest.raises(ValueError, match="chỉ chứa Total/filter metadata"):
        fake_core._get_or_map("Export", total_only, schema, [])

    malformed_business = pd.DataFrame([{"Item": None, "Store": "S1"}])
    assert fake_core._get_or_map("Export", malformed_business, schema, []) == mapping


def test_api_entrypoint_rebinds_input_validator_to_shared_contract():
    from backend import input_contract
    import backend.main_api  # noqa: F401

    assert input_contract.validate_manual_mapping is validate_manual_mapping
