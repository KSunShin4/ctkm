from openpyxl import load_workbook
import pandas as pd

import main


def _task(task_id, source_row_id, **overrides):
    row = {
        "task_id": task_id,
        "promotion_id": "PROMO-1",
        "pog_position_id": "POGPOS-1",
        "item_id": "P1001",
        "item_key": "P1001",
        "item_name": "Product A",
        "campaign_name": "FULL MONTH",
        "campaign_detail": "Giảm 10%",
        "promo_from": pd.Timestamp("2026-09-01"),
        "promo_to": pd.Timestamp("2026-09-30"),
        "program_source": "CTKM",
        "store_id": "S1",
        "pog_code": "POG-A",
        "location_id": "1",
        "scheduled_action": "gắn tem",
        "action": "gắn tem",
        "reason": "trong_thoi_gian_ap_dung",
        "in_report_scope": True,
        "location_completeness": "Đủ vị trí",
        "location_status": "Đủ vị trí",
        "source_row_id": source_row_id,
        "source_row_ids": source_row_id,
        "source_row_count": 1,
        "duplicate_status": "Dòng duy nhất",
        "pog_source_row_id": "POG:Export:10",
        "pog_source_row_ids": "POG:Export:10",
        "pog_source_row_count": 1,
    }
    row.update(overrides)
    return row


def test_work_tracking_rows_preserve_distinct_source_lineages_even_when_task_id_collides():
    frame = pd.DataFrame([
        _task("TASK-1", "CTKM:Monthly:10", item_name="Product A"),
        _task("TASK-1", "CTKM:Monthly:20", item_name="Product A renamed"),
    ])

    result = main.EX._add_work_tracking_columns(frame)

    assert len(result) == 1
    assert result.loc[0, "task_id"] == "TASK-1"
    assert result.loc[0, "source_row_ids"] == "CTKM:Monthly:10 | CTKM:Monthly:20"
    assert result.loc[0, "source_row_count"] == 2
    assert result.loc[0, "pog_source_row_ids"] == "POG:Export:10"
    assert result.loc[0, "pog_source_row_count"] == 1
    assert "Audit hành động" in result.loc[0, "duplicate_status"]
    assert result.loc[0, "processing_status"] == "Chưa xử lý"


def test_distinct_task_ids_remain_distinct_even_for_same_promotion():
    frame = pd.DataFrame([
        _task("TASK-1", "CTKM:Monthly:10", location_id="1", pog_position_id="POGPOS-1"),
        _task("TASK-2", "CTKM:Monthly:10", location_id="2", pog_position_id="POGPOS-2"),
    ])

    result = main.EX._add_work_tracking_columns(frame)

    assert result["task_id"].tolist() == ["TASK-1", "TASK-2"]


def test_rows_without_task_id_are_never_silently_collapsed():
    frame = pd.DataFrame([
        _task(None, "CTKM:Monthly:10"),
        _task("", "CTKM:Monthly:20"),
    ])

    result = main.EX._add_work_tracking_columns(frame)

    assert len(result) == 2


def test_export_keeps_distinct_source_rows_in_both_audit_and_editable_workbook(tmp_path):
    actions = pd.DataFrame([
        _task("TASK-1", "CTKM:Monthly:10", item_name="Product A"),
        _task("TASK-1", "CTKM:Monthly:20", item_name="Product A renamed"),
    ])
    output = tmp_path / "task-identity.xlsx"

    main.EX.export_final_report(
        output,
        pd.DataFrame(),
        pd.DataFrame(),
        pd.DataFrame(),
        [],
        [],
        {"Thang bao cao": "2026-09", "Ngay nghiep vu": pd.Timestamp("2026-09-15")},
        df_fefo_actions=actions,
    )

    workbook = load_workbook(output)

    # One editable physical task, despite two source rows describing it.
    assert workbook["Gắn tem"].max_row == 2
    # Audit retains both source rows for fidelity and investigation.
    assert workbook["Audit hành động"].max_row == 3
    assert workbook["Gắn tem"].data_validations.count == 1
