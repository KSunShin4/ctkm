from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

import main


BUSINESS_DATE = pd.Timestamp("2026-08-15")


def _promotion(item_id="P1001", **overrides):
    row = {
        "item_id": item_id,
        "item_name": f"Name {item_id}",
        "promo_uom": "Hộp",
        "campaign_name": "Monthly",
        "campaign_detail": "Giảm 10%",
        "promo_from": pd.Timestamp("2026-08-01"),
        "promo_to": pd.Timestamp("2026-08-31"),
        "normal_price": 100_000,
        "promo_price": 90_000,
        "channel": "  Kênh LẠ / WEBAPP+Store  ",
        "product_type": "Medicine",
        "program_source": "CTKM",
        "source_sheet": "Monthly",
    }
    row.update(overrides)
    return row


@pytest.mark.parametrize(
    ("promo_from", "promo_to", "business_date", "expected_action", "expected_reason"),
    [
        ("2026-08-20", "2026-08-31", "2026-08-19", "chờ gắn tem", "chua_den_ngay_gan"),
        ("2026-08-20", "2026-08-31", "2026-08-20", "gắn tem", "trong_thoi_gian_ap_dung"),
        ("2026-08-01", "2026-08-31", "2026-08-31", "gắn tem", "trong_thoi_gian_ap_dung"),
        ("2026-08-01", "2026-08-31", "2026-09-01", "tháo tem", "den_ngay_thao"),
        (None, "2026-08-31", "2026-08-15", "review", "thieu_thong_tin_ngay"),
        ("2026-08-31", "2026-08-01", "2026-08-15", "review", "ngay_bat_dau_sau_ngay_ket_thuc"),
    ],
)
def test_action_date_matrix_is_deterministic(
    promo_from, promo_to, business_date, expected_action, expected_reason,
):
    frame = pd.DataFrame([_promotion(
        promo_from=pd.Timestamp(promo_from) if promo_from else pd.NaT,
        promo_to=pd.Timestamp(promo_to) if promo_to else pd.NaT,
    )])

    result = main.RC.build_fefo_action_frame(
        frame, "2026-08", business_date=pd.Timestamp(business_date),
    )

    assert result.loc[0, "scheduled_action"] == expected_action
    assert result.loc[0, "action"] == expected_action
    assert result.loc[0, "reason"] == expected_reason


def test_missing_pog_does_not_change_scheduled_action():
    result = main.RC.build_fefo_action_frame(
        pd.DataFrame([_promotion()]), "2026-08",
        df_pog=pd.DataFrame(columns=["item_id", "store_id", "pog_code", "location_id"]),
        business_date=BUSINESS_DATE,
    )

    assert result.loc[0, "scheduled_action"] == "gắn tem"
    assert result.loc[0, "location_status"] == "Không tìm thấy SKU trong POG"


def test_partial_pog_location_is_not_reported_as_complete():
    pog = pd.DataFrame([{"item_id": "P1001", "store_id": "S1", "pog_code": None, "location_id": "3"}])

    result = main.RC.build_fefo_action_frame(
        pd.DataFrame([_promotion()]), "2026-08", df_pog=pog, business_date=BUSINESS_DATE,
    )

    assert result.loc[0, "scheduled_action"] == "gắn tem"
    assert result.loc[0, "location_status"] == "Thiếu một phần thông tin vị trí"


def test_multiple_pog_positions_are_all_preserved():
    pog = pd.DataFrame([
        {"item_id": "P1001", "store_id": "S1", "pog_code": "A", "location_id": "1", "facing": 2},
        {"item_id": "P1001", "store_id": "S1", "pog_code": "B", "location_id": "4", "facing": 1},
    ])

    result = main.RC.build_fefo_action_frame(
        pd.DataFrame([_promotion()]), "2026-08", df_pog=pog, business_date=BUSINESS_DATE,
    )

    assert result["location_id"].tolist() == ["1", "4"]
    assert set(result["location_status"]) == {"Nhiều vị trí POG"}
    assert set(result["pog_position_count"]) == {2}
    assert "label_quantity" not in result.columns


def test_pog_month_mismatch_warns_without_changing_action():
    pog = pd.DataFrame([{
        "item_id": "P1001", "store_id": "S1", "pog_code": "A", "location_id": "1",
        "pog_snapshot_month": "2026-07",
    }])

    result = main.RC.build_fefo_action_frame(
        pd.DataFrame([_promotion()]), "2026-08", df_pog=pog, business_date=BUSINESS_DATE,
    )

    assert result.loc[0, "scheduled_action"] == "gắn tem"
    assert result.loc[0, "location_status"] == "POG lệch tháng"
    assert result.loc[0, "pog_month_status"] == "POG lệch tháng"


def test_unknown_channel_is_byte_for_byte_preserved_and_does_not_hide_overlap():
    channel_a = "  Kênh LẠ / WEBAPP+Store  "
    channel_b = "MARKETPLACE ???"
    frame = pd.DataFrame([
        _promotion(campaign_name="A", channel=channel_a),
        _promotion(campaign_name="B", channel=channel_b),
    ])

    normalized = main.NM.clean_values(frame)
    assert normalized["channel"].tolist() == [channel_a, channel_b]
    warnings = main.RC.detect_campaign_time_overlaps(normalized)
    result = main.RC.build_fefo_action_frame(
        normalized, "2026-08", business_date=BUSINESS_DATE, overlap_warnings=warnings,
    )

    assert len(warnings) == 1
    assert warnings[0].channel_a == channel_a
    assert warnings[0].channel_b == channel_b
    assert result["channel"].tolist() == [channel_a, channel_b]
    assert set(result["scheduled_action"]) == {"review"}


def test_process_preserves_channel_and_source_excel_row():
    source = pd.DataFrame([{
        "Item": "P1", "Campaign": "A", "From": "2026-08-01", "To": "2026-08-31",
        "Channel": "  Kênh LẠ  ",
    }])
    mapping = {
        "Item": "item_id", "Campaign": "campaign_name", "From": "promo_from",
        "To": "promo_to", "Channel": "channel",
    }

    result = main._process(
        "Monthly", source, main.SCHEMA["fact_fields"], [], True, [], mapping_override=mapping,
    )

    assert result.loc[0, "channel"] == "  Kênh LẠ  "
    assert result.loc[0, "source_row_id"] == "Monthly:2"


def test_overlap_moves_only_affected_campaign_rows_to_review():
    frame = pd.DataFrame([
        _promotion("P1", campaign_name="A"),
        _promotion("P1", campaign_name="B"),
        _promotion("P1", campaign_name="Later", promo_from=pd.Timestamp("2026-09-10"), promo_to=pd.Timestamp("2026-09-20")),
        _promotion("P2", campaign_name="Unrelated"),
    ])
    warnings = main.RC.detect_campaign_time_overlaps(frame)

    result = main.RC.build_fefo_action_frame(
        frame, "2026-08", business_date=BUSINESS_DATE, overlap_warnings=warnings,
    )

    by_campaign = result.set_index("campaign_name")["scheduled_action"].to_dict()
    assert by_campaign == {
        "A": "review", "B": "review", "Later": "chờ gắn tem", "Unrelated": "gắn tem",
    }
    assert "xung_dot_thoi_gian" in result.set_index("campaign_name").loc["A", "reason"]


def test_clearance_same_sku_different_period_is_not_dropped():
    ctkm = pd.DataFrame([_promotion("P1", campaign_name="CTKM")])
    clearance = pd.DataFrame([_promotion(
        "P1", campaign_name="Clearance", promo_from=pd.Timestamp("2026-09-01"),
        promo_to=pd.Timestamp("2026-09-30"), program_source="CLEARANCE",
    )])

    source = main._build_promo_action_source(ctkm, clearance)

    assert source["campaign_name"].tolist() == ["CTKM", "Clearance"]
    assert source["program_source"].tolist() == ["CTKM", "CLEARANCE"]


def test_ctkm_clearance_same_sku_and_period_are_both_kept_for_review():
    frame = main._build_promo_action_source(
        pd.DataFrame([_promotion("P1", campaign_name="CTKM")]),
        pd.DataFrame([_promotion("P1", campaign_name="Clearance", program_source="CLEARANCE")]),
    )
    warnings = main.RC.detect_campaign_time_overlaps(frame)

    actions = main.RC.build_fefo_action_frame(
        frame, "2026-08", business_date=BUSINESS_DATE, overlap_warnings=warnings,
    )

    assert len(actions) == 2
    assert set(actions["program_source"]) == {"CTKM", "CLEARANCE"}
    assert set(actions["scheduled_action"]) == {"review"}
    assert actions["source_row_id"].nunique() == 2
    assert any(value.startswith("CTKM:") for value in actions["source_row_id"])
    assert any(value.startswith("CLEARANCE:") for value in actions["source_row_id"])


def test_exact_duplicates_are_consolidated_with_lineage():
    duplicate = _promotion(source_sheet="Monthly")
    frame = pd.DataFrame([duplicate, duplicate])

    result = main.RC.build_fefo_action_frame(
        frame, "2026-08", business_date=BUSINESS_DATE,
    )

    assert len(result) == 1
    assert result.loc[0, "source_row_count"] == 2
    assert len(result.loc[0, "source_row_ids"].split(" | ")) == 2
    assert result.loc[0, "duplicate_status"] == "Đã gom dòng trùng hoàn toàn"


def test_action_frame_preserves_prices_and_business_columns():
    result = main.RC.build_fefo_action_frame(
        pd.DataFrame([_promotion()]), "2026-08", business_date=BUSINESS_DATE,
    )

    assert result.loc[0, "normal_price"] == 100_000
    assert result.loc[0, "promo_price"] == 90_000
    assert result.loc[0, "product_type"] == "Medicine"
    assert result.loc[0, "channel"] == "  Kênh LẠ / WEBAPP+Store  "


def test_promotion_removed_on_first_day_of_month_is_in_month_scope():
    frame = pd.DataFrame([_promotion(
        promo_from=pd.Timestamp("2026-07-01"), promo_to=pd.Timestamp("2026-07-31"),
    )])

    relevant, operational = main.TF.filter_ctkm_by_month(frame, "2026-08")

    assert len(relevant) == 1
    assert len(operational) == 1
    assert relevant.loc[0, "_month_scope_reason"] == "ngay_thao_trong_thang"


def test_workbook_exports_all_action_groups_and_location_issues(tmp_path):
    actions = pd.DataFrame([
        {**_promotion("P1"), "scheduled_action": "gắn tem", "action": "gắn tem", "reason": "trong_thoi_gian_ap_dung", "in_report_scope": True, "task_id": "T1", "location_status": "Đủ vị trí"},
        {**_promotion("P2"), "scheduled_action": "tháo tem", "action": "tháo tem", "reason": "den_ngay_thao", "in_report_scope": True, "task_id": "T2", "location_status": "Đủ vị trí"},
        {**_promotion("P3"), "scheduled_action": "chờ gắn tem", "action": "chờ gắn tem", "reason": "chua_den_ngay_gan", "in_report_scope": True, "task_id": "T3", "location_status": "Không tìm thấy SKU trong POG"},
        {**_promotion("P4"), "scheduled_action": "review", "action": "review", "reason": "thieu_thong_tin_ngay", "in_report_scope": True, "task_id": "T4", "location_status": "Thiếu một phần thông tin vị trí"},
        {**_promotion("P5"), "scheduled_action": "chờ gắn tem", "action": "chờ gắn tem", "reason": "chua_den_ngay_gan", "in_report_scope": False, "task_id": "T5", "location_status": "Đủ vị trí"},
    ])
    output = tmp_path / "report.xlsx"

    main.EX.export_final_report(
        output, pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), [], [],
        {"Thang bao cao": "2026-08"}, df_fefo_actions=actions,
    )

    workbook = load_workbook(output)
    assert {"Gắn tem", "Tháo tem", "Chờ gắn", "Cần review", "Audit hành động"}.issubset(workbook.sheetnames)
    assert workbook["Gắn tem"].max_row == 2
    assert workbook["Tháo tem"].max_row == 2
    assert workbook["Chờ gắn"].max_row == 2
    assert workbook["Cần review"].max_row == 2
    assert workbook["Audit hành động"].max_row == 6
    assert "thieu_thong_tin_ngay" in [cell.value for cell in workbook["Cần review"][2]]
    assert workbook["CTKM thieu vi tri ke"].max_row == 3
    review_ids = [cell.value for cell in workbook["Cần review"][2]]
    assert "P1" not in review_ids and "P2" not in review_ids
    wait_values = [cell.value for row in workbook["Chờ gắn"].iter_rows(min_row=2) for cell in row]
    assert "P5" not in wait_values
    location_headers = [cell.value for cell in workbook["CTKM thieu vi tri ke"][1]]
    location_task_col = next(idx for idx, value in enumerate(location_headers, start=1) if value and "task_id" in str(value))
    location_task_ids = {
        workbook["CTKM thieu vi tri ke"].cell(row, location_task_col).value
        for row in range(2, workbook["CTKM thieu vi tri ke"].max_row + 1)
    }
    assert {"T3", "T4"} == location_task_ids
    for sheet_name in ("Gắn tem", "Tháo tem", "Chờ gắn", "Cần review"):
        assert workbook[sheet_name].data_validations.count == 1


def test_pog_only_wording_is_neutral(tmp_path):
    output = tmp_path / "neutral.xlsx"
    main.EX.export_final_report(
        output, pd.DataFrame(), pd.DataFrame(), pd.DataFrame([{"item_id": "P9"}]),
        [], [], {"Thang bao cao": "2026-08"},
    )

    workbook = load_workbook(output)
    assert "POG không có CTKM" in workbook.sheetnames
    guide_text = " ".join(str(cell.value or "") for row in workbook["HƯỚNG DẪN"] for cell in row)
    assert "tem nhầm" not in guide_text.lower()
