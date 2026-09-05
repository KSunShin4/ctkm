from openpyxl import load_workbook
import pandas as pd

import main


BUSINESS_DATE = pd.Timestamp("2026-09-15")


def _promo(item_id, *, name, detail):
    return {
        "item_id": item_id,
        "item_name": name,
        "promo_uom": "Hộp" if str(item_id).upper().startswith("P") else None,
        "campaign_name": "September",
        "campaign_detail": detail,
        "promo_from": pd.Timestamp("2026-09-01"),
        "promo_to": pd.Timestamp("2026-09-30"),
        "normal_price": 100000,
        "promo_price": 90000,
        "key_channel": "Offline",
        "channel": "Offline",
        "program_source": "CTKM",
        "source_sheet": "Monthly",
    }


def _pog(item_id):
    return {
        "item_id": item_id,
        "store_id": "S1",
        "pog_code": "A",
        "location_id": "1",
        "facing": 1,
        "pog_snapshot_month": "2026-09",
    }


def test_explicit_bill_is_valid_but_pog_not_applicable_while_bill_text_on_sku_still_uses_pog():
    promotions = pd.DataFrame([
        _promo("Bill", name="Invoice campaign", detail="Discount on total customer bill"),
        _promo(
            "P1001",
            name="SKU with PWP",
            detail="Customer purchase bill 150k, can buy product at discount 50%",
        ),
        _promo("P2001", name="Missing POG SKU", detail="Regular SKU promotion"),
    ])
    pog = pd.DataFrame([_pog("P1001")])

    actions = main.RC.build_fefo_action_frame(
        promotions,
        "2026-09",
        df_pog=pog,
        business_date=BUSINESS_DATE,
    )
    by_item = actions.set_index("item_key")

    bill = by_item.loc["BILL"]
    assert bill["scheduled_action"] == "gắn tem"
    assert bill["reason"] == "trong_thoi_gian_ap_dung"
    assert bool(bill["pog_applicable"]) is False
    assert bill["location_status"] == "Không áp dụng POG"
    assert bill["pog_month_status"] == "Không áp dụng POG"
    assert bill["location_completeness"] == "Đủ vị trí"
    assert bool(bill["location_needs_review"]) is False
    assert bill["pog_position_count"] == 0

    pwp = by_item.loc["P1001"]
    assert pwp["campaign_detail"] == "Customer purchase bill 150k, can buy product at discount 50%"
    assert bool(pwp["pog_applicable"]) is True
    assert pwp["location_completeness"] == "Đủ vị trí"
    assert pwp["location_status"] == "Đủ vị trí"
    assert bool(pwp["location_needs_review"]) is False

    missing = by_item.loc["P2001"]
    assert bool(missing["pog_applicable"]) is True
    assert missing["location_completeness"] == "Không tìm thấy SKU trong POG"
    assert bool(missing["location_needs_review"]) is True


def test_bill_remains_separate_reconciliation_category_without_becoming_pog_error():
    ctkm = main.NM.clean_values(pd.DataFrame([
        {"item_id": "Bill", "campaign_name": "Invoice campaign"},
        {"item_id": "P1001", "campaign_name": "SKU campaign"},
    ]))
    pog = main.NM.clean_values(pd.DataFrame([_pog("P1001")]))

    joined = main.JN.join_ctkm_pog(ctkm, pog)
    report = main.RC.build_reconciliation_report(
        joined.df_matched,
        joined.df_ctkm_only,
        joined.df_pog_only,
        [],
    )

    assert len(joined.df_matched) == 1
    assert joined.df_matched.loc[0, "item_key"] == "P1001"
    assert len(joined.df_ctkm_only) == 1
    assert joined.df_ctkm_only.loc[0, "item_key"] == "BILL"
    assert report.non_sku_ctkm_count == 1
    assert report.invalid_item_ctkm_count == 0
    assert report.missing_item_ctkm_count == 0


def test_export_does_not_count_bill_as_missing_pog_but_keeps_pog_applicability_visible(tmp_path):
    promotions = pd.DataFrame([
        _promo("Bill", name="Invoice campaign", detail="Discount on total customer bill"),
        _promo(
            "P1001",
            name="SKU with PWP",
            detail="Customer purchase bill 150k, can buy product at discount 50%",
        ),
        _promo("P2001", name="Missing POG SKU", detail="Regular SKU promotion"),
    ])
    actions = main.RC.build_fefo_action_frame(
        promotions,
        "2026-09",
        df_pog=pd.DataFrame([_pog("P1001")]),
        business_date=BUSINESS_DATE,
    )
    output = tmp_path / "bill-pog-applicability.xlsx"
    summary = {
        "Thang bao cao": "2026-09",
        "Ngay nghiep vu": BUSINESS_DATE,
        "So dong CTKM thieu vi tri ke": 0,
        "So item CTKM thieu vi tri ke": 0,
        "So dong POG khong co CTKM phu hop": 0,
        "So item POG khong co CTKM phu hop": 0,
    }

    main.EX.export_final_report(
        output_path=output,
        df_matched=pd.DataFrame(),
        df_ctkm_only=pd.DataFrame(),
        df_pog_only=pd.DataFrame(),
        overlap_warnings=[],
        mapping_logs=[],
        reconciliation_summary=summary,
        df_fefo_actions=actions,
        df_pog_review=pd.DataFrame(),
    )

    # Only the real P-code without a POG position is a location issue.
    assert summary["So dong KM thieu vi tri POG"] == 1
    assert summary["So item KM thieu vi tri POG"] == 1
    assert summary["So dong CTKM Monthly thieu vi tri POG"] == 1

    workbook = load_workbook(output, data_only=True)
    issue_values = [
        cell.value for row in workbook["CTKM thieu vi tri ke"].iter_rows() for cell in row
    ]
    assert "P2001" in issue_values
    assert "Bill" not in issue_values
    assert "Invoice campaign" not in issue_values

    action = workbook["Gắn tem"]
    headers = [cell.value for cell in action[1]]
    item_header = main.EX._BILINGUAL_HEADERS["item_id"]
    applicable_header = main.EX._BILINGUAL_HEADERS["pog_applicable"]
    status_header = main.EX._BILINGUAL_HEADERS["location_status"]
    item_col = headers.index(item_header) + 1
    applicable_col = headers.index(applicable_header) + 1
    status_col = headers.index(status_header) + 1

    bill_row = next(
        row for row in range(2, action.max_row + 1)
        if action.cell(row=row, column=item_col).value == "Bill"
    )
    assert action.cell(row=bill_row, column=applicable_col).value is False
    assert action.cell(row=bill_row, column=status_col).value == "Không áp dụng POG"


def test_bill_rows_do_not_create_overlap_conflicts_with_each_other():
    frame = pd.DataFrame([
        _promo("Bill", name="Invoice A", detail="Bill offer A"),
        _promo("Bill", name="Invoice B", detail="Bill offer B"),
    ])
    frame.loc[1, "campaign_name"] = "September B"

    warnings = main.RC.detect_campaign_time_overlaps(frame)

    assert warnings == []
