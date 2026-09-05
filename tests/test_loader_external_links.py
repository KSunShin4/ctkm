import importlib.util
from pathlib import Path
import sys

import pytest


def _loader_module():
    root = Path(__file__).resolve().parents[1]
    modules_dir = root / "modules"
    if "modules" not in sys.modules:
        import types
        pkg = types.ModuleType("modules")
        pkg.__path__ = [str(modules_dir)]
        sys.modules["modules"] = pkg

    config_spec = importlib.util.spec_from_file_location(
        "modules.config_loader", modules_dir / "config_loader.py"
    )
    config_module = importlib.util.module_from_spec(config_spec)
    sys.modules[config_spec.name] = config_module
    config_spec.loader.exec_module(config_module)

    spec = importlib.util.spec_from_file_location("modules.loader_external_link_test", modules_dir / "1_loader.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize("is_pog_file", [False, True])
def test_shared_loader_disables_external_link_parsing(tmp_path, monkeypatch, is_pog_file):
    from openpyxl import Workbook
    import openpyxl

    loader = _loader_module()
    source = tmp_path / ("pog.xlsx" if is_pog_file else "ctkm-or-clearance.xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["STOREID", "POGCODE", "ITEMID", "LOCATIONID"])
    sheet.append(["S1", "POG1", "P1", 1])
    workbook.save(source)

    real_load_workbook = openpyxl.load_workbook
    calls = []

    def recording_load_workbook(*args, **kwargs):
        calls.append(dict(kwargs))
        return real_load_workbook(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", recording_load_workbook)

    loaded = loader.load_workbook_raw(
        str(source),
        is_pog_file=is_pog_file,
        selected_sheet_names=["Export"],
    )

    assert "Export" in loaded.sheets
    assert len(loaded.sheets["Export"].df) == 1
    assert calls
    assert calls[0].get("keep_links") is False
    assert calls[0].get("read_only") is True
    assert calls[0].get("data_only") is True


def test_second_identical_load_reuses_parsed_workbook_once(tmp_path, monkeypatch):
    from openpyxl import Workbook
    import openpyxl

    loader = _loader_module()
    source = tmp_path / "ctkm.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["ITEMID", "Campaign", "From", "To"])
    sheet.append(["P1", "Promo", "2026-09-01", "2026-09-30"])
    workbook.save(source)

    real_load_workbook = openpyxl.load_workbook
    calls = 0

    def recording_load_workbook(*args, **kwargs):
        nonlocal calls
        calls += 1
        return real_load_workbook(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", recording_load_workbook)

    first = loader.load_workbook_raw(
        str(source), is_pog_file=False, selected_sheet_names=["Export"]
    )
    second = loader.load_workbook_raw(
        str(source), is_pog_file=False, selected_sheet_names=["Export"]
    )

    assert calls == 1
    assert second is first
    assert len(second.sheets["Export"].df) == 1


def test_pog_snapshot_is_detected_during_single_data_stream(tmp_path):
    from openpyxl import Workbook

    loader = _loader_module()
    source = tmp_path / "pog.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    # A single 9/1/2026 value is intentionally ambiguous (MM/DD vs DD/MM).
    # Use the real export-style range so the whole range disambiguates September.
    sheet.append(["Applied filters: DATEID 9/1/2026 - 9/30/2026"])
    sheet.append(["STOREID", "POGCODE", "ITEMID", "LOCATIONID"])
    sheet.append(["S1", "POG1", "P1", 1])
    workbook.save(source)

    loaded = loader.load_workbook_raw(
        str(source), is_pog_file=True, selected_sheet_names=["Export"]
    )

    assert loaded.pog_snapshot_month == "2026-09"
    assert len(loaded.sheets["Export"].df) == 1
