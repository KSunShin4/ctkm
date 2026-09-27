from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import main
from backend.final_output_contract import validate_ctkm_workbook
from backend.slim_output_contract import SLIM_SHEET_ORDER, install_slim_output_contract


def _actions():
    rows = [
        {
            "task_id": "T-GAN",
            "item_id": "P1001",
            "item_name": "Gan",
            "key_channel": "Offline",
            "channel": "OFFLINE",
            "campaign_name": "A",
            "promo_from": pd.Timestamp("2026-09-01"),
            "promo_to": pd.Timestamp("2026-09-30"),
            "scheduled_action": "gắn tem",
            "action": "gắn tem",
            "in_report_scope": True,
            "location_completeness": "Đủ vị trí",
            "location_needs_review": False,
            "review_required": True,
            "review_reason": "xung_dot_chuong_trinh",
            "review_severity": "warning",
            "review_blocks_action": False,
            "program_source": "CTKM",
            "extra__NOTE - Cần lưu ý": "  RAW NOTE  ",
        },
        {
            "task_id": "T-THAO",
            "item_id": "P1002",
            "item_name": "Thao",
            "key_channel": "Online",
            "channel": "Online",
            "campaign_name": "B",
            "promo_from": pd.Timestamp("2026-09-01"),
            "promo_to": pd.Timestamp("2026-09-04"),
            "scheduled_action": "tháo tem",
            "action": "tháo tem",
            "in_report_scope": True,
            "location_completeness": "Đủ vị trí",
            "location_needs_review": False,
            "review_required": False,
            "review_reason": "",
            "program_source": "CTKM",
            "extra__NOTE - Cần lưu ý": "",
        },
        {
            "task_id": "T-WAIT",
            "item_id": "P1003",
            "item_name": "Wait",
            "key_channel": "Others",
            "channel": "Offline/WEB/APP",
            "campaign_name": "C",
            "promo_from": pd.Timestamp("2026-09-27"),
            "promo_to": pd.Timestamp("2026-09-27"),
            "scheduled_action": "chờ gắn tem",
            "action": "chờ gắn tem",
            "in_report_scope": True,
            "location_completeness": "Không tìm thấy SKU trong POG",
            "location_needs_review": True,
            "review_required": False,
            "review_reason": "",
            "program_source": "CTKM",
            "extra__NOTE - Cần lưu ý": "",
        },
    ]
    return pd.DataFrame(rows)


def _install_on_fake_backend():
    core = SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace()),
        _main_mod=SimpleNamespace(EX=main.EX, _core=main._core),
    )
    install_slim_output_contract(core)
    return core


def test_slim_backend_export_keeps_operational_reconciliation_and_audit_sheets(tmp_path):
    original_export = main.EX.export_final_report
    try:
        core = _install_on_fake_backend()
        output = tmp_path / "slim.xlsx"
        actions = _actions()
        summary = {
            "Thang bao cao": "2026-09",
            "Ngay nghiep vu": pd.Timestamp("2026-09-05"),
            "So dong KM thieu vi tri POG": 1,
            "So dong can kiem tra vi tri": 1,
            "So dong cho gan": 1,
            "So dong can review": 1,
            "So dong POG khong co KM phu hop": 0,
        }

        main.EX.export_final_report(
            output_path=output,
            df_matched=pd.DataFrame([{"item_id": "P1001"}]),
            df_ctkm_only=pd.DataFrame(),
            df_pog_only=pd.DataFrame(),
            overlap_warnings=[],
            mapping_logs=[],
            reconciliation_summary=summary,
            df_fefo_actions=actions,
            df_pog_review=pd.DataFrame(),
        )

        workbook = load_workbook(output, data_only=False)
        assert workbook.sheetnames == [
            name for name in SLIM_SHEET_ORDER if name not in {"POG metadata", "KM metadata"}
        ]
        assert "Canh bao trung thoi gian" not in workbook.sheetnames
        assert "Mapping log" in workbook.sheetnames
        assert "Tom tat" not in workbook.sheetnames
        assert "Xác nhận CTKM Clearance" not in workbook.sheetnames

        assert workbook["Gắn tem"].max_row == 2
        assert workbook["Tháo tem"].max_row == 2
        assert workbook["Chờ gắn"].max_row == 2
        assert workbook["Cần review"].max_row == 2
        assert workbook["CTKM thieu vi tri ke"].max_row == 2
        assert workbook["POG cần kiểm tra"].max_row == 2
        assert workbook["Audit hành động"].max_row == 4
        assert workbook["POG không có CTKM"].sheet_state == "hidden"

        headers = [cell.value for cell in workbook["Audit hành động"][1]]
        raw_col = headers.index("extra__NOTE - Cần lưu ý") + 1
        assert workbook["Audit hành động"].cell(2, raw_col).value == "  RAW NOTE  "
        workbook.close()

        validate_ctkm_workbook(output, summary)
        assert core.app.state.ctkm_output_profile == "production-slim"
    finally:
        main.EX.export_final_report = original_export


def test_slim_backend_tong_hop_uses_monthly_plus_clearance_and_keeps_monthly_audit(tmp_path):
    original_export = main.EX.export_final_report
    try:
        _install_on_fake_backend()
        output = tmp_path / "all-km-slim.xlsx"
        actions = pd.DataFrame([
            {
                "item_id": "P1001", "item_key": "P1001", "program_source": "CTKM",
                "campaign_name": "Monthly", "promo_from": pd.Timestamp("2026-09-01"),
                "promo_to": pd.Timestamp("2026-09-30"), "scheduled_action": "gắn tem",
                "action": "gắn tem", "in_report_scope": True,
                "location_completeness": "Đủ vị trí", "location_needs_review": False,
            },
            {
                "item_id": "P2001", "item_key": "P2001", "program_source": "CLEARANCE",
                "campaign_name": "Clearance", "promo_from": pd.Timestamp("2026-09-01"),
                "promo_to": pd.Timestamp("2026-09-30"), "scheduled_action": "gắn tem",
                "action": "gắn tem", "in_report_scope": True,
                "location_completeness": "Đủ vị trí", "location_needs_review": False,
            },
        ])
        monthly_matched = pd.DataFrame([
            {"item_id": "P1001", "pog_code": "A"},
        ])
        monthly_pog_only = pd.DataFrame([
            {"item_id": "P2001", "pog_code": "B"},
            {"item_id": "P3001", "pog_code": "C"},
        ])
        summary = {
            "Thang bao cao": "2026-09",
            "Ngay nghiep vu": pd.Timestamp("2026-09-17"),
            "So dong KM thieu vi tri POG": 0,
            "So dong can kiem tra vi tri": 0,
            "So dong cho gan": 0,
            "So dong can review": 0,
            "So dong POG khong co CTKM phu hop": 2,
            "So item POG khong co CTKM phu hop": 2,
        }

        main.EX.export_final_report(
            output_path=output,
            df_matched=monthly_matched,
            df_ctkm_only=pd.DataFrame(),
            df_pog_only=monthly_pog_only,
            overlap_warnings=[],
            mapping_logs=[],
            reconciliation_summary=summary,
            df_fefo_actions=actions,
            df_pog_review=pd.DataFrame(),
        )

        workbook = load_workbook(output, data_only=True)

        def values(sheet_name):
            return {
                cell.value
                for row in workbook[sheet_name].iter_rows()
                for cell in row
                if cell.value is not None
            }

        assert {"P1001", "P2001"}.issubset(values("Tong hop"))
        assert "P3001" not in values("Tong hop")
        assert "P2001" not in values("POG không có KM")
        assert "P3001" in values("POG không có KM")
        assert {"P2001", "P3001"}.issubset(values("POG không có Monthly CTKM"))
        assert "P2001" not in values("POG không có CTKM")
        assert "P3001" in values("POG không có CTKM")
        assert workbook["POG không có CTKM"].sheet_state == "hidden"
        workbook.close()

        assert summary["So item Clearance bo sung vao Tong hop"] == 1
        assert summary["So item KM+POG thieu trong Tong hop"] == 0
        assert summary["So dong POG khong co KM phu hop"] == 1
        assert summary["So dong POG khong co CTKM phu hop"] == 1
        validate_ctkm_workbook(output, summary)
    finally:
        main.EX.export_final_report = original_export


def test_slim_backend_retains_hidden_source_metadata_and_invalid_pog_rows(tmp_path):
    original_export = main.EX.export_final_report
    try:
        _install_on_fake_backend()
        output = tmp_path / "source-audit-slim.xlsx"
        actions = _actions().iloc[:1].copy()

        confirmation = pd.DataFrame()
        confirmation.attrs["fact_metadata"] = pd.DataFrame([
            {"raw_label": "  Grand Total  ", "raw_value": "\xa0"}
        ])
        pog_review = pd.DataFrame([
            {"item_id": "", "item_name": "  POG row without item code  ", "source_row_id": "POG:Export:99"}
        ])
        pog_review.attrs["pog_metadata"] = pd.DataFrame([
            {"raw_filter": "Applied filter DATEID", "raw_value": "  keep me  "}
        ])
        summary = {
            "Thang bao cao": "2026-09",
            "Ngay nghiep vu": pd.Timestamp("2026-09-17"),
            "So dong KM thieu vi tri POG": 0,
            "So dong can kiem tra vi tri": 0,
            "So dong cho gan": 0,
            "So dong can review": 1,
            "So dong POG khong co KM phu hop": 0,
        }

        main.EX.export_final_report(
            output_path=output,
            df_matched=pd.DataFrame([{"item_id": "P1001"}]),
            df_ctkm_only=pd.DataFrame(),
            df_pog_only=pd.DataFrame(),
            overlap_warnings=[],
            mapping_logs=[],
            reconciliation_summary=summary,
            df_fefo_actions=actions,
            df_ctkm_clearance_confirmation=confirmation,
            df_pog_review=pog_review,
        )

        workbook = load_workbook(output, data_only=True)
        assert workbook["POG metadata"].sheet_state == "hidden"
        assert workbook["KM metadata"].sheet_state == "hidden"
        assert workbook["POG metadata"]["B2"].value == "  keep me  "
        assert workbook["KM metadata"]["A2"].value == "  Grand Total  "
        assert workbook["KM metadata"]["B2"].value == "\xa0"
        pog_review_values = [
            cell.value for row in workbook["POG cần review"].iter_rows() for cell in row
        ]
        assert "  POG row without item code  " in pog_review_values
        workbook.close()

        validate_ctkm_workbook(output, summary)
    finally:
        main.EX.export_final_report = original_export


def test_full_library_export_remains_available_outside_backend_contract():
    assert hasattr(main.EX, "_original_export_final_report")



def test_slim_location_review_does_not_fallback_to_missing_location_rows(tmp_path):
    original_export = main.EX.export_final_report
    try:
        _install_on_fake_backend()
        output = tmp_path / "independent-location-axes.xlsx"
        actions = pd.DataFrame([{
            "task_id": "T-MISSING-LOCATION",
            "item_id": "P9001",
            "item_name": "Missing location",
            "campaign_name": "A",
            "promo_from": pd.Timestamp("2026-09-01"),
            "promo_to": pd.Timestamp("2026-09-30"),
            "scheduled_action": "gắn tem",
            "action": "gắn tem",
            "in_report_scope": True,
            "location_completeness": "Không tìm thấy SKU trong POG",
            # Deliberately no location_needs_review column: review membership is unknown.
            "review_required": False,
            "program_source": "CTKM",
        }])
        summary = {
            "Thang bao cao": "2026-09",
            "Ngay nghiep vu": pd.Timestamp("2026-09-15"),
        }

        main.EX.export_final_report(
            output_path=output,
            df_matched=pd.DataFrame(),
            df_ctkm_only=pd.DataFrame(),
            df_pog_only=pd.DataFrame(),
            overlap_warnings=[],
            mapping_logs=[],
            reconciliation_summary=summary,
            df_fefo_actions=actions,
            df_pog_review=pd.DataFrame(),
        )

        workbook = load_workbook(output, data_only=True)
        missing_values = {
            cell.value for row in workbook["CTKM thieu vi tri ke"].iter_rows()
            for cell in row if cell.value is not None
        }
        review_values = {
            cell.value for row in workbook["POG cần kiểm tra"].iter_rows()
            for cell in row if cell.value is not None
        }
        assert "P9001" in missing_values
        assert "P9001" not in review_values
        assert "Không có vị trí POG nào cần review." in review_values
        workbook.close()
    finally:
        main.EX.export_final_report = original_export



def test_final_validator_rejects_missing_total_or_primary_action_row(tmp_path):
    original_export = main.EX.export_final_report
    try:
        _install_on_fake_backend()
        output = tmp_path / "completeness.xlsx"
        actions = _actions()
        summary = {
            "Thang bao cao": "2026-09",
            "Ngay nghiep vu": pd.Timestamp("2026-09-05"),
            "So dong KM thieu vi tri POG": 1,
            "So dong can kiem tra vi tri": 1,
            "So dong cho gan": 1,
            "So dong can review": 1,
            "So dong POG khong co KM phu hop": 0,
        }
        main.EX.export_final_report(
            output_path=output,
            df_matched=pd.DataFrame([{"item_id": "P1001"}]),
            df_ctkm_only=pd.DataFrame(),
            df_pog_only=pd.DataFrame(),
            overlap_warnings=[],
            mapping_logs=[],
            reconciliation_summary=summary,
            df_fefo_actions=actions,
            df_pog_review=pd.DataFrame(),
        )
        validate_ctkm_workbook(output, summary)

        workbook = load_workbook(output)
        workbook["TỔNG HỢP - TEM"].delete_rows(2, 1)
        workbook.save(output)
        workbook.close()
        with pytest.raises(RuntimeError, match="TỔNG HỢP - TEM"):
            validate_ctkm_workbook(output, summary)

        main.EX.export_final_report(
            output_path=output,
            df_matched=pd.DataFrame([{"item_id": "P1001"}]),
            df_ctkm_only=pd.DataFrame(),
            df_pog_only=pd.DataFrame(),
            overlap_warnings=[],
            mapping_logs=[],
            reconciliation_summary=summary,
            df_fefo_actions=actions,
            df_pog_review=pd.DataFrame(),
        )
        workbook = load_workbook(output)
        workbook["Gắn tem"].delete_rows(2, 1)
        workbook.save(output)
        workbook.close()
        with pytest.raises(RuntimeError, match="Gắn tem"):
            validate_ctkm_workbook(output, summary)
    finally:
        main.EX.export_final_report = original_export


def test_final_validator_rejects_same_task_lineage_with_different_business_values(tmp_path):
    original_export = main.EX.export_final_report
    try:
        _install_on_fake_backend()
        output = tmp_path / "silent-collapse.xlsx"
        actions = pd.DataFrame([
            {
                "task_id": "TASK-COLLISION",
                "source_row_id": "CTKM:Monthly:10",
                "item_id": "P1001",
                "item_name": "Product A",
                "campaign_name": "A",
                "scheduled_action": "gắn tem",
                "action": "gắn tem",
                "in_report_scope": True,
                "location_completeness": "Đủ vị trí",
                "location_needs_review": False,
                "review_required": False,
                "program_source": "CTKM",
            },
            {
                "task_id": "TASK-COLLISION",
                "source_row_id": "CTKM:Monthly:10",
                "item_id": "P1001",
                "item_name": "Product B",
                "campaign_name": "A",
                "scheduled_action": "gắn tem",
                "action": "gắn tem",
                "in_report_scope": True,
                "location_completeness": "Đủ vị trí",
                "location_needs_review": False,
                "review_required": False,
                "program_source": "CTKM",
            },
        ])
        summary = {
            "Thang bao cao": "2026-09",
            "Ngay nghiep vu": pd.Timestamp("2026-09-05"),
            "So dong KM thieu vi tri POG": 0,
            "So dong can kiem tra vi tri": 0,
            "So dong cho gan": 0,
            "So dong can review": 0,
            "So dong POG khong co KM phu hop": 0,
        }
        main.EX.export_final_report(
            output_path=output,
            df_matched=pd.DataFrame(),
            df_ctkm_only=pd.DataFrame(),
            df_pog_only=pd.DataFrame(),
            overlap_warnings=[],
            mapping_logs=[],
            reconciliation_summary=summary,
            df_fefo_actions=actions,
            df_pog_review=pd.DataFrame(),
        )
        with pytest.raises(RuntimeError, match="cùng task/lineage"):
            validate_ctkm_workbook(output, summary)
    finally:
        main.EX.export_final_report = original_export



def test_selected_output_columns_only_trim_operational_views_not_audit(tmp_path):
    original_export = main.EX.export_final_report
    try:
        _install_on_fake_backend()
        output = tmp_path / "selected-output.xlsx"
        actions = _actions().copy()
        actions["extra__NOTE - Cần lưu ý"] = "  raw note  "
        actions["extra__POG__FIXELNAME"] = "  Kệ A  "
        actions["source_row_id"] = "CTKM:Monthly:8"
        actions["source_row_ids"] = "CTKM:Monthly:8"
        actions["source_row_count"] = 1

        summary = {
            "Thang bao cao": "2026-09",
            "Ngay nghiep vu": pd.Timestamp("2026-09-05"),
            "So dong KM thieu vi tri POG": 1,
            "So dong can kiem tra vi tri": 1,
            "So dong cho gan": 1,
            "So dong can review": 1,
            "So dong POG khong co KM phu hop": 0,
        }
        main.EX.export_final_report(
            output_path=output,
            df_matched=pd.DataFrame(),
            df_ctkm_only=pd.DataFrame(),
            df_pog_only=pd.DataFrame(),
            overlap_warnings=[],
            mapping_logs=[],
            reconciliation_summary=summary,
            selected_output_columns=["campaign_name"],
            df_fefo_actions=actions,
            df_pog_review=pd.DataFrame(),
        )

        workbook = load_workbook(output, data_only=True)
        try:
            total_headers = [cell.value for cell in workbook["TỔNG HỢP - TEM"][1]]
            action_headers = [cell.value for cell in workbook["Gắn tem"][1]]
            audit_headers = [cell.value for cell in workbook["Audit hành động"][1]]

            # Required user-facing columns stay pinned even when not selected.
            for required in [
                "KEY Channel", "Channel", "Mã sản phẩm",
                "Tên sản phẩm", "Đơn vị sản phẩm (Promo UoM)",
            ]:
                assert required in total_headers
                assert required in action_headers

            assert "Campaign" in total_headers
            assert "Campaign" in action_headers

            # Source/audit fidelity is never trimmed by presentation selection.
            assert "extra__NOTE - Cần lưu ý" in audit_headers
            assert "extra__POG__FIXELNAME" in audit_headers
            assert "Dòng nguồn chính (source_row_id)" in audit_headers

            # Action sheets keep underlying columns for identity/audit, but
            # presentation selection hides unselected canonical/raw fields.
            assert "extra__POG__FIXELNAME" in action_headers
            raw_idx = action_headers.index("extra__POG__FIXELNAME") + 1
            assert workbook["Gắn tem"].column_dimensions[get_column_letter(raw_idx)].hidden is True

            assert "From" in action_headers
            from_idx = action_headers.index("From") + 1
            assert workbook["Gắn tem"].column_dimensions[get_column_letter(from_idx)].hidden is True
            campaign_idx = action_headers.index("Campaign") + 1
            assert workbook["Gắn tem"].column_dimensions[get_column_letter(campaign_idx)].hidden is False
        finally:
            workbook.close()
    finally:
        main.EX.export_final_report = original_export



def test_slim_workbook_keeps_mapping_provenance_labels(tmp_path):
    original_export = main.EX.export_final_report
    try:
        _install_on_fake_backend()
        output = tmp_path / "mapping-provenance.xlsx"
        actions = _actions().iloc[:1].copy()
        summary = {
            "Thang bao cao": "2026-09",
            "Ngay nghiep vu": pd.Timestamp("2026-09-17"),
            "So dong KM thieu vi tri POG": 0,
            "So dong can kiem tra vi tri": 0,
            "So dong cho gan": 0,
            "So dong can review": 1,
            "So dong POG khong co KM phu hop": 0,
        }
        mapping_logs = [
            {
                "sheet_name": "Monthly",
                "mapping": {"SKU": "item_id", "Campaign": "campaign_name"},
                "confidence": {"SKU": 1.0, "Campaign": 1.0},
                "unmapped": ["Note"],
                "low_confidence_fields": [],
                "source": "user-confirmed",
            },
            {
                "sheet_name": "POG",
                "mapping": {"ITEMID": "item_id", "STOREID": "store_id"},
                "confidence": {"ITEMID": 1.0, "STOREID": 1.0},
                "unmapped": [],
                "low_confidence_fields": [],
                "source": "saved-preset",
            },
        ]

        main.EX.export_final_report(
            output_path=output,
            df_matched=pd.DataFrame([{"item_id": "P1001"}]),
            df_ctkm_only=pd.DataFrame(),
            df_pog_only=pd.DataFrame(),
            overlap_warnings=[],
            mapping_logs=mapping_logs,
            reconciliation_summary=summary,
            df_fefo_actions=actions,
            df_pog_review=pd.DataFrame(),
        )

        workbook = load_workbook(output, data_only=True)
        try:
            ws = workbook["Mapping log"]
            headers = [cell.value for cell in ws[1]]
            source_col = headers.index("nguon") + 1
            canonical_col = headers.index("canonical_field") + 1
            values = [
                (ws.cell(row, canonical_col).value, ws.cell(row, source_col).value)
                for row in range(2, ws.max_row + 1)
            ]
            assert ("item_id", "user-confirmed") in values
            assert ("campaign_name", "user-confirmed") in values
            assert ("item_id", "saved-preset") in values
            assert ("(KHONG MAP DUOC)", "user-confirmed") in values
        finally:
            workbook.close()
    finally:
        main.EX.export_final_report = original_export
