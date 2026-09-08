import asyncio
from io import BytesIO

from fastapi import UploadFile
from openpyxl import Workbook

from backend import main_api
from backend import tag_audit_business_date_contract as business_date
from backend import tag_audit_shelf_contract as shelf_contract


HEADERS = [
    "Phạm vi áp dụng (KEY Channel)", "Kênh chương trình (Channel)",
    "Mã sản phẩm", "Tên sản phẩm", "Đơn vị sản phẩm (Promo UoM)",
    "POG", "Vị trí", "Facing", "Công việc", "Ngày gắn", "Ngày tháo",
    "Campaign", "Scheme detail", "From", "To", "Giá thường", "Giá KM",
    "Nguồn", "Lý do / Cảnh báo", "Product type", "Cửa hàng",
]


def _workbook_bytes():
    workbook = Workbook()
    workbook.remove(workbook.active)
    guide = workbook.create_sheet("HƯỚNG DẪN")
    guide.append(["DANH SÁCH THAO TÁC TEM KHUYẾN MÃI"])
    guide.append(["Tháng báo cáo", "2026-09"])
    guide.append(["Ngày nghiệp vụ", "2026-09-07"])
    for name in ("Gắn tem", "Tháo tem", "Chờ gắn"):
        workbook.create_sheet(name).append(HEADERS)
    workbook["Gắn tem"].append([
        "All channel", "ALL CHANNEL", "P25860", "HOTCHLAND SLEEP NOW", "Hộp",
        "POG-A", "KỆ-01", 1, "snapshot action", "2026-09-06", "2026-09-14",
        "GIẢM 15%", "GIẢM 15%", "2026-09-06", "2026-09-13",
        457000, 388450, "CTKM", "", "NORMAL", "S1",
    ])
    workbook["Gắn tem"].append([
        "All channel", "ALL CHANNEL", "P25175", "FELIX GINKGO BILOBA", "Hộp",
        "POG-A", "KỆ-01", 2, "snapshot action", "2026-09-06", "2026-09-16",
        "MUA 1 TẶNG 1", "MUA 1 TẶNG 1", "2026-09-06", "2026-09-15",
        500000, 500000, "CTKM", "", "NORMAL", "S1",
    ])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _route(path: str, method: str):
    matches = [
        route for route in main_api.app.routes
        if getattr(route, "path", None) == path
        and method in (getattr(route, "methods", None) or set())
    ]
    assert len(matches) == 1
    return matches[0]


def test_shelf_routes_are_store_facing_not_admin_only():
    for path, method in (
        ("/api/tag-audit/shelf/session", "POST"),
        ("/api/tag-audit/shelf/session/{session_id}", "GET"),
    ):
        route = _route(path, method)
        dependency_calls = [getattr(dep, "call", None) for dep in route.dependant.dependencies]
        assert main_api.require_admin not in dependency_calls


def test_shelf_routes_precede_static_frontend_when_present():
    routes = list(main_api.app.routes)
    mount_indexes = [
        index for index, route in enumerate(routes)
        if route.__class__.__name__ == "Mount"
        and (getattr(route, "name", None) == "frontend" or getattr(route, "path", None) in {"", "/"})
    ]
    if not mount_indexes:
        return
    mount_index = mount_indexes[0]
    assert routes.index(_route("/api/tag-audit/shelf/session", "POST")) < mount_index
    assert routes.index(_route("/api/tag-audit/shelf/session/{session_id}", "GET")) < mount_index


def test_public_session_parses_active_tasks_and_groups_small_pog(monkeypatch):
    monkeypatch.setattr(business_date, "current_vn_business_date", lambda: "2026-09-07")
    upload = UploadFile(filename="actions.xlsx", file=BytesIO(_workbook_bytes()))

    payload = asyncio.run(main_api.create_tag_audit_shelf_session(upload))

    assert payload["action_count"] == 2
    assert payload["attach_count"] == 2
    assert payload["remove_count"] == 0
    assert payload["shelf_count"] == 1
    assert payload["shelves"][0]["shelf_label"] == "POG-A"
    assert payload["shelves"][0]["location_count"] == 1
    assert payload["shelves"][0]["locations"] == ["KỆ-01"]
    assert payload["shelves"][0]["sku_count"] == 2
    assert [item["item_id"] for item in payload["items"]] == ["P25860", "P25175"]
    assert payload["progress_scope"] == "gan_and_thao"
    assert payload["context_id"].startswith("sctx_")
    assert payload["workbook_scope_id"].startswith("book_")
    assert payload["store_scope_id"].startswith("store_")
    restored = asyncio.run(main_api.read_tag_audit_shelf_session(payload["session_id"]))
    assert restored["context_id"] == payload["context_id"]


def test_same_month_workbook_surfaces_due_removal_tasks_on_later_date(monkeypatch):
    monkeypatch.setattr(business_date, "current_vn_business_date", lambda: "2026-09-16")
    upload = UploadFile(filename="actions.xlsx", file=BytesIO(_workbook_bytes()))

    payload = asyncio.run(main_api.create_tag_audit_shelf_session(upload))

    assert payload["action_count"] == 2
    assert payload["attach_count"] == 0
    assert payload["remove_count"] == 2
    assert all(item["expected_action"] == "thao_tem" for item in payload["items"])
    assert all(item["promotion_key"].startswith("promo_") for item in payload["items"])
    assert all(item["removal_key"].startswith("remove_") for item in payload["items"])


def test_small_pog_is_one_work_unit_even_with_multiple_locations():
    items = shelf_contract._decorate_items([
        {"task_key":"a1","check_key":"a1","item_id":"P1","pog_code":"POG-A","location_id":"L1","facing":1,"expected_action":"gan_tem"},
        {"task_key":"a2","check_key":"a2","item_id":"P1","pog_code":"POG-A","location_id":"L2","facing":2,"expected_action":"gan_tem"},
        {"task_key":"b1","check_key":"b1","item_id":"P2","pog_code":"POG-A","location_id":"L2","facing":3,"expected_action":"thao_tem"},
    ], max_group_actions=30)
    shelves = shelf_contract._build_shelves(items)

    assert {item["shelf_key"] for item in items} == {"POG::POG-A"}
    assert shelves[0]["shelf_label"] == "POG-A"
    assert shelves[0]["action_count"] == 3
    assert shelves[0]["attach_count"] == 2
    assert shelves[0]["remove_count"] == 1
    assert shelves[0]["sku_count"] == 2
    assert shelves[0]["location_count"] == 2
    assert shelves[0]["locations"] == ["L1", "L2"]


def test_large_pog_is_split_by_location_only_after_threshold():
    source = []
    for index in range(31):
        source.append({
            "task_key":f"t{index}", "check_key":f"t{index}", "item_id":f"P{index:05d}",
            "pog_code":"POG-BIG", "location_id":"L1" if index < 16 else "L2",
            "facing":index, "expected_action":"gan_tem",
        })
    items = shelf_contract._decorate_items(source, max_group_actions=30)
    shelves = shelf_contract._build_shelves(items)

    assert {item["shelf_key"] for item in items} == {"POG::POG-BIG::LOC::L1", "POG::POG-BIG::LOC::L2"}
    assert [shelf["action_count"] for shelf in shelves] == [16, 15]


def test_items_follow_natural_location_and_facing_order():
    items = shelf_contract._decorate_items([
        {"task_key":"c","check_key":"c","item_id":"P3","pog_code":"POG-A","location_id":"KỆ-10","facing":1,"expected_action":"gan_tem"},
        {"task_key":"b","check_key":"b","item_id":"P2","pog_code":"POG-A","location_id":"KỆ-2","facing":2,"expected_action":"gan_tem"},
        {"task_key":"a","check_key":"a","item_id":"P1","pog_code":"POG-A","location_id":"KỆ-2","facing":1,"expected_action":"thao_tem"},
    ], max_group_actions=30)
    assert [(item["location_id"], item["facing"], item["item_id"]) for item in items] == [
        ("KỆ-2", 1, "P1"),
        ("KỆ-2", 2, "P2"),
        ("KỆ-10", 1, "P3"),
    ]


def test_missing_pog_and_location_are_grouped_explicitly():
    items = shelf_contract._decorate_items([
        {"task_key":"a","check_key":"a","item_id":"P1","pog_code":"","location_id":"","expected_action":"gan_tem"},
        {"task_key":"b","check_key":"b","item_id":"P2","pog_code":None,"location_id":None,"expected_action":"gan_tem"},
    ])
    shelves = shelf_contract._build_shelves(items)
    assert len(shelves) == 1
    assert shelves[0]["shelf_key"] == "__NO_LOCATION__"
    assert shelves[0]["shelf_label"] == "Chưa có vị trí kệ"
    assert shelves[0]["action_count"] == 2
