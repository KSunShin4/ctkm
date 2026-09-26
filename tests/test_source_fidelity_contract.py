import pandas as pd
from openpyxl import load_workbook

import main


def _promotion(row_id: str, note: str):
    return {
        "item_id": "P29598",
        "item_key": "P29598",
        "item_name": "Nestle Nutren Junior Sua Bot Cho Tre (1-10Y) Lon 400G",
        "promo_uom": "Lon",
        "key_channel": None,
        "channel": "All Channel",
        "campaign_name": "Clearance Isstop 0-2",
        "campaign_detail": "Discount 67000",
        "promo_from": 46266,
        "promo_to": 46295,
        "normal_price": 353000,
        "promo_price": 286000,
        "product_type": "H&B",
        "program_source": "CLEARANCE",
        "source_sheet": "TỔNG HỢP CTKM UPDATE",
        "source_row_id": row_id,
        "source_excel_row": int(row_id.rsplit(":", 1)[-1]),
        # Representative real-file passthrough columns. These are not part of
        # action classification, but must remain exact in Audit hành động.
        "extra__NOTE - Cần lưu ý": note,
        "extra__New P code": "\xa0",
        "extra__Unit 2": "\xa0",
        "extra__RSP (+VAT)": 353000,
        "extra__Saving per unit": 67000,
        "extra__Out": 0.08,
        "extra__Support type for customer": "Cash discount amount",
        "extra__Big cat": "Health and Beauty",
        "extra__Cat 2": "Nutrition Food",
    }


def test_action_audit_restores_exact_source_passthrough_columns():
    source = pd.DataFrame([
        _promotion("CLEARANCE:TỔNG HỢP CTKM UPDATE:406", "  giữ nguyên khoảng trắng  ")
    ])

    actions = main.RC.build_fefo_action_frame(
        source,
        "2026-09",
        business_date=pd.Timestamp("2026-09-17"),
    )

    assert len(actions) == 1
    row = actions.iloc[0]
    assert row["extra__NOTE - Cần lưu ý"] == "  giữ nguyên khoảng trắng  "
    assert row["extra__New P code"] == "\xa0"
    assert row["extra__Unit 2"] == "\xa0"
    assert row["extra__RSP (+VAT)"] == 353000
    assert row["extra__Saving per unit"] == 67000
    assert row["extra__Out"] == 0.08
    assert row["extra__Support type for customer"] == "Cash discount amount"
    assert row["extra__Big cat"] == "Health and Beauty"
    assert row["extra__Cat 2"] == "Nutrition Food"


def test_different_raw_extra_values_are_not_silently_collapsed():
    source = pd.DataFrame([
        _promotion("CLEARANCE:TỔNG HỢP CTKM UPDATE:406", "NOTE A"),
        _promotion("CLEARANCE:TỔNG HỢP CTKM UPDATE:407", "NOTE B"),
    ])

    actions = main.RC.build_fefo_action_frame(
        source,
        "2026-09",
        business_date=pd.Timestamp("2026-09-17"),
    )

    assert len(actions) == 2
    assert actions["extra__NOTE - Cần lưu ý"].tolist() == ["NOTE A", "NOTE B"]
    assert actions["source_row_id"].tolist() == [
        "CLEARANCE:TỔNG HỢP CTKM UPDATE:406",
        "CLEARANCE:TỔNG HỢP CTKM UPDATE:407",
    ]


def test_exact_duplicate_raw_rows_keep_all_lineage_and_one_exact_value():
    first = _promotion("CLEARANCE:TỔNG HỢP CTKM UPDATE:406", "SAME")
    second = dict(first)
    second["source_row_id"] = "CLEARANCE:TỔNG HỢP CTKM UPDATE:407"
    second["source_excel_row"] = 407
    source = pd.DataFrame([first, second])

    actions = main.RC.build_fefo_action_frame(
        source,
        "2026-09",
        business_date=pd.Timestamp("2026-09-17"),
    )

    assert len(actions) == 1
    row = actions.iloc[0]
    assert row["extra__NOTE - Cần lưu ý"] == "SAME"
    assert row["source_row_count"] == 2
    assert set(part.strip() for part in row["source_row_ids"].split("|")) == {
        "CLEARANCE:TỔNG HỢP CTKM UPDATE:406",
        "CLEARANCE:TỔNG HỢP CTKM UPDATE:407",
    }



def _pog(row_id: str, *, fixel_name: str, invent_unit=3):
    return {
        "item_id": "P29598",
        "item_name": "POG Product Name",
        "store_id": "SGPMC1885",
        "pog_code": "BABY2.1885",
        "location_id": 4,
        "facing": 2,
        "pog_snapshot_month": "2026-09",
        "source_sheet": "Export",
        "source_row_id": row_id,
        "source_excel_row": int(row_id.rsplit(":", 1)[-1]),
        "extra__FIXELNAME": fixel_name,
        "extra__TOTALUNIT": 12,
        "extra__INVENTUNIT": invent_unit,
        "extra__PROMOTIONTYPE": "  RAW PROMO TYPE  ",
    }


def test_pog_raw_values_survive_action_and_audit_without_false_deduplication(tmp_path):
    source = pd.DataFrame([
        _promotion("CTKM:Monthly:10", "Monthly raw note"),
    ])
    source.loc[0, "program_source"] = "CTKM"
    source.loc[0, "source_sheet"] = "Monthly"
    source.loc[0, "source_row_id"] = "CTKM:Monthly:10"
    source.loc[0, "source_excel_row"] = 10

    pog = pd.DataFrame([
        _pog("POG:Export:20", fixel_name="  Kệ A  ", invent_unit=3),
        _pog("POG:Export:21", fixel_name="  Kệ B  ", invent_unit=4),
    ])

    actions = main.RC.build_fefo_action_frame(
        source,
        "2026-09",
        pog,
        business_date=pd.Timestamp("2026-09-17"),
    )

    # Same physical shelf coordinates, but different raw source data are not
    # silently collapsed.  Physical task identity remains the same.
    assert len(actions) == 2
    assert actions["pog_position_id"].nunique() == 1
    assert actions["task_id"].nunique() == 1
    assert set(actions["pog_source_row_id"]) == {"POG:Export:20", "POG:Export:21"}
    assert actions["pog_position_count"].tolist() == [1, 1]
    assert actions["extra__POG__FIXELNAME"].tolist() == ["  Kệ A  ", "  Kệ B  "]
    assert actions["extra__POG__INVENTUNIT"].tolist() == [3, 4]
    assert actions["extra__POG__PROMOTIONTYPE"].tolist() == [
        "  RAW PROMO TYPE  ", "  RAW PROMO TYPE  "
    ]

    output = tmp_path / "pog-source-fidelity.xlsx"
    main.EX.export_final_report(
        output,
        pd.DataFrame(), pd.DataFrame(), pd.DataFrame(),
        [], [],
        {"Thang bao cao": "2026-09", "Ngay nghiep vu": pd.Timestamp("2026-09-17")},
        df_fefo_actions=actions,
    )

    workbook = load_workbook(output, data_only=True)
    try:
        audit = workbook["Audit hành động"]
        headers = [cell.value for cell in audit[1]]
        fixel_col = headers.index("extra__POG__FIXELNAME") + 1
        invent_col = headers.index("extra__POG__INVENTUNIT") + 1
        assert [audit.cell(row, fixel_col).value for row in (2, 3)] == ["  Kệ A  ", "  Kệ B  "]
        assert [audit.cell(row, invent_col).value for row in (2, 3)] == [3, 4]
    finally:
        workbook.close()


def test_complete_pog_rows_only_collapse_when_all_raw_source_values_match():
    source = pd.DataFrame([_promotion("CTKM:Monthly:10", "Monthly raw note")])
    source.loc[0, "program_source"] = "CTKM"
    first = _pog("POG:Export:20", fixel_name="Same", invent_unit=3)
    second = dict(first)
    second["source_row_id"] = "POG:Export:21"
    second["source_excel_row"] = 21

    actions = main.RC.build_fefo_action_frame(
        source,
        "2026-09",
        pd.DataFrame([first, second]),
        business_date=pd.Timestamp("2026-09-17"),
    )

    assert len(actions) == 1
    assert actions.loc[0, "pog_source_row_count"] == 2
    assert set(actions.loc[0, "pog_source_row_ids"].split(" | ")) == {
        "POG:Export:20", "POG:Export:21"
    }
    assert actions.loc[0, "extra__POG__FIXELNAME"] == "Same"
