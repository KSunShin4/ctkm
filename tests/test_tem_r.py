import importlib.util
from pathlib import Path
import sys

import pandas as pd


def _module():
    path = Path(__file__).resolve().parents[1] / "modules" / "10_tem_r.py"
    spec = importlib.util.spec_from_file_location("tem_r", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_stock_two_can_represent_itself_and_one_adjacent_empty_location():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 3, "ITEMID": "P3", "ITEMNAME": "Three", "TỒN": 2},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 4, "ITEMID": "P4", "ITEMNAME": "Four", "TỒN": 0},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 5, "ITEMID": "P5", "ITEMNAME": "Five", "TỒN": 0},
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert pd.isna(plan.rows.loc[0, "TEM R"])
    assert plan.rows.loc[1, "TEM R"] == "P3"
    assert plan.rows.loc[2, "TEM R"] == "KHÔNG CÓ SẢN PHẨM THAY THẾ"
    assert len(plan.exceptions) == 1


def test_short_shelf_stock_is_reserved_for_the_nearest_adjacent_location():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "CVS_NOMILK_1B_180X70_LI", "POGCODE": "CVS8.NEW.LI", "FIXELID": "1/6", "LOCATIONID": 1, "ITEMID": "P25988", "ITEMNAME": "Mycospring", "TỒN": 0},
        {"STOREID": "CVS_NOMILK_1B_180X70_LI", "POGCODE": "CVS8.NEW.LI", "FIXELID": "1/6", "LOCATIONID": 2, "ITEMID": "P25990", "ITEMNAME": "Mycospring Chuoi", "TỒN": 0},
        {"STOREID": "CVS_NOMILK_1B_180X70_LI", "POGCODE": "CVS8.NEW.LI", "FIXELID": "1/6", "LOCATIONID": 3, "ITEMID": "P22814", "ITEMNAME": "Tracybee", "TỒN": 0},
        {"STOREID": "CVS_NOMILK_1B_180X70_LI", "POGCODE": "CVS8.NEW.LI", "FIXELID": "1/6", "LOCATIONID": 4, "ITEMID": "P05464", "ITEMNAME": "Xuan Nguyen", "TỒN": 2},
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert plan.rows.loc[0, "TEM R"] == "KHÔNG CÓ SẢN PHẨM THAY THẾ"
    assert plan.rows.loc[1, "TEM R"] == "KHÔNG CÓ SẢN PHẨM THAY THẾ"
    assert plan.rows.loc[2, "TEM R"] == "P05464"
    assert pd.isna(plan.rows.loc[3, "TEM R"])
    assert plan.rows.loc[3, "TEM_R_STATUS"] == "Còn tồn - không cần Tem R"


def test_shelf_with_zero_stock_is_fully_rearranged():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 1, "ITEMID": "P1", "ITEMNAME": "One", "TỒN": 0},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 2, "ITEMID": "P2", "ITEMNAME": "Two", "TỒN": 1},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 3, "ITEMID": "P3", "ITEMNAME": "Three", "TỒN": 3},
    ])

    plan = tem_r.build_tem_r_plan(frame, max_location_distance=2)

    assert plan.rows["TEM R"].tolist() == ["P2", "P3", "P3"]


def test_split_replacement_group_triggers_full_shelf_rearrangement():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 1, "ITEMID": "P1", "ITEMNAME": "One", "TỒN": 0},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 2, "ITEMID": "PB", "ITEMNAME": "B", "TỒN": 1},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 3, "ITEMID": "PA", "ITEMNAME": "A", "TỒN": 3},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 4, "ITEMID": "P2", "ITEMNAME": "Two", "TỒN": 0},
    ])

    plan = tem_r.build_tem_r_plan(frame, max_location_distance=2)

    assert plan.rows["TEM R"].tolist() == ["PB", "PA", "PA", "PA"]
    assert plan.rows["TEM R"].tolist().count("PA") == 3


def test_repeated_pog_rows_do_not_add_the_same_store_stock_twice():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 1, "ITEMID": "PA", "ITEMNAME": "A", "TỒN": 3},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 2, "ITEMID": "PA", "ITEMNAME": "A", "TỒN": 3},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 3, "ITEMID": "P0", "ITEMNAME": "Zero", "TỒN": 0},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 4, "ITEMID": "P9", "ITEMNAME": "Zero", "TỒN": 0},
    ])

    plan = tem_r.build_tem_r_plan(frame)

    # Tồn PA là 3 (không phải 3 + 3); hai vị trí gốc nên chỉ thay thêm một vị trí.
    assert plan.rows["TEM R"].dropna().tolist().count("PA") == 1
    assert plan.rows.loc[3, "TEM R"] == "KHÔNG CÓ SẢN PHẨM THAY THẾ"


def test_user_can_choose_a_differently_named_stock_column():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 1,
         "ITEMID": "P1", "ITEMNAME": "One", "On hand": 2},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 2,
         "ITEMID": "P2", "ITEMNAME": "Two", "On hand": 0},
    ])

    plan = tem_r.build_tem_r_plan(frame, stock_column="On hand")

    assert plan.rows.loc[1, "TEM R"] == "P1"
    assert "On hand" in plan.audit.columns


def test_blank_stock_is_reported_as_an_exception():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 1,
         "ITEMID": "P1", "ITEMNAME": "One", "TỒN": None},
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert len(plan.exceptions) == 1
    assert plan.rows.loc[0, "TEM R"] == "BỎ QUA - KHÔNG CÓ DỮ LIỆU TỒN KHO"
    assert plan.exceptions.iloc[0]["reason"] == "Bỏ qua - không có dữ liệu tồn kho"


def test_na_stock_text_is_reported_as_an_exception_instead_of_zero_stock():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 1,
         "ITEMID": "P1", "ITEMNAME": "One", "TỒN": "NA"},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 2,
         "ITEMID": "P2", "ITEMNAME": "Two", "TỒN": "#N/A"},
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert plan.rows["TEM R"].tolist() == [
        "BỎ QUA - KHÔNG CÓ DỮ LIỆU TỒN KHO",
        "BỎ QUA - TỒN KHO LỖI #N/A",
    ]
    assert plan.exceptions["reason"].tolist() == [
        "Bỏ qua - không có dữ liệu tồn kho", "Bỏ qua - tồn kho lỗi #N/A",
    ]


def test_workbook_skips_cached_formula_error_and_continues(tmp_path):
    tem_r = _module()
    from openpyxl import Workbook
    import re
    import zipfile

    source = tmp_path / "formula-error.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "TỒN"])
    sheet.append(["S1", "G1", "1/1", 1, "P1", "One", "=NA()"])
    sheet.append(["S1", "G1", "1/1", 2, "P2", "Two", 2])
    sheet.append(["S1", "G1", "1/1", 3, "P3", "Three", 0])
    workbook.save(source)

    with zipfile.ZipFile(source, "r") as zin:
        members = {name: zin.read(name) for name in zin.namelist()}
    patched, replacements = re.subn(
        br'<c r="G2"[^>]*><f>NA\(\)</f><v\s*/></c>',
        b'<c r="G2" t="e"><f>NA()</f><v>#N/A</v></c>',
        members["xl/worksheets/sheet1.xml"],
    )
    assert replacements == 1
    members["xl/worksheets/sheet1.xml"] = patched
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, content in members.items():
            zout.writestr(name, content)

    plan = tem_r.create_tem_r_workbook(source, output)

    assert plan.rows.loc[0, "TEM R"] == "BỎ QUA - TỒN KHO LỖI #N/A"
    assert plan.rows.loc[0, "TEM_R_STATUS"] == "Bỏ qua - tồn kho lỗi #N/A"
    assert plan.rows.loc[2, "TEM R"] == "P2"
    assert plan.exceptions.iloc[0].to_dict() == {
        "row": 2, "ITEMID": "P1", "reason": "Bỏ qua - tồn kho lỗi #N/A",
    }

    from openpyxl import load_workbook
    result = load_workbook(output, data_only=False)
    export = result["Export"]
    assert export.cell(2, 8).value == "BỎ QUA - TỒN KHO LỖI #N/A"
    assert export.cell(4, 8).value == "P2"
    assert export.cell(2, 7).fill.fgColor.rgb == "00FCE4D6"
    assert result["Tem R ngoại lệ"].cell(2, 1).value == 2


def test_workbook_skips_formula_without_cached_value_and_processes_other_rows(tmp_path):
    tem_r = _module()
    from openpyxl import Workbook, load_workbook

    source = tmp_path / "formula-without-cache.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "TỒN"])
    sheet.append(["S1", "G1", "1/1", 1, "P1", "One", "=1+1"])
    sheet.append(["S1", "G1", "1/1", 2, "P2", "Two", 2])
    sheet.append(["S1", "G1", "1/1", 3, "P3", "Three", 0])
    workbook.save(source)

    plan = tem_r.create_tem_r_workbook(source, output)

    assert plan.rows.loc[0, "TEM R"] == "BỎ QUA - KHÔNG CÓ DỮ LIỆU TỒN KHO"
    assert plan.rows.loc[2, "TEM R"] == "P2"
    assert plan.exceptions.iloc[0]["row"] == 2
    export = load_workbook(output, data_only=False)["Export"]
    assert export.cell(2, 8).value == "BỎ QUA - KHÔNG CÓ DỮ LIỆU TỒN KHO"
    assert export.cell(4, 8).value == "P2"


def test_fractional_stock_below_one_is_treated_as_zero():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 1,
         "ITEMID": "P1", "ITEMNAME": "One", "TỒN": 0.166667},
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert plan.rows.loc[0, "_stock"] == 0
    assert plan.rows.loc[0, "TEM R"] == "KHÔNG CÓ SẢN PHẨM THAY THẾ"
    assert plan.rows.loc[0, "TEM_R_STATUS"] == "Kệ không còn đủ tồn để thay"


def test_earlier_product_with_enough_stock_fills_following_positions():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 1, "ITEMID": "P1", "ITEMNAME": "One", "TỒN": 5},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 3, "ITEMID": "P3", "ITEMNAME": "Three", "TỒN": 2},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 4, "ITEMID": "P4", "ITEMNAME": "Four", "TỒN": 0},
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert plan.rows["TEM R"].isna().iloc[:2].all()
    assert plan.rows.loc[2, "TEM R"] == "P3"


def test_workbook_accepts_user_mapping_for_required_columns(tmp_path):
    tem_r = _module()
    from openpyxl import Workbook, load_workbook

    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["Cửa hàng", "Mã POG", "Kệ", "Vị trí", "Mã hàng", "Tên hàng", "Tồn thực tế"])
    sheet.append(["S1", "G1", "1/1", 1, "P1", "One", 2])
    sheet.append(["S1", "G1", "1/1", 2, "P2", "Two", 0])
    workbook.save(source)

    mapping = dict(zip(tem_r.REQUIRED_COLUMNS, ["Cửa hàng", "Mã POG", "Kệ", "Vị trí", "Mã hàng", "Tên hàng"]))
    plan = tem_r.create_tem_r_workbook(
        source, output, stock_column="Tồn thực tế", column_mapping=mapping,
    )

    assert plan.rows.loc[1, "TEM R"] == "P1"
    result = load_workbook(output, data_only=True)["Export"]
    assert result.cell(2, 8).value is None
    assert result.cell(3, 8).value == "P1"


def test_inspection_handles_workbook_without_dimension_metadata(tmp_path):
    tem_r = _module()
    from openpyxl import Workbook
    import zipfile

    source = tmp_path / "missing-dimension.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "TỒN"])
    sheet.append(["S1", "G1", "1/1", 1, "P1", "One", 2])
    workbook.save(source)

    with zipfile.ZipFile(source, "r") as zin:
        members = {name: zin.read(name) for name in zin.namelist()}
    import re
    members["xl/worksheets/sheet1.xml"] = re.sub(
        br"<dimension[^>]*/>", b"", members["xl/worksheets/sheet1.xml"]
    )
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, content in members.items():
            zout.writestr(name, content)

    result = tem_r.inspect_tem_r_workbook(source)

    assert result["sheet_name"] == "Export"
    assert result["suggested_item_column"] == "ITEMID"
    assert result["suggested_stock_column"] == "TỒN"


def test_inspection_rejects_duplicate_headers(tmp_path):
    tem_r = _module()
    from openpyxl import Workbook

    source = tmp_path / "duplicate-headers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ITEMID", "ITEMID", "TỒN"])
    sheet.append(["P1", "P1", 1])
    workbook.save(source)

    try:
        tem_r.inspect_tem_r_workbook(source)
    except ValueError as exc:
        assert "tên cột bị trùng" in str(exc)
    else:
        raise AssertionError("duplicate headers should be rejected")


def test_shelf_without_zero_stock_has_no_tem_r():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 1, "ITEMID": "P1", "ITEMNAME": "One", "TỒN": 3},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 2, "ITEMID": "P2", "ITEMNAME": "Two", "TỒN": 2},
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert plan.rows["TEM R"].isna().all()


def test_stock_units_fill_positions_on_the_same_fixture():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 1, "ITEMID": "P1", "ITEMNAME": "One", "TỒN": 3},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 4, "ITEMID": "P4", "ITEMNAME": "Four", "TỒN": 0},
    ])

    plan = tem_r.build_tem_r_plan(frame, max_location_distance=3)

    assert plan.rows.loc[1, "TEM R"] == "P1"
    assert plan.rows.loc[1, "TEM_R_STATUS"] == "Tồn 0 - mã thay đã nằm liền nhóm"


def test_one_stock_unit_covers_only_its_own_position():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 1, "ITEMID": "P1", "ITEMNAME": "One", "TỒN": 1},
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 2, "ITEMID": "P2", "ITEMNAME": "Two", "TỒN": 0},
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert pd.isna(plan.rows.loc[0, "TEM R"])
    assert plan.rows.loc[1, "TEM R"] == "KHÔNG CÓ SẢN PHẨM THAY THẾ"
    assert len(plan.exceptions) == 1


def test_writes_explicit_message_when_pog_has_no_replacement():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID": "S1", "POGCODE": "G1", "FIXELID": "1/1", "LOCATIONID": 1, "ITEMID": "P1", "ITEMNAME": "One", "TỒN": 0},
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert plan.rows.loc[0, "TEM R"] == "KHÔNG CÓ SẢN PHẨM THAY THẾ"
    assert plan.rows.loc[0, "TEM_R_STATUS"] == "Kệ không còn đủ tồn để thay"


def test_same_shelf_replacement_is_not_rejected_by_distance():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID":"S1", "POGCODE":"G1", "FIXELID":"1/1", "LOCATIONID":1,
         "ITEMID":"P1", "ITEMNAME":"One", "TỒN":3},
        {"STOREID":"S1", "POGCODE":"G1", "FIXELID":"1/1", "LOCATIONID":4,
         "ITEMID":"P4", "ITEMNAME":"Four", "TỒN":0},
    ])

    # max_location_distance remains accepted for API compatibility, but distance
    # is only a priority signal now; same-shelf stock must remain eligible.
    plan = tem_r.build_tem_r_plan(frame, max_location_distance=1)

    assert plan.rows.loc[1, "TEM R"] == "P1"


def test_same_shelf_uses_nearest_available_item_first():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID":"S1", "POGCODE":"G1", "FIXELID":"1/1", "LOCATIONID":1,
         "ITEMID":"PA", "ITEMNAME":"A", "TỒN":2},
        {"STOREID":"S1", "POGCODE":"G1", "FIXELID":"1/1", "LOCATIONID":3,
         "ITEMID":"P0", "ITEMNAME":"Zero", "TỒN":0},
        {"STOREID":"S1", "POGCODE":"G1", "FIXELID":"1/1", "LOCATIONID":4,
         "ITEMID":"PB", "ITEMNAME":"B", "TỒN":2},
    ])

    plan = tem_r.build_tem_r_plan(frame, max_location_distance=1)

    assert plan.rows.loc[1, "TEM R"] == "PB"


def test_real_confect_case_uses_far_same_shelf_stock_and_groups_split_item():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID":"CONFECT", "POGCODE":"CTCONF", "FIXELID":"1/5", "LOCATIONID":2,
         "ITEMID":"P08844", "ITEMNAME":"Huro Kẹo", "TỒN":34},
        {"STOREID":"CONFECT", "POGCODE":"CTCONF", "FIXELID":"1/5", "LOCATIONID":3,
         "ITEMID":"P29216", "ITEMNAME":"Kefir Snack", "TỒN":0},
        {"STOREID":"CONFECT", "POGCODE":"CTCONF", "FIXELID":"1/5", "LOCATIONID":4,
         "ITEMID":"P29217", "ITEMNAME":"Kefir Snack", "TỒN":1},
        {"STOREID":"CONFECT", "POGCODE":"CTCONF", "FIXELID":"1/5", "LOCATIONID":5,
         "ITEMID":"P25749", "ITEMNAME":"Pororo", "TỒN":0},
        {"STOREID":"CONFECT", "POGCODE":"CTCONF", "FIXELID":"1/5", "LOCATIONID":6,
         "ITEMID":"P25748", "ITEMNAME":"Pororo", "TỒN":0},
    ])

    plan = tem_r.build_tem_r_plan(frame, max_location_distance=1)
    codes = plan.rows["TEM R"].tolist()

    assert "KHÔNG CÓ SẢN PHẨM THAY THẾ" not in codes
    assert codes == ["P08844", "P08844", "P08844", "P08844", "P29217"]
    assert codes.count("P08844") == 4


def test_missing_shelf_key_is_quarantined_instead_of_grouped():
    tem_r = _module()
    frame = pd.DataFrame([
        {"STOREID":None, "POGCODE":"G1", "FIXELID":"1/1", "LOCATIONID":1,
         "ITEMID":"P1", "ITEMNAME":"One", "TỒN":3},
        {"STOREID":None, "POGCODE":"G1", "FIXELID":"1/1", "LOCATIONID":2,
         "ITEMID":"P2", "ITEMNAME":"Two", "TỒN":0},
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert plan.rows["TEM R"].isna().all()
    assert plan.exceptions["reason"].tolist() == ["Thiếu mã cửa hàng", "Thiếu mã cửa hàng"]


def test_stock_parser_supports_common_localized_numbers():
    tem_r = _module()

    assert tem_r._as_stock("1,000") == 1000
    assert tem_r._as_stock("1.000") == 1000
    assert tem_r._as_stock("2,5") == 2
    assert tem_r._as_stock("0.166") == 0
    assert tem_r._as_stock("0,5") == 0
    assert tem_r._as_stock("0.999") == 0
    assert tem_r._as_stock("1,5") == 1
    assert tem_r._as_stock("2.9") == 2
    assert tem_r._as_stock("1,234.5") == 1234
