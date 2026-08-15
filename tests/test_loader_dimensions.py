import importlib.util
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook


def _loader_module():
    import modules.config_loader  # ensure the relative import has a package

    path = Path(__file__).resolve().parents[1] / "modules" / "1_loader.py"
    spec = importlib.util.spec_from_file_location("modules.loader_dimensions_test", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_loader_ignores_inflated_xfd_dimension(tmp_path):
    source = tmp_path / "inflated.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ITEMID", "ITEMNAME"])
    sheet.append(["P1", "One"])
    workbook.save(source)

    with zipfile.ZipFile(source, "r") as zin:
        members = {name: zin.read(name) for name in zin.namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
        b'ref="A1:B2"', b'ref="A1:XFD2"'
    )
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, content in members.items():
            zout.writestr(name, content)

    loaded = _loader_module().load_workbook_raw(str(source))

    assert list(loaded.sheets["Sheet"].df.columns) == ["ITEMID", "ITEMNAME"]
    assert loaded.sheets["Sheet"].df.iloc[0].tolist() == ["P1", "One"]


def test_loader_rejects_fake_xlsx_with_clear_error(tmp_path):
    source = tmp_path / "fake.xlsx"
    source.write_text("not an excel workbook", encoding="utf-8")

    try:
        _loader_module().load_workbook_raw(str(source))
    except ValueError as exc:
        assert "không phải XLSX hợp lệ" in str(exc)
    else:
        raise AssertionError("fake XLSX should be rejected")


def test_loader_finds_header_below_old_six_row_limit(tmp_path):
    source = tmp_path / "deep-header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for _ in range(10):
        sheet.append([])
    sheet.append(["ITEMID", "ITEMNAME", "STOREID"])
    sheet.append(["P1", "One", "S1"])
    workbook.save(source)

    loaded = _loader_module().load_workbook_raw(str(source))

    assert loaded.sheets["Sheet"].header_row_index == 10
    assert loaded.sheets["Sheet"].df.iloc[0].tolist() == ["P1", "One", "S1"]


def test_loader_materializes_only_selected_sheet(tmp_path):
    source = tmp_path / "multiple.xlsx"
    workbook = Workbook()
    chosen = workbook.active
    chosen.title = "Chosen"
    chosen.append(["ITEMID", "ITEMNAME"])
    chosen.append(["P1", "One"])
    ignored = workbook.create_sheet("Ignored")
    ignored.append(["BROKEN", "DATA"])
    ignored.append(["x", "y"])
    workbook.save(source)

    module = _loader_module()
    assert module.inspect_workbook_sheets(str(source)) == ["Chosen", "Ignored"]
    loaded = module.load_workbook_raw(str(source), selected_sheet_names=["Chosen"])

    assert list(loaded.sheets) == ["Chosen"]
    assert loaded.sheets["Chosen"].df.iloc[0].tolist() == ["P1", "One"]
