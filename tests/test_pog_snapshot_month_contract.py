from openpyxl import Workbook

import main


def test_standard_pharmacity_dateid_range_detects_month_first_format():
    row = [
        "Applied filters:\nFIXELNAME does not contain 'NONPOG'\n"
        "MODELAPPLYSTORE is VTB\nDATEID 9/1/2026 - 9/30/2026"
    ]

    assert main.L._extract_pog_snapshot_month_from_values(row) == "2026-09"


def test_day_first_dateid_range_is_supported_when_range_disambiguates_it():
    row = ["Applied filters: DATEID 1/9/2026 - 30/9/2026"]

    assert main.L._extract_pog_snapshot_month_from_values(row) == "2026-09"


def test_ambiguous_dateid_does_not_guess_the_snapshot_month():
    # 2/3 can mean either 2 March or 3 February. With no disambiguating day in
    # the range, choosing either would silently attach the wrong month to all POG rows.
    row = ["Applied filters: DATEID 2/3/2026 - 2/3/2026"]

    assert main.L._extract_pog_snapshot_month_from_values(row) is None


def test_snapshot_month_comes_from_trailing_sheet_metadata_not_filename(tmp_path):
    source = tmp_path / "data (3).xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append([
        "STOREID", "POGCODE", "ITEMID", "ITEMNAME", "LOCATIONID", "TOTALFACING",
    ])
    sheet.append(["S1", "A", "P1001", "Product", 1, 2])
    sheet.append([
        "Applied filters:\nFIXELNAME does not contain 'NONPOG'\n"
        "MODELAPPLYSTORE is VTB\nDATEID 9/1/2026 - 9/30/2026"
    ])
    workbook.save(source)

    loaded = main.L.load_workbook_raw(
        str(source), is_pog_file=True, selected_sheet_names=["Export"]
    )

    assert loaded.pog_snapshot_month == "2026-09"
    assert "Export" in loaded.sheets


def test_snapshot_alignment_uses_detected_metadata_month(tmp_path):
    source = tmp_path / "opaque-name.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["STOREID", "POGCODE", "ITEMID", "LOCATIONID"])
    sheet.append(["S1", "A", "P1001", 1])
    sheet.append(["DATEID 9/1/2026 - 9/30/2026"])
    workbook.save(source)

    loaded = main.L.load_workbook_raw(
        str(source), is_pog_file=True, selected_sheet_names=["Export"]
    )
    alignment = main.TF.check_snapshot_alignment("2026-09", loaded.pog_snapshot_month)

    assert alignment.is_aligned is True
    assert alignment.warning_message is None
