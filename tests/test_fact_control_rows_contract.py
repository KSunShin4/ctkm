from openpyxl import load_workbook
import pandas as pd

import main


BUSINESS_DATE = pd.Timestamp("2026-09-15")


def _raw_detail_sheet():
    return pd.DataFrame([
        {
            "Item": "P1001",
            "Product": "Valid product",
            "Promo UoM": "Hộp",
            "Channel": "Offline",
            "Campaign": "Combo Baby",
            "From": pd.Timestamp("2026-09-01"),
            "To": pd.Timestamp("2026-09-30"),
            "Normal": 100000,
            "Promo": 90000,
        },
        {
            # Genuine malformed business row: missing item code, but all other
            # business anchors are present. It must remain reviewable.
            "Item": None,
            "Product": "Missing-code product",
            "Promo UoM": "Hộp",
            "Channel": "Offline",
            "Campaign": "Broken campaign",
            "From": pd.Timestamp("2026-09-01"),
            "To": pd.Timestamp("2026-09-30"),
            "Normal": 120000,
            "Promo": 100000,
        },
        {
            # Shape seen in optional Monthly combo/detail tabs: aggregate prices
            # can be populated even though there is no SKU/date/location context.
            "Item": None,
            "Product": None,
            "Promo UoM": None,
            "Channel": None,
            "Campaign": 'Combo "Mẹ tròn - Con vuông" Total',
            "From": None,
            "To": None,
            "Normal": 1542000,
            "Promo": 1184100,
        },
        {
            "Item": None,
            "Product": None,
            "Promo UoM": None,
            "Channel": None,
            "Campaign": "Grand Total",
            "From": None,
            "To": None,
            "Normal": 5407000,
            "Promo": 4545275,
        },
    ])


def _mapping():
    return {
        "Item": "item_id",
        "Product": "item_name",
        "Promo UoM": "promo_uom",
        "Channel": "channel",
        "Campaign": "campaign_name",
        "From": "promo_from",
        "To": "promo_to",
        "Normal": "normal_price",
        "Promo": "promo_price",
    }


def _processed_detail_sheet():
    warnings = []
    result = main._process(
        "Combo detail",
        _raw_detail_sheet(),
        main.SCHEMA["fact_fields"],
        [],
        True,
        warnings,
        mapping_override=_mapping(),
        source_excel_rows=[2, 3, 4, 5],
        source_prefix="CTKM",
    )
    return result, warnings


def test_fact_subtotals_are_separated_without_hiding_real_missing_item_rows():
    processed, warnings = _processed_detail_sheet()

    assert len(processed) == 2
    assert processed["item_id"].tolist()[0] == "P1001"
    assert pd.isna(processed["item_id"].tolist()[1])
    assert processed["item_name"].tolist() == ["Valid product", "Missing-code product"]

    metadata = processed.attrs.get("fact_metadata")
    assert metadata is not None
    assert len(metadata) == 2
    assert set(metadata["fact_control_kind"]) == {"subtotal", "grand_total"}
    assert set(metadata["source_row_id"]) == {"CTKM:Combo detail:4", "CTKM:Combo detail:5"}
    assert set(metadata["program_source"]) == {"CTKM"}
    assert metadata.loc[
        metadata["fact_control_kind"].eq("subtotal"), "campaign_name"
    ].iloc[0] == 'Combo "Mẹ tròn - Con vuông" Total'
    assert metadata.loc[
        metadata["fact_control_kind"].eq("grand_total"), "normal_price"
    ].iloc[0] == 5407000

    assert len(warnings) == 1
    assert "Tách 2 dòng subtotal/control" in warnings[0]
    assert "KM metadata" in warnings[0]

    actions = main.RC.build_fefo_action_frame(
        processed,
        "2026-09",
        business_date=BUSINESS_DATE,
    )
    assert len(actions) == 2
    by_name = actions.set_index("item_name")["scheduled_action"].to_dict()
    assert by_name == {
        "Valid product": "gắn tem",
        "Missing-code product": "review",
    }
    assert actions.set_index("item_name").loc["Missing-code product", "reason"] == "thieu_ma_san_pham"
    assert not actions["campaign_name"].astype(str).str.endswith("Total").any()


def test_fact_subtotals_remain_source_auditable_but_not_user_work_rows(tmp_path):
    processed, _warnings = _processed_detail_sheet()
    confirmation = main._build_ctkm_clearance_confirmation(processed, None)
    metadata = confirmation.attrs.get("fact_metadata")
    assert metadata is not None and len(metadata) == 2

    actions = main.RC.build_fefo_action_frame(
        processed,
        "2026-09",
        business_date=BUSINESS_DATE,
    )
    output = tmp_path / "fact-control-rows.xlsx"
    summary = {
        "Thang bao cao": "2026-09",
        "Ngay nghiep vu": BUSINESS_DATE,
        "So dong CTKM thieu vi tri ke": 0,
        "So item CTKM thieu vi tri ke": 0,
        "So dong POG khong co CTKM phu hop": 0,
        "So item POG khong co CTKM phu hop": 0,
    }

    main.EX.export_final_report(
        output_path=output,
        df_matched=pd.DataFrame(),
        df_ctkm_only=pd.DataFrame(),
        df_pog_only=pd.DataFrame(),
        overlap_warnings=[],
        mapping_logs=[],
        reconciliation_summary=summary,
        df_fefo_actions=actions,
        df_ctkm_clearance_confirmation=confirmation,
        df_pog_review=pd.DataFrame(),
    )

    workbook = load_workbook(output, data_only=True)

    review_values = [
        cell.value for row in workbook["Cần review"].iter_rows() for cell in row
    ]
    assert "Missing-code product" in review_values
    assert 'Combo "Mẹ tròn - Con vuông" Total' not in review_values
    assert "Grand Total" not in review_values

    confirmation_values = [
        cell.value for row in workbook["Xác nhận CTKM Clearance"].iter_rows() for cell in row
    ]
    assert "Valid product" in confirmation_values
    assert "Missing-code product" in confirmation_values
    assert 'Combo "Mẹ tròn - Con vuông" Total' not in confirmation_values
    assert "Grand Total" not in confirmation_values

    assert "KM metadata" in workbook.sheetnames
    audit = workbook["KM metadata"]
    assert audit.sheet_state == "hidden"
    audit_values = [cell.value for row in audit.iter_rows() for cell in row]
    assert "CTKM:Combo detail:4" in audit_values
    assert "CTKM:Combo detail:5" in audit_values
    assert 'Combo "Mẹ tròn - Con vuông" Total' in audit_values
    assert "Grand Total" in audit_values
    assert "subtotal" in audit_values
    assert "grand_total" in audit_values
    assert 1542000 in audit_values
    assert 5407000 in audit_values
