from copy import deepcopy
import asyncio
import math
import pytest
from pathlib import Path

from openpyxl import load_workbook
from fastapi import HTTPException

from backend import main_api
from backend.scheduling_contract import export_schedule_to_template




def _style_signature(cell):
    def color(value):
        if value is None:
            return None
        return (value.type, value.rgb, value.indexed, value.theme, value.tint)

    def side(value):
        if value is None:
            return None
        return (value.style, color(value.color))

    font = cell.font
    fill = cell.fill
    border = cell.border
    alignment = cell.alignment
    protection = cell.protection
    return (
        (font.name, font.sz, font.b, font.i, font.u, font.strike, color(font.color), font.vertAlign, font.scheme),
        (fill.fill_type, color(fill.fgColor), color(fill.bgColor)),
        (
            side(border.left), side(border.right), side(border.top), side(border.bottom),
            side(border.diagonal), border.diagonalUp, border.diagonalDown,
        ),
        (
            alignment.horizontal, alignment.vertical, alignment.text_rotation,
            alignment.wrap_text, alignment.shrink_to_fit, alignment.indent, alignment.readingOrder,
        ),
        cell.number_format,
        (protection.locked, protection.hidden),
    )


def _route(path, method):
    matches = [
        route for route in main_api.app.routes
        if getattr(route, "path", None) == path
        and method in (getattr(route, "methods", set()) or set())
    ]
    assert len(matches) == 1
    return matches[0]


def _config():
    employees = [
        {"id": "phuong", "name": "Phương", "staff_code": "800603", "role": "STORE_MANAGER", "active": True},
        {"id": "mi", "name": "My", "staff_code": "825982", "role": "EMPLOYEE", "active": True},
        {"id": "nhan", "name": "Nhàn", "staff_code": "827370", "role": "EMPLOYEE", "active": True},
        {"id": "vu", "name": "Vũ", "staff_code": "825433", "role": "EMPLOYEE", "active": True},
        {"id": "ly", "name": "Ly", "staff_code": "827389", "role": "EMPLOYEE", "active": True},
    ]
    days = [f"2026-09-{day:02d}" for day in range(14, 21)]
    return {
        "store_label": "1885",
        "week_start": days[0],
        "employees": employees,
        "staffing": [{"date": day, "morning": 3 if i == 0 else 2, "afternoon": 2 if i == 0 else 3} for i, day in enumerate(days)],
        "receiving_days": [],
        "requests": [],
        "fixed_weekday_rules": [
            {"employee_id": "phuong", "weekday": 0, "shift": "MORNING", "level": "HARD"},
            {"employee_id": "phuong", "weekday": 1, "shift": "AFTERNOON", "level": "HARD"},
        ],
        "employee_rules": [],
        "pairing_rules": [],
        "morning_time_overrides": [],
    }


def _assignments():
    schedule = {
        "2026-09-14": (["phuong", "mi", "nhan"], ["vu", "ly"]),
        "2026-09-15": (["mi", "vu"], ["phuong", "nhan", "ly"]),
        "2026-09-16": (["phuong", "ly"], ["mi", "nhan", "vu"]),
        "2026-09-17": (["mi", "vu"], ["phuong", "nhan", "ly"]),
        "2026-09-18": (["phuong", "ly"], ["mi", "nhan", "vu"]),
        "2026-09-19": (["mi", "vu"], ["phuong", "nhan", "ly"]),
        "2026-09-20": (["phuong", "ly"], ["mi", "nhan", "vu"]),
    }
    result = []
    all_ids = ["phuong", "mi", "nhan", "vu", "ly"]
    for date, (morning, afternoon) in schedule.items():
        for employee_id in all_ids:
            if employee_id in morning:
                start = "07:00" if date == "2026-09-14" and employee_id == "phuong" else "06:00"
                result.append({"employee_id": employee_id, "date": date, "shift": "MORNING", "start_time": start, "receiving": False})
            elif employee_id in afternoon:
                result.append({"employee_id": employee_id, "date": date, "shift": "AFTERNOON", "start_time": "14:00", "receiving": False})
            else:
                result.append({"employee_id": employee_id, "date": date, "shift": "OFF", "start_time": None, "receiving": False})
    return result


def test_scheduling_routes_are_installed_before_static_frontend_and_require_admin():
    assert main_api.app.state.staff_scheduling_contract_v1 is True
    paths = [
        ("/api/schedules/generate", "POST"),
        ("/api/schedules/validate", "POST"),
        ("/api/schedules", "POST"),
        ("/api/schedules", "GET"),
        ("/api/schedules/storage", "GET"),
        ("/api/schedules/{schedule_id}", "GET"),
        ("/api/schedules/{schedule_id}", "PATCH"),
        ("/api/schedules/{schedule_id}/finalize", "POST"),
        ("/api/schedules/{schedule_id}/export.xlsx", "GET"),
    ]
    routes = list(main_api.app.routes)
    mount_index = next(i for i, route in enumerate(routes) if route.__class__.__name__ == "Mount" and getattr(route, "name", None) == "frontend") if any(route.__class__.__name__ == "Mount" and getattr(route, "name", None) == "frontend" for route in routes) else len(routes)
    for path, method in paths:
        route = _route(path, method)
        dependencies = [getattr(dep, "call", None) for dep in route.dependant.dependencies]
        assert main_api.require_admin in dependencies
        assert routes.index(route) < mount_index
    assert _route("/schedule", "GET")


def test_exact_template_export_preserves_layout_and_replaces_only_schedule_values(tmp_path, monkeypatch):
    source_template = Path(__file__).resolve().parents[1] / "backend" / "templates" / "1885_Lich_lam_viec_template.xlsx"
    original = load_workbook(source_template)
    original_ws = original[original.sheetnames[0]]
    original_styles = {(row, col): _style_signature(original_ws.cell(row, col)) for row in range(1, 11) for col in range(1, 9)}
    original_merges = sorted(str(item) for item in original_ws.merged_cells.ranges)
    original_widths = {key: value.width for key, value in original_ws.column_dimensions.items()}
    original_heights = {key: value.height for key, value in original_ws.row_dimensions.items()}
    original.close()

    monkeypatch.setattr(main_api, "_ROOT", tmp_path)
    item = {"id": "test-schedule", "version": 1, "config": _config(), "assignments": _assignments()}
    path, filename = export_schedule_to_template(main_api, item)

    assert filename.endswith(".xlsx")
    workbook = load_workbook(path)
    try:
        ws = workbook[workbook.sheetnames[0]]
        assert ws["A1"].value == "Lịch làm việc 1885 ngày 14/09/26 đến 20/09/26"
        assert [ws.cell(2, col).value for col in range(2, 9)] == [14, 15, 16, 17, 18, 19, 20]
        assert [ws.cell(row, 1).value for row in range(4, 11)] == ["6H-14H", "6H-14H", "7H-15H", "14H-22H", "14H-22H", "14H-22H", "OFF"]
        assert ws["B6"].value == "Phương - 800603"
        assert ws["B4"].value in {"My - 825982", "Nhàn - 827370"}
        assert sorted(str(item) for item in ws.merged_cells.ranges) == original_merges
        assert {(row, col): _style_signature(ws.cell(row, col)) for row in range(1, 11) for col in range(1, 9)} == original_styles
        exported_widths = {key: value.width for key, value in ws.column_dimensions.items()}
        exported_heights = {key: value.height for key, value in ws.row_dimensions.items()}
        assert exported_widths.keys() == original_widths.keys()
        assert exported_heights.keys() == original_heights.keys()
        for key, expected in original_widths.items():
            assert math.isclose(exported_widths[key], expected, rel_tol=0.0, abs_tol=1e-12)
        for key, expected in original_heights.items():
            if expected is None:
                assert exported_heights[key] is None
            else:
                assert math.isclose(exported_heights[key], expected, rel_tol=0.0, abs_tol=1e-12)
        # Template content outside the scheduling table is deliberately untouched.
        assert ws["E17"].value == "s"
    finally:
        workbook.close()


def test_export_uses_saved_assignments_without_rerunning_optimizer(tmp_path, monkeypatch):
    monkeypatch.setattr(main_api, "_ROOT", tmp_path)
    assignments = _assignments()
    edited = deepcopy(assignments)
    # Manual Sunday swap: Mi moves to morning and Phương moves to afternoon.
    # This stays valid but is intentionally chosen here, not by the optimizer.
    for row in edited:
        if row["date"] != "2026-09-20":
            continue
        if row["employee_id"] == "mi":
            row.update(shift="MORNING", start_time="06:00", receiving=False)
        elif row["employee_id"] == "phuong":
            row.update(shift="AFTERNOON", start_time="14:00", receiving=False)
    item = {"id": "manual", "version": 4, "config": _config(), "assignments": edited}

    path, _ = export_schedule_to_template(main_api, item)

    workbook = load_workbook(path, data_only=False)
    try:
        ws = workbook[workbook.sheetnames[0]]
        sunday_morning = {ws["H4"].value, ws["H5"].value}
        sunday_afternoon = {ws["H7"].value, ws["H8"].value, ws["H9"].value}
        assert sunday_morning == {"My - 825982", "Ly - 827389"}
        assert "Phương - 800603" in sunday_afternoon
    finally:
        workbook.close()

def test_schedule_storage_reports_local_ephemeral_by_default(tmp_path, monkeypatch):
    monkeypatch.setattr(main_api, "_ROOT", tmp_path)
    monkeypatch.delenv("SCHEDULE_GCS_BUCKET", raising=False)
    monkeypatch.delenv("SCHEDULE_STORE_PATH", raising=False)
    monkeypatch.delenv("SCHEDULE_STORE_DURABLE", raising=False)

    info = asyncio.run(main_api.schedule_storage("admin"))

    assert info["mode"] == "local"
    assert info["durable"] is False
    assert "ephemeral" in info["warning"]


def test_schedule_update_rejects_stale_version_and_final_requires_explicit_reopen(tmp_path, monkeypatch):
    monkeypatch.setattr(main_api, "_ROOT", tmp_path)
    monkeypatch.delenv("SCHEDULE_GCS_BUCKET", raising=False)
    monkeypatch.setenv("SCHEDULE_STORE_PATH", str(tmp_path / "schedules.json"))

    created = asyncio.run(main_api.schedule_create({
        "config": _config(),
        "assignments": _assignments(),
        "variant": "manual",
    }, "admin"))
    schedule_id = created["id"]
    assert created["version"] == 1

    updated = asyncio.run(main_api.schedule_update(schedule_id, {
        "config": _config(),
        "assignments": _assignments(),
        "variant": "manual",
        "expected_version": 1,
    }, "admin"))
    assert updated["version"] == 2

    with pytest.raises(HTTPException) as stale:
        asyncio.run(main_api.schedule_update(schedule_id, {
            "config": _config(),
            "assignments": _assignments(),
            "expected_version": 1,
        }, "admin"))
    assert stale.value.status_code == 409

    final = asyncio.run(main_api.schedule_finalize(schedule_id, {"expected_version": 2}, "admin"))
    assert final["status"] == "FINAL"
    assert final["version"] == 3

    with pytest.raises(HTTPException) as locked:
        asyncio.run(main_api.schedule_update(schedule_id, {
            "config": _config(),
            "assignments": _assignments(),
            "expected_version": 3,
        }, "admin"))
    assert locked.value.status_code == 409
    assert "FINAL" in str(locked.value.detail)

    reopened = asyncio.run(main_api.schedule_update(schedule_id, {
        "config": _config(),
        "assignments": _assignments(),
        "expected_version": 3,
        "reopen": True,
    }, "admin"))
    assert reopened["status"] == "DRAFT"
    assert reopened["version"] == 4

