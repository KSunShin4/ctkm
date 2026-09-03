from openpyxl import load_workbook
import pandas as pd

import main


BUSINESS_DATE = pd.Timestamp("2026-09-15")


def _promotion(item_id):
    return {
        "item_id": item_id,
        "item_name": f"Product {item_id}",
        "promo_uom": "Hộp",
        "campaign_name": "September",
        "campaign_detail": "Giảm 10%",
        "promo_from": pd.Timestamp("2026-09-01"),
        "promo_to": pd.Timestamp("2026-09-30"),
        "normal_price": 100000,
        "promo_price": 90000,
        "program_source": "CTKM",
        "source_sheet": "Monthly",
    }


def _pog(item_id, *, pog_code="A", location_id="1", snapshot="2026-09"):
    return {
        "item_id": item_id,
        "store_id": "S1",
        "pog_code": pog_code,
        "location_id": location_id,
        "facing": 1,
        "pog_snapshot_month": snapshot,
    }


def test_location_state_axes_do_not_reclassify_calendar_action():
    promotions = pd.DataFrame([
        _promotion("P1001"),  # one complete current-month position
        _promotion("P2001"),  # two complete current-month positions
        _promotion("P3001"),  # complete but stale snapshot
        _promotion("P4001"),  # complete but unknown snapshot month
        _promotion("P5001"),  # partial coordinates
        _promotion("P6001"),  # no POG match
    ])
    pog = pd.DataFrame([
        _pog("P1001"),
        _pog("P2001", pog_code="B", location_id="1"),
        _pog("P2001", pog_code="C", location_id="2"),
        _pog("P3001", snapshot="2026-08"),
        _pog("P4001", snapshot=None),
        _pog("P5001", pog_code=None, location_id="3"),
    ])

    actions = main.RC.build_fefo_action_frame(
        promotions,
        "2026-09",
        df_pog=pog,
        business_date=BUSINESS_DATE,
    )

    # Calendar action is orthogonal to POG quality. Every promotion is active.
    assert set(actions["scheduled_action"]) == {"gắn tem"}
    assert set(actions["action"]) == {"gắn tem"}

    by_item = {item: group.reset_index(drop=True) for item, group in actions.groupby("item_id")}

    current = by_item["P1001"].iloc[0]
    assert current["location_completeness"] == "Đủ vị trí"
    assert current["pog_month_status"] == "POG đúng tháng"
    assert current["location_needs_review"] is False or current["location_needs_review"] == False
    assert current["location_status"] == "Đủ vị trí"

    multiple = by_item["P2001"]
    assert len(multiple) == 2
    assert set(multiple["location_completeness"]) == {"Đủ vị trí"}
    assert set(multiple["pog_month_status"]) == {"POG đúng tháng"}
    assert set(multiple["location_needs_review"]) == {False}
    assert set(multiple["location_status"]) == {"Nhiều vị trí POG"}
    assert multiple["task_id"].nunique() == 2

    stale = by_item["P3001"].iloc[0]
    assert stale["location_completeness"] == "Đủ vị trí"
    assert stale["pog_month_status"] == "POG lệch tháng"
    assert bool(stale["location_needs_review"]) is True
    assert stale["location_status"] == "POG lệch tháng"

    unknown = by_item["P4001"].iloc[0]
    assert unknown["location_completeness"] == "Đủ vị trí"
    assert unknown["pog_month_status"] == "Không rõ tháng POG"
    assert bool(unknown["location_needs_review"]) is True
    # The legacy one-field status must never report a false-safe "Đủ vị trí".
    assert unknown["location_status"] == "Không rõ tháng POG"

    partial = by_item["P5001"].iloc[0]
    assert partial["location_completeness"] == "Thiếu một phần thông tin vị trí"
    assert bool(partial["location_needs_review"]) is True

    missing = by_item["P6001"].iloc[0]
    assert missing["location_completeness"] == "Không tìm thấy SKU trong POG"
    assert missing["pog_month_status"] == "Không có snapshot POG"
    assert bool(missing["location_needs_review"]) is True
    assert missing["location_status"] == "Không tìm thấy SKU trong POG"


def test_export_routes_location_review_without_turning_it_into_business_review(tmp_path):
    promotions = pd.DataFrame([
        _promotion("P1001"),
        _promotion("P2001"),
        _promotion("P3001"),
        _promotion("P4001"),
        _promotion("P5001"),
        _promotion("P6001"),
    ])
    pog = pd.DataFrame([
        _pog("P1001"),
        _pog("P2001", pog_code="B", location_id="1"),
        _pog("P2001", pog_code="C", location_id="2"),
        _pog("P3001", snapshot="2026-08"),
        _pog("P4001", snapshot=None),
        _pog("P5001", pog_code=None, location_id="3"),
    ])
    actions = main.RC.build_fefo_action_frame(
        promotions,
        "2026-09",
        df_pog=pog,
        business_date=BUSINESS_DATE,
    )

    output = tmp_path / "location-state-contract.xlsx"
    summary = {
        "Thang bao cao": "2026-09",
        "Ngay nghiep vu": BUSINESS_DATE,
    }
    main.EX.export_final_report(
        output_path=output,
        df_matched=pd.DataFrame(),
        df_ctkm_only=pd.DataFrame(),
        df_pog_only=pd.DataFrame(),
        overlap_warnings=[],
        mapping_logs=[],
        reconciliation_summary=summary,
        df_fefo_actions=actions,
    )

    workbook = load_workbook(output, data_only=True)

    # Seven physical tasks: P2001 legitimately has two POG positions.
    gan = workbook["Gắn tem"]
    assert gan.max_row == 8

    gan_values = [cell.value for row in gan.iter_rows() for cell in row]
    for item_id in ("P1001", "P2001", "P3001", "P4001", "P5001", "P6001"):
        assert item_id in gan_values

    # POG-quality review is a separate axis, not the business/date Cần review sheet.
    business_review_values = [
        cell.value for row in workbook["Cần review"].iter_rows() for cell in row
    ]
    for item_id in ("P3001", "P4001", "P5001", "P6001"):
        assert item_id not in business_review_values

    location_review = workbook["POG cần kiểm tra"]
    location_review_values = [cell.value for row in location_review.iter_rows() for cell in row]
    for item_id in ("P3001", "P4001", "P5001", "P6001"):
        assert item_id in location_review_values
    assert "P1001" not in location_review_values
    assert "P2001" not in location_review_values

    location_issue_values = [
        cell.value for row in workbook["CTKM thieu vi tri ke"].iter_rows() for cell in row
    ]
    assert "P5001" in location_issue_values
    assert "P6001" in location_issue_values
    assert "P3001" not in location_issue_values
    assert "P4001" not in location_issue_values

    # Snapshot quality is visible next to location status and highlighted on the
    # operational sheet so staff cannot miss a task that needs POG verification.
    headers = [cell.value for cell in gan[1]]
    item_header = main.EX._BILINGUAL_HEADERS["item_id"]
    month_header = main.EX._BILINGUAL_HEADERS["pog_month_status"]
    review_header = main.EX._BILINGUAL_HEADERS["location_needs_review"]
    item_col = headers.index(item_header) + 1
    month_col = headers.index(month_header) + 1
    review_col = headers.index(review_header) + 1

    p4001_row = next(
        row for row in range(2, gan.max_row + 1)
        if gan.cell(row=row, column=item_col).value == "P4001"
    )
    assert gan.cell(row=p4001_row, column=month_col).value == "Không rõ tháng POG"
    assert gan.cell(row=p4001_row, column=review_col).value is True
    assert (gan.cell(row=p4001_row, column=month_col).fill.fgColor.rgb or "").endswith("FFC7CE")
    assert (gan.cell(row=p4001_row, column=review_col).fill.fgColor.rgb or "").endswith("FFC7CE")
