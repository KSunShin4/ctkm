from openpyxl import load_workbook
import pandas as pd

import main


BUSINESS_DATE = pd.Timestamp("2026-09-20")


def _promo(*, key_channel, channel, campaign, detail, price, start="2026-09-01", source_row_id=None, item_name="Herbal shampoo"):
    row = {
        "item_id": "P28944",
        "item_name": item_name,
        "promo_uom": "Chai",
        "key_channel": key_channel,
        "channel": channel,
        "campaign_name": campaign,
        "campaign_detail": detail,
        "promo_from": pd.Timestamp(start),
        "promo_to": pd.Timestamp("2026-09-30"),
        "normal_price": 189000,
        "promo_price": price,
        "product_type": "H&B",
        "program_source": "CTKM",
        "source_sheet": "Monthly",
    }
    if source_row_id is not None:
        row["source_row_id"] = source_row_id
    return row


def _pog():
    return pd.DataFrame([{
        "item_id": "P28944",
        "store_id": "S1",
        "pog_code": "HAIR1.N",
        "location_id": "7",
        "facing": 2,
        "pog_snapshot_month": "2026-09",
        "source_row_id": "POG:Export:100",
    }])


def _realistic_cross_scope_promotions():
    # Mirrors the real September Monthly pattern for P28944:
    # an All-channel regular discount starts 16 Sep while an Offline PWP runs all month.
    return pd.DataFrame([
        _promo(
            key_channel="All channel",
            channel="All Channel",
            campaign="Full month",
            detail="Discount 25%",
            price=141750,
            start="2026-09-16",
            source_row_id="CTKM:Monthly:150",
        ),
        _promo(
            key_channel="Offline",
            channel="Offline",
            campaign="PWP",
            detail="Customer purchase bill 150k, can buy product at discount 50%",
            price=94500,
            start="2026-09-01",
            source_row_id="CTKM:Monthly:151",
        ),
    ])


def test_distinct_cross_scope_promotions_keep_distinct_tasks_and_flag_same_position():
    promotions = _realistic_cross_scope_promotions()
    warnings = main.RC.detect_campaign_time_overlaps(promotions)

    # Existing reviewed contract: different channel scopes warn but do not block.
    assert len(warnings) == 1
    assert warnings[0].conflict_type == "potential_cross_channel_overlap"
    assert warnings[0].is_blocking is False

    actions = main.RC.build_fefo_action_frame(
        promotions,
        "2026-09",
        df_pog=_pog(),
        business_date=BUSINESS_DATE,
        overlap_warnings=warnings,
    )

    assert len(actions) == 2
    assert set(actions["scheduled_action"]) == {"gắn tem"}
    assert actions["promotion_id"].nunique() == 2
    assert actions["task_id"].nunique() == 2
    assert actions["pog_position_id"].nunique() == 1
    assert set(actions["promotion_count_at_position"]) == {2}
    assert set(actions["has_multiple_promotions_at_position"]) == {True}

    # Raw scope values remain source-faithful and visible; diagnostics are derived only.
    assert set(actions["key_channel"]) == {"All channel", "Offline"}
    assert set(actions["channel"]) == {"All Channel", "Offline"}
    assert set(actions["promo_price"]) == {141750, 94500}


def test_same_promotion_from_multiple_source_rows_counts_once_at_position():
    promotions = pd.DataFrame([
        _promo(
            key_channel="All channel", channel="All Channel",
            campaign="Full month", detail="Discount 25%", price=141750,
            start="2026-09-16", source_row_id="CTKM:Monthly:150",
            item_name="Herbal shampoo",
        ),
        _promo(
            key_channel="All channel", channel="All Channel",
            campaign="Full month", detail="Discount 25%", price=141750,
            start="2026-09-16", source_row_id="CTKM:Monthly:999",
            item_name="Herbal shampoo renamed",
        ),
    ])

    actions = main.RC.build_fefo_action_frame(
        promotions,
        "2026-09",
        df_pog=_pog(),
        business_date=BUSINESS_DATE,
        overlap_warnings=[],
    )

    # Descriptive source differences preserve both audit rows but share one stable promotion/task.
    assert len(actions) == 2
    assert actions["promotion_id"].nunique() == 1
    assert actions["task_id"].nunique() == 1
    assert actions["pog_position_id"].nunique() == 1
    assert set(actions["promotion_count_at_position"]) == {1}
    assert set(actions["has_multiple_promotions_at_position"]) == {False}


def test_missing_pog_does_not_create_a_fake_same_position_collision():
    actions = main.RC.build_fefo_action_frame(
        _realistic_cross_scope_promotions(),
        "2026-09",
        df_pog=pd.DataFrame(columns=["item_id", "store_id", "pog_code", "location_id"]),
        business_date=BUSINESS_DATE,
        overlap_warnings=[],
    )

    assert set(actions["location_completeness"]) == {"Không tìm thấy SKU trong POG"}
    assert set(actions["promotion_count_at_position"]) == {0}
    assert set(actions["has_multiple_promotions_at_position"]) == {False}


def test_action_sheet_surfaces_same_position_warning_without_turning_it_into_review(tmp_path):
    promotions = _realistic_cross_scope_promotions()
    warnings = main.RC.detect_campaign_time_overlaps(promotions)
    actions = main.RC.build_fefo_action_frame(
        promotions,
        "2026-09",
        df_pog=_pog(),
        business_date=BUSINESS_DATE,
        overlap_warnings=warnings,
    )
    output = tmp_path / "same-position-promotions.xlsx"
    summary = {
        "Thang bao cao": "2026-09",
        "Ngay nghiep vu": BUSINESS_DATE,
        "So dong CTKM thieu vi tri ke": 0,
        "So item CTKM thieu vi tri ke": 0,
        "So dong POG khong co CTKM phu hop": 0,
        "So item POG khong co CTKM phu hop": 0,
    }

    main.EX.export_final_report(
        output_path=output,
        df_matched=pd.DataFrame(),
        df_ctkm_only=pd.DataFrame(),
        df_pog_only=pd.DataFrame(),
        overlap_warnings=warnings,
        mapping_logs=[],
        reconciliation_summary=summary,
        df_fefo_actions=actions,
        df_pog_review=pd.DataFrame(),
    )

    workbook = load_workbook(output, data_only=True)
    action = workbook["Gắn tem"]
    review = workbook["Cần review"]

    headers = [cell.value for cell in action[1]]
    count_header = main.EX._BILINGUAL_HEADERS["promotion_count_at_position"]
    flag_header = main.EX._BILINGUAL_HEADERS["has_multiple_promotions_at_position"]
    assert count_header in headers
    assert flag_header in headers

    count_col = headers.index(count_header) + 1
    flag_col = headers.index(flag_header) + 1
    assert action.max_row == 3
    for row in range(2, 4):
        assert action.cell(row=row, column=count_col).value == 2
        assert action.cell(row=row, column=flag_col).value is True
        # Yellow diagnostic, not red error/review styling.
        assert action.cell(row=row, column=flag_col).fill.fgColor.rgb.endswith("FFF2CC")

    review_values = [cell.value for row in review.iter_rows() for cell in row]
    assert "P28944" not in review_values
