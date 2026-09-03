from types import SimpleNamespace

import pandas as pd
from openpyxl import Workbook, load_workbook

from backend import export_memory_contract as contract
from modules.facade_contracts import run_post_facade_hooks


def test_combined_postprocess_opens_workbook_once_and_preserves_metadata(tmp_path, monkeypatch):
    path = tmp_path / "report.xlsx"
    workbook = Workbook()
    guide = workbook.active
    guide.title = "HƯỚNG DẪN"
    guide["A6"] = "Thiếu vị trí"
    summary = workbook.create_sheet("Tom tat")
    summary.append(["Chi so", "Gia tri"])
    summary.append(["So dong CTKM thieu vi tri ke", 2])
    workbook.save(path)

    pog_metadata = pd.DataFrame([
        {"ITEMID": "Total", "pog_control_kind": "summary_total"},
    ])
    fact_metadata = pd.DataFrame([
        {"item_id": "Combo Total", "fact_control_kind": "summary_total"},
    ])

    real_load = contract.load_workbook
    calls = []

    def counted_load(*args, **kwargs):
        calls.append(1)
        return real_load(*args, **kwargs)

    monkeypatch.setattr(contract, "load_workbook", counted_load)
    contract._postprocess_workbook(
        path, pog_metadata=pog_metadata, fact_metadata=fact_metadata
    )

    assert len(calls) == 1
    result = load_workbook(path)
    assert result["POG metadata"].sheet_state == "hidden"
    assert result["KM metadata"].sheet_state == "hidden"
    assert result["Tom tat"].row_dimensions[2].hidden is True
    assert result["HƯỚNG DẪN"]["A6"].value == "Thiếu vị trí KM"
    result.close()


def test_export_memory_hook_is_idempotent_and_survives_facade_rebuild(monkeypatch):
    def base_export(*_args, **_kwargs):
        return "report.xlsx"

    exporter = SimpleNamespace(
        export_final_report=base_export,
        _append_pog_metadata_audit=lambda *_args, **_kwargs: None,
    )
    main_core = SimpleNamespace(EX=exporter)
    facade = SimpleNamespace(_core=main_core, EX=exporter)
    api_core = SimpleNamespace(
        _main_mod=facade,
        app=SimpleNamespace(state=SimpleNamespace()),
    )

    postprocess_calls = []
    monkeypatch.setattr(
        contract,
        "_postprocess_workbook",
        lambda output, **kwargs: postprocess_calls.append((output, kwargs)) or output,
    )
    monkeypatch.setattr(contract, "_release_heap", lambda: None)
    monkeypatch.setattr(contract, "_log_memory", lambda _stage: None)

    contract.install_export_memory_contract(api_core)
    first_wrapper = exporter.export_final_report
    assert getattr(first_wrapper, "_ctkm_export_memory_contract_v1", False)
    assert api_core.app.state.export_memory_contract_v1 is True

    contract.install_export_memory_contract(api_core)
    assert exporter.export_final_report is first_wrapper

    assert exporter.export_final_report("report.xlsx") == "report.xlsx"
    assert len(postprocess_calls) == 1

    # Simulate main.py rebuilding the raw facade surface under another alias.
    exporter.export_final_report = base_export
    run_post_facade_hooks(main_core, facade)
    assert getattr(exporter.export_final_report, "_ctkm_export_memory_contract_v1", False)
    assert exporter.export_final_report is not first_wrapper
