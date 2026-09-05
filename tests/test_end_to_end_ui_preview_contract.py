from openpyxl import Workbook

from backend.preview_contract import build_visible_preview


def test_preview_treats_no_overlap_note_as_zero_and_hides_nondata_sheets(tmp_path):
    path = tmp_path / "preview.xlsx"
    wb = Workbook()
    guide = wb.active
    guide.title = "HƯỚNG DẪN"
    guide["A1"] = "Guide"

    warning = wb.create_sheet("Canh bao trung thoi gian")
    warning.append(["Ghi chu"])
    warning.append(["Khong phat hien campaign nao trung thoi gian."])

    action = wb.create_sheet("Gắn tem")
    action.append(["item_id", "task_id"])
    action.append(["P1001", "TASK-1"])
    action.column_dimensions["B"].hidden = True

    metadata = wb.create_sheet("POG metadata")
    metadata.append(["raw"])
    metadata.append(["Total"])
    metadata.sheet_state = "hidden"
    wb.save(path)
    wb.close()

    preview = build_visible_preview(path, rows=100)["sheets"]
    assert "HƯỚNG DẪN" not in preview
    assert "POG metadata" not in preview
    assert preview["Canh bao trung thoi gian"]["total_rows"] == 0
    assert preview["Canh bao trung thoi gian"]["rows"] == []
    assert preview["Gắn tem"]["headers"] == ["item_id"]
    assert preview["Gắn tem"]["rows"] == [["P1001"]]
    assert preview["Gắn tem"]["total_rows"] == 1


def test_frontend_compatibility_text_matches_current_backend_contracts():
    source = open("frontend/src/main.jsx", encoding="utf-8").read()
    assert "Có cảnh báo mapping. Hãy kiểm tra mapping bên dưới." in source
    assert "Hệ thống đang kiểm tra mapping cột" in source
    assert "Workbook CTKM tháng" in source
    assert "POG không có KM phù hợp" in source
    assert "Tối đa 100 dòng hiển thị mỗi sheet" in source
