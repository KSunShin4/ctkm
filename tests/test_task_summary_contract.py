from openpyxl import load_workbook
import pandas as pd

import main


def _action(task_id, source_row_id, *, scheduled_action, program_source="CTKM", in_scope=True):
    return {
        "task_id": task_id,
        "promotion_id": f"PROMO-{task_id}",
        "pog_position_id": "POGPOS-1",
        "item_id": "P1001",
        "item_key": "P1001",
        "program_source": program_source,
        "scheduled_action": scheduled_action,
        "action": scheduled_action,
        "in_report_scope": in_scope,
        "location_completeness": "Đủ vị trí",
        "location_status": "Đủ vị trí",
        "source_row_id": source_row_id,
        "source_row_ids": source_row_id,
        "source_row_count": 1,
        "duplicate_status": "Dòng duy nhất",
    }


def _summary():
    return {
        "Thang bao cao": "2026-09",
        "Ngay nghiep vu": pd.Timestamp("2026-09-15"),
        # Deliberately wrong raw-row counts. Export adapter must replace only
        # task-oriented metrics with the same unique-task semantics as work sheets.
        "So dong tem tu Clearance": 99,
        "So dong cho gan": 99,
        "So dong can review": 99,
        # Audit/location metrics keep their existing row semantics.
        "So dong CTKM thieu vi tri ke": 7,
        "So dong can kiem tra vi tri": 9,
    }


def test_operational_summary_counts_unique_in_scope_tasks_only():
    actions = pd.DataFrame([
        _action("TASK-WAIT", "CLEARANCE:Monthly:10", scheduled_action="chờ gắn tem", program_source="CLEARANCE"),
        _action("TASK-WAIT", "CLEARANCE:Monthly:20", scheduled_action="chờ gắn tem", program_source="CLEARANCE"),
        _action("TASK-REVIEW", "CTKM:Monthly:30", scheduled_action="review"),
        _action("TASK-REVIEW", "CTKM:Monthly:40", scheduled_action="review"),
        _action("TASK-OUT", "CLEARANCE:Monthly:50", scheduled_action="chờ gắn tem", program_source="CLEARANCE", in_scope=False),
    ])
    summary = _summary()

    main._performance.align_task_summary_with_work_rows(summary, actions)

    assert summary["So dong tem tu Clearance"] == 1
    assert summary["So dong cho gan"] == 1
    assert summary["So dong can review"] == 1
    assert summary["So dong CTKM thieu vi tri ke"] == 7
    assert summary["So dong can kiem tra vi tri"] == 9


def test_export_mutates_pipeline_summary_before_writing_tom_tat(tmp_path):
    actions = pd.DataFrame([
        _action("TASK-WAIT", "CLEARANCE:Monthly:10", scheduled_action="chờ gắn tem", program_source="CLEARANCE"),
        _action("TASK-WAIT", "CLEARANCE:Monthly:20", scheduled_action="chờ gắn tem", program_source="CLEARANCE"),
        _action("TASK-REVIEW", "CTKM:Monthly:30", scheduled_action="review"),
        _action("TASK-REVIEW", "CTKM:Monthly:40", scheduled_action="review"),
    ])
    summary = _summary()
    output = tmp_path / "summary-task-counts.xlsx"

    main.EX.export_final_report(
        output,
        pd.DataFrame(),
        pd.DataFrame(),
        pd.DataFrame(),
        [],
        [],
        summary,
        df_fefo_actions=actions,
    )

    # The same dict is returned by run_pipeline, so API/backend sees these values.
    assert summary["So dong tem tu Clearance"] == 1
    assert summary["So dong cho gan"] == 1
    assert summary["So dong can review"] == 1

    workbook = load_workbook(output, data_only=True)
    tom_tat = {
        workbook["Tom tat"].cell(row=row, column=1).value:
        workbook["Tom tat"].cell(row=row, column=2).value
        for row in range(2, workbook["Tom tat"].max_row + 1)
    }
    assert tom_tat["So dong tem tu Clearance"] == 1
    assert tom_tat["So dong cho gan"] == 1
    assert tom_tat["So dong can review"] == 1
    assert tom_tat["So dong CTKM thieu vi tri ke"] == 7
    assert tom_tat["So dong can kiem tra vi tri"] == 9

    # Workbook work sheets use the exact same task semantics as the summary.
    assert workbook["Chờ gắn"].max_row == 2
    assert workbook["Cần review"].max_row == 2
