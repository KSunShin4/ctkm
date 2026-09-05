from io import BytesIO
import importlib.util
from pathlib import Path
import sys


def _module():
    path = Path(__file__).resolve().parents[1] / "modules" / "11_expiration_checker.py"
    spec = importlib.util.spec_from_file_location("expiration_workbook_hardening", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _workbook_bytes():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(["ITEMID", "BATCHNUMER", "EXPDATE", "ITEMNAME", "ONHANDNEAREXPIRATION"])
    sheet.append(["P1", "LOT1", "31/12/2027", "One", 2])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_expiration_loader_disables_external_links(monkeypatch):
    module = _module()
    import openpyxl

    real_load_workbook = openpyxl.load_workbook
    calls = []

    def recording_load_workbook(*args, **kwargs):
        calls.append(dict(kwargs))
        return real_load_workbook(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", recording_load_workbook)

    frame, sheet_name = module.load_expiration_export(_workbook_bytes())

    assert sheet_name == "Export"
    assert frame.loc[0, "ITEMID"] == "P1"
    assert calls
    assert calls[0].get("keep_links") is False
    assert calls[0].get("read_only") is True
    assert calls[0].get("data_only") is True


def test_expiration_loader_rejects_fake_xlsx():
    module = _module()

    try:
        module.load_expiration_export(b"not an xlsx")
    except ValueError as exc:
        assert "không phải XLSX hợp lệ" in str(exc)
    else:
        raise AssertionError("fake XLSX should be rejected")
