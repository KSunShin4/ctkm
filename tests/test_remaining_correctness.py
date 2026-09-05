from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

import main


BUSINESS_DATE = pd.Timestamp("2026-08-15")


def _promo(item_id="P1001", **overrides):
    row = {
        "item_id": item_id,
        "item_name": "Product",
        "promo_uom": "Hộp",
        "campaign_name": "Campaign A",
        "campaign_detail": "Giảm 10%",
        "promo_from": pd.Timestamp("2026-08-01"),
        "promo_to": pd.Timestamp("2026-08-31"),
        "normal_price": 100_000,
        "promo_price": 90_000,
        "channel": "OFFLINE",
        "program_source": "CTKM",
        "source_sheet": "Monthly",
    }
    row.update(overrides)
    return row


def _pog(item_id="P1001", **overrides):
    row = {
        "item_id": item_id,
        "store_id": "S1",
        "pog_code": "A",
        "location_id": "1",
        "facing": 2,
        "pog_snapshot_month": "2026-08",
    }
    row.update(overrides)
    return row


def _actions(promotions, pog=None, warnings=None, business_date=BUSINESS_DATE):
    frame = pd.DataFrame(promotions)
    return main.RC.build_fefo_action_frame(
        frame,
        "2026-08",
        None if pog is None else pd.DataFrame(pog),
        business_date=business_date,
        overlap_warnings=warnings,
    )


def test_missing_item_and_campaign_are_preserved_with_priority_review_reasons():
    result = _actions([
        _promo(None, campaign_name=None, promo_from=pd.NaT),
        _promo("P2", campaign_name=None, promo_from=pd.NaT),
    ])

    assert len(result) == 2
    assert result["scheduled_action"].tolist() == ["review", "review"]
    assert result["reason"].tolist() == ["thieu_ma_san_pham", "thieu_thong_tin_ngay"]
    assert result["review_blocks_action"].astype(bool).tolist() == [True, True]


def test_normalizer_keeps_missing_item_and_creates_nonmatching_item_key():
    source = pd.DataFrame({"item_id": [None, " p1001 "], "campaign_name": ["Missing", "Valid"]})

    result = main.NM.clean_values(source)

    assert len(result) == 2
    assert pd.isna(result.loc[0, "item_key"])
    assert result.loc[1, "item_key"] == "P1001"
    assert result.loc[1, "item_id"] == " p1001 "


def test_null_ctkm_and_null_pog_keys_never_match_and_bad_pog_is_reviewable():
    joined = main.JN.join_ctkm_pog(
        pd.DataFrame([{"item_id": None, "item_key": None, "campaign_name": "Bad"}]),
        pd.DataFrame([{"item_id": None, "item_key": None, "store_id": "S1"}]),
    )

    assert joined.df_matched.empty
    assert len(joined.df_ctkm_only) == 1
    assert len(joined.df_pog_review) == 1
    assert joined.df_pog_review.loc[0, "reason"] == "thieu_ma_san_pham_pog"


def test_item_key_matches_case_and_spaces_but_keeps_original_ids():
    ctkm = main.NM.clean_values(pd.DataFrame([{"item_id": " p1001 ", "campaign_name": "A"}]))
    pog = main.NM.clean_values(pd.DataFrame([{"item_id": "P1001", "store_id": "S1"}]))

    joined = main.JN.join_ctkm_pog(ctkm, pog)

    assert len(joined.df_matched) == 1
    assert joined.df_matched.loc[0, "item_key"] == "P1001"
    assert joined.df_matched.loc[0, "item_id__ctkm"] == " p1001 "
    assert joined.df_matched.loc[0, "item_id__pog"] == "P1001"


def test_cross_channel_overlap_is_warning_only_and_preserves_scheduled_actions():
    frame = pd.DataFrame([
        _promo(campaign_name="A", channel="OFFLINE"),
        _promo(campaign_name="B", channel="WEBAPP"),
    ])
    warnings = main.RC.detect_campaign_time_overlaps(frame)
    result = _actions(frame.to_dict("records"), warnings=warnings)

    assert len(warnings) == 1
    assert warnings[0].conflict_type == "potential_cross_channel_overlap"
    assert warnings[0].is_blocking is False
    assert result["scheduled_action"].tolist() == ["gắn tem", "gắn tem"]


def test_same_program_different_channel_is_not_a_blocking_conflict():
    frame = pd.DataFrame([
        _promo(channel="OFFLINE"),
        _promo(channel="WEBAPP"),
    ])
    warnings = main.RC.detect_campaign_time_overlaps(frame)

    result = _actions(frame.to_dict("records"), warnings=warnings)

    assert len(warnings) == 1
    assert warnings[0].is_blocking is False
    assert result["channel"].tolist() == ["OFFLINE", "WEBAPP"]
    assert set(result["scheduled_action"]) == {"gắn tem"}


def test_same_channel_different_offer_keeps_action_and_marks_review_warning():
    frame = pd.DataFrame([
        _promo("P1", campaign_name="A", promo_price=90_000),
        _promo("P1", campaign_name="B", promo_price=80_000),
        _promo("P1", campaign_name="Later", promo_from=pd.Timestamp("2026-09-01"), promo_to=pd.Timestamp("2026-09-30")),
        _promo("P2", campaign_name="Other"),
    ])
    warnings = main.RC.detect_campaign_time_overlaps(frame)
    result = _actions(frame.to_dict("records"), warnings=warnings)

    assert len(warnings) == 1
    assert warnings[0].conflict_type == "blocking_conflict"
    assert warnings[0].is_blocking is True
    by_campaign = result.set_index("campaign_name")
    assert by_campaign["scheduled_action"].to_dict() == {
        "A": "gắn tem", "B": "gắn tem", "Later": "chờ gắn tem", "Other": "gắn tem"
    }
    assert bool(by_campaign.loc["A", "review_required"]) is True
    assert bool(by_campaign.loc["B", "review_required"]) is True
    assert bool(by_campaign.loc["A", "review_blocks_action"]) is False
    assert bool(by_campaign.loc["B", "review_blocks_action"]) is False


def test_same_channel_same_offer_overlapping_different_periods_keep_actions():
    frame = pd.DataFrame([
        _promo("P1", promo_from=pd.Timestamp("2026-08-01"), promo_to=pd.Timestamp("2026-08-20")),
        _promo("P1", promo_from=pd.Timestamp("2026-08-10"), promo_to=pd.Timestamp("2026-08-31")),
        _promo("P1", promo_from=pd.Timestamp("2026-09-01"), promo_to=pd.Timestamp("2026-09-30")),
        _promo("P2"),
    ])

    warnings = main.RC.detect_campaign_time_overlaps(frame)
    result = _actions(frame.to_dict("records"), warnings=warnings)

    assert len(warnings) == 1
    assert warnings[0].conflict_type == "blocking_conflict"
    assert warnings[0].is_blocking is True
    assert result["scheduled_action"].tolist() == ["gắn tem", "gắn tem", "chờ gắn tem", "gắn tem"]
    assert result["review_required"].astype(bool).tolist() == [True, True, False, False]
    assert not result.loc[:1, "review_blocks_action"].astype(bool).any()


def test_ctkm_clearance_overlap_requires_review_without_suppressing_action():
    frame = main._build_promo_action_source(
        pd.DataFrame([_promo(channel="OFFLINE")]),
        pd.DataFrame([_promo(campaign_name="Clearance", channel="WEBAPP")]),
    )
    warnings = main.RC.detect_campaign_time_overlaps(frame)
    result = _actions(frame.to_dict("records"), warnings=warnings)

    assert len(warnings) == 1
    assert warnings[0].is_blocking is True
    assert warnings[0].conflict_type == "blocking_conflict"
    assert set(result["scheduled_action"]) == {"gắn tem"}
    assert result["review_required"].astype(bool).all()
    assert not result["review_blocks_action"].astype(bool).any()


def test_promotion_and_task_ids_are_deterministic_under_input_reordering():
    promos = [_promo("P1", campaign_name="A"), _promo("P2", campaign_name="B")]
    pog = [_pog("P1", location_id="1"), _pog("P2", location_id="2")]

    first = _actions(promos, pog)
    second = _actions(list(reversed(promos)), list(reversed(pog)))

    first_ids = first.set_index("item_key")[["promotion_id", "task_id"]].to_dict("index")
    second_ids = second.set_index("item_key")[["promotion_id", "task_id"]].to_dict("index")
    assert first_ids == second_ids


def test_action_and_location_are_part_of_task_id_but_business_date_is_not():
    promotion = _promo()
    two_locations = [_pog(location_id="1"), _pog(location_id="2")]
    attach = _actions([promotion], two_locations, business_date=pd.Timestamp("2026-08-15"))
    same_action_other_day = _actions([promotion], two_locations, business_date=pd.Timestamp("2026-08-16"))
    remove = _actions([promotion], two_locations, business_date=pd.Timestamp("2026-09-01"))

    assert attach["promotion_id"].nunique() == 1
    assert attach["task_id"].nunique() == 2
    assert set(attach["task_id"]) == set(same_action_other_day["task_id"])
    assert set(attach["task_id"]).isdisjoint(set(remove["task_id"]))


def test_exact_duplicate_pog_creates_one_task_with_pog_lineage():
    duplicated = _pog(source_row_id="POG:Sheet:6")
    second = _pog(source_row_id="POG:Sheet:9")

    result = _actions([_promo()], [duplicated, second])

    assert len(result) == 1
    assert result.loc[0, "pog_source_row_count"] == 2
    assert result.loc[0, "pog_source_row_ids"] == "POG:Sheet:6 | POG:Sheet:9"
    assert result.loc[0, "pog_duplicate_status"] == "Đã gom dòng POG trùng hoàn toàn"
    assert result.loc[0, "pog_position_count"] == 1


def test_pog_facing_is_part_of_stable_task_identity():
    pog = [_pog(facing=1), _pog(facing=2)]

    first = _actions([_promo()], pog)
    reordered = _actions([_promo()], list(reversed(pog)))

    assert first["task_id"].nunique() == 2
    assert first.set_index("facing")["task_id"].to_dict() == reordered.set_index("facing")["task_id"].to_dict()


def test_pog_snapshot_month_is_part_of_task_identity():
    result = _actions([_promo()], [
        _pog(pog_snapshot_month="2026-07"),
        _pog(pog_snapshot_month="2026-08"),
    ])

    assert len(result) == 2
    assert result["task_id"].nunique() == 2
    assert result["pog_position_id"].nunique() == 2


@pytest.mark.parametrize("value", [None, pd.NA, np.nan, "", "   "])
def test_is_sku_handles_missing_values_without_boolean_ambiguity(value):
    assert main.RC.is_sku(value) is False


def test_missing_item_survives_join_and_report_without_counting_as_bill():
    joined = main.JN.join_ctkm_pog(
        pd.DataFrame([{"item_id": pd.NA, "item_key": pd.NA, "campaign_name": "Broken"}]),
        pd.DataFrame([{"item_id": "P1", "store_id": "S1"}]),
    )

    report = main.RC.build_reconciliation_report(
        joined.df_matched, joined.df_ctkm_only, joined.df_pog_only, [],
    )
    actions = _actions([_promo(pd.NA)])

    assert report.non_sku_ctkm_count == 0
    assert report.missing_item_ctkm_count == 1
    assert actions.loc[0, "scheduled_action"] == "review"
    assert actions.loc[0, "reason"] == "thieu_ma_san_pham"


def test_report_distinguishes_bill_from_invalid_item_codes():
    joined = main.JN.join_ctkm_pog(
        pd.DataFrame([
            {"item_id": "Bill", "item_key": "BILL"},
            {"item_id": "UNKNOWN", "item_key": "UNKNOWN"},
        ]),
        pd.DataFrame([{"item_id": "P1", "store_id": "S1"}]),
    )

    report = main.RC.build_reconciliation_report(
        joined.df_matched, joined.df_ctkm_only, joined.df_pog_only, [],
    )
    actions = _actions([_promo("UNKNOWN")])

    assert report.non_sku_ctkm_count == 1
    assert report.invalid_item_ctkm_count == 1
    assert report.missing_item_ctkm_count == 0
    assert actions.loc[0, "scheduled_action"] == "review"
    assert actions.loc[0, "reason"] == "ma_san_pham_khong_hop_le"


def test_multiple_complete_locations_are_not_missing_or_review_errors():
    result = _actions([_promo()], [_pog(location_id="1"), _pog(location_id="2")])

    assert set(result["location_completeness"]) == {"Đủ vị trí"}
    assert set(result["has_multiple_pog_positions"]) == {True}
    assert set(result["location_needs_review"]) == {False}


def test_pog_month_mismatch_is_not_missing_but_needs_location_review():
    result = _actions([_promo()], [_pog(pog_snapshot_month="2026-07")])

    assert result.loc[0, "scheduled_action"] == "gắn tem"
    assert result.loc[0, "location_completeness"] == "Đủ vị trí"
    assert result.loc[0, "pog_month_status"] == "POG lệch tháng"
    assert bool(result.loc[0, "location_needs_review"]) is True


def test_loader_keeps_true_excel_rows_with_nonfirst_header_and_garbage(tmp_path):
    source = tmp_path / "rows.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Report title"])
    sheet.append([])
    sheet.append(["Applied filters: ignored"])
    sheet.append([])
    sheet.append(["Item", "Campaign", "From", "To"])
    sheet.append(["P1", "A", "2026-08-01", "2026-08-31"])
    sheet.append([])
    sheet.append(["P2", "B", "2026-08-01", "2026-08-31"])
    workbook.save(source)

    loaded = main.L.load_workbook_raw(source, selected_sheet_names=["Data"])
    loaded_sheet = loaded.sheets["Data"]

    assert loaded_sheet.header_row_index == 4
    assert loaded_sheet.source_excel_rows == [6, 8]
    processed = main._process(
        "Data", loaded_sheet.df, main.SCHEMA["fact_fields"], [], True, [],
        mapping_override={
            "Item": "item_id", "Campaign": "campaign_name",
            "From": "promo_from", "To": "promo_to",
        },
        source_excel_rows=loaded_sheet.source_excel_rows,
        source_prefix="CTKM",
    )
    assert processed["source_row_id"].tolist() == ["CTKM:Data:6", "CTKM:Data:8"]


def test_unpivot_keeps_excel_row_and_creates_distinct_logical_ids():
    frame = pd.DataFrame([{
        "item_id": "P1", "source_excel_row": 8, "source_row_id": "Monthly:8",
        "campaign_detail__1": "A", "campaign_detail__2": "B",
        "promo_price__1": 90, "promo_price__2": 80,
    }])

    result, _ = main.NM.unpivot_multi_promo(frame)

    assert result["source_excel_row"].tolist() == [8, 8]
    assert result["promo_slot_id"].tolist() == [1, 2]
    assert result["logical_row_id"].nunique() == 2
    assert result["source_row_id"].tolist() == ["Monthly:8:slot-1", "Monthly:8:slot-2"]


def test_export_uses_business_date_for_remaining_days_and_pog_review(tmp_path):
    promotions = [_promo(), _promo(None, campaign_name="Broken")]
    actions = _actions(promotions, [
        _pog(source_row_id="POG:Sheet:6"),
        _pog(source_row_id="POG:Sheet:9"),
    ])
    pog_review = pd.DataFrame([{"item_id": None, "reason": "thieu_ma_san_pham_pog", "source_excel_row": 9}])
    output = tmp_path / "business-date.xlsx"

    main.EX.export_final_report(
        output, pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), [], [],
        {"Thang bao cao": "2026-08", "Ngay nghiep vu": BUSINESS_DATE},
        df_fefo_actions=actions, df_pog_review=pog_review,
    )

    workbook = load_workbook(output, data_only=False)
    assert "POG cần review" in workbook.sheetnames
    assert workbook["Gắn tem"].max_row == 2
    assert workbook["Cần review"].max_row == 2
    assert workbook["Audit hành động"].max_row == 3
    guide_values = [cell.value for row in workbook["HƯỚNG DẪN"] for cell in row]
    assert "Ngày nghiệp vụ" in guide_values
    assert BUSINESS_DATE.to_pydatetime() in guide_values
    values = [cell.value for row in workbook["Gắn tem"] for cell in row]
    assert 16 in values
    assert not any(isinstance(value, str) and "TODAY()" in value for value in values)
    all_values = [cell.value for sheet in workbook for row in sheet for cell in row]
    assert not any(
        isinstance(value, str) and ("#REF!" in value or "#NAME?" in value)
        for value in all_values
    )
    assert "thieu_ma_san_pham" in all_values
    assert "POG:Sheet:6 | POG:Sheet:9" in all_values
    assert workbook["Gắn tem"].data_validations.count == 1
    assert workbook["Gắn tem"].auto_filter.ref is not None
    summary_values = [cell.value for row in workbook["Tom tat"] for cell in row]
    assert "Ngay nghiep vu" in summary_values


def test_integrated_pipeline_workbook_has_conflicts_stable_tasks_and_missing_summary(tmp_path):
    ctkm_path = tmp_path / "ctkm.xlsx"
    pog_path = tmp_path / "pog.xlsx"
    output_path = tmp_path / "integrated.xlsx"

    ctkm_book = Workbook()
    ctkm_sheet = ctkm_book.active
    ctkm_sheet.title = "CTKM"
    ctkm_sheet.append([
        "Item", "Name", "UoM", "Campaign", "Detail", "From", "To",
        "Normal", "Promo", "Channel",
    ])
    ctkm_sheet.append(["P1", "Product 1", "Hộp", "Same", "Giảm 10%", "2026-08-01", "2026-08-20", 100, 90, "OFFLINE"])
    ctkm_sheet.append(["P1", "Product 1", "Hộp", "Same", "Giảm 10%", "2026-08-10", "2026-08-31", 100, 90, "OFFLINE"])
    ctkm_sheet.append(["P2", "Product 2", "Hộp", "Other", "Giảm 5%", "2026-08-01", "2026-08-31", 100, 95, "OFFLINE"])
    ctkm_sheet.append([None, "Broken", "Hộp", "Missing code", "Giảm 5%", "2026-08-01", "2026-08-31", 100, 95, "OFFLINE"])
    ctkm_book.save(ctkm_path)

    pog_book = Workbook()
    pog_sheet = pog_book.active
    pog_sheet.title = "POG"
    pog_sheet.append(["Applied filters DATEID 8/31/2026"])
    pog_sheet.append(["Item", "Store", "POG", "Location", "Facing"])
    pog_sheet.append(["P1", "S1", "A", "1", 1])
    pog_sheet.append(["P1", "S1", "A", "1", 2])
    pog_sheet.append(["P2", "S1", "B", "2", 1])
    pog_book.save(pog_path)

    result = main.run_pipeline(
        ctkm_filepath=ctkm_path,
        pog_filepath=pog_path,
        target_month="2026-08",
        output_path=output_path,
        ctkm_sheet_name="CTKM",
        pog_sheet_name="POG",
        business_date="2026-08-15",
        ctkm_mapping_override={
            "Item": "item_id", "Name": "item_name", "UoM": "promo_uom",
            "Campaign": "campaign_name", "Detail": "campaign_detail",
            "From": "promo_from", "To": "promo_to", "Normal": "normal_price",
            "Promo": "promo_price", "Channel": "channel",
        },
        pog_mapping_override={
            "Item": "item_id", "Store": "store_id", "POG": "pog_code",
            "Location": "location_id", "Facing": "facing",
        },
    )

    workbook = load_workbook(result.output_filepath, data_only=False)
    audit = workbook["Audit hành động"]
    headers = [cell.value for cell in audit[1]]
    task_column = headers.index("Mã công việc (task_id)") + 1
    task_ids = [audit.cell(row=row, column=task_column).value for row in range(2, audit.max_row + 1)]
    review_values = [cell.value for row in workbook["Cần review"] for cell in row]
    warning_values = [cell.value for row in workbook["Canh bao trung thoi gian"] for cell in row]
    summary_values = [cell.value for row in workbook["Tom tat"] for cell in row]

    assert audit.max_row == 7
    assert len(task_ids) == len(set(task_ids)) == 6
    assert "thieu_ma_san_pham" in review_values
    assert "blocking_conflict" in warning_values
    assert True in warning_values
    assert "So dong CTKM thieu ma san pham" in summary_values
    assert result.summary["So dong CTKM thieu ma san pham"] == 1
    assert result.summary["  Trong do CTKM theo hoa don"] == 0
