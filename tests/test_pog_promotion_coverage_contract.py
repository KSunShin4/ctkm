from openpyxl import load_workbook
import pandas as pd

import main
from modules.all_km_reconciliation import build_all_km_reconciliation_frames


def _action(item_id, source, *, in_scope=True):
    return {
        "task_id": f"TASK-{source}-{item_id}",
        "promotion_id": f"PROMO-{source}-{item_id}",
        "pog_position_id": f"POGPOS-{item_id}",
        "item_id": item_id,
        "item_key": item_id,
        "item_name": f"NAME-{item_id}",
        "program_source": source,
        "scheduled_action": "gắn tem",
        "action": "gắn tem",
        "in_report_scope": in_scope,
        "location_completeness": "Đủ vị trí",
        "location_status": "Đủ vị trí",
        "campaign_name": f"CAMPAIGN-{source}-{item_id}",
    }


def _frames():
    # P1001 is covered by Monthly and therefore appears in the Monthly matched set.
    matched = pd.DataFrame([
        {"item_id": "P1001", "item_key": "P1001", "store_id": "S1", "pog_code": "A"},
    ])

    # P2001 has no Monthly campaign but does have an in-scope Clearance campaign.
    # P3001 has no promotion at all and appears at two physical POG positions.
    pog_only_monthly = pd.DataFrame([
        {"item_id": "P2001", "item_key": "P2001", "store_id": "S1", "pog_code": "B", "location_id": "1"},
        {"item_id": "P3001", "item_key": "P3001", "store_id": "S1", "pog_code": "C", "location_id": "1"},
        {"item_id": "P3001", "item_key": "P3001", "store_id": "S1", "pog_code": "D", "location_id": "2"},
    ])

    actions = pd.DataFrame([
        _action("P1001", "CTKM"),
        _action("P2001", "CLEARANCE"),
        _action("P4001", "CLEARANCE", in_scope=False),
    ])
    return matched, pog_only_monthly, actions


def _summary():
    return {
        "Thang bao cao": "2026-09",
        "Ngay nghiep vu": pd.Timestamp("2026-09-15"),
        "So dong khop (Matched)": 1,
        "So item duy nhat khop": 1,
        # Historical values are Monthly-only reconciliation results.
        "So dong POG khong co CTKM phu hop": 3,
        "So item POG khong co CTKM phu hop": 2,
    }


def test_clearance_only_pog_item_is_added_to_all_km_tong_hop():
    matched, pog_only_monthly, actions = _frames()

    matched_all, missing_all_km, added_keys = build_all_km_reconciliation_frames(
        main._core, matched, pog_only_monthly, actions
    )

    assert set(matched_all["item_id"]) == {"P1001", "P2001"}
    assert set(missing_all_km["item_id"]) == {"P3001"}
    assert len(missing_all_km) == 2
    assert added_keys == {"P2001"}

    clearance_row = matched_all[matched_all["item_id"].eq("P2001")].iloc[0]
    assert clearance_row["program_source"] == "CLEARANCE"
    assert clearance_row["reconciliation_basis"] == "KM + POG"


def test_clearance_only_pog_item_is_not_missing_from_all_promotion_coverage():
    matched, pog_only_monthly, actions = _frames()
    summary = _summary()

    main._align_pog_promotion_coverage_summary(
        summary, matched, pog_only_monthly, actions
    )

    # Monthly reconciliation is preserved explicitly.
    assert summary["So dong POG khong co CTKM Monthly phu hop"] == 3
    assert summary["So item POG khong co CTKM Monthly phu hop"] == 2

    # All-KM coverage removes the Clearance-only P2001 from the missing set.
    assert summary["So dong POG khong co KM phu hop"] == 2
    assert summary["So item POG khong co KM phu hop"] == 1
    assert summary["So item POG co KM phu hop"] == 2
    assert summary["So item POG tong hop de doi chieu KM"] == 3

    # Existing frontend releases read these historical keys, so keep their
    # operational value aligned with all-KM coverage until the UI migrates.
    assert summary["So dong POG khong co CTKM phu hop"] == 2
    assert summary["So item POG khong co CTKM phu hop"] == 1


def test_export_tong_hop_is_all_km_and_monthly_audit_remains_explicit(tmp_path):
    matched, pog_only_monthly, actions = _frames()
    summary = _summary()
    output = tmp_path / "pog-promotion-coverage.xlsx"

    main.EX.export_final_report(
        output_path=output,
        df_matched=matched,
        df_ctkm_only=pd.DataFrame(),
        df_pog_only=pog_only_monthly,
        overlap_warnings=[],
        mapping_logs=[],
        reconciliation_summary=summary,
        df_fefo_actions=actions,
    )

    workbook = load_workbook(output, data_only=True)

    # Tong hop is now the actual Monthly + Clearance reconciliation surface.
    tong_hop_values = [cell.value for row in workbook["Tong hop"].iter_rows() for cell in row]
    assert "P1001" in tong_hop_values
    assert "P2001" in tong_hop_values
    assert "P3001" not in tong_hop_values

    # User-facing all-KM missing audit no longer falsely lists the Clearance SKU.
    all_km_sheet = workbook["POG không có KM"]
    all_km_values = [cell.value for row in all_km_sheet.iter_rows() for cell in row]
    assert "P2001" not in all_km_values
    assert "P3001" in all_km_values

    # Historical Monthly-only evidence remains available under an explicit name.
    monthly_sheet = workbook["POG không có Monthly CTKM"]
    monthly_values = [cell.value for row in monthly_sheet.iter_rows() for cell in row]
    assert "P2001" in monthly_values
    assert "P3001" in monthly_values

    # Compatibility sheet remains for old clients/contracts but is hidden and
    # carries all-KM missing semantics, not the misleading Monthly-only rows.
    compat = workbook["POG không có CTKM"]
    compat_values = [cell.value for row in compat.iter_rows() for cell in row]
    assert compat.sheet_state == "hidden"
    assert "P2001" not in compat_values
    assert "P3001" in compat_values

    assert summary["So dong khop (Matched)"] == 2
    assert summary["So item duy nhat khop"] == 2
    assert summary["So item Clearance bo sung vao Tong hop"] == 1
    assert summary["So item KM+POG thieu trong Tong hop"] == 0
    assert summary["So dong POG khong co KM phu hop"] == 2
    assert summary["So item POG khong co KM phu hop"] == 1
    assert summary["So dong POG khong co CTKM Monthly phu hop"] == 3
    assert summary["So item POG khong co CTKM Monthly phu hop"] == 2

    tom_tat = workbook["Tom tat"]
    rows = {
        tom_tat.cell(row=row, column=1).value: (tom_tat.cell(row=row, column=2).value, row)
        for row in range(2, tom_tat.max_row + 1)
    }
    assert rows["So item KM+POG thieu trong Tong hop"][0] == 0
    assert rows["So item Clearance bo sung vao Tong hop"][0] == 1

    # Misleading legacy labels are hidden from the normal workbook summary view.
    legacy_row = rows["So dong POG khong co CTKM phu hop"][1]
    assert tom_tat.row_dimensions[legacy_row].hidden is True
