import importlib.util
from pathlib import Path
import sys
from types import SimpleNamespace

import pandas as pd
from openpyxl import Workbook, load_workbook

from backend.tem_r_source_fidelity_contract import (
    install_tem_r_source_fidelity_contract,
)


def _module():
    path = Path(__file__).resolve().parents[1] / "modules" / "10_tem_r.py"
    name = "tem_r_source_fidelity_test"
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def _install(tem_r):
    core = SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace()),
        _tem_r_mod=tem_r,
    )
    install_tem_r_source_fidelity_contract(core)
    return core


def _row(item_id, stock, *, excel_row=None, location=1):
    result = {
        "STOREID": " S1 ",
        "POGCODE": " G1 ",
        "FIXELID": " 1/1 ",
        "LOCATIONID": location,
        "ITEMID": item_id,
        "ITEMNAME": "Name",
        "TỒN": stock,
    }
    if excel_row is not None:
        result["_excel_row"] = excel_row
    return result


def test_tem_r_plan_keeps_raw_itemid_and_exposes_normalized_technical_key():
    tem_r = _module()
    _install(tem_r)
    frame = pd.DataFrame([
        _row(" p00123 ", 2, location=1),
        _row("P00999", 0, location=2),
    ])
    original = frame.copy(deep=True)

    plan = tem_r.build_tem_r_plan(frame)

    pd.testing.assert_frame_equal(frame, original)
    assert plan.rows.loc[0, "ITEMID"] == " p00123 "
    assert plan.rows.loc[0, "_item_key"] == "P00123"
    assert plan.audit.loc[0, "ITEMID"] == " p00123 "
    # Derived representative may use normalized identity; raw source stays raw.
    assert plan.rows.loc[1, "TEM R"] == "P00123"


def test_tem_r_exception_keeps_exact_raw_itemid_by_excel_row():
    tem_r = _module()
    _install(tem_r)
    frame = pd.DataFrame([
        _row(" p00123 ", None, excel_row=17),
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert plan.rows.loc[0, "ITEMID"] == " p00123 "
    assert plan.audit.loc[0, "ITEMID"] == " p00123 "
    assert plan.exceptions.iloc[0]["row"] == 17
    assert plan.exceptions.iloc[0]["ITEMID"] == " p00123 "


def test_repeated_header_is_removed_without_mutating_other_raw_itemids():
    tem_r = _module()
    _install(tem_r)
    frame = pd.DataFrame([
        _row("ITEMID", 0, location=1),
        _row(" p00123 ", 1, location=2),
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert len(plan.rows) == 1
    assert plan.rows.iloc[0]["ITEMID"] == " p00123 "
    assert plan.rows.iloc[0]["_item_key"] == "P00123"


def test_create_tem_r_workbook_keeps_raw_itemid_in_source_and_audit(tmp_path):
    tem_r = _module()
    _install(tem_r)
    source = tmp_path / "tem-r-source.xlsx"
    output = tmp_path / "tem-r-output.xlsx"

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Export"
    ws.append(["STOREID", "POGCODE", "FIXELID", "LOCATIONID", "ITEMID", "ITEMNAME", "TỒN"])
    ws.append(["S1", "G1", "1/1", 1, " p00123 ", "A", 2])
    ws.append(["S1", "G1", "1/1", 2, "P00999", "B", 0])
    workbook.save(source)
    workbook.close()

    plan = tem_r.create_tem_r_workbook(source, output)

    assert plan.rows.loc[0, "ITEMID"] == " p00123 "
    result = load_workbook(output, data_only=False)
    try:
        assert result["Export"].cell(2, 5).value == " p00123 "
        audit = result["Tem R phân bổ"]
        headers = [cell.value for cell in audit[1]]
        item_col = headers.index("ITEMID") + 1
        assert audit.cell(2, item_col).value == " p00123 "
    finally:
        result.close()


def test_contract_installation_is_idempotent():
    tem_r = _module()
    core = _install(tem_r)
    once = tem_r.build_tem_r_plan
    install_tem_r_source_fidelity_contract(core)
    assert tem_r.build_tem_r_plan is once


def test_tem_r_core_itself_preserves_raw_itemid_without_backend_wrapper():
    tem_r = _module()
    frame = pd.DataFrame([
        _row(" p00123 ", 2, location=1),
        _row("P00999", 0, location=2),
    ])
    original = frame.copy(deep=True)

    # Call the core allocator directly so this invariant cannot depend on the
    # backend installation order or compatibility wrapper.
    plan = tem_r._core.build_tem_r_plan(frame)

    pd.testing.assert_frame_equal(frame, original)
    assert plan.rows.loc[0, "ITEMID"] == " p00123 "
    assert plan.rows.loc[0, "_item_key"] == "P00123"
    assert plan.audit.loc[0, "ITEMID"] == " p00123 "
    assert plan.rows.loc[1, "TEM R"] == "P00123"


def test_source_fidelity_wrapper_does_not_replace_core_technical_item_key():
    tem_r = _module()
    _install(tem_r)
    frame = pd.DataFrame([
        _row(" p00123 ", 2, location=1),
        _row("P00999", 0, location=2),
    ])

    plan = tem_r.build_tem_r_plan(frame)

    assert plan.rows.loc[0, "ITEMID"] == " p00123 "
    assert plan.rows.loc[0, "_item_key"] == "P00123"
