# -*- coding: utf-8 -*-
"""Public CTKM orchestrator entrypoint with fidelity contract installed eagerly."""
from __future__ import annotations

import importlib
import importlib.util
from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook

from modules.facade_contracts import run_post_facade_hooks
from modules.item_code_identity import build_transition_identity, canonical_series

_CORE_NAME = "_ctkm_main_core"
_CORE_PATH = Path(__file__).with_name("main_core.py")

if _CORE_NAME in sys.modules:
    _core = sys.modules[_CORE_NAME]
else:
    _spec = importlib.util.spec_from_file_location(_CORE_NAME, _CORE_PATH)
    if _spec is None or _spec.loader is None:
        raise ImportError(f"Không nạp được CTKM core: {_CORE_PATH}")
    _core = importlib.util.module_from_spec(_spec)
    sys.modules[_CORE_NAME] = _core
    _spec.loader.exec_module(_core)

_performance = importlib.import_module("modules.ctkm_performance")
_performance.install_performance_patches(_core)

# ``main.py`` is loaded as ``ctkm_main`` by the API and may also be imported as
# ``main`` by tests/CLI tooling in the same process. Keep one immutable baseline
# after the performance contract so repeated facade execution replaces wrappers
# instead of wrapping wrappers from a previous module name/import order.
_baseline = getattr(_core, "_ctkm_main_facade_baseline", None)
if not isinstance(_baseline, dict):
    _baseline = {
        "infer_standard_fact_mapping": _core._infer_standard_fact_mapping,
        "get_or_map": _core._get_or_map,
        "process": _core._process,
        "build_confirmation": _core._build_ctkm_clearance_confirmation,
        "export_final_report": _core.EX.export_final_report,
    }
    _core._ctkm_main_facade_baseline = _baseline

_original_infer_standard_fact_mapping = _baseline["infer_standard_fact_mapping"]
_original_get_or_map = _baseline["get_or_map"]
_original_process = _baseline["process"]
_original_build_confirmation = _baseline["build_confirmation"]


def _infer_standard_fact_mapping(df):
    mapping = _original_infer_standard_fact_mapping(df)

    if not mapping:
        normalized_columns = {
            _core._normalized_header(column): column for column in df.columns
        }
        campaign_detail_source = normalized_columns.get("campaign detail")
        if "campaign" not in normalized_columns and campaign_detail_source is not None:
            inference_frame = df.rename(columns={campaign_detail_source: "Campaign"})
            mapping = _original_infer_standard_fact_mapping(inference_frame)
            if "Campaign" in mapping:
                campaign_target = mapping.pop("Campaign")
                mapping[campaign_detail_source] = campaign_target

    if not mapping:
        return mapping
    return _performance.preserve_both_channel_mappings(
        mapping, df, normalize_header=_core._normalized_header,
    )


def _infer_standard_pog_mapping(df):
    """Map the stable Pharmacity POG export headers without Gemini."""
    normalized_columns = {}
    for column in df.columns:
        normalized_columns.setdefault(_core._normalized_header(column), column)

    aliases = {
        "item_id": ("itemid", "item id", "item code"),
        "item_name": ("itemname", "item name", "product name"),
        "store_id": ("storeid", "store id"),
        "pog_code": ("pogcode", "pog code"),
        "location_id": ("locationid", "location id"),
        "facing": ("totalfacing", "total facing", "facing"),
    }
    mapping = {}
    for target, target_aliases in aliases.items():
        for alias in target_aliases:
            source = normalized_columns.get(alias)
            if source is not None:
                mapping[source] = target
                break

    if not {"item_id", "store_id"}.issubset(set(mapping.values())):
        return {}
    return mapping


def _nonblank_samples(series, limit=40):
    values = []
    for value in series:
        if value is None:
            continue
        try:
            if pd.isna(value):
                continue
        except (TypeError, ValueError):
            pass
        text = str(value).strip()
        if text:
            values.append(value)
        if len(values) >= limit:
            break
    return values


def _pattern_ratio(series, pattern, limit=40):
    values = _nonblank_samples(series, limit)
    if len(values) < 3:
        return 0.0
    matched = sum(bool(pattern.fullmatch(str(value).strip().upper())) for value in values)
    return matched / len(values)


def _date_sample_value(value):
    if isinstance(value, (pd.Timestamp,)):
        return value.normalize()
    text = str(value).strip()
    if not (
        re.fullmatch(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}(?:\s+.*)?", text)
        or re.fullmatch(r"\d{1,2}[-/.]\d{1,2}[-/.]\d{4}(?:\s+.*)?", text)
    ):
        return None
    parsed = pd.to_datetime(text, errors="coerce", dayfirst=not text[:4].isdigit())
    if pd.isna(parsed) or not (2020 <= parsed.year <= 2035):
        return None
    return parsed.normalize()


def _date_ratio(series, limit=40):
    values = _nonblank_samples(series, limit)
    if len(values) < 3:
        return 0.0
    return sum(_date_sample_value(value) is not None for value in values) / len(values)


def _unique_pattern_candidate(df, excluded, pattern, minimum=0.75):
    candidates = []
    for column in df.columns:
        if column in excluded:
            continue
        ratio = _pattern_ratio(df[column], pattern)
        if ratio >= minimum:
            candidates.append((column, ratio))
    if len(candidates) != 1:
        return None
    return candidates[0]


def _required_fact_header_partial(df):
    normalized = {}
    for column in df.columns:
        normalized.setdefault(_core._normalized_header(column), column)

    mapping = {}
    trusted = set()

    def claim(target, aliases):
        for alias in aliases:
            source = normalized.get(alias)
            if source is not None and target not in mapping.values():
                mapping[source] = target
                trusted.add(source)
                return

    claim("item_id", ("item code", "item code 2", "itemid"))
    claim("campaign_name", ("campaign",))
    if "campaign_name" not in mapping.values():
        claim("campaign_name", ("campaign detail",))
    claim("promo_from", ("from",))
    claim("promo_to", ("to",))
    return mapping, trusted


def _required_pog_header_partial(df):
    normalized = {}
    for column in df.columns:
        normalized.setdefault(_core._normalized_header(column), column)
    mapping = {}
    trusted = set()

    aliases = {
        "item_id": ("itemid", "item id", "item code"),
        "store_id": ("storeid", "store id"),
    }
    for target, names in aliases.items():
        for alias in names:
            source = normalized.get(alias)
            if source is not None:
                mapping[source] = target
                trusted.add(source)
                break
    return mapping, trusted


def _infer_local_value_mapping(df, kind):
    """Return a complete required mapping candidate only when local evidence is strong enough to evaluate."""
    if df is None or df.empty:
        return None

    confidence = {}
    if kind == "fact":
        mapping, trusted = _required_fact_header_partial(df)
        used = set(mapping)

        if "item_id" not in mapping.values():
            item = _unique_pattern_candidate(
                df, used, re.compile(r"P\d{3,10}", re.IGNORECASE)
            )
            if item is not None:
                source, score = item
                mapping[source] = "item_id"
                confidence[source] = round(float(score), 4)
                used.add(source)

        missing_dates = {
            target for target in ("promo_from", "promo_to")
            if target not in mapping.values()
        }
        if missing_dates == {"promo_from", "promo_to"}:
            date_candidates = []
            for column in df.columns:
                if column in used:
                    continue
                ratio = _date_ratio(df[column])
                if ratio >= 0.75:
                    date_candidates.append((column, ratio))
            if len(date_candidates) == 2:
                (left, left_ratio), (right, right_ratio) = date_candidates
                pairs = []
                for left_value, right_value in zip(df[left], df[right]):
                    a = _date_sample_value(left_value)
                    b = _date_sample_value(right_value)
                    if a is not None and b is not None:
                        pairs.append((a, b))
                    if len(pairs) >= 40:
                        break
                if len(pairs) >= 3:
                    left_before = sum(a <= b for a, b in pairs) / len(pairs)
                    right_before = sum(b <= a for a, b in pairs) / len(pairs)
                    if left_before >= 0.75 and left_before > right_before:
                        mapping[left], mapping[right] = "promo_from", "promo_to"
                        score = min(left_ratio, right_ratio, left_before)
                        confidence[left] = confidence[right] = round(float(score), 4)
                    elif right_before >= 0.75 and right_before > left_before:
                        mapping[right], mapping[left] = "promo_from", "promo_to"
                        score = min(left_ratio, right_ratio, right_before)
                        confidence[left] = confidence[right] = round(float(score), 4)

        required = {"item_id", "campaign_name", "promo_from", "promo_to"}
        if not required.issubset(set(mapping.values())):
            return None
        return {
            "mapping": mapping,
            "confidence": confidence,
            "trusted_sources": sorted(trusted),
            "source": "local-value-fact",
        }

    if kind == "pog":
        mapping, trusted = _required_pog_header_partial(df)
        used = set(mapping)
        if "item_id" not in mapping.values():
            item = _unique_pattern_candidate(
                df, used, re.compile(r"P\d{3,10}", re.IGNORECASE)
            )
            if item is not None:
                source, score = item
                mapping[source] = "item_id"
                confidence[source] = round(float(score), 4)
                used.add(source)
        if "store_id" not in mapping.values():
            store = _unique_pattern_candidate(
                df, used, re.compile(r"(?:[A-Z]{1,4})?PMC\d{3,6}", re.IGNORECASE)
            )
            if store is not None:
                source, score = store
                mapping[source] = "store_id"
                confidence[source] = round(float(score), 4)

        if not {"item_id", "store_id"}.issubset(set(mapping.values())):
            return None
        return {
            "mapping": mapping,
            "confidence": confidence,
            "trusted_sources": sorted(trusted),
            "source": "local-value-pog",
        }

    return None


def _mapping_result(sheet_name, mapping, df, schema_section, source):
    result = _core._build_manual_mapping_result(
        sheet_name, mapping, df, schema_section
    )
    entry = _core.MC.mapping_result_to_dict_safe(result)
    entry["source"] = source
    return result, {**entry, "sheet_name": sheet_name}


def _get_or_map(sheet_name, df, schema_section, mapping_logs):
    is_fact_schema = (
        "campaign_name" in schema_section
        and "promo_from" in schema_section
        and "promo_to" in schema_section
    )
    if is_fact_schema:
        mapping = _infer_standard_fact_mapping(df)
        if mapping:
            result, log_entry = _mapping_result(
                sheet_name, mapping, df, schema_section, "standard-fact"
            )
            mapping_logs.append(log_entry)
            return result.mapping

    is_pog_schema = (
        "store_id" in schema_section
        and "pog_code" in schema_section
        and "campaign_name" not in schema_section
    )
    if is_pog_schema:
        mapping = _infer_standard_pog_mapping(df)
        if mapping:
            result, log_entry = _mapping_result(
                sheet_name, mapping, df, schema_section, "standard-pog"
            )
            mapping_logs.append(log_entry)
            return result.mapping

    return _original_get_or_map(sheet_name, df, schema_section, mapping_logs)


def _blank_fact_value(value):
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(value, str) and not value.strip()


def _fact_control_kind(row) -> str | None:
    """Recognize only unmistakable subtotal/footer rows from optional fact tabs."""
    if not _blank_fact_value(row.get("item_key")) or not _blank_fact_value(row.get("item_id")):
        return None

    business_anchors = (
        "item_name", "promo_uom", "promo_from", "promo_to",
        "key_channel", "channel", "campaign_detail", "product_type",
    )
    if any(not _blank_fact_value(row.get(column)) for column in business_anchors):
        return None

    markers = []
    for value in row.values:
        if not isinstance(value, str):
            continue
        text = " ".join(value.strip().lower().split())
        if text:
            markers.append(text)
    if any(text in {"grand total", "total", "tổng", "tổng cộng"} for text in markers):
        return "grand_total"
    if any(text.endswith(" total") for text in markers):
        return "subtotal"
    return None


def _separate_fact_control_rows(df):
    if df is None or df.empty:
        return df, pd.DataFrame(columns=[] if df is None else df.columns)
    kinds = df.apply(_fact_control_kind, axis=1)
    control_mask = kinds.notna()
    business = df.loc[~control_mask].copy().reset_index(drop=True)
    metadata = df.loc[control_mask].copy().reset_index(drop=True)
    if not metadata.empty:
        metadata["fact_control_kind"] = kinds.loc[control_mask].to_list()
    return business, metadata


def _process(*args, **kwargs):
    """Process mapped facts while keeping subtotal rows outside business logic."""
    result = _original_process(*args, **kwargs)
    is_fact = kwargs.get("is_fact")
    if is_fact is None and len(args) > 4:
        is_fact = args[4]
    if not is_fact or result is None or result.empty:
        return result

    business, metadata = _separate_fact_control_rows(result)
    if metadata.empty:
        return business

    source_prefix = kwargs.get("source_prefix")
    if source_prefix is None and len(args) > 9:
        source_prefix = args[9]
    if source_prefix:
        metadata["program_source"] = str(source_prefix).upper()

    business.attrs["fact_metadata"] = metadata
    warnings = kwargs.get("all_warnings")
    if warnings is None and len(args) > 5:
        warnings = args[5]
    sheet_name = kwargs.get("sheet_name")
    if sheet_name is None and args:
        sheet_name = args[0]
    if isinstance(warnings, list):
        warnings.append(
            f"[Sheet '{sheet_name}'] Tách {len(metadata)} dòng subtotal/control khỏi xử lý nghiệp vụ; "
            "giữ nguyên trong sheet ẩn 'KM metadata'."
        )
    return business


def _build_ctkm_clearance_confirmation(df_ctkm, df_clearance):
    metadata_frames = []
    for frame in (df_ctkm, df_clearance):
        if frame is None:
            continue
        metadata = getattr(frame, "attrs", {}).get("fact_metadata")
        if metadata is not None and not metadata.empty:
            metadata_frames.append(metadata.copy())

    result = _original_build_confirmation(df_ctkm, df_clearance)
    if metadata_frames:
        result.attrs["fact_metadata"] = pd.concat(metadata_frames, ignore_index=True, sort=False)
    return result


def _comparison_identity(frame):
    return build_transition_identity(
        frame,
        key_func=_core.NM.make_item_key,
        is_sku=_core.RC.is_sku,
    )


def _comparison_item_series(frame, identity):
    return canonical_series(
        frame, identity, key_func=_core.NM.make_item_key,
    )


def _frame_item_keys(frame, identity=None):
    """Return comparison product keys without changing source item values."""
    if frame is None or frame.empty:
        return set()
    identity = identity or _comparison_identity(frame)
    series = _comparison_item_series(frame, identity)
    return {value for value in series.dropna() if str(value).strip()}


def _location_issue_rows(df_fefo_actions):
    if df_fefo_actions is None:
        return pd.DataFrame()
    actions = df_fefo_actions.copy()
    if "in_report_scope" in actions.columns:
        actions = actions[actions["in_report_scope"].fillna(False).astype(bool)].copy()
    if actions.empty:
        return actions
    completeness = actions.get(
        "location_completeness",
        actions.get("location_status", pd.Series("Đủ vị trí", index=actions.index)),
    )
    return actions[~completeness.eq("Đủ vị trí")].copy()


def _align_location_summary_with_program_sources(summary, df_fefo_actions):
    """Align location issue counts with product-transition comparison identity."""
    if not isinstance(summary, dict) or df_fefo_actions is None:
        return summary

    identity = _comparison_identity(df_fefo_actions)
    issues = _location_issue_rows(df_fefo_actions)
    issue_keys = _comparison_item_series(issues, identity)
    sources = issues.get("program_source", pd.Series(index=issues.index, dtype=object))
    ctkm_mask = sources.eq("CTKM")
    clearance_mask = sources.eq("CLEARANCE")

    total_rows = len(issues)
    total_items = issue_keys.nunique(dropna=True)
    ctkm_items = issue_keys[ctkm_mask].nunique(dropna=True)
    clearance_items = issue_keys[clearance_mask].nunique(dropna=True)

    summary["So dong KM thieu vi tri POG"] = total_rows
    summary["So item KM thieu vi tri POG"] = total_items
    summary["So dong CTKM Monthly thieu vi tri POG"] = int(ctkm_mask.sum())
    summary["So item CTKM Monthly thieu vi tri POG"] = ctkm_items
    summary["So dong Clearance thieu vi tri POG"] = int(clearance_mask.sum())
    summary["So item Clearance thieu vi tri POG"] = clearance_items

    if "So dong CTKM thieu vi tri ke" in summary:
        summary["So dong CTKM thieu vi tri ke"] = total_rows
    if "So item CTKM thieu vi tri ke" in summary:
        summary["So item CTKM thieu vi tri ke"] = total_items
    return summary


def _align_cross_source_summary(summary, df_fefo_actions):
    """Count CTKM/Clearance shared products across safe source/new-code pairs."""
    if not isinstance(summary, dict) or df_fefo_actions is None:
        return summary
    actions = df_fefo_actions.copy()
    if "in_report_scope" in actions.columns:
        actions = actions[actions["in_report_scope"].fillna(False).astype(bool)].copy()
    if actions.empty:
        if "So item trung CTKM va Clearance" in summary:
            summary["So item trung CTKM va Clearance"] = 0
        return summary

    identity = _comparison_identity(actions)
    keys = _comparison_item_series(actions, identity)
    sources = actions.get("program_source", pd.Series(index=actions.index, dtype=object))
    ctkm_keys = set(keys[sources.eq("CTKM")].dropna())
    clearance_keys = set(keys[sources.eq("CLEARANCE")].dropna())
    if "So item trung CTKM va Clearance" in summary:
        summary["So item trung CTKM va Clearance"] = len(ctkm_keys & clearance_keys)
    return summary


def _align_matched_unique_summary(summary, df_matched, df_fefo_actions):
    if not isinstance(summary, dict) or "So item duy nhat khop" not in summary:
        return summary
    identity = _comparison_identity(df_fefo_actions)
    summary["So item duy nhat khop"] = len(_frame_item_keys(df_matched, identity))
    return summary


def _align_pog_promotion_coverage_summary(
    summary, df_matched, df_pog_only, df_fefo_actions
):
    """Separate Monthly reconciliation from all-promotion POG coverage.

    Product sets are compared through safe source/new-code identity so a POG row
    already migrated to New P code is not reported as promotion-less merely
    because the operational action intentionally retains the source item_id.
    """
    if not isinstance(summary, dict):
        return summary

    identity = _comparison_identity(df_fefo_actions)
    matched_keys = _frame_item_keys(df_matched, identity)
    monthly_pog_only_keys = _frame_item_keys(df_pog_only, identity)
    all_pog_keys = matched_keys | monthly_pog_only_keys

    actions = df_fefo_actions.copy() if df_fefo_actions is not None else pd.DataFrame()
    if not actions.empty and "in_report_scope" in actions.columns:
        actions = actions[actions["in_report_scope"].fillna(False).astype(bool)].copy()
    km_keys = _frame_item_keys(actions, identity)

    missing_km_keys = all_pog_keys - km_keys
    covered_km_keys = all_pog_keys & km_keys

    missing_rows = 0
    if df_pog_only is not None and not df_pog_only.empty:
        row_keys = _comparison_item_series(df_pog_only, identity)
        missing_rows = int(row_keys.isin(missing_km_keys).sum())

    monthly_row_fallback = len(df_pog_only) if df_pog_only is not None else 0
    monthly_rows = int(summary.get(
        "So dong POG khong co CTKM Monthly phu hop",
        summary.get("So dong POG khong co CTKM phu hop", monthly_row_fallback),
    ))
    monthly_items = int(summary.get(
        "So item POG khong co CTKM Monthly phu hop",
        len(monthly_pog_only_keys),
    ))
    summary["So dong POG khong co CTKM Monthly phu hop"] = monthly_rows
    summary["So item POG khong co CTKM Monthly phu hop"] = monthly_items

    summary["So dong POG khong co KM phu hop"] = missing_rows
    summary["So item POG khong co KM phu hop"] = len(missing_km_keys)
    summary["So item POG co KM phu hop"] = len(covered_km_keys)
    summary["So item POG tong hop de doi chieu KM"] = len(all_pog_keys)

    if "So dong POG khong co CTKM phu hop" in summary:
        summary["So dong POG khong co CTKM phu hop"] = missing_rows
    if "So item POG khong co CTKM phu hop" in summary:
        summary["So item POG khong co CTKM phu hop"] = len(missing_km_keys)
    return summary


def _append_fact_metadata_audit(output_path, metadata):
    if metadata is None or metadata.empty:
        return
    path = Path(output_path)
    if not path.is_file():
        return
    workbook = load_workbook(path)
    sheet_name = "KM metadata"
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    ws = workbook.create_sheet(sheet_name)
    columns = list(metadata.columns)
    ws.append(columns)
    for _, row in metadata.iterrows():
        values = []
        for column in columns:
            value = row[column]
            try:
                missing = bool(pd.isna(value))
            except (TypeError, ValueError):
                missing = False
            if missing:
                values.append(None)
            elif isinstance(value, pd.Timestamp):
                values.append(value.to_pydatetime())
            else:
                values.append(value)
        ws.append(values)
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
    ws.sheet_state = "hidden"
    workbook.save(path)


def _clarify_location_workbook(output_path):
    """Hide misleading legacy summary aliases while preserving workbook schema."""
    path = Path(output_path)
    if not path.is_file():
        return
    workbook = load_workbook(path)
    changed = False

    if "Tom tat" in workbook.sheetnames:
        ws = workbook["Tom tat"]
        legacy = {
            "So dong CTKM thieu vi tri ke",
            "So item CTKM thieu vi tri ke",
            "So dong POG khong co CTKM phu hop",
            "So item POG khong co CTKM phu hop",
        }
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value in legacy:
                ws.row_dimensions[row].hidden = True
                changed = True

    if "HƯỚNG DẪN" in workbook.sheetnames:
        ws = workbook["HƯỚNG DẪN"]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == "Thiếu vị trí":
                    cell.value = "Thiếu vị trí KM"
                    changed = True

    if changed:
        workbook.save(path)


_core._infer_standard_fact_mapping = _infer_standard_fact_mapping
_core._infer_standard_pog_mapping = _infer_standard_pog_mapping
_core._get_or_map = _get_or_map
_core._process = _process
_core._build_ctkm_clearance_confirmation = _build_ctkm_clearance_confirmation

_task_aligned_export = _baseline["export_final_report"]


def _export_with_location_summary(*args, **kwargs):
    summary = kwargs.get("reconciliation_summary")
    if summary is None and len(args) > 6:
        summary = args[6]
    actions = kwargs.get("df_fefo_actions")
    if actions is None and len(args) > 9:
        actions = args[9]
    confirmation = kwargs.get("df_ctkm_clearance_confirmation")
    if confirmation is None and len(args) > 10:
        confirmation = args[10]
    df_matched = kwargs.get("df_matched")
    if df_matched is None and len(args) > 1:
        df_matched = args[1]
    df_pog_only = kwargs.get("df_pog_only")
    if df_pog_only is None and len(args) > 3:
        df_pog_only = args[3]

    _align_location_summary_with_program_sources(summary, actions)
    _align_cross_source_summary(summary, actions)
    _align_matched_unique_summary(summary, df_matched, actions)
    _align_pog_promotion_coverage_summary(summary, df_matched, df_pog_only, actions)
    output_path = _task_aligned_export(*args, **kwargs)
    _clarify_location_workbook(output_path)
    metadata = (
        getattr(confirmation, "attrs", {}).get("fact_metadata")
        if confirmation is not None else None
    )
    _append_fact_metadata_audit(output_path, metadata)
    return output_path


_core.EX.export_final_report = _export_with_location_summary

# Backend/runtime contracts can register hooks that must survive a second facade
# execution under another module name. Reapply them only after this facade has
# rebuilt its raw wrappers, then export the hardened core surface below.
_post_facade_hooks = run_post_facade_hooks(_core, sys.modules.get(__name__))

for _name in dir(_core):
    if not _name.startswith("__"):
        globals()[_name] = getattr(_core, _name)

if __name__ == "__main__":
    _core._cli()
