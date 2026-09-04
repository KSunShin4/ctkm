# -*- coding: utf-8 -*-
"""Reduce CPU cost of CTKM workbook styling without changing workbook semantics.

The historical exporter styled each worksheet in several full passes and created
many equivalent OpenPyXL style objects per cell. That is tolerable on a desktop
but unnecessarily expensive on a very small Render CPU. This contract replaces
only presentation helpers in the shared exporter core:

* one main pass computes widths, row styling, action colors, shelf bands and
  review/collision highlights;
* immutable style objects are reused instead of recreated for every cell;
* the later shelf-color pass becomes a no-op because its final visual state has
  already been applied in the main pass;
* action/confirmation input helpers reuse protection/fill objects.

No business dataframe, sheet name, row count, mapping, task identity or source
value is changed by this module.
"""
from __future__ import annotations

import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


_ACTION_SHEETS = {"Gắn tem", "Tháo tem", "Chờ gắn", "Cần review"}
_COLOR_MAP = {
    "Tong hop": "4F81BD",
    "CTKM thieu vi tri ke": "C0504D",
    "Gắn tem": "2E7D32",
    "Tháo tem": "F28C00",
    "Chờ gắn": "C79B00",
    "Cần review": "7030A0",
    "Audit hành động": "5B5EA6",
    "POG không có CTKM": "9BBB59",
    "POG cần review": "C65911",
    "POG cần kiểm tra": "C0504D",
    "Canh bao trung thoi gian": "C79B00",
    "Xác nhận CTKM Clearance": "00A6A6",
    "Mapping log": "5B5EA6",
    "Tom tat": "6F42C1",
}
_DATE_HEADERS = {"promo_from", "promo_to", "ngay_gan_tem", "ngay_thao_tem"}
_REVIEW_FIELDS = (
    "location_completeness", "pog_month_status",
    "location_status", "location_needs_review",
)
_COLLISION_FIELDS = (
    "promotion_count_at_position", "has_multiple_promotions_at_position",
)


def _truthy(value) -> bool:
    return value is True or str(value).strip().lower() in {"true", "1", "yes"}


def _column_values(df, name: str, length: int):
    if df is None or getattr(df, "empty", True) or name not in df.columns:
        return [None] * length
    values = df[name].tolist()
    if len(values) < length:
        values.extend([None] * (length - len(values)))
    return values


def _optimized_professional_style(ws, sheet_name: str, df: pd.DataFrame | None = None):
    if ws is None:
        return

    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    if max_col <= 0:
        return

    thin = Side(border_style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_alignment = Alignment(vertical="center", wrap_text=True)
    alt_fill = PatternFill(fill_type="solid", fgColor="F7F9FC")
    white_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    attach_fill = PatternFill(fill_type="solid", fgColor="E8F5E9")
    remove_fill = PatternFill(fill_type="solid", fgColor="FFF3E0")
    shelf_fills = (
        PatternFill(fill_type="solid", fgColor="EAF3FF"),
        PatternFill(fill_type="solid", fgColor="FFF4E5"),
    )
    review_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")
    review_font = Font(bold=True, color="9C0006")
    warning_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    warning_font = Font(bold=True, color="7F6000")

    ws.freeze_panes = "C2" if sheet_name in _ACTION_SHEETS else "A2"
    ws.auto_filter.ref = ws.dimensions if max_row >= 1 and max_col >= 1 else None
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.sheet_properties.outlinePr.summaryBelow = True

    headers = [ws.cell(row=1, column=index).value for index in range(1, max_col + 1)]
    header_text = ["" if value is None else str(value) for value in headers]
    widths = [max(10, len(value)) for value in header_text]

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border

    data_rows = max(0, max_row - 1)
    action_values = _column_values(df, "action", data_rows)
    location_values = _column_values(
        df,
        "location_completeness" if df is not None and "location_completeness" in getattr(df, "columns", ())
        else "location_status",
        data_rows,
    )
    review_values = _column_values(df, "location_needs_review", data_rows)
    collision_values = _column_values(df, "has_multiple_promotions_at_position", data_rows)
    store_values = _column_values(df, "store_id", data_rows)
    pog_values = _column_values(df, "pog_code", data_rows)

    action_col = next(
        (index for index, value in enumerate(header_text, start=1) if value.lower() == "action"),
        None,
    )
    completeness_header = (
        "location_completeness" if "location_completeness" in header_text
        else "location_status" if "location_status" in header_text else None
    )
    completeness_col = header_text.index(completeness_header) + 1 if completeness_header else None
    review_columns = [header_text.index(field) + 1 for field in _REVIEW_FIELDS if field in header_text]
    collision_columns = [header_text.index(field) + 1 for field in _COLLISION_FIELDS if field in header_text]

    shelf_color_index: dict[tuple[str, str], int] = {}

    for offset, row_idx in enumerate(range(2, max_row + 1)):
        base_fill = alt_fill if row_idx % 2 == 0 else white_fill
        action_fill = None
        if sheet_name in _ACTION_SHEETS and offset < len(action_values):
            action_text = "" if action_values[offset] is None else str(action_values[offset]).lower()
            if "gắn" in action_text:
                action_fill = attach_fill
            elif "tháo" in action_text:
                action_fill = remove_fill

        shelf_fill = None
        if offset < len(store_values) and offset < len(pog_values):
            store = "" if pd.isna(store_values[offset]) else str(store_values[offset]).strip()
            pog = "" if pd.isna(pog_values[offset]) else str(pog_values[offset]).strip()
            if store and pog:
                shelf = (store, pog)
                if shelf not in shelf_color_index:
                    shelf_color_index[shelf] = len(shelf_color_index) % len(shelf_fills)
                shelf_fill = shelf_fills[shelf_color_index[shelf]]

        row_fill = shelf_fill or action_fill or base_fill
        shelf_row = shelf_fill is not None
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            value = cell.value
            if value is not None:
                text = str(value)
                if len(text) > widths[col_idx - 1]:
                    widths[col_idx - 1] = len(text)
                cell.alignment = data_alignment
                cell.border = border
                cell.fill = row_fill
                if isinstance(value, str) and value.startswith("="):
                    cell.number_format = "0"
                header = header_text[col_idx - 1].lower()
                if header in _DATE_HEADERS:
                    if isinstance(value, pd.Timestamp):
                        cell.number_format = "dd/mm/yyyy"
                    elif isinstance(value, str):
                        try:
                            pd.to_datetime(value)
                            cell.number_format = "dd/mm/yyyy"
                        except Exception:
                            pass
            elif shelf_row:
                # Original shelf pass colors the complete Excel row, including blanks.
                cell.fill = row_fill

        if completeness_col is not None and offset < len(location_values):
            value = location_values[offset]
            if isinstance(value, str) and value != "Đủ vị trí":
                cell = ws.cell(row=row_idx, column=completeness_col)
                cell.fill = review_fill
                cell.font = review_font

        if offset < len(collision_values) and _truthy(collision_values[offset]):
            for column in collision_columns:
                cell = ws.cell(row=row_idx, column=column)
                cell.fill = warning_fill
                cell.font = warning_font

        if offset < len(review_values) and _truthy(review_values[offset]):
            for column in review_columns:
                cell = ws.cell(row=row_idx, column=column)
                cell.fill = review_fill
                cell.font = review_font

        ws.row_dimensions[row_idx].height = 22

    ws.sheet_properties.tabColor = _COLOR_MAP.get(sheet_name, "4F81BD")
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = min(60, max(12, width + 2))


def _shelf_pass_already_applied(_ws, _df) -> None:
    """Shelf/review/collision final fills are folded into the main style pass."""
    return None


def _optimized_action_inputs(ws) -> None:
    headers = [str(cell.value or "") for cell in ws[1]]
    if "processing_status" not in headers:
        return
    status_col = headers.index("processing_status") + 1
    editable = [name for name in ["processing_status", "processed_by", "processed_at", "staff_note"] if name in headers]
    validation = DataValidation(
        type="list",
        formula1='"Chưa xử lý,Đã xong,Không tìm thấy,Không thực hiện được"',
        allow_blank=False,
    )
    validation.error = "Hãy chọn một trạng thái trong danh sách."
    validation.errorTitle = "Trạng thái không hợp lệ"
    ws.add_data_validation(validation)

    yellow_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    unlocked = Protection(locked=False)
    if ws.max_row >= 2:
        status_letter = get_column_letter(status_col)
        validation.add(f"{status_letter}2:{status_letter}{ws.max_row}")
        full_range = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"
        ws.conditional_formatting.add(
            full_range,
            FormulaRule(
                formula=[f'${status_letter}2="Đã xong"'],
                fill=PatternFill(fill_type="solid", fgColor="E2F0D9"),
            ),
        )
        ws.conditional_formatting.add(
            full_range,
            FormulaRule(
                formula=[f'${status_letter}2="Không tìm thấy"'],
                fill=PatternFill(fill_type="solid", fgColor="FFC7CE"),
            ),
        )

    editable_indexes = [(name, headers.index(name) + 1) for name in editable]
    for row_idx in range(2, ws.max_row + 1):
        for name, col_idx in editable_indexes:
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.protection = unlocked
            cell.fill = yellow_fill
            if name == "processed_at":
                cell.number_format = "dd/mm/yyyy hh:mm"
    ws.protection.autoFilter = True
    ws.protection.sort = True
    ws.protection.sheet = True


def _optimized_confirmation_inputs(ws) -> None:
    headers = [str(cell.value or "") for cell in ws[1]]
    if "xac_nhan_ap_dung" not in headers:
        return
    confirm_col = headers.index("xac_nhan_ap_dung") + 1
    validation = DataValidation(
        type="list",
        formula1='"Chưa xác nhận,Áp dụng CTKM,Áp dụng Clearance,Áp dụng cả hai,Không áp dụng"',
        allow_blank=False,
    )
    validation.error = "Hãy chọn một trạng thái xác nhận trong danh sách."
    validation.errorTitle = "Trạng thái không hợp lệ"
    ws.add_data_validation(validation)

    yellow_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    unlocked = Protection(locked=False)
    if ws.max_row >= 2:
        letter = get_column_letter(confirm_col)
        validation.add(f"{letter}2:{letter}{ws.max_row}")
        note_col = headers.index("ghi_chu_xac_nhan") + 1
        for row_idx in range(2, ws.max_row + 1):
            for col_idx in (confirm_col, note_col):
                cell = ws.cell(row_idx, col_idx)
                cell.fill = yellow_fill
                cell.protection = unlocked
    ws.protection.autoFilter = True
    ws.protection.sort = True
    ws.protection.sheet = True


def install_export_speed_contract(core_module) -> None:
    if getattr(core_module.app.state, "export_speed_contract_v1", False):
        return
    main_core = getattr(core_module._main_mod, "_core", None)
    exporter = getattr(main_core, "EX", None) if main_core is not None else None
    exporter_core = getattr(exporter, "_core", None) if exporter is not None else None
    if exporter_core is None:
        raise RuntimeError("Export-speed contract không tìm thấy exporter core")

    _optimized_professional_style._ctkm_export_speed_contract_v1 = True
    _shelf_pass_already_applied._ctkm_export_speed_contract_v1 = True
    exporter_core._apply_professional_style = _optimized_professional_style
    exporter_core._apply_alternating_shelf_colors = _shelf_pass_already_applied
    exporter_core._configure_action_inputs = _optimized_action_inputs
    exporter_core._configure_confirmation_inputs = _optimized_confirmation_inputs

    # Keep the facade surface coherent for diagnostics/tests. The core exporter
    # function resolves helpers from exporter_core globals at runtime.
    exporter._apply_professional_style = _optimized_professional_style
    exporter._apply_alternating_shelf_colors = _shelf_pass_already_applied
    exporter._configure_action_inputs = _optimized_action_inputs
    exporter._configure_confirmation_inputs = _optimized_confirmation_inputs
    core_module.app.state.export_speed_contract_v1 = True
