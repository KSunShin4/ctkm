# -*- coding: utf-8 -*-
"""Render-oriented CTKM workbook export without sacrificing source auditability.

The production profile avoids expensive diagnostic styling, but it must not
silently remove reconciliation/source evidence.  Operational sheets stay compact;
raw passthrough columns remain available in ``Audit hành động`` and source/control
metadata remains hidden for audit.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from modules.all_km_reconciliation import (
    _frame_keys,
    _identity,
    build_all_km_reconciliation_frames,
)
from modules.ctkm_performance import align_task_summary_with_work_rows
from backend.fast_slim_excel import write_slim_workbook


SLIM_SHEET_ORDER = [
    "HƯỚNG DẪN",
    "TỔNG HỢP - TEM",
    "Gắn tem",
    "Tháo tem",
    "Chờ gắn",
    "Cần review",
    "CTKM thieu vi tri ke",
    "POG cần kiểm tra",
    "Tong hop",
    "POG không có KM",
    "POG không có Monthly CTKM",
    "POG cần review",
    "Audit hành động",
    "Mapping log",
    "POG không có CTKM",
    "POG metadata",
    "KM metadata",
]


def _truthy(value) -> bool:
    return value is True or str(value).strip().lower() in {"true", "1", "yes"}


def _in_scope(actions: pd.DataFrame) -> pd.Series:
    if "in_report_scope" not in actions.columns:
        return pd.Series(True, index=actions.index)
    return actions["in_report_scope"].fillna(False).astype(bool)


def _review_mask(actions: pd.DataFrame, exporter) -> pd.Series:
    if "review_required" in actions.columns:
        return actions["review_required"].map(_truthy)
    field = "scheduled_action" if "scheduled_action" in actions.columns else "action"
    if field not in actions.columns:
        return pd.Series(False, index=actions.index)
    return actions[field].apply(exporter._normalize_action_text).eq("review")


def _prepare_frames(exporter, df_fefo_actions: pd.DataFrame | None):
    if df_fefo_actions is None or df_fefo_actions.empty:
        base = pd.DataFrame(columns=[
            "task_id", "item_id", "item_name", "promo_uom", "program_source",
            "campaign_name", "campaign_detail", "promo_from", "promo_to",
            "normal_price", "promo_price", "key_channel", "channel", "product_type",
            "source_sheet", "scheduled_action", "action", "reason", "target_month",
            "location_status", "review_required", "review_reason",
        ])
        actions = base.copy()
    else:
        actions = df_fefo_actions.copy()

    for required_column in ("key_channel", "channel", "item_id", "item_name", "promo_uom"):
        if required_column not in actions.columns:
            actions[required_column] = pd.NA

    scope = _in_scope(actions)
    action_field = "scheduled_action" if "scheduled_action" in actions.columns else "action"
    normalized = (
        actions[action_field].apply(exporter._normalize_action_text)
        if action_field in actions.columns
        else pd.Series("", index=actions.index)
    )

    gan = actions[scope & normalized.isin({"gan tem", "gan tem fefo"})].copy()
    thao = actions[scope & normalized.isin({"thao tem", "thao tem fefo", "thao rem fefo"})].copy()
    wait = actions[scope & normalized.eq("cho gan tem")].copy()
    review = actions[scope & _review_mask(actions, exporter)].copy()

    if "location_completeness" in actions.columns:
        location_issues = actions[scope & ~actions["location_completeness"].eq("Đủ vị trí")].copy()
    elif "location_status" in actions.columns:
        location_issues = actions[scope & ~actions["location_status"].eq("Đủ vị trí")].copy()
    else:
        location_issues = pd.DataFrame(columns=actions.columns)

    if "location_needs_review" in actions.columns:
        location_review = actions[
            scope & actions["location_needs_review"].fillna(False).astype(bool)
        ].copy()
    else:
        # Completeness and review are independent location-quality axes.
        # Missing/partial coordinates belong to CTKM thieu vi tri ke, but without
        # the canonical review flag there is no evidence that a row belongs to
        # POG cần kiểm tra. Never manufacture review rows by copying the issue set.
        location_review = pd.DataFrame(columns=actions.columns)

    for frame in (gan, thao, wait, review):
        if "so_ngay_con_lai" not in frame.columns:
            frame["so_ngay_con_lai"] = ""

    gan = exporter._prepare_action_sheet(exporter._add_work_tracking_columns(gan))
    thao = exporter._prepare_action_sheet(exporter._add_work_tracking_columns(thao))
    wait = exporter._prepare_action_sheet(exporter._add_work_tracking_columns(wait))
    review = exporter._prepare_action_sheet(exporter._add_work_tracking_columns(review))
    location_issues = exporter._prepare_action_sheet(location_issues)
    location_review = exporter._prepare_action_sheet(location_review)
    audit = exporter._prepare_action_sheet(actions)
    total = exporter._build_total_action_frame(actions)
    return total, gan, thao, wait, review, location_issues, location_review, audit


def _lightweight_audit_style(ws, exporter) -> None:
    """Keep audit/reconciliation usable without expensive per-cell body styling."""
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
    fill = PatternFill(fill_type="solid", fgColor="5B5EA6")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, cell in enumerate(ws[1], start=1):
        header = str(cell.value or "")
        width = getattr(exporter, "_COLUMN_WIDTHS", {}).get(
            header, min(30, max(12, len(header) + 2))
        )
        ws.column_dimensions[exporter.get_column_letter(index)].width = width
    exporter._apply_bilingual_headers(ws)


def _trim_guide_to_slim_profile(workbook) -> None:
    """Remove only cards for user-facing sheets still omitted by slim profile."""
    ws = workbook["HƯỚNG DẪN"]
    rows_to_delete = []
    for row_idx in range(1, ws.max_row + 1):
        if ws.cell(row_idx, 1).value in {"Xác nhận CTKM / Clearance"}:
            rows_to_delete.append(row_idx)
    for row_idx in reversed(rows_to_delete):
        ws.delete_rows(row_idx, 1)
    ws["D10"] = (
        "Bản production slim: giảm styling/diagnostic nặng nhưng vẫn giữ "
        "reconciliation và source audit để không mất dữ liệu nguồn."
    )


def _safe_frame(exporter, frame: pd.DataFrame | None, message: str) -> pd.DataFrame:
    source = frame if frame is not None else pd.DataFrame()
    return exporter._safe(source, message)


def _metadata_from(frame: pd.DataFrame | None, attr_name: str) -> pd.DataFrame | None:
    if frame is None:
        return None
    metadata = getattr(frame, "attrs", {}).get(attr_name)
    return metadata.copy() if isinstance(metadata, pd.DataFrame) and not metadata.empty else None


def _write_hidden_metadata(writer, sheet_name: str, frame: pd.DataFrame | None) -> None:
    if frame is None or frame.empty:
        return
    frame.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.sheets[sheet_name]
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
    ws.sheet_state = "hidden"


def _update_all_km_summary(
    main_core,
    summary: dict | None,
    matched_all: pd.DataFrame,
    monthly_pog_only: pd.DataFrame | None,
    missing_all_km: pd.DataFrame,
    actions: pd.DataFrame | None,
    added_keys: set[str],
) -> None:
    if not isinstance(summary, dict):
        return

    scoped_actions = actions.copy() if actions is not None else pd.DataFrame()
    if not scoped_actions.empty and "in_report_scope" in scoped_actions.columns:
        scoped_actions = scoped_actions[
            scoped_actions["in_report_scope"].fillna(False).astype(bool)
        ].copy()
    identity_input = scoped_actions if not scoped_actions.empty else matched_all
    identity = _identity(main_core, identity_input)

    monthly_keys = set(_frame_keys(main_core, identity, monthly_pog_only).dropna())
    matched_keys = set(_frame_keys(main_core, identity, matched_all).dropna())
    missing_keys = set(_frame_keys(main_core, identity, missing_all_km).dropna())

    summary["So dong khop (Matched)"] = len(matched_all)
    summary["So item duy nhat khop"] = len(matched_keys)
    summary["So dong POG khong co CTKM Monthly phu hop"] = (
        0 if monthly_pog_only is None else len(monthly_pog_only)
    )
    summary["So item POG khong co CTKM Monthly phu hop"] = len(monthly_keys)
    summary["So dong POG khong co KM phu hop"] = len(missing_all_km)
    summary["So item POG khong co KM phu hop"] = len(missing_keys)
    summary["So item Clearance bo sung vao Tong hop"] = len(added_keys)
    summary["So item KM+POG thieu trong Tong hop"] = 0

    # Legacy API aliases stay aligned with the canonical all-KM meaning.
    if "So dong POG khong co CTKM phu hop" in summary:
        summary["So dong POG khong co CTKM phu hop"] = len(missing_all_km)
    if "So item POG khong co CTKM phu hop" in summary:
        summary["So item POG khong co CTKM phu hop"] = len(missing_keys)


_SELECTABLE_PRESENTATION_FIELDS = {
    "item_id", "item_name", "promo_uom", "key_channel", "channel",
    "campaign_name", "campaign_detail", "promo_from", "promo_to",
    "normal_price", "promo_price", "product_type",
    "store_id", "pog_code", "location_id", "facing",
}
_PINNED_PRESENTATION_FIELDS = {
    "key_channel", "channel", "item_id", "item_name", "promo_uom",
}


def _apply_selected_output_visibility(ws, exporter, selected: set[str]) -> None:
    if not selected:
        return
    for index, cell in enumerate(ws[1], start=1):
        header = str(cell.value or "").strip()
        if header in _PINNED_PRESENTATION_FIELDS:
            ws.column_dimensions[exporter.get_column_letter(index)].hidden = False
            continue
        if header in _SELECTABLE_PRESENTATION_FIELDS:
            ws.column_dimensions[exporter.get_column_letter(index)].hidden = (
                header.lower() not in selected
            )


def _mapping_log_frame(mapping_logs) -> pd.DataFrame:
    rows = []
    for log in mapping_logs or []:
        sheet_name = log.get("sheet_name", "")
        source = log.get("source", "")
        confidence = log.get("confidence") or {}
        low = set(log.get("low_confidence_fields") or [])
        mapping = log.get("mapping") or {}
        for source_column, canonical_field in mapping.items():
            rows.append({
                "sheet": sheet_name,
                "cot_goc": source_column,
                "canonical_field": canonical_field,
                "confidence": confidence.get(source_column),
                "can_review": source_column in low,
                "nguon": source,
            })
        for source_column in log.get("unmapped") or []:
            rows.append({
                "sheet": sheet_name,
                "cot_goc": source_column,
                "canonical_field": "(KHONG MAP DUOC)",
                "confidence": confidence.get(source_column),
                "can_review": True,
                "nguon": source,
            })
    if rows:
        return pd.DataFrame(rows)
    return pd.DataFrame({"Ghi chu": ["Chưa có mapping log."]})


def install_slim_output_contract(core_module) -> None:
    if getattr(core_module.app.state, "slim_output_contract_v1", False):
        return

    main_mod = core_module._main_mod
    main_core = getattr(main_mod, "_core", main_mod)
    exporter = main_mod.EX

    def export_final_report_slim(
        output_path,
        df_matched,
        df_ctkm_only,
        df_pog_only,
        overlap_warnings,
        mapping_logs,
        reconciliation_summary,
        snapshot_alignment_warning=None,
        selected_output_columns=None,
        df_fefo_actions=None,
        df_ctkm_clearance_confirmation=None,
        df_pog_review=None,
    ) -> str:
        # These are intentionally unused in slim presentation, but do not discard
        # df_matched/df_pog_only: they are required for all-KM reconciliation.
        del df_ctkm_only, overlap_warnings
        del snapshot_alignment_warning

        matched_all, missing_all_km, added_keys = build_all_km_reconciliation_frames(
            main_core, df_matched, df_pog_only, df_fefo_actions
        )
        _update_all_km_summary(
            main_core,
            reconciliation_summary,
            matched_all,
            df_pog_only,
            missing_all_km,
            df_fefo_actions,
            added_keys,
        )

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        align_task_summary_with_work_rows(reconciliation_summary, df_fefo_actions)
        total, gan, thao, wait, review, location_issues, location_review, audit = _prepare_frames(
            exporter, df_fefo_actions
        )

        # selected_output_columns is presentation-only; all workbook data columns
        # remain present so audit/identity validation cannot be weakened by a UI choice.
        selected = {
            str(column).strip().lower()
            for column in (selected_output_columns or [])
            if str(column).strip()
        }

        pog_review = df_pog_review.copy() if df_pog_review is not None else pd.DataFrame()
        pog_metadata = _metadata_from(df_pog_review, "pog_metadata")
        km_metadata = _metadata_from(df_ctkm_clearance_confirmation, "fact_metadata")

        return write_slim_workbook(
            output_path=output_path,
            exporter=exporter,
            total=total,
            gan=gan,
            thao=thao,
            wait=wait,
            review=review,
            location_issues=location_issues,
            location_review=location_review,
            matched_all=matched_all,
            missing_all_km=missing_all_km,
            monthly_pog_only=df_pog_only,
            pog_review=pog_review,
            audit=audit,
            mapping_log=_mapping_log_frame(mapping_logs),
            reconciliation_summary=reconciliation_summary,
            selected=selected,
            pog_metadata=pog_metadata,
            km_metadata=km_metadata,
        )

    exporter.export_final_report = export_final_report_slim
    # main_core.run_pipeline holds the same exporter module object, so replacing
    # the module callable is sufficient for all backend CTKM jobs.
    core_module.app.state.slim_output_contract_v1 = True
    core_module.app.state.ctkm_output_profile = "production-slim"
