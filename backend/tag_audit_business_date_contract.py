# -*- coding: utf-8 -*-
"""Business-date contract for the physical shelf-tag scanner.

The exported workbook is a monthly source snapshot, while Gắn/Tháo/Chờ is a
business-date decision. A store must therefore be able to reuse a same-month
workbook on later days without trusting the action sheet that happened to be
correct on export day.

This layer preserves the workbook's original values and action-sheet identity for
audit, then creates shallow scanner-only row copies whose ``expected_action`` is
derived from ``promo_from`` / ``promo_to`` as of the current Viet Nam date.
Known report-month mismatches remain blocking because the monthly source itself is
then incompatible; a merely older ``business_date`` is informational, not a block.
Parsed workbooks must contain both report-month and business-date metadata; missing,
incomplete, or unreadable metadata is fail-closed instead of trusting snapshot actions.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO
import re
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


_GUIDE_SHEET = "HƯỚNG DẪN"
_BUSINESS_DATE_KEY = "_report_business_date"
_REPORT_MONTH_KEY = "_report_month"
_SNAPSHOT_ACTION_KEY = "_snapshot_expected_action"
_SCANNER_DATE_KEY = "_scanner_action_date"
_SCANNER_REASON_KEY = "_scanner_action_reason"
_METADATA_STATUS_KEY = "_report_metadata_status"
_METADATA_VALID = "valid"
_METADATA_BLOCKING = {"missing", "incomplete", "invalid"}
_VN_TZ = ZoneInfo("Asia/Ho_Chi_Minh")


def _copyable_source(source):
    if isinstance(source, (bytes, bytearray)):
        return BytesIO(bytes(source))
    getvalue = getattr(source, "getvalue", None)
    if callable(getvalue):
        try:
            return BytesIO(getvalue())
        except Exception:
            pass
    return source


def _iso_date(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value)
            if isinstance(converted, datetime):
                return converted.date().isoformat()
            if isinstance(converted, date):
                return converted.isoformat()
        except Exception:
            return None
    text = str(value).strip()
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return None


def read_report_metadata(source) -> dict:
    """Read exporter metadata without depending on fixed cell coordinates."""
    workbook = None
    try:
        workbook = load_workbook(_copyable_source(source), data_only=True, read_only=True)
        if _GUIDE_SHEET not in workbook.sheetnames:
            return {}
        ws = workbook[_GUIDE_SHEET]
        metadata: dict[str, str] = {}
        for label, value in ws.iter_rows(min_row=1, max_row=24, min_col=1, max_col=2, values_only=True):
            key = str(label or "").strip()
            if key == "Tháng báo cáo":
                month = str(value or "").strip()
                if re.fullmatch(r"\d{4}-\d{2}", month):
                    metadata["report_month"] = month
            elif key == "Ngày nghiệp vụ":
                parsed = _iso_date(value)
                if parsed:
                    metadata["business_date"] = parsed
        return metadata
    finally:
        if workbook is not None:
            workbook.close()


def current_vn_business_date() -> str:
    return datetime.now(_VN_TZ).date().isoformat()


def metadata_from_rows(rows: list[dict]) -> dict:
    for row in rows or []:
        business_date = str(row.get(_BUSINESS_DATE_KEY) or "").strip()
        report_month = str(row.get(_REPORT_MONTH_KEY) or "").strip()
        status = str(row.get(_METADATA_STATUS_KEY) or "").strip().lower()
        if business_date or report_month or status:
            if not status:
                status = _METADATA_VALID if business_date and report_month else "legacy_unmarked"
            return {
                "business_date": business_date or None,
                "report_month": report_month or None,
                "metadata_status": status,
            }
    return {
        "business_date": None,
        "report_month": None,
        "metadata_status": "legacy_unmarked",
    }


def _metadata_block_reason(rows: list[dict], today: str) -> str | None:
    metadata = metadata_from_rows(rows)
    status = metadata.get("metadata_status")
    if status in _METADATA_BLOCKING:
        return "missing_report_metadata"
    report_month = metadata.get("report_month")
    if report_month and report_month != str(today)[:7]:
        return "report_month_mismatch"
    return None


def derive_scanner_action(row: dict, business_date: str) -> tuple[str, str]:
    """Derive today's scanner action without mutating source business values."""
    today_text = _iso_date(business_date)
    promo_from_text = _iso_date(row.get("promo_from"))
    promo_to_text = _iso_date(row.get("promo_to"))
    if not today_text:
        return "review", "scanner_business_date_invalid"
    if not promo_from_text or not promo_to_text:
        return "review", "missing_or_invalid_promo_dates"

    today = date.fromisoformat(today_text)
    promo_from = date.fromisoformat(promo_from_text)
    promo_to = date.fromisoformat(promo_to_text)
    if promo_from > promo_to:
        return "review", "promo_from_after_promo_to"
    if today < promo_from:
        return "cho_gan", "before_attach_date"
    if today <= promo_to:
        return "gan_tem", "promotion_active"
    if today >= promo_to + timedelta(days=1):
        return "thao_tem", "removal_date_reached"
    return "review", "scanner_action_undetermined"


def scanner_rows_for_current_date(rows: list[dict], business_date: str | None = None) -> list[dict]:
    """Return scanner-only shallow copies reclassified for the requested date."""
    # Parsed workbooks are fail-closed unless both required metadata fields are
    # valid. Direct synthetic rows used by lower-level tests remain legacy-unmarked.
    metadata = metadata_from_rows(rows)
    if metadata.get("metadata_status") in _METADATA_BLOCKING:
        return [dict(row) for row in rows or []]
    if not metadata.get("business_date") and not metadata.get("report_month"):
        return [dict(row) for row in rows or []]

    today = _iso_date(business_date) if business_date is not None else current_vn_business_date()
    today = today or current_vn_business_date()
    derived: list[dict] = []
    for source_row in rows or []:
        row = dict(source_row)
        row[_SNAPSHOT_ACTION_KEY] = source_row.get("expected_action")
        action, reason = derive_scanner_action(source_row, today)
        row["expected_action"] = action
        row[_SCANNER_DATE_KEY] = today
        row[_SCANNER_REASON_KEY] = reason
        derived.append(row)
    return derived


def _known_report_month_mismatch(rows: list[dict], today: str) -> bool:
    return _metadata_block_reason(rows, today) == "report_month_mismatch"


def install_tag_audit_business_date_contract(tag_audit_module) -> None:
    if getattr(tag_audit_module, "_business_date_contract_v1", False):
        return

    original_parse = tag_audit_module._parse_action_workbook
    original_context_payload = tag_audit_module._context_payload
    original_compare = tag_audit_module.compare_observation

    def parse_with_report_metadata(source):
        metadata_source = _copyable_source(source)
        rows = original_parse(source)
        metadata_status = _METADATA_VALID
        try:
            metadata = read_report_metadata(metadata_source)
        except Exception:
            metadata = {}
            metadata_status = "invalid"
        business_date = metadata.get("business_date")
        report_month = metadata.get("report_month")
        if metadata_status != "invalid":
            if business_date and report_month:
                metadata_status = _METADATA_VALID
            elif business_date or report_month:
                metadata_status = "incomplete"
            else:
                metadata_status = "missing"
        for row in rows:
            row[_BUSINESS_DATE_KEY] = business_date
            row[_REPORT_MONTH_KEY] = report_month
            row[_METADATA_STATUS_KEY] = metadata_status
        return rows

    def context_payload_with_business_date(cache: dict) -> dict:
        payload = original_context_payload(cache)
        rows = cache.get("rows") or []
        metadata = metadata_from_rows(rows)
        today = current_vn_business_date()
        business_date = metadata.get("business_date")
        report_month = metadata.get("report_month")
        payload["report_month"] = report_month
        payload["business_date"] = business_date
        payload["scanner_business_date"] = today
        block_reason = _metadata_block_reason(rows, today)
        payload["business_date_stale"] = bool(business_date and business_date != today)
        payload["business_date_blocked"] = block_reason is not None
        payload["business_date_block_reason"] = block_reason
        payload["scanner_reclassifies_actions"] = (
            metadata.get("metadata_status") == _METADATA_VALID
        )
        return payload

    def compare_with_current_business_date(rows, observed, manual_item_id=""):
        today = current_vn_business_date()
        metadata = metadata_from_rows(rows)
        report_month = metadata.get("report_month")
        block_reason = _metadata_block_reason(rows, today)
        if block_reason == "missing_report_metadata":
            raise tag_audit_module.HTTPException(
                status_code=409,
                detail=(
                    "Workbook kiểm tem thiếu hoặc không đọc được metadata Tháng báo cáo/Ngày nghiệp vụ. "
                    "Hãy chạy hoặc tải workbook CTKM do hệ thống xuất trước khi kiểm tem."
                ),
            )
        if block_reason == "report_month_mismatch":
            raise tag_audit_module.HTTPException(
                status_code=409,
                detail=(
                    f"Workbook kiểm tem thuộc tháng {report_month}; tháng kiểm hiện tại là {today[:7]}. "
                    "Hãy chạy hoặc tải workbook CTKM đúng tháng trước khi kiểm tem."
                ),
            )

        scanner_rows = scanner_rows_for_current_date(rows, today)
        result = original_compare(scanner_rows, observed, manual_item_id=manual_item_id)
        expected = result.get("expected") or {}
        if str(expected.get("expected_action") or "") == "review":
            result["verdict"] = {
                "code": "review_business_date",
                "title": "Cần kiểm tra ngày CTKM",
                "detail": "Dữ liệu ngày chương trình thiếu hoặc không hợp lệ nên chưa thể xác định Gắn/Tháo/Chờ.",
                "severity": "warning",
            }
            mismatches = list(result.get("mismatches") or [])
            if "business_date" not in mismatches:
                mismatches.append("business_date")
            result["mismatches"] = mismatches
        return result

    tag_audit_module._parse_action_workbook = parse_with_report_metadata
    tag_audit_module._context_payload = context_payload_with_business_date
    tag_audit_module.compare_observation = compare_with_current_business_date
    tag_audit_module.read_tag_audit_report_metadata = read_report_metadata
    tag_audit_module.tag_audit_current_business_date = current_vn_business_date
    tag_audit_module.derive_scanner_action = derive_scanner_action
    tag_audit_module.scanner_rows_for_current_date = scanner_rows_for_current_date
    tag_audit_module._business_date_contract_v1 = True
