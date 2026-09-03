# -*- coding: utf-8 -*-
"""Contracts for Tem R and other auxiliary workflows.

CTKM already validates inputs before queue acceptance and verifies its final
workbook before publishing it. Tem R historically accepted stale sheet/mapping
configuration and only discovered it inside the worker; its inspect endpoint also
left temporary directories behind. This module brings Tem R to the same lifecycle
without changing its allocation algorithm.
"""
from __future__ import annotations

import asyncio
from datetime import datetime
import json
from pathlib import Path
import shutil
import time
import uuid

import openpyxl
import pandas as pd
from fastapi import Depends, File, Form, HTTPException, UploadFile
from typing import Optional


def _remove_tree(path: Path) -> None:
    try:
        shutil.rmtree(path, ignore_errors=True)
    except OSError:
        pass


def _replace_route(core_module, path: str, method: str, endpoint) -> None:
    for route in core_module.app.routes:
        if getattr(route, "path", None) != path:
            continue
        if method not in (getattr(route, "methods", None) or set()):
            continue
        route.endpoint = endpoint
        dependant = getattr(route, "dependant", None)
        if dependant is not None:
            dependant.call = endpoint
        return
    raise RuntimeError(f"Không tìm thấy route {method} {path} để cài auxiliary contract")


def _parse_mapping(raw: Optional[str]) -> dict | None:
    try:
        mapping = json.loads(raw) if raw else None
    except json.JSONDecodeError as exc:
        raise ValueError(f"Ánh xạ cột Tem R không phải JSON hợp lệ: {exc.msg}") from exc
    if mapping is not None and not isinstance(mapping, dict):
        raise ValueError("Ánh xạ cột Tem R phải là JSON object")
    return mapping


def _resolve_tem_r_selection(
    tem_r_mod,
    source_path: Path,
    *,
    stock_column: str,
    column_mapping: dict | None,
    sheet_name: str | None,
    header_row: int | None,
) -> tuple[str, int, dict[str, str], dict]:
    """Resolve and validate sheet/header/mapping before a queue slot is accepted."""
    inspection = tem_r_mod.inspect_tem_r_workbook(source_path)
    options = [
        option for option in inspection.get("sheet_options", [])
        if option.get("columns")
    ]
    if not options:
        raise ValueError("File Tem R không có sheet dữ liệu khả dụng.")

    selected_name = str(sheet_name or "").strip()
    if selected_name:
        selected = next((item for item in options if item.get("sheet_name") == selected_name), None)
        if selected is None:
            available = ", ".join(str(item.get("sheet_name")) for item in options)
            raise ValueError(
                f"Không tìm thấy sheet Tem R '{selected_name}'. Sheet khả dụng: {available}"
            )
    elif len(options) == 1:
        selected = options[0]
        selected_name = str(selected.get("sheet_name") or "")
    else:
        available = ", ".join(str(item.get("sheet_name")) for item in options)
        raise ValueError(
            "File Tem R có nhiều sheet dữ liệu. Hãy chọn rõ sheet trước khi phân bổ: "
            + available
        )

    detected_header = selected.get("header_row")
    if detected_header is None:
        raise ValueError(f"Không xác định được dòng tiêu đề của sheet '{selected_name}'.")
    detected_header = int(detected_header)
    if header_row is not None and int(header_row) != detected_header:
        raise ValueError(
            f"Dòng tiêu đề đã chọn ({header_row}) không còn khớp file hiện tại "
            f"(phát hiện dòng {detected_header}). Hãy đọc lại file Tem R."
        )

    columns = [str(value) for value in selected.get("columns", [])]
    stock = str(stock_column or "").strip()
    if not stock:
        raise ValueError("Hãy chọn cột tồn kho")
    if stock not in columns:
        raise ValueError(f"Không tìm thấy cột tồn kho đã chọn: {stock}")

    required = tuple(tem_r_mod.REQUIRED_COLUMNS)
    mapping = dict(column_mapping or {})
    if not mapping:
        # Backward-compatible exact-header path: no inference beyond exact names.
        mapping = {canonical: canonical for canonical in required if canonical in columns}

    missing_targets = [canonical for canonical in required if not str(mapping.get(canonical) or "").strip()]
    if missing_targets:
        raise ValueError("Hãy xác nhận đủ cột Tem R: " + ", ".join(missing_targets))

    resolved = {canonical: str(mapping[canonical]).strip() for canonical in required}
    missing_sources = sorted({source for source in resolved.values() if source not in columns})
    if missing_sources:
        raise ValueError("Không tìm thấy cột đã chọn: " + ", ".join(missing_sources))
    duplicate_sources = sorted({
        source for source in resolved.values()
        if list(resolved.values()).count(source) > 1
    })
    if duplicate_sources:
        raise ValueError("Mỗi trường Tem R phải dùng một cột riêng: " + ", ".join(duplicate_sources))
    if stock in resolved.values():
        canonical = next(name for name, source in resolved.items() if source == stock)
        raise ValueError(
            f"Cột tồn kho '{stock}' đang đồng thời được dùng cho {canonical}. "
            "Hãy chọn hai cột khác nhau."
        )

    # Inspect already validates the header. Also require at least one real ITEMID
    # below it so a title/control-only sheet is rejected before queue acceptance.
    workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=True, keep_links=False)
    try:
        worksheet = workbook[selected_name]
        item_col = columns.index(resolved["ITEMID"]) + 1
        has_item = False
        for row in worksheet.iter_rows(
            min_row=detected_header + 1,
            min_col=item_col,
            max_col=item_col,
            values_only=True,
        ):
            value = row[0] if row else None
            text = "" if value is None else str(value).strip()
            if text and text.upper() not in {"ITEMID", "TOTAL", "GRAND TOTAL", "TỔNG", "TONG"}:
                has_item = True
                break
        if not has_item:
            raise ValueError(
                f"Sheet Tem R '{selected_name}' không có dòng sản phẩm ITEMID để phân bổ."
            )
    finally:
        workbook.close()

    return selected_name, detected_header, resolved, inspection


def validate_tem_r_workbook(
    output_path: str | Path,
    plan,
    *,
    sheet_name: str,
    header_row: int,
) -> None:
    """Verify the saved workbook matches the allocation plan before publication."""
    path = Path(output_path)
    if not path.is_file():
        raise RuntimeError("Tem R kết thúc nhưng không tạo được file kết quả.")

    workbook = None
    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=False, keep_links=False)
        required_sheets = {sheet_name, "Tem R phân bổ", "Tem R ngoại lệ"}
        missing = sorted(required_sheets - set(workbook.sheetnames))
        if missing:
            raise RuntimeError("Workbook Tem R thiếu sheet bắt buộc: " + ", ".join(missing))

        source = workbook[sheet_name]
        headers = [
            "" if source.cell(header_row, column).value is None
            else str(source.cell(header_row, column).value).strip()
            for column in range(1, (source.max_column or 0) + 1)
        ]
        if "TEM R" not in headers:
            raise RuntimeError("Workbook Tem R không có cột 'TEM R' sau khi xử lý.")
        tem_r_col = headers.index("TEM R") + 1

        if "_excel_row" in plan.rows.columns:
            for _, row in plan.rows.iterrows():
                excel_row = int(row["_excel_row"])
                expected_value = row.get("TEM R")
                expected = None if pd.isna(expected_value) else str(expected_value)
                actual_value = source.cell(excel_row, tem_r_col).value
                actual = None if actual_value is None else str(actual_value)
                if expected != actual:
                    raise RuntimeError(
                        f"Workbook Tem R không khớp kế hoạch tại dòng Excel {excel_row}: "
                        f"plan={expected!r}, file={actual!r}"
                    )

        audit = workbook["Tem R phân bổ"]
        exceptions = workbook["Tem R ngoại lệ"]
        audit_rows = max(0, (audit.max_row or 0) - 1)
        exception_rows = max(0, (exceptions.max_row or 0) - 1)
        if audit_rows != len(plan.audit):
            raise RuntimeError(
                f"Sheet 'Tem R phân bổ' không khớp plan: file={audit_rows}, plan={len(plan.audit)}"
            )
        if exception_rows != len(plan.exceptions):
            raise RuntimeError(
                f"Sheet 'Tem R ngoại lệ' không khớp plan: file={exception_rows}, plan={len(plan.exceptions)}"
            )
    except RuntimeError:
        raise
    except Exception as exc:
        raise RuntimeError(f"Workbook Tem R không đọc được: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _tem_r_summary(plan) -> dict:
    return {
        "tem_r_total_rows": len(plan.rows),
        "tem_r_replaced_rows": int(plan.rows["TEM_R_REPRESENTATIVE"].notna().sum()),
        "tem_r_skipped_stock_rows": int(plan.rows["_stock"].isna().sum()),
        "tem_r_exceptions": len(plan.exceptions),
    }


def install_auxiliary_workflows_contract(core_module) -> None:
    if getattr(core_module.app.state, "auxiliary_workflows_contract_v1", False):
        return

    tem_r_mod = core_module._tem_r_mod

    async def inspect_tem_r_file(
        tem_r_file: UploadFile = File(...),
        _admin: str = Depends(core_module.require_admin),
    ):
        if not tem_r_file.filename or not tem_r_file.filename.lower().endswith(".xlsx"):
            raise HTTPException(status_code=422, detail="tem_r_file phải là file .xlsx")
        content = await core_module._read_upload_limited(
            tem_r_file, core_module._MAX_TEM_R_BYTES, "File Tem R"
        )
        job_dir = core_module._TMP_DIR / f"inspect_tem_r_{uuid.uuid4().hex[:8]}"
        job_dir.mkdir(parents=True, exist_ok=True)
        source_path = job_dir / "tem_r.xlsx"
        source_path.write_bytes(content)
        try:
            result = await asyncio.to_thread(tem_r_mod.inspect_tem_r_workbook, source_path)
            result["required_columns"] = list(tem_r_mod.REQUIRED_COLUMNS)
            for option in result.get("sheet_options", []):
                columns = [str(column) for column in option.get("columns", [])]
                lookup = {column.strip().upper(): column for column in columns}
                option["suggested_mapping"] = {
                    canonical: lookup.get(canonical, "") for canonical in tem_r_mod.REQUIRED_COLUMNS
                }
            columns = [str(column) for column in result.get("columns", [])]
            lookup = {column.strip().upper(): column for column in columns}
            result["suggested_mapping"] = {
                canonical: lookup.get(canonical, "") for canonical in tem_r_mod.REQUIRED_COLUMNS
            }
            usable = [item for item in result.get("sheet_options", []) if item.get("columns")]
            result["requires_sheet_selection"] = len(usable) > 1
            return result
        except HTTPException:
            raise
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        except Exception as exc:
            raise HTTPException(status_code=500, detail=f"Lỗi đọc cột file Tem R: {exc}") from exc
        finally:
            _remove_tree(job_dir)

    async def enqueue_tem_r(
        tem_r_file: UploadFile = File(...),
        max_location_distance: float = Form(1),
        stock_column: str = Form(...),
        column_mapping: Optional[str] = Form(None),
        sheet_name: Optional[str] = Form(None),
        header_row: Optional[int] = Form(None),
        _admin: str = Depends(core_module.require_admin),
    ):
        core_module._cleanup_old_jobs()
        if not tem_r_file.filename or not tem_r_file.filename.lower().endswith(".xlsx"):
            raise HTTPException(status_code=422, detail="tem_r_file phải là file .xlsx")
        if max_location_distance <= 0:
            raise HTTPException(status_code=422, detail="Khoảng cách vị trí tối đa phải lớn hơn 0")
        content = await core_module._read_upload_limited(
            tem_r_file, core_module._MAX_TEM_R_BYTES, "File Tem R"
        )
        try:
            mapping = _parse_mapping(column_mapping)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc

        job_id = uuid.uuid4().hex[:8]
        job_dir = core_module._TMP_DIR / f"tem_r_{job_id}"
        job_dir.mkdir(parents=True, exist_ok=True)
        source_path = job_dir / "tem_r.xlsx"
        source_path.write_bytes(content)
        output_path = core_module._OUT_DIR / f"tem_r_{job_id}.xlsx"

        try:
            resolved_sheet, resolved_header, resolved_mapping, _inspection = await asyncio.to_thread(
                _resolve_tem_r_selection,
                tem_r_mod,
                source_path,
                stock_column=stock_column,
                column_mapping=mapping,
                sheet_name=sheet_name,
                header_row=header_row,
            )
        except Exception as exc:
            _remove_tree(job_dir)
            raise HTTPException(status_code=422, detail=f"Preflight Tem R không đạt: {exc}") from exc

        job = core_module._new_job(job_id, "tem_r", output_path, output_path.name)
        job["sheet_name"] = resolved_sheet
        job["header_row"] = resolved_header

        def work(current):
            current.update(progress=30, message="Đang phân bổ Tem R")
            try:
                plan = tem_r_mod.create_tem_r_workbook(
                    source_path,
                    output_path,
                    max_location_distance=max_location_distance,
                    stock_column=str(stock_column).strip(),
                    column_mapping=resolved_mapping,
                    sheet_name=resolved_sheet,
                    header_row=resolved_header,
                )
                validate_tem_r_workbook(
                    output_path, plan, sheet_name=resolved_sheet, header_row=resolved_header
                )
                return _tem_r_summary(plan), plan.exceptions["reason"].tolist()
            finally:
                source_path.unlink(missing_ok=True)

        core_module._submit_job(job_id, "tem_r", work)
        return core_module._job_public(job)

    async def run_tem_r(
        tem_r_file: UploadFile = File(...),
        max_location_distance: float = Form(1),
        stock_column: str = Form(...),
        column_mapping: Optional[str] = Form(None),
        sheet_name: Optional[str] = Form(None),
        header_row: Optional[int] = Form(None),
        _admin: str = Depends(core_module.require_admin),
    ):
        """Legacy synchronous endpoint with the same preflight/output contract."""
        if not tem_r_file.filename or not tem_r_file.filename.lower().endswith(".xlsx"):
            raise HTTPException(status_code=422, detail="tem_r_file phải là file .xlsx")
        if max_location_distance <= 0:
            raise HTTPException(status_code=422, detail="Khoảng cách vị trí tối đa phải lớn hơn 0")
        content = await core_module._read_upload_limited(
            tem_r_file, core_module._MAX_TEM_R_BYTES, "File Tem R"
        )
        try:
            mapping = _parse_mapping(column_mapping)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc

        job_id = uuid.uuid4().hex[:8]
        job_dir = core_module._TMP_DIR / f"tem_r_{job_id}"
        job_dir.mkdir(parents=True, exist_ok=True)
        source_path = job_dir / "tem_r.xlsx"
        source_path.write_bytes(content)
        output_path = core_module._OUT_DIR / f"tem_r_{job_id}.xlsx"
        acquired = False
        try:
            resolved_sheet, resolved_header, resolved_mapping, _inspection = await asyncio.to_thread(
                _resolve_tem_r_selection,
                tem_r_mod,
                source_path,
                stock_column=stock_column,
                column_mapping=mapping,
                sheet_name=sheet_name,
                header_row=header_row,
            )
            core_module._heavy_queue.reserve(job_id, "tem_r")
            await asyncio.to_thread(core_module._heavy_queue.acquire, job_id, "tem_r")
            acquired = True
            plan = await asyncio.to_thread(
                tem_r_mod.create_tem_r_workbook,
                source_path,
                output_path,
                max_location_distance=max_location_distance,
                stock_column=str(stock_column).strip(),
                column_mapping=resolved_mapping,
                sheet_name=resolved_sheet,
                header_row=resolved_header,
            )
            validate_tem_r_workbook(
                output_path, plan, sheet_name=resolved_sheet, header_row=resolved_header
            )
            summary = _tem_r_summary(plan)
            job = core_module._new_job(job_id, "tem_r", output_path, output_path.name)
            job.update(
                status="completed", progress=100, message="Hoàn tất",
                summary=summary, warnings=plan.exceptions["reason"].tolist(), error=None,
                created_timestamp=time.time(),
            )
            return {
                "job_id": job_id,
                "summary": summary,
                "warnings": job["warnings"],
            }
        except HTTPException:
            Path(output_path).unlink(missing_ok=True)
            raise
        except RuntimeError as exc:
            Path(output_path).unlink(missing_ok=True)
            status = 503 if "Hàng đợi" in str(exc) else 500
            raise HTTPException(status_code=status, detail=str(exc)) from exc
        except ValueError as exc:
            Path(output_path).unlink(missing_ok=True)
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        except Exception as exc:
            Path(output_path).unlink(missing_ok=True)
            raise HTTPException(status_code=500, detail=f"Lỗi xử lý Tem R: {exc}") from exc
        finally:
            if acquired:
                core_module._heavy_queue.release(job_id)
            else:
                core_module._heavy_queue.cancel(job_id)
            _remove_tree(job_dir)

    _replace_route(core_module, "/api/inspect-tem-r", "POST", inspect_tem_r_file)
    _replace_route(core_module, "/api/jobs/tem-r", "POST", enqueue_tem_r)
    _replace_route(core_module, "/api/tem-r", "POST", run_tem_r)
    core_module.inspect_tem_r_file = inspect_tem_r_file
    core_module.enqueue_tem_r = enqueue_tem_r
    core_module.run_tem_r = run_tem_r
    core_module._resolve_tem_r_selection = _resolve_tem_r_selection
    core_module._validate_tem_r_workbook = validate_tem_r_workbook
    core_module.app.state.auxiliary_workflows_contract_v1 = True
