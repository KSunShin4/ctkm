# -*- coding: utf-8 -*-
"""Reduce peak RSS of the production CTKM Excel export path.

The historical facade chain created the workbook once and then reopened the full
OpenPyXL workbook up to three times to append POG metadata, clarify legacy labels,
and append KM metadata. On a 512 MiB Render instance those full reload/save passes
can materially increase peak RSS after an already style-heavy export.

This contract keeps the business/exporter implementation unchanged, suppresses
only those post-save reopen helpers in the hardened API path, and performs all
three mutations in one combined workbook pass. It also emits small RSS checkpoints
to Render logs so a future OOM can be localized without exposing business data.
"""
from __future__ import annotations

import ctypes
import gc
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from modules.facade_contracts import register_post_facade_hook


_HOOK_NAME = "export-memory-contract-v1"


def _memory_mb() -> tuple[float | None, float | None]:
    """Return current/peak RSS on Linux without an optional dependency."""
    current = None
    peak = None
    try:
        for line in Path("/proc/self/status").read_text(encoding="utf-8").splitlines():
            if line.startswith("VmRSS:"):
                current = float(line.split()[1]) / 1024.0
            elif line.startswith("VmHWM:"):
                peak = float(line.split()[1]) / 1024.0
    except (OSError, ValueError, IndexError):
        pass
    return current, peak


def _log_memory(stage: str) -> None:
    current, peak = _memory_mb()
    if current is None and peak is None:
        return
    current_text = "?" if current is None else f"{current:.1f}"
    peak_text = "?" if peak is None else f"{peak:.1f}"
    print(f"[ctkm-memory] {stage}: rss_mb={current_text} peak_mb={peak_text}", flush=True)


def _release_heap() -> None:
    """Collect Python cycles and return free glibc arenas when supported."""
    gc.collect()
    try:
        libc = ctypes.CDLL("libc.so.6")
        trim = getattr(libc, "malloc_trim", None)
        if trim is not None:
            trim(0)
    except (OSError, AttributeError):
        pass


def _is_missing(value: Any) -> bool:
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _metadata_value(value: Any):
    if _is_missing(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def _replace_metadata_sheet(workbook, sheet_name: str, metadata) -> bool:
    if metadata is None or getattr(metadata, "empty", True):
        return False
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    ws = workbook.create_sheet(sheet_name)
    columns = list(metadata.columns)
    ws.append(columns)
    for row in metadata.itertuples(index=False, name=None):
        ws.append([_metadata_value(value) for value in row])
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
    ws.sheet_state = "hidden"
    return True


def _clarify_legacy_workbook(workbook) -> bool:
    changed = False
    if "Tom tat" in workbook.sheetnames:
        ws = workbook["Tom tat"]
        legacy = {
            "So dong CTKM thieu vi tri ke",
            "So item CTKM thieu vi tri ke",
            "So dong POG khong co CTKM phu hop",
            "So item POG khong co CTKM phu hop",
        }
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value in legacy:
                if not ws.row_dimensions[row].hidden:
                    ws.row_dimensions[row].hidden = True
                    changed = True

    if "HƯỚNG DẪN" in workbook.sheetnames:
        ws = workbook["HƯỚNG DẪN"]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == "Thiếu vị trí":
                    cell.value = "Thiếu vị trí KM"
                    changed = True
    return changed


def _postprocess_workbook(output_path, *, pog_metadata=None, fact_metadata=None) -> str:
    """Apply every post-export workbook mutation with one load/save cycle."""
    path = Path(output_path)
    if not path.is_file():
        return str(output_path)

    _log_memory("before-postprocess-open")
    workbook = load_workbook(path)
    try:
        changed = False
        changed |= _replace_metadata_sheet(workbook, "POG metadata", pog_metadata)
        changed |= _replace_metadata_sheet(workbook, "KM metadata", fact_metadata)
        changed |= _clarify_legacy_workbook(workbook)
        if changed:
            workbook.save(path)
    finally:
        try:
            workbook.close()
        except Exception:
            pass
        del workbook
        _release_heap()
    _log_memory("after-postprocess-close")
    return str(output_path)


def _arg(args, kwargs, name: str, position: int):
    if name in kwargs:
        return kwargs.get(name)
    return args[position] if len(args) > position else None


def _noop_post_export(*_args, **_kwargs) -> None:
    return None


def _apply_export_memory_contract(core_module, facade_module=None) -> None:
    main_core = getattr(core_module._main_mod, "_core", None)
    if main_core is None:
        raise RuntimeError("Export-memory contract không tìm thấy _ctkm_main_core")

    exporter = getattr(main_core, "EX", None)
    if exporter is None:
        raise RuntimeError("Export-memory contract không tìm thấy exporter module")

    # The exporter facade wrapper resolves this helper dynamically from its own
    # module globals. Suppress its standalone POG workbook reopen; the combined
    # pass below appends the same metadata later.
    if hasattr(exporter, "_append_pog_metadata_audit"):
        exporter._append_pog_metadata_audit = _noop_post_export

    current_export = getattr(exporter, "export_final_report", None)
    if not callable(current_export):
        raise RuntimeError("Export-memory contract không tìm thấy export_final_report")
    if getattr(current_export, "_ctkm_export_memory_contract_v1", False):
        return

    # main.py's location/fact postprocessors are looked up through the wrapper's
    # globals at call time. Disable exactly those reopen helpers for this current
    # facade generation; a post-facade hook repeats this after any reload.
    export_globals = getattr(current_export, "__globals__", None)
    if isinstance(export_globals, dict):
        if "_clarify_location_workbook" in export_globals:
            export_globals["_clarify_location_workbook"] = _noop_post_export
        if "_append_fact_metadata_audit" in export_globals:
            export_globals["_append_fact_metadata_audit"] = _noop_post_export

    def memory_export(*args, **kwargs):
        df_pog_review = _arg(args, kwargs, "df_pog_review", 11)
        confirmation = _arg(args, kwargs, "df_ctkm_clearance_confirmation", 10)
        pog_metadata = (
            getattr(df_pog_review, "attrs", {}).get("pog_metadata")
            if df_pog_review is not None else None
        )
        fact_metadata = (
            getattr(confirmation, "attrs", {}).get("fact_metadata")
            if confirmation is not None else None
        )

        _log_memory("before-base-export")
        output_path = current_export(*args, **kwargs)
        _log_memory("after-base-export")

        # At this point the ExcelWriter/OpenPyXL workbook from the base exporter
        # is unreachable. Reclaim it before opening the final workbook once.
        _release_heap()
        _log_memory("after-base-export-trim")
        return _postprocess_workbook(
            output_path,
            pog_metadata=pog_metadata,
            fact_metadata=fact_metadata,
        )

    memory_export.__name__ = getattr(current_export, "__name__", "export_final_report")
    memory_export.__doc__ = getattr(current_export, "__doc__", None)
    memory_export.__wrapped__ = current_export
    memory_export._ctkm_export_memory_contract_v1 = True
    exporter.export_final_report = memory_export

    # Keep the currently retained facade surface synchronized with the shared
    # exporter module after a main.py alias reload.
    for facade in (core_module._main_mod, facade_module):
        if facade is None:
            continue
        facade_exporter = getattr(facade, "EX", None)
        if facade_exporter is exporter:
            facade_exporter.export_final_report = memory_export


def install_export_memory_contract(core_module) -> None:
    if getattr(core_module.app.state, "export_memory_contract_v1", False):
        return

    main_core = getattr(core_module._main_mod, "_core", None)
    if main_core is None:
        raise RuntimeError("Export-memory contract không tìm thấy _ctkm_main_core")

    def reapply(_main_core, facade_module=None) -> None:
        _apply_export_memory_contract(core_module, facade_module)

    register_post_facade_hook(main_core, _HOOK_NAME, reapply)
    _apply_export_memory_contract(core_module, core_module._main_mod)
    core_module.app.state.export_memory_contract_v1 = True
