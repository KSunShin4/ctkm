# -*- coding: utf-8 -*-
"""Shared semantics for visible Excel data rows and empty placeholders."""
from __future__ import annotations

from openpyxl.utils import get_column_letter


_EMPTY_NOTE_PREFIXES = (
    "không có",
    "khong co",
    "không phát hiện",
    "khong phat hien",
    "chưa có",
    "chua co",
)


def visible_column_indexes(ws) -> list[int]:
    indexes: list[int] = []
    for index in range(1, (ws.max_column or 0) + 1):
        dimension = ws.column_dimensions.get(get_column_letter(index))
        if dimension is not None and bool(dimension.hidden):
            continue
        indexes.append(index)
    return indexes


def visible_nonempty_rows(ws, column_indexes: list[int], *, start_row: int = 2):
    for row_index in range(start_row, (ws.max_row or 0) + 1):
        dimension = ws.row_dimensions.get(row_index)
        if dimension is not None and bool(dimension.hidden):
            continue
        values = [ws.cell(row=row_index, column=index).value for index in column_indexes]
        if not any(value is not None and str(value).strip() for value in values):
            continue
        yield values


def is_empty_placeholder(headers: list[str], rows: list[list[object]]) -> bool:
    """Recognize exporter note rows that represent an operational zero count."""
    if len(headers) != 1 or headers[0].strip().lower() not in {"ghi chu", "ghi chú"}:
        return False
    if len(rows) != 1:
        return False
    value = "" if rows[0][0] is None else str(rows[0][0]).strip().lower()
    return value.startswith(_EMPTY_NOTE_PREFIXES)


def visible_data_row_count(ws) -> int:
    columns = visible_column_indexes(ws)
    headers = [
        "" if ws.cell(row=1, column=index).value is None
        else str(ws.cell(row=1, column=index).value)
        for index in columns
    ]
    rows = list(visible_nonempty_rows(ws, columns))
    return 0 if is_empty_placeholder(headers, rows) else len(rows)
