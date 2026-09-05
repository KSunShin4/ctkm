# -*- coding: utf-8 -*-
"""Preview the final workbook exactly as users can see it in Excel.

The exporter deliberately hides technical columns, metadata sheets and a few
legacy summary rows. Preview uses the same shared visible-row semantics as final
artifact validation so badges and workbook cards cannot disagree.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from backend.workbook_semantics import (
    is_empty_placeholder as _is_empty_placeholder,
    visible_column_indexes as _visible_column_indexes,
    visible_nonempty_rows as _visible_nonempty_rows,
)


_NON_TABULAR_SHEETS = {"HƯỚNG DẪN"}
_ACTION_SHEETS = {"Gắn tem", "Tháo tem", "Chờ gắn", "Cần review"}
_RECONCILIATION_SHEETS = {
    "Tong hop", "CTKM thieu vi tri ke", "POG cần kiểm tra",
    "POG không có CTKM", "POG cần review",
}
_AUDIT_SHEETS = {"Audit hành động", "Canh bao trung thoi gian"}
_CONFIRMATION_SHEETS = {"Xác nhận CTKM Clearance"}
_SUMMARY_SHEETS = {"Tom tat"}


def sheet_kind(sheet_name: str) -> str:
    if sheet_name in _ACTION_SHEETS:
        return "action"
    if sheet_name in _RECONCILIATION_SHEETS:
        return "reconciliation"
    if sheet_name in _AUDIT_SHEETS:
        return "audit"
    if sheet_name in _CONFIRMATION_SHEETS:
        return "confirmation"
    if sheet_name in _SUMMARY_SHEETS:
        return "summary"
    return "data"


def build_visible_preview(path: str | Path, rows: int = 50) -> dict:
    """Return only visible, tabular workbook content.

    ``total_rows`` counts visible non-empty data rows. Empty-sheet placeholders
    created by exporter are reported as zero rows. Hidden technical columns/rows
    and hidden metadata sheets are intentionally excluded.
    """
    limit = int(rows)
    if limit < 1 or limit > 500:
        raise ValueError("rows phải nằm trong khoảng 1..500")

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        sheets = {}
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            if ws.sheet_state != "visible" or sheet_name in _NON_TABULAR_SHEETS:
                continue

            column_indexes = _visible_column_indexes(ws)
            headers = [
                "" if ws.cell(row=1, column=index).value is None
                else str(ws.cell(row=1, column=index).value)
                for index in column_indexes
            ]
            raw_rows = list(_visible_nonempty_rows(ws, column_indexes))
            if _is_empty_placeholder(headers, raw_rows):
                total_rows = 0
                preview_rows = []
            else:
                total_rows = len(raw_rows)
                preview_rows = raw_rows[:limit]

            sheets[sheet_name] = {
                "headers": headers,
                "rows": [
                    ["" if value is None else str(value) for value in row]
                    for row in preview_rows
                ],
                "total_rows": total_rows,
                "sheet_kind": sheet_kind(sheet_name),
            }
        return {"sheets": sheets}
    finally:
        workbook.close()


def install_preview_contract(core_module) -> None:
    """Replace only the preview endpoint call while retaining FastAPI dependencies."""
    if getattr(core_module.app.state, "visible_preview_contract", False):
        return

    def visible_preview_result(job_id: str, rows: int = 50, _admin: str | None = None):
        job = core_module._jobs.get(job_id)
        if not job:
            raise core_module.HTTPException(
                status_code=404, detail=f"Khong tim thay job '{job_id}'"
            )

        status_fn = getattr(core_module, "_effective_job_status", None)
        status = status_fn(job) if callable(status_fn) else str(job.get("status") or "").lower()
        ready_fn = getattr(core_module, "_job_output_ready", None)
        if callable(ready_fn) and not ready_fn(job):
            if status in {"queued", "running"}:
                raise core_module.HTTPException(
                    status_code=409, detail="Tác vụ chưa hoàn tất; chưa thể xem trước file."
                )
            if status == "failed":
                raise core_module.HTTPException(
                    status_code=409, detail="Tác vụ thất bại; không có file hợp lệ để xem trước."
                )

        path = Path(job["output_path"])
        if not path.is_file():
            raise core_module.HTTPException(
                status_code=404, detail="File kết quả đã hết hạn hoặc không còn tồn tại."
            )
        if rows < 1 or rows > 500:
            raise core_module.HTTPException(
                status_code=422, detail="rows phải nằm trong khoảng 1..500"
            )
        try:
            preview = build_visible_preview(path, rows)
            return {
                "job_id": job_id,
                "month": job.get("month"),
                "summary": job.get("summary") or {},
                **preview,
            }
        except core_module.HTTPException:
            raise
        except Exception as exc:
            raise core_module.HTTPException(
                status_code=500, detail=f"Loi doc file ket qua: {exc}"
            ) from exc

    replaced = False
    for route in core_module.app.routes:
        if getattr(route, "path", None) != "/api/preview/{job_id}":
            continue
        if "GET" not in (getattr(route, "methods", None) or set()):
            continue
        route.endpoint = visible_preview_result
        dependant = getattr(route, "dependant", None)
        if dependant is not None:
            dependant.call = visible_preview_result
        replaced = True
        break

    if not replaced:
        raise RuntimeError("Không tìm thấy route GET /api/preview/{job_id} để cài preview contract")

    core_module.preview_result = visible_preview_result
    core_module.app.state.visible_preview_contract = True
