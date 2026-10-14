# -*- coding: utf-8 -*-
"""Stateless staff-scheduling API and exact-template Excel export."""
from __future__ import annotations

from copy import copy
from datetime import datetime, timedelta
from io import BytesIO
import re
from urllib.parse import quote

from fastapi import Body, Depends, HTTPException
from fastapi.responses import FileResponse, Response
from openpyxl import load_workbook

from modules.scheduling_engine import (
    SchedulingError,
    generate_schedule,
    normalize_config,
    validate_schedule,
)

_TEMPLATE_NAME = "1885_Lich_lam_viec_template.xlsx"
_SAFE_FILENAME = re.compile(r"[^0-9A-Za-zÀ-ỹ._-]+")


def _template_path():
    from pathlib import Path
    return Path(__file__).resolve().parent / "templates" / _TEMPLATE_NAME


def _validation_payload(body: dict) -> tuple[dict, list[dict]]:
    if not isinstance(body, dict):
        raise HTTPException(status_code=422, detail="Payload không hợp lệ")
    config = body.get("config") if isinstance(body.get("config"), dict) else body
    assignments = body.get("assignments")
    if not isinstance(assignments, list):
        raise HTTPException(status_code=422, detail="Thiếu assignments")
    return config, assignments


def _validate_or_422(config: dict, assignments: list[dict]) -> dict:
    try:
        normalized = normalize_config(config)
        validation = validate_schedule(normalized, assignments)
    except SchedulingError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    return {"normalized": normalized, "validation": validation}


def _display_name(employee: dict) -> str:
    name = employee.get("name", "")
    code = employee.get("staff_code", "")
    return f"{name} - {code}" if code else str(name)


def _export_filename(config: dict) -> str:
    start = datetime.fromisoformat(config["week_start"]).strftime("%d-%m")
    end = datetime.fromisoformat(config["week_end"]).strftime("%d_%m_%Y")
    label = _SAFE_FILENAME.sub("_", str(config.get("store_label") or "lich")).strip("_") or "lich"
    return f"{label}_Lich_lam_viec_{start}-{end}.xlsx"


def _write_group_to_slots(ws, column: int, rows: list[dict], slot_rows: list[int], employee_by_id: dict) -> None:
    """Write every employee into the fixed template without blocking export.

    If the user intentionally puts more people into a shift than the template has
    physical rows for, the overflow names are joined in the last slot instead of
    rejecting the workbook.
    """
    names = [_display_name(employee_by_id[row["employee_id"]]) for row in rows]
    if not names or not slot_rows:
        return
    for index, row_number in enumerate(slot_rows):
        if index < len(slot_rows) - 1:
            ws.cell(row_number, column).value = names[index] if index < len(names) else None
            continue
        remaining = names[index:]
        ws.cell(row_number, column).value = "\n".join(remaining) if remaining else None
        if len(remaining) > 1:
            cell = ws.cell(row_number, column)
            alignment = copy(cell.alignment)
            alignment.wrapText = True
            cell.alignment = alignment


def export_schedule_to_template(core_module, item: dict) -> tuple[bytes, str]:
    """Render exactly the current submitted schedule; never save or re-optimize it."""
    del core_module  # Compatibility with the previous helper signature.
    template = _template_path()
    if not template.is_file():
        raise HTTPException(status_code=500, detail="Thiếu template Excel lịch làm việc 1885")

    config = item.get("config") or {}
    assignments = item.get("assignments") or []
    checked = _validate_or_422(config, assignments)
    validation = checked["validation"]
    normalized = checked["normalized"]
    employee_by_id = {employee["id"]: employee for employee in normalized["employees"]}
    order = {employee["id"]: index for index, employee in enumerate(normalized["employees"])}
    rows_by_day = {entry["date"]: [] for entry in normalized["staffing"]}
    for row in validation["assignments"]:
        rows_by_day[row["date"]].append(row)

    output_name = _export_filename(normalized)
    workbook = load_workbook(template)
    try:
        ws = workbook[workbook.sheetnames[0]]
        start_date = datetime.fromisoformat(normalized["week_start"])
        end_date = datetime.fromisoformat(normalized["week_end"])
        ws["A1"] = (
            f"Lịch làm việc {normalized.get('store_label') or '1885'} ngày "
            f"{start_date.strftime('%d/%m/%y')} đến {end_date.strftime('%d/%m/%y')}"
        )
        weekdays = ["thứ 2", "thứ 3", "thứ 4", "thứ 5", "thứ 6", "thứ 7", "chủ nhật"]
        day_keys = []
        for offset in range(7):
            current = start_date.date() + timedelta(days=offset)
            day_keys.append(current.isoformat())
            column = 2 + offset
            ws.cell(2, column).value = current.day
            ws.cell(3, column).value = weekdays[offset]
            for row_number in range(4, 11):
                ws.cell(row_number, column).value = None

        for offset, day_key in enumerate(day_keys):
            column = 2 + offset
            day_rows = sorted(
                rows_by_day.get(day_key, []),
                key=lambda row: order.get(row["employee_id"], 999),
            )
            m06 = [
                row for row in day_rows
                if row["shift"] == "MORNING" and row.get("start_time") == "06:00"
            ]
            m07 = [
                row for row in day_rows
                if row["shift"] == "MORNING" and row.get("start_time") == "07:00"
            ]
            afternoon = [row for row in day_rows if row["shift"] == "AFTERNOON"]
            off = [row for row in day_rows if row["shift"] == "OFF"]

            _write_group_to_slots(ws, column, m06, [4, 5], employee_by_id)
            _write_group_to_slots(ws, column, m07, [6], employee_by_id)
            _write_group_to_slots(ws, column, afternoon, [7, 8, 9], employee_by_id)
            if off:
                ws.cell(10, column).value = ", ".join(
                    f"{employee_by_id[row['employee_id']]['name']} off" for row in off
                )

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue(), output_name
    finally:
        workbook.close()


def _insert_before_frontend(core_module, path: str, endpoint, methods: list[str], **kwargs) -> None:
    core_module.app.add_api_route(path, endpoint, methods=methods, **kwargs)
    routes = core_module.app.router.routes
    route = next(
        item for item in reversed(routes)
        if getattr(item, "path", None) == path
        and set(methods).issubset(getattr(item, "methods", set()) or set())
    )
    routes.remove(route)
    mount_index = next(
        (
            index for index, item in enumerate(routes)
            if item.__class__.__name__ == "Mount"
            and (getattr(item, "name", None) == "frontend" or getattr(item, "path", None) in {"", "/"})
        ),
        len(routes),
    )
    routes.insert(mount_index, route)


def install_scheduling_contract(core_module) -> None:
    if getattr(core_module.app.state, "staff_scheduling_contract_v2", False):
        return

    async def schedule_page():
        index_path = core_module._FRONTEND_DIST / "index.html"
        if not index_path.is_file():
            raise HTTPException(status_code=404, detail="Frontend xếp lịch chưa được build")
        return FileResponse(index_path)

    async def schedule_generate(
        payload: dict = Body(...),
        _admin: str = Depends(core_module.require_admin),
    ):
        try:
            return generate_schedule(payload, variants=3)
        except SchedulingError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc

    async def schedule_validate(
        payload: dict = Body(...),
        _admin: str = Depends(core_module.require_admin),
    ):
        config, assignments = _validation_payload(payload)
        return _validate_or_422(config, assignments)["validation"]

    async def schedule_export(
        payload: dict = Body(...),
        _admin: str = Depends(core_module.require_admin),
    ):
        config, assignments = _validation_payload(payload)
        content, filename = export_schedule_to_template(
            core_module,
            {"config": config, "assignments": assignments},
        )
        encoded = quote(filename)
        return Response(
            content=content,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
        )

    bindings = (
        ("/schedule", schedule_page, ["GET"], {"include_in_schema": False}),
        ("/api/schedules/generate", schedule_generate, ["POST"], {}),
        ("/api/schedules/validate", schedule_validate, ["POST"], {}),
        ("/api/schedules/export.xlsx", schedule_export, ["POST"], {}),
    )
    for path, endpoint, methods, kwargs in bindings:
        _insert_before_frontend(core_module, path, endpoint, methods, **kwargs)

    for name, value in {
        "schedule_page": schedule_page,
        "schedule_generate": schedule_generate,
        "schedule_validate": schedule_validate,
        "schedule_export": schedule_export,
    }.items():
        setattr(core_module, name, value)

    core_module.app.state.staff_scheduling_contract_v1 = True
    core_module.app.state.staff_scheduling_contract_v2 = True
