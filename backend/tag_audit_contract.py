# -*- coding: utf-8 -*-
"""Single-tag shelf-label audit for store operations.

The feature is intentionally split into two trust domains:

* Gemini Vision is optional and may only extract text printed on one physical tag.
* Every operational verdict is deterministic and compares the extracted observation
  with the CTKM workbook produced by this application.

Barcode values are collected as supporting evidence only. They are never treated
as an ITEMID unless a future, company-owned barcode-to-ITEMID mapping is supplied.
"""
from __future__ import annotations

import asyncio
import base64
from datetime import date, datetime
from difflib import SequenceMatcher
from hashlib import sha256
from io import BytesIO
import json
import math
import os
from pathlib import Path
import re
import threading
import time
import unicodedata

from fastapi import Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
import requests

from modules.config_loader import GEMINI_CONFIG
from modules.channel_scope import scope_identity_key


_MAX_IMAGE_BYTES = max(1, int(os.getenv("MAX_TAG_AUDIT_IMAGE_MB", "8"))) * 1024 * 1024
_ACTION_SHEETS = {
    "Gắn tem": "gan_tem",
    "Tháo tem": "thao_tem",
    "Chờ gắn": "cho_gan",
}

_DISPLAY_TO_KEY = {
    "KEY Channel": "key_channel",
    "Channel": "channel",
    "Mã sản phẩm": "item_id",
    "Tên sản phẩm": "item_name",
    "Đơn vị sản phẩm (Promo UoM)": "promo_uom",
    "POG": "pog_code",
    "Vị trí": "location_id",
    "Facing": "facing",
    "Công việc": "action",
    "Ngày gắn": "ngay_gan_tem",
    "Ngày tháo": "ngay_thao_tem",
    "Campaign": "campaign_name",
    "Scheme detail": "campaign_detail",
    "From": "promo_from",
    "To": "promo_to",
    "Giá thường": "normal_price",
    "Giá KM": "promo_price",
    "Nguồn": "program_source",
    "Lý do / Cảnh báo": "reason",
    "Product type": "product_type",
    "Cửa hàng": "store_id",
}
_CANONICAL_KEYS = set(_DISPLAY_TO_KEY.values())

_CONTEXT_LOCK = threading.RLock()
_CONTEXT_CACHE: dict = {
    "source_key": None,
    "source": None,
    "job_id": None,
    "month": None,
    "rows": None,
}
_VISION_LOCK = threading.Lock()
_VISION_LAST_CALL = 0.0


def _blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return not str(value).strip()


def _cell_value(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def _norm_text(value) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _norm_item(value) -> str:
    return str(value or "").strip().upper()


def _parse_money(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return int(round(float(value)))
    digits = re.sub(r"[^0-9]", "", str(value))
    return int(digits) if digits else None


def _parse_iso_date(value):
    if _blank(value):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return None


def _similarity(left, right) -> float:
    a, b = _norm_text(left), _norm_text(right)
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, a, b).ratio()


def _header_key(header) -> str | None:
    text = str(header or "").strip()
    if text in _DISPLAY_TO_KEY:
        return _DISPLAY_TO_KEY[text]
    if text in _CANONICAL_KEYS:
        return text
    return None


def _parse_action_workbook(source) -> list[dict]:
    workbook = None
    try:
        workbook = load_workbook(source, data_only=True, read_only=True)
        rows: list[dict] = []
        found_sheet = False
        for sheet_name, expected_action in _ACTION_SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                continue
            found_sheet = True
            ws = workbook[sheet_name]
            iterator = ws.iter_rows(values_only=True)
            headers = next(iterator, None)
            if not headers:
                continue
            keys = [_header_key(header) for header in headers]
            if "item_id" not in keys:
                raise ValueError(f"Sheet '{sheet_name}' thiếu cột Mã sản phẩm")
            for values in iterator:
                record = {}
                for index, key in enumerate(keys):
                    if key is None or index >= len(values):
                        continue
                    record[key] = _cell_value(values[index])
                if _blank(record.get("item_id")):
                    continue
                record["expected_action"] = expected_action
                record["source_sheet"] = sheet_name
                rows.append(record)
        if not found_sheet:
            raise ValueError(
                "Workbook không có các sheet Gắn tem / Tháo tem / Chờ gắn. "
                "Hãy dùng file kết quả CTKM do hệ thống xuất."
            )

        # The same business promotion may be repeated only because it has several
        # POG positions. A physical tag identifies the promotion, not one facing.
        deduped: list[dict] = []
        seen: set[tuple] = set()
        for record in rows:
            key = (
                _norm_item(record.get("item_id")),
                str(record.get("key_channel") or ""),
                str(record.get("channel") or ""),
                str(record.get("campaign_name") or ""),
                str(record.get("campaign_detail") or ""),
                _parse_iso_date(record.get("promo_from")),
                _parse_iso_date(record.get("promo_to")),
                _parse_money(record.get("promo_price")),
                record.get("expected_action"),
            )
            if key in seen:
                continue
            seen.add(key)
            deduped.append(record)
        if not deduped:
            raise ValueError("Workbook không có dòng hành động CTKM để kiểm tem")
        return deduped
    finally:
        if workbook is not None:
            workbook.close()


def _latest_completed_job(core_module):
    candidates = []
    for job in core_module._jobs.values():
        if job.get("kind") != "ctkm":
            continue
        status_fn = getattr(core_module, "_effective_job_status", None)
        status = status_fn(job) if callable(status_fn) else job.get("status")
        path = Path(str(job.get("output_path") or ""))
        if status != "completed" or not path.is_file():
            continue
        candidates.append(job)
    if not candidates:
        return None
    return max(
        candidates,
        key=lambda job: float(job.get("created_timestamp") or 0),
    )


def _context_payload(cache: dict) -> dict:
    rows = cache.get("rows") or []
    return {
        "ready": bool(rows),
        "source": cache.get("source"),
        "job_id": cache.get("job_id"),
        "month": cache.get("month"),
        "row_count": len(rows),
        "item_count": len({_norm_item(row.get("item_id")) for row in rows}),
        "vision_available": bool(os.getenv("GEMINI_API_KEY")),
        "barcode_role": "supporting_only",
    }


def _cache_latest_job(core_module, *, force: bool = False) -> dict:
    job = _latest_completed_job(core_module)
    if job is None:
        raise HTTPException(
            status_code=409,
            detail=(
                "Chưa có workbook CTKM khả dụng trong phiên Cloud Run này. "
                "Hãy chạy CTKM hoặc tải workbook kết quả vào màn hình Kiểm tem."
            ),
        )
    path = Path(str(job["output_path"]))
    stat = path.stat()
    source_key = f"job:{job['job_id']}:{stat.st_mtime_ns}:{stat.st_size}"
    with _CONTEXT_LOCK:
        if not force and _CONTEXT_CACHE.get("source_key") == source_key:
            return dict(_CONTEXT_CACHE)
    rows = _parse_action_workbook(path)
    cache = {
        "source_key": source_key,
        "source": "latest-job",
        "job_id": job.get("job_id"),
        "month": job.get("month"),
        "rows": rows,
    }
    with _CONTEXT_LOCK:
        _CONTEXT_CACHE.update(cache)
        return dict(_CONTEXT_CACHE)


def _current_context(core_module) -> dict:
    with _CONTEXT_LOCK:
        cached = dict(_CONTEXT_CACHE)
    if cached.get("rows") and cached.get("source") == "uploaded-workbook":
        return cached
    try:
        return _cache_latest_job(core_module)
    except HTTPException:
        if cached.get("rows"):
            return cached
        raise


def _vision_spacing() -> None:
    global _VISION_LAST_CALL
    minimum = max(0.0, float(os.getenv("TAG_AUDIT_MIN_SECONDS_BETWEEN_OCR", "1.0")))
    with _VISION_LOCK:
        now = time.monotonic()
        wait = minimum - (now - _VISION_LAST_CALL)
        if wait > 0:
            time.sleep(wait)
        _VISION_LAST_CALL = time.monotonic()


def _extract_tag_with_gemini(image: bytes, mime_type: str, month: str | None) -> dict:
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise RuntimeError("GEMINI_API_KEY chưa được cấu hình; hãy nhập các trường tem thủ công.")

    target_month = str(month or "").strip()
    prompt = f"""Đọc DUY NHẤT một tem giá/khuyến mãi Pharmacity trong ảnh.
Bạn chỉ là OCR/extractor, KHÔNG được kết luận tem đúng hay sai và KHÔNG được suy đoán dữ liệu CTKM.
Nếu một trường không đọc chắc chắn thì trả chuỗi rỗng.

Cần trích xuất:
- item_id: mã sản phẩm in trên tem dạng P + chữ số nếu nhìn thấy (ví dụ P29134).
- barcode: dãy số nằm dưới barcode nếu đọc chắc chắn.
- product_name: tên sản phẩm trên tem.
- promo_text: nội dung ưu đãi như MUA 1 TẶNG 1, GIẢM 25%, MUA SP2 GIẢM 50%.
- normal_price: giá thường, chỉ chữ số, không dấu phân cách và không VND.
- promo_price: giá khuyến mãi nổi bật, chỉ chữ số.
- date_text: nguyên văn phần KM/ngày in trên tem.
- promo_from, promo_to: YYYY-MM-DD nếu xác định chắc chắn. Tháng báo cáo hiện tại là {target_month or '(không rõ)'}; chỉ dùng năm của tháng báo cáo để hoàn chỉnh ngày DD/MM trên tem, không tự suy diễn ngày khác.
- raw_text: tối đa 500 ký tự văn bản quan trọng đọc được trên tem.
- confidence: độ tin cậy tổng thể 0..1.

Barcode chỉ là dữ liệu quan sát phụ, không đồng nghĩa với ITEMID."""

    schema = {
        "type": "object",
        "properties": {
            "item_id": {"type": "string"},
            "barcode": {"type": "string"},
            "product_name": {"type": "string"},
            "promo_text": {"type": "string"},
            "normal_price": {"type": "string"},
            "promo_price": {"type": "string"},
            "date_text": {"type": "string"},
            "promo_from": {"type": "string"},
            "promo_to": {"type": "string"},
            "raw_text": {"type": "string"},
            "confidence": {"type": "number", "minimum": 0, "maximum": 1},
        },
        "required": [
            "item_id", "barcode", "product_name", "promo_text", "normal_price",
            "promo_price", "date_text", "promo_from", "promo_to", "raw_text", "confidence",
        ],
    }
    url = GEMINI_CONFIG["api_endpoint"].format(model=GEMINI_CONFIG["model"])
    payload = {
        "system_instruction": {
            "parts": [{"text": "Trích xuất OCR trung thực từ ảnh; không quyết định nghiệp vụ."}]
        },
        "contents": [{
            "role": "user",
            "parts": [
                {"text": prompt},
                {"inline_data": {
                    "mime_type": mime_type,
                    "data": base64.b64encode(image).decode("ascii"),
                }},
            ],
        }],
        "generationConfig": {
            "temperature": 0.0,
            "maxOutputTokens": 2048,
            "responseMimeType": "application/json",
            "responseSchema": schema,
        },
    }
    _vision_spacing()
    response = requests.post(
        url,
        headers={"x-goog-api-key": api_key},
        json=payload,
        timeout=max(10, int(GEMINI_CONFIG.get("timeout_seconds", 30))),
    )
    if response.status_code == 429:
        raise RuntimeError("Gemini đang giới hạn lượt OCR. Chờ một chút rồi quét lại.")
    try:
        response.raise_for_status()
    except requests.RequestException as exc:
        raise RuntimeError(f"OCR Gemini thất bại: HTTP {response.status_code}") from exc
    data = response.json()
    candidates = data.get("candidates") or []
    if not candidates:
        raise RuntimeError("OCR Gemini không trả về kết quả")
    parts = candidates[0].get("content", {}).get("parts", [])
    text = "".join(part.get("text", "") for part in parts if isinstance(part, dict)).strip()
    text = text.strip("`").strip()
    if text.lower().startswith("json"):
        text = text[4:].strip()
    try:
        result = json.loads(text)
    except json.JSONDecodeError as exc:
        raise RuntimeError("OCR Gemini trả về dữ liệu không đọc được") from exc
    if not isinstance(result, dict):
        raise RuntimeError("OCR Gemini trả về sai định dạng")
    result["item_id"] = _norm_item(result.get("item_id"))
    result["barcode"] = re.sub(r"\D", "", str(result.get("barcode") or ""))
    result["promo_price"] = _parse_money(result.get("promo_price"))
    result["normal_price"] = _parse_money(result.get("normal_price"))
    result["promo_from"] = _parse_iso_date(result.get("promo_from"))
    result["promo_to"] = _parse_iso_date(result.get("promo_to"))
    try:
        result["confidence"] = max(0.0, min(1.0, float(result.get("confidence") or 0)))
    except (TypeError, ValueError):
        result["confidence"] = 0.0
    return result


def _row_business_key(row: dict) -> tuple:
    return (
        row.get("expected_action"),
        scope_identity_key(row.get("key_channel"), row.get("channel")),
        _parse_iso_date(row.get("promo_from")),
        _parse_iso_date(row.get("promo_to")),
        _parse_money(row.get("promo_price")),
        str(row.get("campaign_name") or ""),
        str(row.get("campaign_detail") or ""),
    )


def _score_candidate(row: dict, observed: dict) -> float:
    score = 0.0
    observed_price = _parse_money(observed.get("promo_price"))
    expected_price = _parse_money(row.get("promo_price"))
    if observed_price is not None:
        score += 5.0 if expected_price == observed_price else -1.0

    observed_from = _parse_iso_date(observed.get("promo_from"))
    observed_to = _parse_iso_date(observed.get("promo_to"))
    expected_from = _parse_iso_date(row.get("promo_from"))
    expected_to = _parse_iso_date(row.get("promo_to"))
    if observed_from:
        score += 2.0 if observed_from == expected_from else -0.5
    if observed_to:
        score += 2.0 if observed_to == expected_to else -0.5

    promo_expected = " ".join(
        str(row.get(key) or "") for key in ("campaign_name", "campaign_detail")
    )
    if observed.get("promo_text"):
        score += 2.5 * _similarity(observed.get("promo_text"), promo_expected)
    if observed.get("product_name"):
        score += 1.5 * _similarity(observed.get("product_name"), row.get("item_name"))
    return score


def _public_expected(row: dict | None):
    if not row:
        return None
    return {
        key: _cell_value(row.get(key)) for key in (
            "key_channel", "channel", "item_id", "item_name", "promo_uom",
            "campaign_name", "campaign_detail", "promo_from", "promo_to",
            "normal_price", "promo_price", "pog_code", "location_id", "facing",
            "reason", "store_id", "source_sheet", "expected_action",
        )
    }


def _verdict(code: str, title: str, detail: str, *, severity: str = "warning") -> dict:
    return {"code": code, "title": title, "detail": detail, "severity": severity}


def compare_observation(rows: list[dict], observed: dict, manual_item_id: str = "") -> dict:
    """Deterministically compare one observed physical tag with CTKM action rows."""
    manual_key = _norm_item(manual_item_id)
    observed_key = _norm_item(observed.get("item_id"))
    item_key = manual_key or observed_key
    identity_source = "manual_item_id" if manual_key else ("ocr_item_id" if observed_key else "")

    if item_key:
        candidates = [row for row in rows if _norm_item(row.get("item_id")) == item_key]
        if not candidates:
            return {
                "verdict": _verdict(
                    "not_found", "Không tìm thấy sản phẩm",
                    f"Mã {item_key} không có trong danh sách hành động CTKM hiện tại.",
                    severity="error",
                ),
                "identity_source": identity_source,
                "observed": observed,
                "expected": None,
                "mismatches": ["item_id"],
            }
    else:
        observed_name = observed.get("product_name")
        if not observed_name:
            return {
                "verdict": _verdict(
                    "review_identity", "Cần xác nhận sản phẩm",
                    "Ảnh chưa đọc chắc mã sản phẩm. Barcode không được dùng thay ITEMID.",
                ),
                "identity_source": "none",
                "observed": observed,
                "expected": None,
                "mismatches": ["identity"],
            }
        best_by_item: dict[str, tuple[float, dict]] = {}
        for row in rows:
            similarity = _similarity(observed_name, row.get("item_name"))
            key = _norm_item(row.get("item_id"))
            if key and (key not in best_by_item or similarity > best_by_item[key][0]):
                best_by_item[key] = (similarity, row)
        ranked_names = sorted(best_by_item.values(), key=lambda pair: pair[0], reverse=True)
        if not ranked_names or ranked_names[0][0] < 0.72:
            return {
                "verdict": _verdict(
                    "not_found", "Không xác định được sản phẩm",
                    "Tên đọc từ tem không đủ gần với dữ liệu CTKM. Hãy nhập Mã sản phẩm để kiểm lại.",
                    severity="error",
                ),
                "identity_source": "name_candidate",
                "observed": observed,
                "expected": None,
                "mismatches": ["identity"],
            }
        candidate_item = _norm_item(ranked_names[0][1].get("item_id"))
        candidates = [row for row in rows if _norm_item(row.get("item_id")) == candidate_item]
        identity_source = "name_candidate"

    scored = sorted(
        ((_score_candidate(row, observed), row) for row in candidates),
        key=lambda pair: pair[0],
        reverse=True,
    )
    best_score, best = scored[0]
    if len(scored) > 1:
        second_score, second = scored[1]
        if abs(best_score - second_score) < 0.75 and _row_business_key(best) != _row_business_key(second):
            return {
                "verdict": _verdict(
                    "review_ambiguous", "Có nhiều chương trình phù hợp",
                    "Ảnh chưa đủ dữ liệu để xác định chính xác tem thuộc chương trình nào.",
                ),
                "identity_source": identity_source,
                "observed": observed,
                "expected": _public_expected(best),
                "mismatches": ["promotion_identity"],
            }

    if identity_source == "name_candidate":
        return {
            "verdict": _verdict(
                "review_identity", "Cần xác nhận mã sản phẩm",
                f"Hệ thống chỉ đang gợi ý {_norm_item(best.get('item_id'))} theo tên; chưa dùng gợi ý này để kết luận.",
            ),
            "identity_source": identity_source,
            "observed": observed,
            "expected": _public_expected(best),
            "mismatches": ["identity"],
        }

    action = best.get("expected_action")
    if action == "thao_tem":
        return {
            "verdict": _verdict(
                "remove_tag", "Phải tháo tem",
                "Tem vẫn đang hiện diện nhưng chương trình này đã đến hành động tháo.",
                severity="error",
            ),
            "identity_source": identity_source,
            "observed": observed,
            "expected": _public_expected(best),
            "mismatches": ["action"],
        }
    if action == "cho_gan":
        return {
            "verdict": _verdict(
                "early_tag", "Gắn tem quá sớm",
                "Tem đã xuất hiện trước thời điểm chương trình được phép gắn.",
                severity="error",
            ),
            "identity_source": identity_source,
            "observed": observed,
            "expected": _public_expected(best),
            "mismatches": ["action"],
        }

    mismatches: list[str] = []
    observed_price = _parse_money(observed.get("promo_price"))
    expected_price = _parse_money(best.get("promo_price"))
    if observed_price is not None and expected_price is not None and observed_price != expected_price:
        mismatches.append("promo_price")

    observed_normal = _parse_money(observed.get("normal_price"))
    expected_normal = _parse_money(best.get("normal_price"))
    if observed_normal is not None and expected_normal is not None and observed_normal != expected_normal:
        mismatches.append("normal_price")

    observed_from = _parse_iso_date(observed.get("promo_from"))
    observed_to = _parse_iso_date(observed.get("promo_to"))
    expected_from = _parse_iso_date(best.get("promo_from"))
    expected_to = _parse_iso_date(best.get("promo_to"))
    if observed_from and expected_from and observed_from != expected_from:
        mismatches.append("promo_from")
    if observed_to and expected_to and observed_to != expected_to:
        mismatches.append("promo_to")

    promo_expected = " ".join(
        str(best.get(key) or "") for key in ("campaign_name", "campaign_detail")
    ).strip()
    if observed.get("promo_text") and promo_expected:
        if _similarity(observed.get("promo_text"), promo_expected) < 0.30:
            mismatches.append("promo_text")

    if mismatches:
        if "promo_price" in mismatches and any(key in mismatches for key in ("promo_from", "promo_to")):
            verdict = _verdict(
                "wrong_price_and_date", "Sai giá và thời gian",
                "Giá khuyến mãi và thời gian trên tem không khớp CTKM.", severity="error",
            )
        elif "promo_price" in mismatches:
            verdict = _verdict(
                "wrong_price", "Sai giá",
                "Giá khuyến mãi trên tem không khớp CTKM.", severity="error",
            )
        elif any(key in mismatches for key in ("promo_from", "promo_to")):
            verdict = _verdict(
                "wrong_date", "Sai thời gian",
                "Thời gian khuyến mãi trên tem không khớp CTKM.", severity="error",
            )
        elif "promo_text" in mismatches:
            verdict = _verdict(
                "wrong_program", "Sai nội dung chương trình",
                "Nội dung ưu đãi đọc được trên tem không khớp chương trình CTKM.", severity="error",
            )
        else:
            verdict = _verdict(
                "wrong_fields", "Có thông tin không khớp",
                "Một hoặc nhiều trường trên tem khác dữ liệu CTKM.", severity="error",
            )
        return {
            "verdict": verdict,
            "identity_source": identity_source,
            "observed": observed,
            "expected": _public_expected(best),
            "mismatches": mismatches,
        }

    missing: list[str] = []
    if expected_price is not None and observed_price is None:
        missing.append("promo_price")
    if expected_from and not observed_from:
        missing.append("promo_from")
    if expected_to and not observed_to:
        missing.append("promo_to")
    if missing:
        return {
            "verdict": _verdict(
                "review_incomplete", "Ảnh chưa đủ rõ",
                "Chưa đọc đủ giá/ngày để xác nhận tem đúng. Chụp gần hơn hoặc nhập trường còn thiếu.",
            ),
            "identity_source": identity_source,
            "observed": observed,
            "expected": _public_expected(best),
            "mismatches": missing,
        }

    return {
        "verdict": _verdict(
            "correct", "Đúng tem",
            "Mã sản phẩm, giá và thời gian đọc được khớp dữ liệu CTKM.", severity="success",
        ),
        "identity_source": identity_source,
        "observed": observed,
        "expected": _public_expected(best),
        "mismatches": [],
    }


def _insert_api_route_before_frontend(core_module, path: str, endpoint, methods: list[str], **kwargs):
    core_module.app.add_api_route(path, endpoint, methods=methods, **kwargs)
    routes = core_module.app.router.routes
    new_route = next(
        route for route in reversed(routes)
        if getattr(route, "path", None) == path
        and set(methods).issubset(getattr(route, "methods", set()) or set())
    )
    routes.remove(new_route)
    mount_index = next(
        (index for index, route in enumerate(routes)
         if getattr(route, "path", None) == "/" and route.__class__.__name__ == "Mount"),
        len(routes),
    )
    routes.insert(mount_index, new_route)


def install_tag_audit_contract(core_module) -> None:
    if getattr(core_module.app.state, "tag_audit_contract_v1", False):
        return

    async def tag_audit_context(_admin: str = Depends(core_module.require_admin)):
        context = _current_context(core_module)
        return _context_payload(context)

    async def load_tag_audit_context(
        workbook: UploadFile = File(...),
        _admin: str = Depends(core_module.require_admin),
    ):
        filename = str(workbook.filename or "")
        if not filename.lower().endswith(".xlsx"):
            raise HTTPException(status_code=422, detail="Workbook kiểm tem phải là file .xlsx")
        content = await core_module._read_upload_limited(
            workbook, core_module._MAX_PIPELINE_BYTES, "Workbook CTKM"
        )
        try:
            rows = await asyncio.to_thread(_parse_action_workbook, BytesIO(content))
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"Workbook kiểm tem không hợp lệ: {exc}") from exc
        cache = {
            "source_key": "upload:" + sha256(content).hexdigest(),
            "source": "uploaded-workbook",
            "job_id": None,
            "month": None,
            "rows": rows,
        }
        with _CONTEXT_LOCK:
            _CONTEXT_CACHE.update(cache)
        return _context_payload(cache)

    async def check_tag(
        image: UploadFile = File(...),
        manual_item_id: str = Form(""),
        barcode_hint: str = Form(""),
        manual_promo_price: str = Form(""),
        manual_normal_price: str = Form(""),
        manual_promo_from: str = Form(""),
        manual_promo_to: str = Form(""),
        manual_promo_text: str = Form(""),
        _admin: str = Depends(core_module.require_admin),
    ):
        context = _current_context(core_module)
        mime_type = str(image.content_type or "").lower()
        if mime_type not in {"image/jpeg", "image/png", "image/webp"}:
            raise HTTPException(status_code=422, detail="Ảnh tem phải là JPEG, PNG hoặc WebP")
        content = await core_module._read_upload_limited(image, _MAX_IMAGE_BYTES, "Ảnh tem")

        vision_used = False
        vision_warning = None
        try:
            observed = await asyncio.to_thread(
                _extract_tag_with_gemini, content, mime_type, context.get("month")
            )
            vision_used = True
        except Exception as exc:
            observed = {
                "item_id": "", "barcode": "", "product_name": "", "promo_text": "",
                "normal_price": None, "promo_price": None, "date_text": "",
                "promo_from": None, "promo_to": None, "raw_text": "", "confidence": 0.0,
            }
            vision_warning = str(exc)

        barcode = re.sub(r"\D", "", str(barcode_hint or ""))
        if barcode:
            observed["barcode"] = barcode
        if manual_promo_price.strip():
            observed["promo_price"] = _parse_money(manual_promo_price)
        if manual_normal_price.strip():
            observed["normal_price"] = _parse_money(manual_normal_price)
        if manual_promo_from.strip():
            observed["promo_from"] = _parse_iso_date(manual_promo_from)
        if manual_promo_to.strip():
            observed["promo_to"] = _parse_iso_date(manual_promo_to)
        if manual_promo_text.strip():
            observed["promo_text"] = manual_promo_text.strip()

        result = compare_observation(context["rows"], observed, manual_item_id=manual_item_id)
        result["context"] = _context_payload(context)
        result["vision_used"] = vision_used
        result["vision_warning"] = vision_warning
        result["barcode_role"] = "supporting_only"
        return result

    _insert_api_route_before_frontend(
        core_module, "/api/tag-audit/context", tag_audit_context, ["GET"],
        response_model=None,
    )
    _insert_api_route_before_frontend(
        core_module, "/api/tag-audit/context", load_tag_audit_context, ["POST"],
        response_model=None,
    )
    _insert_api_route_before_frontend(
        core_module, "/api/tag-audit/check", check_tag, ["POST"],
        response_model=None,
    )
    core_module.tag_audit_context = tag_audit_context
    core_module.load_tag_audit_context = load_tag_audit_context
    core_module.check_tag = check_tag
    core_module.app.state.tag_audit_contract_v1 = True
