# -*- coding: utf-8 -*-
"""End-to-end validation for the final CTKM workbook artifact."""
from __future__ import annotations

from datetime import date, datetime
from collections import Counter
from numbers import Number
from pathlib import Path
import unicodedata

import openpyxl

from backend.workbook_semantics import (
    is_empty_placeholder,
    visible_column_indexes,
    visible_data_row_count,
    visible_nonempty_rows,
)


_FULL_REQUIRED_CTKM_SHEETS = {
    "HƯỚNG DẪN",
    "Gắn tem",
    "Tháo tem",
    "Chờ gắn",
    "Cần review",
    "CTKM thieu vi tri ke",
    "POG cần kiểm tra",
    "Tong hop",
    "POG không có CTKM",
    "POG cần review",
    "Audit hành động",
    "Canh bao trung thoi gian",
    "Xác nhận CTKM Clearance",
    "Tom tat",
    "Mapping log",
}

_SLIM_REQUIRED_CTKM_SHEETS = {
    "HƯỚNG DẪN",
    "TỔNG HỢP - TEM",
    "Gắn tem",
    "Tháo tem",
    "Chờ gắn",
    "Cần review",
    "CTKM thieu vi tri ke",
    "POG cần kiểm tra",
    "Tong hop",
    "POG không có KM",
    "POG không có Monthly CTKM",
    "POG không có CTKM",
    "POG cần review",
    "Audit hành động",
}

_REQUIRED_CTKM_SHEETS = _FULL_REQUIRED_CTKM_SHEETS

_SUMMARY_KEYS_TO_VERIFY = (
    "Thang bao cao",
    "So dong khop (Matched)",
    "So item duy nhat khop",
    "So dong KM thieu vi tri POG",
    "So item KM thieu vi tri POG",
    "So dong can kiem tra vi tri",
    "So canh bao trung thoi gian",
    "So dong cho gan",
    "So dong can review",
    "So dong POG khong co KM phu hop",
    "So item POG khong co KM phu hop",
)

_LEGACY_CANONICAL_PAIRS = (
    ("So dong CTKM thieu vi tri ke", "So dong KM thieu vi tri POG"),
    ("So item CTKM thieu vi tri ke", "So item KM thieu vi tri POG"),
    ("So dong POG khong co CTKM phu hop", "So dong POG khong co KM phu hop"),
    ("So item POG khong co CTKM phu hop", "So item POG khong co KM phu hop"),
)

_SUMMARY_SHEET_COUNT_PAIRS = (
    ("So dong KM thieu vi tri POG", "CTKM thieu vi tri ke"),
    ("So dong can kiem tra vi tri", "POG cần kiểm tra"),
    ("So dong cho gan", "Chờ gắn"),
    ("So dong can review", "Cần review"),
    ("So dong POG khong co KM phu hop", "POG không có KM"),
)

_FULL_GUIDE_TARGETS = {
    "Gắn tem": "Gắn tem",
    "Tháo tem": "Tháo tem",
    "Chờ gắn": "Chờ gắn",
    "Cần review": "Cần review",
    "Thiếu vị trí": "CTKM thieu vi tri ke",
    "Thiếu vị trí KM": "CTKM thieu vi tri ke",
    "Review vị trí": "POG cần kiểm tra",
    "POG cần review": "POG cần review",
    "Xác nhận CTKM / Clearance": "Xác nhận CTKM Clearance",
}

_SLIM_GUIDE_TARGETS = {
    key: value for key, value in _FULL_GUIDE_TARGETS.items()
    if value in {
        "Gắn tem", "Tháo tem", "Chờ gắn", "Cần review",
        "CTKM thieu vi tri ke", "POG cần kiểm tra", "POG cần review",
    }
}
_GUIDE_TARGETS = _FULL_GUIDE_TARGETS

_LEGACY_HIDDEN_SUMMARY_KEYS = {
    "So dong CTKM thieu vi tri ke",
    "So item CTKM thieu vi tri ke",
    "So dong POG khong co CTKM phu hop",
    "So item POG khong co CTKM phu hop",
}


def _summary_from_sheet(ws) -> dict:
    result = {}
    for row in range(2, (ws.max_row or 0) + 1):
        key = ws.cell(row=row, column=1).value
        if key is None or not str(key).strip():
            continue
        result[str(key)] = ws.cell(row=row, column=2).value
    return result


def _summary_row_indexes(ws) -> dict[str, int]:
    result = {}
    for row in range(2, (ws.max_row or 0) + 1):
        key = ws.cell(row=row, column=1).value
        if key is not None and str(key).strip():
            result[str(key)] = row
    return result


def _same_value(expected, actual) -> bool:
    if expected is None and actual is None:
        return True
    if isinstance(expected, Number) and isinstance(actual, Number):
        return float(expected) == float(actual)
    if isinstance(expected, (date, datetime)) and isinstance(actual, (date, datetime)):
        return expected == actual
    return str(expected) == str(actual)


def _workbook_profile(workbook) -> str:
    names = set(workbook.sheetnames)
    if _FULL_REQUIRED_CTKM_SHEETS.issubset(names):
        return "full"
    if _SLIM_REQUIRED_CTKM_SHEETS.issubset(names):
        return "production-slim"
    missing_slim = sorted(_SLIM_REQUIRED_CTKM_SHEETS - names)
    raise RuntimeError(
        "Workbook kết quả thiếu sheet bắt buộc cho profile vận hành: "
        + ", ".join(missing_slim)
    )


_FIELD_ALIASES = {
    "task_id": ("task_id", "Mã công việc (task_id)"),
    "pog_position_id": ("pog_position_id", "Mã vị trí POG ổn định (pog_position_id)"),
    "source_row_id": ("source_row_id", "Dòng nguồn chính (source_row_id)"),
    "source_row_ids": ("source_row_ids", "Các dòng nguồn (source_row_ids)"),
    "logical_row_id": ("logical_row_id",),
    "scheduled_action": ("scheduled_action", "Hành động theo lịch (scheduled_action)"),
    "review_required": ("review_required", "Cần review?"),
    "in_report_scope": ("in_report_scope",),
}


def _normalize_text(value) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value).strip().lower())
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def _truthy_cell(value) -> bool:
    return value is True or _normalize_text(value) in {"true", "1", "yes", "co", "có"}


def _normalized_cell(value):
    if value is None:
        return ("missing", "")
    if isinstance(value, (date, datetime)):
        return ("date", value.isoformat())
    if isinstance(value, Number):
        return ("number", float(value))
    return ("text", str(value))


def _worksheet_records(ws) -> list[dict[str, object]]:
    """Read worksheet records sequentially so production-slim validation can stream."""
    iterator = ws.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration:
        return []
    headers = ["" if value is None else str(value) for value in header_row]
    rows = []
    for row in iterator:
        values = list(row)
        if any(
            value is not None and str(value).strip()
            for value in values
        ):
            rows.append(values)
    if is_empty_placeholder(headers, rows):
        return []
    return [dict(zip(headers, row)) for row in rows]


def _record_value(record: dict[str, object], field: str):
    for header in _FIELD_ALIASES.get(field, (field,)):
        if header in record:
            return record.get(header)
    return None


def _lineage_parts(value) -> list[str]:
    if value is None:
        return []
    return [part.strip() for part in str(value).split("|") if part.strip()]


def _scoped_work_records(audit_records: list[dict[str, object]]) -> list[dict[str, object]]:
    scoped = []
    for record in audit_records:
        scope = _record_value(record, "in_report_scope")
        if scope is not None and not _truthy_cell(scope):
            continue
        scoped.append(record)

    groups: dict[tuple, list[dict[str, object]]] = {}
    for position, record in enumerate(scoped):
        task = _record_value(record, "task_id")
        task = "" if task is None else str(task).strip()
        lineage = []
        lineage.extend(_lineage_parts(_record_value(record, "source_row_ids")))
        lineage.extend(_lineage_parts(_record_value(record, "source_row_id")))
        lineage = list(dict.fromkeys(lineage))
        key = ("task_source", task, tuple(lineage)) if task and lineage else ("row", position)
        groups.setdefault(key, []).append(record)

    return [records[0] for records in groups.values()]


def _projection_counter(records: list[dict[str, object]], headers: list[str]) -> Counter:
    return Counter(
        tuple(_normalized_cell(record.get(header)) for header in headers)
        for record in records
    )


def _action_identity(record: dict[str, object]) -> tuple:
    task = _record_value(record, "task_id")
    logical = _record_value(record, "logical_row_id")
    source = _record_value(record, "source_row_id")
    position = _record_value(record, "pog_position_id")
    return tuple(
        _normalized_cell(value)
        for value in (task, logical if logical not in (None, "") else source, position)
    )


def _validate_export_completeness(workbook) -> None:
    required = {"Audit hành động", "Gắn tem", "Tháo tem", "Chờ gắn", "Cần review"}
    if not required.issubset(workbook.sheetnames):
        return

    audit_records = _worksheet_records(workbook["Audit hành động"])
    work_records = _scoped_work_records(audit_records)

    if "TỔNG HỢP - TEM" in workbook.sheetnames:
        total_ws = workbook["TỔNG HỢP - TEM"]
        total_records = _worksheet_records(total_ws)
        total_headers = [
            str(total_ws.cell(row=1, column=index).value or "")
            for index in range(1, (total_ws.max_column or 0) + 1)
        ]
        audit_headers = set(audit_records[0].keys()) if audit_records else set()
        common_headers = [header for header in total_headers if header in audit_headers]

        # A same-task/same-lineage collapse is only legitimate when all user-facing
        # business values agree. Otherwise the exporter would silently choose one.
        scoped_raw = [
            record for record in audit_records
            if _record_value(record, "in_report_scope") is None
            or _truthy_cell(_record_value(record, "in_report_scope"))
        ]
        grouped: dict[tuple, list[dict[str, object]]] = {}
        for position, record in enumerate(scoped_raw):
            task = str(_record_value(record, "task_id") or "").strip()
            lineage = list(dict.fromkeys([
                *_lineage_parts(_record_value(record, "source_row_ids")),
                *_lineage_parts(_record_value(record, "source_row_id")),
            ]))
            key = ("task_source", task, tuple(lineage)) if task and lineage else ("row", position)
            grouped.setdefault(key, []).append(record)
        for key, records in grouped.items():
            if len(records) <= 1 or not common_headers:
                continue
            if len(_projection_counter(records, common_headers)) != 1:
                raise RuntimeError(
                    "Export completeness invariant failed; cùng task/lineage nhưng "
                    f"dữ liệu nghiệp vụ khác nhau: {key}"
                )

        if len(total_records) != len(work_records):
            raise RuntimeError(
                "Export completeness invariant failed; TỔNG HỢP - TEM không đủ "
                f"work rows: total={len(total_records)}, expected={len(work_records)}"
            )
        if common_headers and (
            _projection_counter(total_records, common_headers)
            != _projection_counter(work_records, common_headers)
        ):
            raise RuntimeError(
                "Export completeness invariant failed; TỔNG HỢP - TEM lệch dữ liệu "
                "so với Audit hành động in-scope."
            )

    scheduled_groups = {
        "Gắn tem": {"gan tem", "gan tem fefo"},
        "Tháo tem": {"thao tem", "thao tem fefo", "thao rem fefo"},
        "Chờ gắn": {"cho gan tem"},
    }
    for sheet_name, accepted in scheduled_groups.items():
        expected = [
            record for record in work_records
            if _normalize_text(_record_value(record, "scheduled_action")) in accepted
        ]
        actual = _worksheet_records(workbook[sheet_name])
        if Counter(map(_action_identity, actual)) != Counter(map(_action_identity, expected)):
            raise RuntimeError(
                f"Export completeness invariant failed; sheet '{sheet_name}' "
                "không khớp scheduled_action trong Audit hành động."
            )

    review_expected = []
    for record in work_records:
        review_value = _record_value(record, "review_required")
        if review_value is not None:
            needs_review = _truthy_cell(review_value)
        else:
            needs_review = _normalize_text(_record_value(record, "scheduled_action")) == "review"
        if needs_review:
            review_expected.append(record)
    review_actual = _worksheet_records(workbook["Cần review"])
    if Counter(map(_action_identity, review_actual)) != Counter(map(_action_identity, review_expected)):
        raise RuntimeError(
            "Export completeness invariant failed; sheet 'Cần review' không khớp "
            "review_required trong Audit hành động."
        )


def validate_ctkm_workbook(path: str | Path, summary: dict | None) -> None:
    """Reject an unreadable/inconsistent full or production-slim workbook."""
    artifact = Path(path)
    if not artifact.is_file():
        raise RuntimeError("Pipeline không tạo được file Excel kết quả.")

    workbook = None
    try:
        # Read production-slim artifacts in streaming mode. Full workbooks still
        # need mutable row/column metadata checks, so they are reopened normally.
        workbook = openpyxl.load_workbook(artifact, data_only=True, read_only=True)
        profile = _workbook_profile(workbook)
        if profile == "full":
            workbook.close()
            workbook = openpyxl.load_workbook(artifact, data_only=True, read_only=False)
        if not workbook.sheetnames or workbook.sheetnames[0] != "HƯỚNG DẪN":
            raise RuntimeError("Workbook kết quả không đặt HƯỚNG DẪN ở sheet đầu tiên.")

        if profile == "full" and workbook["Mapping log"].sheet_state != "hidden":
            raise RuntimeError("Mapping log phải được ẩn trong workbook kết quả.")
        if "POG không có CTKM" in workbook.sheetnames:
            if workbook["POG không có CTKM"].sheet_state != "hidden":
                raise RuntimeError(
                    "POG không có CTKM là compatibility sheet và phải được ẩn."
                )
        for metadata_name in ("POG metadata", "KM metadata"):
            if metadata_name in workbook.sheetnames and workbook[metadata_name].sheet_state != "hidden":
                raise RuntimeError(f"{metadata_name} phải được ẩn trong workbook kết quả.")

        _validate_export_completeness(workbook)

        summary_dict = summary if isinstance(summary, dict) else {}
        if profile == "full":
            summary_ws = workbook["Tom tat"]
            workbook_summary = _summary_from_sheet(summary_ws)
            for key in _SUMMARY_KEYS_TO_VERIFY:
                if key not in summary_dict:
                    continue
                if key not in workbook_summary:
                    raise RuntimeError(f"Tom tat thiếu chỉ số '{key}'.")
                if not _same_value(summary_dict[key], workbook_summary[key]):
                    raise RuntimeError(
                        f"Tom tat không khớp summary cho '{key}': "
                        f"API={summary_dict[key]!r}, Excel={workbook_summary[key]!r}"
                    )

            summary_rows = _summary_row_indexes(summary_ws)
            for legacy_key in _LEGACY_HIDDEN_SUMMARY_KEYS:
                row = summary_rows.get(legacy_key)
                if row is not None and not bool(summary_ws.row_dimensions[row].hidden):
                    raise RuntimeError(
                        f"Tom tat phải ẩn legacy metric '{legacy_key}' để tránh nhãn sai nghĩa."
                    )

        for legacy_key, canonical_key in _LEGACY_CANONICAL_PAIRS:
            if legacy_key in summary_dict and canonical_key in summary_dict:
                if not _same_value(summary_dict[legacy_key], summary_dict[canonical_key]):
                    raise RuntimeError(
                        f"Summary legacy '{legacy_key}' lệch canonical '{canonical_key}'."
                    )

        for summary_key, sheet_name in _SUMMARY_SHEET_COUNT_PAIRS:
            if summary_key not in summary_dict or sheet_name not in workbook.sheetnames:
                continue
            sheet_count = (
                len(_worksheet_records(workbook[sheet_name]))
                if profile == "production-slim"
                else visible_data_row_count(workbook[sheet_name])
            )
            if not _same_value(summary_dict[summary_key], sheet_count):
                raise RuntimeError(
                    f"Summary '{summary_key}' không khớp sheet '{sheet_name}': "
                    f"API={summary_dict[summary_key]!r}, rows={sheet_count}"
                )

        guide_targets = _FULL_GUIDE_TARGETS if profile == "full" else _SLIM_GUIDE_TARGETS
        guide = workbook["HƯỚNG DẪN"]
        seen_targets: set[str] = set()
        for row in range(1, (guide.max_row or 0) + 1):
            label = guide.cell(row=row, column=1).value
            if label not in guide_targets:
                continue
            target = guide_targets[label]
            expected_count = (
                len(_worksheet_records(workbook[target]))
                if profile == "production-slim"
                else visible_data_row_count(workbook[target])
            )
            actual_count = guide.cell(row=row, column=2).value
            if not _same_value(expected_count, actual_count):
                raise RuntimeError(
                    f"HƯỚNG DẪN không khớp sheet '{target}': "
                    f"card={actual_count!r}, rows={expected_count}"
                )
            seen_targets.add(target)

        missing_cards = sorted(set(guide_targets.values()) - seen_targets)
        if missing_cards:
            raise RuntimeError(
                "HƯỚNG DẪN thiếu card cho sheet: " + ", ".join(missing_cards)
            )
    except RuntimeError:
        raise
    except Exception as exc:
        raise RuntimeError(f"Workbook kết quả không đọc được: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()


def install_final_output_contract(core_module) -> None:
    if getattr(core_module.app.state, "final_output_contract_v1", False):
        return

    original_run_pipeline = core_module.run_pipeline

    def validated_run_pipeline(*args, **kwargs):
        result = original_run_pipeline(*args, **kwargs)
        validate_ctkm_workbook(result.output_filepath, result.summary)
        return result

    core_module.run_pipeline = validated_run_pipeline

    from backend import preview_contract

    preview_contract._visible_column_indexes = visible_column_indexes
    preview_contract._visible_nonempty_rows = visible_nonempty_rows
    preview_contract._is_empty_placeholder = is_empty_placeholder

    core_module.app.state.final_output_contract_v1 = True
